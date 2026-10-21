import csv
import io
import json
import math
import random
import re
from datetime import datetime, date

from openpyxl import load_workbook

from core.workflow import transition_task

# 카탈로그 QC 샘플 검수 5~10% 기준
DEFAULT_SAMPLE_RATE = 10
MIN_SAMPLE = 5
PASS_ACCURACY = 95

MAX_ROWS = 5000
MAX_COLS = 50
MAX_SIZE = 10 * 1024 * 1024

STATUS_LABELS = {"open": "미처리", "done": "대기", "hold": "보류", "approved": "승인", "rejected": "반려"}


class CsvError(Exception):
    pass


def parse_csv(raw):
    if len(raw) > MAX_SIZE:
        raise CsvError("파일 크기가 10MB를 초과합니다.")
    # 국내 엑셀 기본 저장 형식 CP949 대응
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp949")
        except UnicodeDecodeError:
            raise CsvError("인코딩을 확인할 수 없습니다. UTF-8 또는 CP949로 저장해 주세요.")
    reader = csv.DictReader(io.StringIO(text))
    headers = reader.fieldnames or []
    if not headers:
        raise CsvError("헤더 행을 읽을 수 없습니다.")
    if len(headers) > MAX_COLS:
        raise CsvError(f"열이 {MAX_COLS}개를 초과합니다.")
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise CsvError("헤더에 빈 열 이름이나 중복된 열 이름이 있습니다.")
    # 열 개수 불일치 행 보정. 초과 값 제외, 부족 값 빈칸 처리
    rows = [{h: r.get(h) or "" for h in headers} for r in reader]
    if not rows:
        raise CsvError("데이터 행이 없습니다.")
    if len(rows) > MAX_ROWS:
        raise CsvError(f"행이 {MAX_ROWS}건을 초과합니다. 요청을 나눠 올려 주세요.")
    return headers, rows


def _xlsx_cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_xlsx(raw):
    if len(raw) > MAX_SIZE:
        raise CsvError("파일 크기가 10MB를 초과합니다.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header_cells = next(rows_iter, None)
    except Exception:
        raise CsvError("엑셀 파일을 열 수 없습니다. 손상되지 않았는지 확인해 주세요.")
    if not header_cells:
        raise CsvError("헤더 행을 읽을 수 없습니다.")
    headers = [_xlsx_cell(c) for c in header_cells]
    if len(headers) > MAX_COLS:
        raise CsvError(f"열이 {MAX_COLS}개를 초과합니다.")
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise CsvError("헤더에 빈 열 이름이나 중복된 열 이름이 있습니다.")
    rows = []
    for cells in rows_iter:
        if not cells or all(c is None for c in cells):
            continue
        vals = [_xlsx_cell(c) for c in cells]
        rows.append({h: (vals[i] if i < len(vals) else "") for i, h in enumerate(headers)})
    if not rows:
        raise CsvError("데이터 행이 없습니다.")
    if len(rows) > MAX_ROWS:
        raise CsvError(f"행이 {MAX_ROWS}건을 초과합니다. 요청을 나눠 올려 주세요.")
    return headers, rows


def parse_item_file(filename, raw):
    name = (filename or "").lower()
    xlsx_magic = raw[:2] == b"PK"
    if name.endswith(".xlsx"):
        if not xlsx_magic:
            raise CsvError("엑셀 파일 형식이 올바르지 않습니다.")
        return parse_xlsx(raw)
    if name.endswith(".csv"):
        if xlsx_magic:
            raise CsvError("CSV가 아닌 파일입니다. 엑셀 파일은 xlsx로 올려 주세요.")
        return parse_csv(raw)
    raise CsvError("CSV 또는 엑셀(xlsx) 파일만 올릴 수 있습니다.")


FIELD_TYPES = ("text", "select", "tags", "number", "date", "phone", "long")


def validate_value(ftype, value):
    if ftype == "number":
        return bool(re.match(r'^-?\d+(\.\d+)?$', value))
    if ftype == "date":
        try:
            datetime.strptime(value, "%Y-%m-%d")
            return True
        except ValueError:
            return False
    if ftype == "phone":
        return bool(re.match(r'^[0-9-]+$', value)) and 9 <= len(value.replace("-", "")) <= 13
    return True


def get_item_spec(row):
    raw = row["item_spec"] if row else None
    return json.loads(raw) if raw else None


def task_item_spec(conn, task_id):
    row = conn.execute(
        "SELECT wr.item_spec FROM tasks t JOIN work_requests wr ON wr.id = t.work_request_id WHERE t.id=?",
        (task_id,),
    ).fetchone()
    return get_item_spec(row)


def assign_items(conn, req_id, task_id, volume):
    if volume <= 0:
        return 0
    rows = conn.execute(
        "SELECT id FROM work_items WHERE request_id=? AND task_id IS NULL AND status='open' "
        "ORDER BY seq LIMIT ?",
        (req_id, volume),
    ).fetchall()
    ids = [r["id"] for r in rows]
    if ids:
        ph = ",".join("?" * len(ids))
        conn.execute(f"UPDATE work_items SET task_id=? WHERE id IN ({ph})", [task_id] + ids)
    return len(ids)


def release_items(conn, task_id):
    conn.execute(
        "UPDATE work_items SET task_id=NULL, status='open', hold_note='' "
        "WHERE task_id=? AND status IN ('open','hold')",
        (task_id,),
    )


def run_sample(conn, task_id, rate, resubmit=False):
    rate = rate or DEFAULT_SAMPLE_RATE
    total_assigned, rnd = conn.execute(
        "SELECT COUNT(*), COALESCE(MAX(round), 0) + 1 FROM work_items WHERE task_id=?", (task_id,)
    ).fetchone()

    if resubmit:
        # 반려 건, 보류 후 되돌린 건은 전수 재검수
        conn.execute(
            "UPDATE work_items SET sampled=1, round=? WHERE task_id=? AND status='done' "
            "AND (review_note != '' OR hold_note != '')",
            (rnd, task_id),
        )
        pool = [r["id"] for r in conn.execute(
            "SELECT id FROM work_items WHERE task_id=? AND status='done' "
            "AND COALESCE(review_note, '') = '' AND COALESCE(hold_note, '') = ''",
            (task_id,),
        ).fetchall()]
    else:
        pool = [r["id"] for r in conn.execute(
            "SELECT id FROM work_items WHERE task_id=? AND status='done'", (task_id,)
        ).fetchall()]

    if not pool:
        return
    if total_assigned < MIN_SAMPLE:
        size = len(pool)
    else:
        size = min(len(pool), max(MIN_SAMPLE, math.ceil(len(pool) * rate / 100)))
    chosen = random.sample(pool, size)
    ph = ",".join("?" * len(chosen))
    conn.execute(f"UPDATE work_items SET sampled=1, round=? WHERE id IN ({ph})", [rnd] + chosen)


def review_cnt(conn, task_id):
    return conn.execute(
        "SELECT COUNT(*) FROM work_items WHERE task_id=? AND (status='hold' OR (status='done' AND sampled=1))",
        (task_id,),
    ).fetchone()[0]


def finalize_review(conn, task_id, user_id):
    # 이전 회차 승인 건 제외. 이번 제출 회차 기준 정확도 산정
    sample_rows = conn.execute(
        "SELECT status FROM work_items WHERE task_id=? AND sampled=1 AND status IN ('approved','rejected') "
        "AND round = (SELECT MAX(round) FROM work_items WHERE task_id=?)",
        (task_id, task_id),
    ).fetchall()
    total_s = len(sample_rows)
    ok_s = sum(1 for r in sample_rows if r["status"] == "approved")
    accuracy = 100.0 if total_s == 0 else ok_s / total_s * 100
    reopened = conn.execute(
        "SELECT COUNT(*) FROM work_items WHERE task_id=? AND status='open'", (task_id,)
    ).fetchone()[0]

    if accuracy < PASS_ACCURACY or reopened:
        transition_task(conn, task_id, 'reject', user_id, 'client', f'샘플 정확도 {accuracy:.0f}%')
        conn.execute("UPDATE work_items SET status='open' WHERE task_id=? AND status='rejected'", (task_id,))
    else:
        transition_task(conn, task_id, 'approve', user_id, 'client')
        conn.execute("UPDATE work_items SET status='approved' WHERE task_id=? AND status='done'", (task_id,))
    conn.commit()
    return accuracy


def export_rows(conn, req_id, item_spec):
    show = item_spec.get("show", [])
    fields = item_spec.get("fields", [])
    headers = ["seq"] + show + [f["label"] for f in fields] + ["검수결과"]
    items = conn.execute(
        "SELECT * FROM work_items WHERE request_id=? ORDER BY seq", (req_id,)
    ).fetchall()
    rows = []
    for it in items:
        src = json.loads(it["source"]) if it["source"] else {}
        res = json.loads(it["result"]) if it["result"] else {}
        row = [it["seq"]] + [src.get(c, "") for c in show]
        for f in fields:
            v = res.get(f["key"], "")
            row.append(", ".join(v) if isinstance(v, list) else v)
        row.append(STATUS_LABELS.get(it["status"], it["status"]))
        rows.append(row)
    return headers, rows
