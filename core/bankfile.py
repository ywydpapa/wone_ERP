import csv
import hashlib
import io
import pathlib
from datetime import date, datetime

from openpyxl import load_workbook

# 프리셋별 필드 헤더명 후보. date, amt_in, amt_out 후보가 모두 매칭되면 그 프리셋으로 인식
PRESETS = {
    "kb": {"date": ["거래일시"], "memo": ["적요"], "party": ["보낸분/받는분"],
           "amt_out": ["출금액"], "amt_in": ["입금액"], "balance": ["잔액"]},
    "shinhan": {"date": ["거래일자"], "memo": ["적요"], "party": ["내용"],
                "amt_out": ["출금"], "amt_in": ["입금"], "balance": ["잔액"]},
    "woori": {"date": ["거래일시"], "memo": ["기재내용"],
              "amt_out": ["찾으신금액"], "amt_in": ["맡기신금액"], "balance": ["잔액"]},
    "hana": {"date": ["거래일시"], "memo": ["적요"],
             "amt_out": ["출금"], "amt_in": ["입금"], "balance": ["잔액"]},
    "nh": {"date": ["거래일자"], "memo": ["거래내용"],
           "amt_out": ["출금금액"], "amt_in": ["입금금액"], "balance": ["거래후잔액"]},
}

REQUIRED_FIELDS = ("date", "amt_in", "amt_out")
DT_FMTS = ("%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S",
           "%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d")


def _norm_header(v):
    s = "" if v is None else str(v).strip()
    for ch in (" ", "\u00b7", "/"):
        s = s.replace(ch, "")
    return s


def read_rows(data, filename):
    ext = pathlib.Path(filename).suffix.lower()
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp949")
        return list(csv.reader(io.StringIO(text)))
    raise ValueError(f"지원하지 않는 파일 형식: {filename}")


def find_header(rows, fields):
    for idx, row in enumerate(rows[:20]):
        cells = {}
        for col, v in enumerate(row):
            norm = _norm_header(v)
            if norm and norm not in cells:
                cells[norm] = col
        cmap = {}
        for field, cands in fields.items():
            for cand in cands:
                col = cells.get(_norm_header(cand))
                if col is not None:
                    cmap[field] = col
                    break
        if all(f in cmap for f in REQUIRED_FIELDS):
            return idx, cmap
    return None


def _cell(row, col):
    if col is None or col >= len(row):
        return None
    return row[col]


def _norm_dt(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        if v.hour or v.minute or v.second:
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s:
        return None
    for fmt in DT_FMTS:
        try:
            dt = datetime.strptime(s, fmt)
        except ValueError:
            continue
        return dt.strftime("%Y-%m-%d %H:%M:%S") if "%H" in fmt else dt.strftime("%Y-%m-%d")
    return None


def _norm_amt(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).replace(",", "").replace(" ", "").strip()
    if s in ("", "-"):
        return 0
    try:
        return int(s)
    except ValueError:
        return 0


def parse(rows, cmap):
    txns = []
    for row in rows:
        at = _norm_dt(_cell(row, cmap.get("date")))
        if at is None:
            continue
        party = str(_cell(row, cmap.get("party")) or "").strip()
        memo = str(_cell(row, cmap.get("memo")) or "").strip() if "memo" in cmap else party
        txns.append({
            "txn_at": at, "memo": memo, "party": party,
            "amt_in": _norm_amt(_cell(row, cmap.get("amt_in"))),
            "amt_out": _norm_amt(_cell(row, cmap.get("amt_out"))),
            "balance": _norm_amt(_cell(row, cmap.get("balance"))) if "balance" in cmap else None,
        })
    return txns


def txn_hash(acct_id, t):
    key = f"{acct_id}|{t['txn_at']}|{t['amt_in']}|{t['amt_out']}|{t['memo']}|{t['balance']}"
    return hashlib.sha1(key.encode()).hexdigest()
