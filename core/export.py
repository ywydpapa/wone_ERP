import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from fastapi.responses import StreamingResponse


# 공통 스타일
HEADER_FONT = Font(name="맑은 고딕", bold=True, size=10)
HEADER_FILL = PatternFill(start_color="E8EEF4", end_color="E8EEF4", fill_type="solid")
CELL_FONT = Font(name="맑은 고딕", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def styled_workbook(title, headers, rows, col_widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title

    # 헤더
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # 데이터
    for ri, row in enumerate(rows, 2):
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"

    # 열 너비
    if col_widths:
        for ci, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    return wb


def excel_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def pdf_response(html_str, filename):
    from weasyprint import HTML
    buf = io.BytesIO()
    HTML(string=html_str).write_pdf(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
