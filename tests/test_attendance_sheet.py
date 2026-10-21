import io

from openpyxl import load_workbook


def _emp_id(db, user_id):
    return db.execute("SELECT id FROM employees WHERE user_id=?", (user_id,)).fetchone()[0]


def _add_log(db, emp_id, work_date, field, old, new, reason, editor_id):
    att_id = db.execute(
        "INSERT INTO attendance (employee_id, work_date, status, clock_in, clock_out) "
        "VALUES (?, ?, 'normal', '09:00:00', '18:00:00')",
        (emp_id, work_date),
    ).lastrowid
    db.execute(
        "INSERT INTO attendance_logs (attendance_id, changed_by, field, old_value, new_value, reason) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (att_id, editor_id, field, old, new, reason),
    )
    return att_id


def test_export_excel_has_change_log_sheet(db, seed, login):
    emp_a = _emp_id(db, seed["a"]["wk"])
    staff_id = seed["staff"]
    _add_log(db, emp_a, "2026-03-05", "clock_in", "09:30", "09:00", "지각 정정", staff_id)
    db.commit()

    c = login("t_staff")
    r = c.get("/hr/attendance/sheet/export?year=2026&month=3&company_id=" + str(seed["a"]["company"]))
    assert r.status_code == 200

    wb = load_workbook(io.BytesIO(r.content))
    assert "수정이력" in wb.sheetnames
    ws = wb["수정이력"]
    header = [c.value for c in ws[1]]
    assert header == ["일자", "사번", "성명", "항목", "변경 전", "변경 후", "사유", "수정자", "수정일시"]
    row = [c.value for c in ws[2]]
    assert row[0] == "2026-03-05"
    assert row[3] == "출근시각"
    assert row[4] == "09:30"
    assert row[5] == "09:00"
    assert row[6] == "지각 정정"


def test_export_change_log_scoped_to_company(db, seed, login):
    emp_a = _emp_id(db, seed["a"]["wk"])
    emp_b = _emp_id(db, seed["b"]["wk"])
    staff_id = seed["staff"]
    _add_log(db, emp_a, "2026-03-05", "status", "late", "normal", "A사 수정", staff_id)
    _add_log(db, emp_b, "2026-03-06", "status", "late", "normal", "B사 수정", staff_id)
    db.commit()

    c = login("t_staff")
    r = c.get("/hr/attendance/sheet/export?year=2026&month=3&company_id=" + str(seed["a"]["company"]))
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["수정이력"]
    reasons = [row[6].value for row in ws.iter_rows(min_row=2)]
    assert "A사 수정" in reasons
    assert "B사 수정" not in reasons


def test_export_pdf_includes_change_log(db, seed, login):
    emp_a = _emp_id(db, seed["a"]["wk"])
    staff_id = seed["staff"]
    _add_log(db, emp_a, "2026-03-05", "overtime_minutes", "0", "60", "연장근무 반영", staff_id)
    db.commit()

    c = login("t_staff")
    r = c.get("/hr/attendance/sheet/export?year=2026&month=3&company_id=" +
              str(seed["a"]["company"]) + "&format=pdf")
    assert r.status_code == 200
    assert b"%PDF" in r.content[:10]
