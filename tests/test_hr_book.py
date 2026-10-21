import io
import sqlite3

from openpyxl import load_workbook

from core.books import INTERNAL
from migrate import run_migrations
from routers.hr_book import _load_ledger
from tests.conftest import csrf


def _wone_emp(db, name="원직원", employee_no="W1"):
    # book 0(원 내부) 직원은 company_id NULL
    return db.execute(
        "INSERT INTO employees (name, company_id, employee_no, hire_date, status) "
        "VALUES (?, NULL, ?, '2024-01-01', 'active')",
        (name, employee_no),
    ).lastrowid


def _delegate(db, company_id, on=True):
    db.execute("UPDATE client_companies SET payroll_delegated=? WHERE id=?", (1 if on else 0, company_id))
    db.commit()


# 마이그레이션 076

def test_migration_076_preserves_ids_and_fk(tmp_path):
    path = str(tmp_path / "pre076.db")
    run_migrations(path, upto=75)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    cid = conn.execute("INSERT INTO client_companies (name) VALUES ('이관테스트')").lastrowid
    conn.execute(
        "INSERT INTO pay_items (id, company_id, code, name) VALUES (700, ?, 'BASE', '기본급')", (cid,)
    )
    emp_id = conn.execute(
        "INSERT INTO employees (name, company_id, hire_date) VALUES ('직원', ?, '2024-01-01')", (cid,)
    ).lastrowid
    ps_id = conn.execute(
        "INSERT INTO payslips (employee_id, pay_year, pay_month, gross_pay, total_deduction, net_pay) "
        "VALUES (?, 2026, 1, 1000, 0, 1000)",
        (emp_id,),
    ).lastrowid
    conn.execute(
        "INSERT INTO payslip_items (id, payslip_id, item_id, amount) VALUES (800, ?, 700, 1000)", (ps_id,)
    )
    conn.commit()
    conn.close()

    run_migrations(path, upto=76)

    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    assert conn.execute("SELECT company_id FROM pay_items WHERE id=700").fetchone()["company_id"] == cid
    assert conn.execute("SELECT item_id FROM payslip_items WHERE id=800").fetchone()["item_id"] == 700
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []

    conn.execute("INSERT INTO pay_items (company_id, code, name) VALUES (0, 'BASE', '기본급')")
    conn.commit()
    assert conn.execute("SELECT 1 FROM pay_items WHERE company_id=0").fetchone()
    conn.close()


def test_client_companies_payroll_delegated_default_0(db):
    cid = db.execute("INSERT INTO client_companies (name) VALUES ('테스트')").lastrowid
    db.commit()
    row = db.execute("SELECT payroll_delegated FROM client_companies WHERE id=?", (cid,)).fetchone()
    assert row["payroll_delegated"] == 0


# book 0(원 내부) 범위

def test_staff_default_book_sees_only_null_company_employees(db, seed, login):
    wone_id = _wone_emp(db)
    db.commit()

    c = login("t_staff")

    html = c.get("/payroll").text
    assert "원직원" in html
    assert "가테스트근로자" not in html

    html = c.get("/payroll/ledger").text
    assert "원직원" in html
    assert "가테스트근로자" not in html

    html = c.get("/withholding").text
    assert "원" in html  # 내부 book 라벨

    html = c.get("/roster").text
    assert "원직원" in html
    assert "가테스트근로자" not in html

    html = c.get("/severance/estimate").text
    assert "퇴직금 추계액" in html


def test_staff_switch_to_non_delegated_company_stays_zero(db, seed, login):
    a = seed["a"]["company"]
    c = login("t_staff")
    c.post("/hr-book", data={"book": a, "next": "/payroll"})
    r = c.get("/payroll")
    assert "가테스트근로자" not in r.text


def test_staff_switch_to_delegated_company_sees_it(db, seed, login):
    a = seed["a"]["company"]
    _delegate(db, a)
    c = login("t_staff")
    c.post("/hr-book", data={"book": a, "next": "/payroll"})
    r = c.get("/payroll")
    assert "가테스트근로자" in r.text


def test_delegation_turned_off_resets_to_zero_next_request(db, seed, login):
    a = seed["a"]["company"]
    _delegate(db, a)
    c = login("t_staff")
    c.post("/hr-book", data={"book": a, "next": "/payroll"})
    assert "가테스트근로자" in c.get("/payroll").text

    _delegate(db, a, on=False)
    r = c.get("/payroll")
    assert "가테스트근로자" not in r.text


def test_staff_book_zero_404_on_client_employee(db, seed, login):
    a = seed["a"]["company"]
    emp_a = db.execute("SELECT id FROM employees WHERE user_id=?", (seed["a"]["wk"],)).fetchone()["id"]
    c = login("t_staff")
    r = c.get(f"/payroll/{emp_a}")
    assert r.status_code == 404


# client 범위

def test_client_hr_sees_own_company_only(db, seed, login):
    b = seed["b"]["company"]
    emp_b = db.execute("SELECT id FROM employees WHERE user_id=?", (seed["b"]["wk"],)).fetchone()["id"]
    c = login("a_mgr")
    r = c.get(f"/payroll/{emp_b}")
    assert r.status_code == 404
    r2 = c.post(f"/payroll/{emp_b}", data={"year": "2026", "month": "1"})
    assert r2.status_code == 404
    r3 = c.get(f"/severance/{emp_b}")
    assert r3.status_code == 404
    r4 = c.get(f"/withholding/employees/{emp_b}")
    assert r4.status_code == 404


def test_client_ops_redirected_from_payroll(db, seed, login):
    db.execute("UPDATE client_users SET client_role='ops' WHERE user_id=?", (seed["b"]["mgr"],))
    db.commit()
    c = login("b_mgr")
    r = c.get("/payroll", follow_redirects=False)
    assert r.status_code == 303


def test_worker_redirected_from_payroll(db, seed, login):
    c = login("a_wk")
    r = c.get("/payroll", follow_redirects=False)
    assert r.status_code == 303


# 삭제된 구 URL

def test_old_urls_gone(db, seed, login):
    c = login("a_mgr")
    assert c.get("/client/payroll", follow_redirects=False).status_code == 404
    assert c.get("/client/withholding", follow_redirects=False).status_code == 404

    cs = login("t_staff")
    assert cs.get("/hr/payroll/ledger", follow_redirects=False).status_code == 404
    assert cs.get("/hr/withholding", follow_redirects=False).status_code == 404


# /hr-book open-redirect 방지

def test_hr_book_rejects_open_redirect(db, seed, login):
    c = login("t_staff")
    r = c.post("/hr-book", data={"book": "0", "next": "//evil.com"}, follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"] == "/payroll"

    r2 = c.post("/hr-book", data={"book": "0", "next": "http://evil.com"}, follow_redirects=False)
    assert r2.status_code == 303
    assert r2.headers["location"] == "/payroll"


# 플랫폼 고객사 수정 - 급여 대행 위탁 체크박스

def _company_form_data(name="체크박스테스트"):
    return {
        "name": name, "business_no": "", "representative": "", "biz_type": "", "biz_item": "",
        "address": "", "contact_name": "", "contact_phone": "", "contact_email": "",
        "contract_start": "", "contract_end": "", "contract_amount": "", "contract_content": "",
        "status": "active", "start_time": "09:00", "end_time": "18:00", "break_min": "60",
        "rest_day": "6", "off_day": "5", "employee_count": "0",
    }


def test_platform_company_edit_saves_delegation_checkbox(db, seed, login):
    a = seed["a"]["company"]
    c = login("t_staff")

    data = _company_form_data()
    data["payroll_delegated"] = "on"
    c.post(f"/platform/companies/{a}", data=data, follow_redirects=False)
    assert db.execute("SELECT payroll_delegated FROM client_companies WHERE id=?", (a,)).fetchone()["payroll_delegated"] == 1

    data2 = _company_form_data()
    c.post(f"/platform/companies/{a}", data=data2, follow_redirects=False)
    assert db.execute("SELECT payroll_delegated FROM client_companies WHERE id=?", (a,)).fetchone()["payroll_delegated"] == 0


# 직원 화면에서 급여 마스터 필드 제거

def test_hr_employee_edit_no_base_salary(db, seed, login):
    emp_a = db.execute("SELECT id FROM employees WHERE user_id=?", (seed["a"]["wk"],)).fetchone()["id"]
    c = login("t_staff")
    html = c.get(f"/hr/employees/{emp_a}/edit").text
    assert 'name="base_salary"' not in html
    assert 'name="tax_rate"' not in html
    assert 'name="bank_account"' not in html


# 급여대장 야간/휴일근로시간 분리

def test_ledger_export_splits_night_and_holiday_hours(db):
    emp_id = _wone_emp(db, "야휴테스트", "W9")
    # 2026-03-09(월) 21~23시 근무 -> 야간 1시간
    db.execute(
        "INSERT INTO attendance (employee_id, work_date, status, clock_in, clock_out, work_minutes) "
        "VALUES (?, '2026-03-09', 'normal', '21:00:00', '23:00:00', 120)",
        (emp_id,),
    )
    # 2026-03-01(일, 휴무일) 9~18시 근무 -> 휴일 8시간
    db.execute(
        "INSERT INTO attendance (employee_id, work_date, status, clock_in, clock_out, work_minutes) "
        "VALUES (?, '2026-03-01', 'normal', '09:00:00', '18:00:00', 480)",
        (emp_id,),
    )
    db.commit()

    rows, att_map = _load_ledger(db, 2026, 3, INTERNAL)
    row = next(r for r in rows if r["id"] == emp_id)
    att = att_map[emp_id]
    night_hours = att["night_mins"] // 60
    holiday_hours = att["holiday_mins"] // 60
    assert night_hours == 1
    assert holiday_hours == 8


def test_ledger_excel_has_night_holiday_columns(db, login):
    emp_id = _wone_emp(db, "엑셀야휴", "W10")
    db.execute(
        "INSERT INTO attendance (employee_id, work_date, status, clock_in, clock_out, work_minutes) "
        "VALUES (?, '2026-03-09', 'normal', '21:00:00', '23:00:00', 120)",
        (emp_id,),
    )
    db.execute(
        "INSERT INTO attendance (employee_id, work_date, status, clock_in, clock_out, work_minutes) "
        "VALUES (?, '2026-03-01', 'normal', '09:00:00', '18:00:00', 480)",
        (emp_id,),
    )
    db.commit()

    c = login("t_staff")
    r = c.get("/payroll/ledger/export?year=2026&month=3&reason=테스트조회")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["급여대장"]
    header = [cell.value for cell in ws[2]]
    assert "야간근로시간" in header
    assert "휴일근로시간" in header
    night_col = header.index("야간근로시간") + 1
    holiday_col = header.index("휴일근로시간") + 1
    target_row = next(row for row in ws.iter_rows(min_row=3) if row[1].value == "엑셀야휴")
    assert target_row[night_col - 1].value == 1
    assert target_row[holiday_col - 1].value == 8
