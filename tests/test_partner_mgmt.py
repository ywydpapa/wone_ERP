import io

import pytest
from openpyxl import Workbook, load_workbook

from core import credit, partner


def _slip(db, cid, uid, day, pid, debit, credit_amt, no):
    doc = db.execute(
        """INSERT INTO erp_docs (user_id, doc_type, status, slip_date, slip_no, doc_number, company_id, title)
           VALUES (?,'expense','done',?,?,?,?,'t')""",
        (uid, day, no, no, cid),
    ).lastrowid
    db.execute(
        """INSERT INTO slip_lines (doc_id, line_no, account_name, account_code, debit, credit, partner_id)
           VALUES (?,1,'외상매출금','110',?,?,?)""",
        (doc, debit, credit_amt, pid),
    )


@pytest.fixture
def ar(db, seed):
    # 9/1 매출 100만, 9/20 수금 30만. 미수 70만
    cid, uid = seed["a"]["company"], seed["a"]["mgr"]
    pid = db.execute(
        """INSERT INTO partners (code, name, company_id, partner_type, credit_limit, credit_days)
           VALUES ('C1','가고객',?,'customer',1000000,30)""", (cid,)
    ).lastrowid
    _slip(db, cid, uid, "2026-09-01", pid, 1000000, 0, "S1")
    _slip(db, cid, uid, "2026-09-20", pid, 0, 300000, "S2")
    item = db.execute(
        "INSERT INTO items (item_code, item_name, unit_price, company_id) VALUES ('I1','품목',100000,?)", (cid,)
    ).lastrowid
    db.commit()
    return {"cid": cid, "pid": pid, "item": item}


def _modes(db, cid, limit, days):
    db.execute("UPDATE client_companies SET credit_limit_mode=?, credit_days_mode=? WHERE id=?", (limit, days, cid))
    db.commit()


def test_credit_off_by_default(db, ar):
    assert credit.check(db, ar["cid"], ar["pid"], 5000000, "2026-10-08") is None


def test_credit_limit_over(db, ar):
    _modes(db, ar["cid"], "warn", "off")
    assert credit.check(db, ar["cid"], ar["pid"], 300000, "2026-10-08") is None
    r = credit.check(db, ar["cid"], ar["pid"], 400000, "2026-10-08")
    assert not r["block"]
    assert "초과금액: 100,000원" in r["msg"]


def test_credit_days_offsets_later_collection(db, ar):
    _modes(db, ar["cid"], "off", "block")
    # 기준일 9/7 이전 미수 100만, 이후 수금 30만. 부족 70만
    r = credit.check(db, ar["cid"], ar["pid"], 1, "2026-10-07")
    assert r["block"]
    assert "2026.09.07.까지의 미수금: 1,000,000원" in r["msg"]
    assert "여신 부족 미수잔액: 700,000원" in r["msg"]
    # 기준일이 매출일 이전이면 해당 없음
    assert credit.check(db, ar["cid"], ar["pid"], 1, "2026-09-30") is None


def _quote(c, ar, qty, **extra):
    return c.post("/quotations/save", data={
        "qt_id": "0", "partner_id": str(ar["pid"]), "quote_date": "2026-10-08",
        "line_nums": "1", "item_id_1": str(ar["item"]), "qty_1": str(qty), "price_1": "100000", **extra,
    }, follow_redirects=False)


def test_quote_save_warn_then_confirm(db, ar, login):
    _modes(db, ar["cid"], "warn", "off")
    c = login("a_mgr")
    r = _quote(c, ar, 3)
    assert r.status_code == 409
    assert "여신한도를 초과합니다" in r.json()["confirm"]
    assert db.execute("SELECT COUNT(*) FROM quotations").fetchone()[0] == 0
    assert _quote(c, ar, 3, credit_ok="1").status_code == 303


def test_quote_save_block_ignores_confirm(db, ar, login):
    _modes(db, ar["cid"], "block", "off")
    r = _quote(login("a_mgr"), ar, 3, credit_ok="1")
    assert r.status_code == 400
    assert "저장할 수 없습니다" in r.json()["notice"]
    assert db.execute("SELECT COUNT(*) FROM quotations").fetchone()[0] == 0


def test_partner_save_new_fields(db, seed, login):
    a, b = seed["a"]["company"], seed["b"]["company"]
    g = db.execute("INSERT INTO partner_groups (company_id, name) VALUES (?, '서울')", (a,)).lastrowid
    pg = db.execute("INSERT INTO price_groups (company_id, name, rate) VALUES (?, '도매가', -15)", (a,)).lastrowid
    other = db.execute("INSERT INTO price_groups (company_id, name, rate) VALUES (?, '타사', -50)", (b,)).lastrowid
    f = db.execute("INSERT INTO partner_fields (company_id, label, kind) VALUES (?, '계약금액', 'num')", (a,)).lastrowid
    db.commit()

    r = login("a_mgr").post("/partners/save", data={
        "name": "신규", "group_id": str(g), "sale_pg_id": str(pg), "buy_pg_id": str(other),
        "credit_limit": "5,000,000", "credit_days": "60", f"f_{f}": "1,200",
    }, follow_redirects=False)
    assert r.status_code == 303
    p = db.execute("SELECT * FROM partners WHERE name='신규'").fetchone()
    assert (p["group_id"], p["sale_pg_id"], p["buy_pg_id"]) == (g, pg, None)
    assert (p["credit_limit"], p["credit_days"]) == (5000000, 60)
    assert p["extra"] == f'{{"{f}": "1200"}}'


def test_api_partners_rates(db, seed, login):
    a = seed["a"]["company"]
    pg = db.execute("INSERT INTO price_groups (company_id, name, rate) VALUES (?, '도매가', -15)", (a,)).lastrowid
    db.execute("INSERT INTO partners (code, name, company_id, sale_pg_id) VALUES ('P1','가',?,?)", (a, pg))
    db.commit()
    rows = login("a_mgr").get("/api/partners").json()
    assert rows[0]["sale_rate"] == -15
    assert rows[0]["buy_rate"] == 0


def test_price_group_delete_resets_partners(db, seed, login):
    a = seed["a"]["company"]
    c = login("a_mgr")
    c.post("/partners/price-groups/save", data={"name": "도매가", "rate": "-15"})
    pg = db.execute("SELECT id FROM price_groups WHERE company_id=?", (a,)).fetchone()["id"]
    db.execute("INSERT INTO partners (code, name, company_id, sale_pg_id, buy_pg_id) VALUES ('P1','가',?,?,?)",
               (a, pg, pg))
    db.commit()
    c.post(f"/partners/price-groups/{pg}/delete")
    p = db.execute("SELECT sale_pg_id, buy_pg_id FROM partners WHERE code='P1'").fetchone()
    assert (p["sale_pg_id"], p["buy_pg_id"]) == (None, None)


def test_settings_other_company_group_untouched(db, seed, login):
    g = db.execute("INSERT INTO partner_groups (company_id, name) VALUES (?, '남의그룹')", (seed["b"]["company"],)).lastrowid
    db.commit()
    login("a_mgr").post(f"/partners/groups/{g}/delete")
    assert db.execute("SELECT 1 FROM partner_groups WHERE id=?", (g,)).fetchone()


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_add_and_update(db, seed, login):
    a = seed["a"]["company"]
    db.execute("INSERT INTO partner_groups (company_id, name) VALUES (?, '서울')", (a,))
    db.execute("INSERT INTO partners (code, name, company_id) VALUES ('00007','기존',?)", (a,))
    db.commit()
    data = _xlsx([
        ["거래처코드", "거래처명", "유형", "사업자번호", "거래처그룹", "여신한도"],
        ["00007", "기존수정", "고객", 2208162517, "서울", 3000000],
        [None, "신규", None, None, None, None],
    ])
    r = login("a_mgr").post("/partners/import", files={"file": ("p.xlsx", data)})
    assert "신규 1건, 수정 1건" in r.text
    old = db.execute("SELECT * FROM partners WHERE code='00007'").fetchone()
    assert (old["name"], old["partner_type"], old["biz_no"], old["credit_limit"]) == ("기존수정", "customer", "220-81-62517", 3000000)
    assert old["group_id"] is not None
    assert db.execute("SELECT code FROM partners WHERE name='신규'").fetchone()["code"] == "00008"


def test_import_error_saves_nothing(db, seed, login):
    data = _xlsx([
        ["거래처코드", "거래처명", "거래처그룹", "여신한도"],
        ["A1", "정상", None, None],
        ["A2", "", "없는그룹", "abc"],
    ])
    before = db.execute("SELECT COUNT(*) FROM partners").fetchone()[0]
    r = login("a_mgr").post("/partners/import", files={"file": ("p.xlsx", data)})
    assert "3행: 거래처명 누락, 없는 거래처그룹(없는그룹), 여신한도 숫자 오류(abc)" in r.text
    assert db.execute("SELECT COUNT(*) FROM partners").fetchone()[0] == before


def test_import_template_roundtrip(db, seed, login):
    db.execute("INSERT INTO partners (code, name, company_id, partner_type) VALUES ('00001','가',?,'both')",
               (seed["a"]["company"],))
    db.commit()
    c = login("a_mgr")
    ws = load_workbook(io.BytesIO(c.get("/partners/import/template").content)).active
    assert ws["A2"].value == "00001"
    assert ws["C2"].value == "공급/고객"
    buf = io.BytesIO()
    ws.parent.save(buf)
    r = c.post("/partners/import", files={"file": ("p.xlsx", buf.getvalue())})
    assert "신규 0건, 수정 1건" in r.text


def test_slip_entry_collection_redirects_to_settlement_new(db, ar, login):
    r = login("a_mgr").get(f"/slip_entry?partner_id={ar['pid']}&kind=수금", follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"] == f"/settlements/new?kind=%EC%88%98%EA%B8%88&partner_id={ar['pid']}"


def test_detail_lists_sales(db, ar, login):
    db.execute(
        """INSERT INTO quotations (quote_number, partner_id, quote_date, total_amount, tax_amount, company_id)
           VALUES ('QT-1',?,'2026-10-01',1000,100,?)""", (ar["pid"], ar["cid"])
    )
    db.commit()
    r = login("a_mgr").get(f"/partners/{ar['pid']}")
    assert "QT-1" in r.text
    assert "1,100" in r.text
    assert "잔여 300,000" in r.text


def test_import_bank_contact_and_biz_no(db, seed, login):
    data = _xlsx([
        ["거래처명", "사업자번호", "담당자명", "담당자이메일", "은행", "계좌번호"],
        ["계좌", "220-81-62517", "김담당", "k@x.kr", "국민", "111-222"],
        ["오류", "220-81-62518", None, None, None, None],
    ])
    r = login("a_mgr").post("/partners/import", files={"file": ("p.xlsx", data)})
    assert "3행: 사업자번호 오류(220-81-62518)" in r.text
    data = _xlsx([
        ["거래처명", "사업자번호", "담당자명", "담당자이메일", "은행", "계좌번호"],
        ["계좌", "2208162517", "김담당", "k@x.kr", "국민", "111-222"],
    ])
    login("a_mgr").post("/partners/import", files={"file": ("p.xlsx", data)})
    p = db.execute("SELECT * FROM partners WHERE name='계좌'").fetchone()
    assert p["biz_no"] == "220-81-62517"
    bank = db.execute("SELECT bank, account_no FROM partner_banks WHERE partner_id=? AND is_default=1",
                      (p["id"],)).fetchone()
    assert tuple(bank) == ("국민", "111-222")
    c = db.execute("SELECT name, email FROM partner_contacts WHERE partner_id=? AND is_primary=1",
                   (p["id"],)).fetchone()
    assert tuple(c) == ("김담당", "k@x.kr")


def test_po_skips_ar_credit(db, ar):
    _modes(db, ar["cid"], "block", "off")
    assert partner.guard(db, ar["cid"], ar["pid"], 5000000, "2026-10-08", "견적")["block"]
    assert partner.guard(db, ar["cid"], ar["pid"], 5000000, "2026-10-08", "발주", ar=False) is None


def _tax_inv(c, ar, **extra):
    return c.post("/tax-invoices", data={
        "kind": "매출", "tax_type": "과세", "evidence": "세금계산서", "issue_date": "2026-10-08",
        "partner_id": str(ar["pid"]), "acct": "401", "pay": "110",
        "line_count": "1", "item_name_1": "품목", "item_qty_1": "1", "item_price_1": "500000",
        "item_supply_1": "500000", "item_vat_1": "50000", **extra,
    }, follow_redirects=False)


def test_tax_invoice_credit_warn_then_confirm(db, ar, login):
    _modes(db, ar["cid"], "warn", "off")
    c = login("a_mgr")
    r = _tax_inv(c, ar)
    assert r.status_code == 409
    assert "저장하겠습니까?" in r.json()["confirm"]
    assert db.execute("SELECT COUNT(*) FROM tax_invoices").fetchone()[0] == 0
    assert _tax_inv(c, ar, credit_ok="1").status_code == 303


def test_tax_invoice_credit_block(db, ar, login):
    _modes(db, ar["cid"], "block", "off")
    r = _tax_inv(login("a_mgr"), ar, credit_ok="1")
    assert r.status_code == 400
    assert "저장할 수 없습니다" in r.json()["notice"]
    assert db.execute("SELECT COUNT(*) FROM tax_invoices").fetchone()[0] == 0


def test_import_without_code_column_autonumbers(db, seed, login):
    data = _xlsx([["거래처명", "유형"], ["코드없음", "고객"]])
    login("a_mgr").post("/partners/import", files={"file": ("p.xlsx", data)})
    code = db.execute("SELECT code FROM partners WHERE name='코드없음'").fetchone()[0]
    assert code.isdigit() and len(code) == 5


def test_aging_keeps_unmatched_legacy_name(db, seed, ar, login):
    doc = db.execute(
        """INSERT INTO erp_docs (user_id, doc_type, status, slip_date, slip_no, doc_number, company_id, title)
           VALUES (?,'expense','done','2026-08-01','S9','S9',?,'t')""", (seed["a"]["mgr"], ar["cid"])
    ).lastrowid
    db.execute(
        """INSERT INTO slip_lines (doc_id, line_no, account_name, account_code, debit, credit, partner)
           VALUES (?,1,'외상매출금','110',250000,0,'폐업한옛거래처')""", (doc,))
    db.commit()
    c = login("a_mgr")
    for basis in ("date", "due"):
        r = c.get(f"/ledger/aging?view=ar&basis={basis}")
        assert "폐업한옛거래처" in r.text
        assert "250,000" in r.text
