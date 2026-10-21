import io
import json
import re
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from core.data_work import (
    CsvError, parse_csv, parse_xlsx, parse_item_file, validate_value,
    run_sample, finalize_review,
)
from tests.conftest import csrf


def _csv_bytes(headers, rows, encoding="utf-8-sig"):
    lines = [",".join(headers)]
    for r in rows:
        lines.append(",".join(str(v) for v in r))
    text = "\r\n".join(lines) + "\r\n"
    return text.encode(encoding)


def _xlsx_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_csv_utf8_sig():
    raw = _csv_bytes(["상호", "주소"], [["가게1", "서울"], ["가게2", "부산"]], "utf-8-sig")
    headers, rows = parse_csv(raw)
    assert headers == ["상호", "주소"]
    assert rows[0]["상호"] == "가게1"


def test_parse_csv_cp949():
    raw = _csv_bytes(["상호", "주소"], [["가게1", "서울"]], "cp949")
    headers, rows = parse_csv(raw)
    assert headers == ["상호", "주소"]
    assert rows[0]["상호"] == "가게1"


def test_parse_csv_row_limit():
    raw = _csv_bytes(["a"], [[i] for i in range(5001)])
    with pytest.raises(CsvError):
        parse_csv(raw)


def test_parse_csv_rejects_dup_header_and_pads_short_row():
    with pytest.raises(CsvError):
        parse_csv(_csv_bytes(["상호", "상호"], [["a", "b"]]))
    headers, rows = parse_csv("상호,주소\r\n가게1\r\n가게2,서울,초과\r\n".encode("utf-8"))
    assert rows[0] == {"상호": "가게1", "주소": ""}
    assert rows[1] == {"상호": "가게2", "주소": "서울"}


def test_parse_csv_col_limit():
    headers = [f"c{i}" for i in range(51)]
    raw = _csv_bytes(headers, [["x"] * 51])
    with pytest.raises(CsvError):
        parse_csv(raw)


def _new_req(c, tok, title="데이터입력"):
    raw = _csv_bytes(["상호", "주소", "비고"], [["가게1", "서울", "x"], ["가게2", "부산", "y"], ["가게3", "대구", "z"]])
    r = c.post("/client/requests", data={
        "title": title, "task_type": "data_entry", "volume": "1",
        "priority": "normal", "output_format": "excel", "data_work": "1",
        "csrf_token": tok,
    }, files={"item_csv": ("items.csv", raw, "text/csv")}, follow_redirects=False)
    assert r.status_code == 303
    req_id = int(r.headers["location"].split("/")[3])
    return req_id


def _save_spec(c, req_id, tok, show=("상호", "주소"), sample_rate=10):
    r = c.post(f"/client/requests/{req_id}/item-spec", data={
        "show": list(show),
        "field_label": ["업종"], "field_type": ["text"], "field_options": [""], "field_required": ["1"],
        "sample_rate": str(sample_rate),
        "csrf_token": tok,
    }, follow_redirects=False)
    assert r.status_code == 303
    return r


def test_create_request_builds_work_items_show_only(db, seed, login):
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    req_id = _new_req(c, tok)

    rows = db.execute("SELECT * FROM work_items WHERE request_id=? ORDER BY seq", (req_id,)).fetchall()
    assert len(rows) == 3
    assert [r["seq"] for r in rows] == [1, 2, 3]

    # 빈 열이 없으면 표시 체크는 전부 기본 선택, 입력 항목은 빈 행 하나만 제공
    html = c.get(f"/client/requests/{req_id}/item-spec").text
    assert html.count('name="show"') == 3
    assert 'checked' in html
    assert html.count('name="field_label"') == 1

    _save_spec(c, req_id, csrf(c, f"/client/requests/{req_id}/item-spec"))

    wr = db.execute("SELECT item_spec, volume FROM work_requests WHERE id=?", (req_id,)).fetchone()
    spec = json.loads(wr["item_spec"])
    assert spec["show"] == ["상호", "주소"]
    assert wr["volume"] == 3

    src = json.loads(db.execute("SELECT source FROM work_items WHERE request_id=? AND seq=1", (req_id,)).fetchone()["source"])
    assert set(src.keys()) == {"상호", "주소"}
    assert "비고" not in src


def test_task_create_allocates_in_seq_order_and_keeps_on_unassign(db, seed, login):
    a = seed["a"]["company"]
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    req_id = _new_req(c, tok)
    _save_spec(c, req_id, csrf(c, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()

    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "1차배정", "requested_volume": "2", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]

    emp_id = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_id), "csrf_token": tok,
    }, follow_redirects=False)

    assigned = db.execute(
        "SELECT seq FROM work_items WHERE task_id=? ORDER BY seq", (task_id,)
    ).fetchall()
    assert [r["seq"] for r in assigned] == [1, 2]

    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": "", "csrf_token": tok,
    }, follow_redirects=False)
    kept = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE task_id=?", (task_id,)).fetchone()["n"]
    assert kept == 2

    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_id), "csrf_token": tok,
    }, follow_redirects=False)
    kept = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE task_id=?", (task_id,)).fetchone()["n"]
    assert kept == 2


def test_task_create_caps_volume_and_skips_erp_sync(db, seed, login):
    c = login("a_mgr")
    req_id = _new_req(c, csrf(c, "/client/requests/new"))
    _save_spec(c, req_id, csrf(c, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()

    c.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "초과", "requested_volume": "10", "erp_module": "slip_purchase",
        "csrf_token": csrf(c, f"/client/requests/{req_id}"),
    }, follow_redirects=False)
    t = db.execute("SELECT id, requested_volume, erp_module FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()
    assert t["requested_volume"] == 3
    assert t["erp_module"] == ""

    r = c.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "남은 건 없음", "requested_volume": "1",
        "csrf_token": csrf(c, f"/client/requests/{req_id}"),
    }, follow_redirects=False)
    assert "no_items_left" in r.headers["location"]
    assert db.execute("SELECT COUNT(*) AS n FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["n"] == 1


def test_submit_blocked_until_item_spec_and_cancel_purges_rows(db, seed, login):
    c = login("a_mgr")
    req_id = _new_req(c, csrf(c, "/client/requests/new"))

    r = c.post(f"/client/requests/{req_id}/submit", data={
        "csrf_token": csrf(c, f"/client/requests/{req_id}"),
    }, follow_redirects=False)
    assert r.headers["location"].endswith("/item-spec")
    assert db.execute("SELECT status FROM work_requests WHERE id=?", (req_id,)).fetchone()["status"] == "draft"

    c.post(f"/client/requests/{req_id}/cancel", data={
        "csrf_token": csrf(c, f"/client/requests/{req_id}"),
    }, follow_redirects=False)
    left = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE request_id=?", (req_id,)).fetchone()["n"]
    assert left == 0


def test_cancel_releases_items(db, seed, login):
    a = seed["a"]["company"]
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    req_id = _new_req(c, tok)
    _save_spec(c, req_id, csrf(c, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()

    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "취소될업무", "requested_volume": "2", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_id = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_id), "csrf_token": tok,
    }, follow_redirects=False)

    tok = csrf(c, f"/client/requests/{req_id}")
    c.post(f"/client/tasks/{task_id}/cancel", data={"csrf_token": tok}, follow_redirects=False)

    pool = db.execute(
        "SELECT COUNT(*) AS n FROM work_items WHERE request_id=? AND task_id IS NULL AND status='open'", (req_id,)
    ).fetchone()["n"]
    assert pool == 3


def test_cross_company_cannot_see_other_request(db, seed, login):
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)

    cb = login("b_mgr")
    r = cb.get(f"/client/requests/{req_id}", follow_redirects=False)
    assert r.status_code in (303, 307)
    assert r.headers["location"] != f"/client/requests/{req_id}"


def test_worker_cannot_open_other_workers_task_work_screen(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()

    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    cb = login("b_wk")
    r = cb.get(f"/task/{task_id}/work", follow_redirects=False)
    assert r.status_code in (303, 307)
    assert r.headers["location"] != f"/task/{task_id}/work"


def test_staff_view_does_not_expose_item_source(db, seed, login):
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))

    staff = login("t_staff")
    r = staff.get(f"/platform/work-requests/{req_id}")
    assert "가게1" not in r.text
    assert "서울" not in r.text


def _mk_task_items(db, req_id, task_id, n, done=True):
    for seq in range(1, n + 1):
        db.execute(
            "INSERT INTO work_items (request_id, task_id, seq, source, status) VALUES (?,?,?,?,?)",
            (req_id, task_id, seq, "{}", "done" if done else "open"),
        )
    db.commit()


def test_sample_min_five_when_small(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'in_progress')",
        (req_id,),
    ).lastrowid
    _mk_task_items(db, req_id, task_id, 20)
    run_sample(db, task_id, 10)
    sampled = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE task_id=? AND sampled=1", (task_id,)).fetchone()["n"]
    assert sampled == 5


def test_sample_full_when_under_five_assigned(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'in_progress')",
        (req_id,),
    ).lastrowid
    _mk_task_items(db, req_id, task_id, 3)
    run_sample(db, task_id, 10)
    sampled = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE task_id=? AND sampled=1", (task_id,)).fetchone()["n"]
    assert sampled == 3


def test_finalize_review_rejects_below_accuracy_and_reopens(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'review')",
        (req_id,),
    ).lastrowid
    for seq, status in [(1, "approved"), (2, "approved"), (3, "rejected")]:
        db.execute(
            "INSERT INTO work_items (request_id, task_id, seq, source, status, sampled) VALUES (?,?,?,?,?,1)",
            (req_id, task_id, seq, "{}", status),
        )
    db.commit()
    acc = finalize_review(db, task_id, seed["a"]["mgr"])
    assert round(acc) == 67
    task = db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()
    assert task["status"] == "returned"
    reopened = db.execute("SELECT status FROM work_items WHERE task_id=? AND seq=3", (task_id,)).fetchone()
    assert reopened["status"] == "open"


def test_finalize_review_approves_when_accuracy_ok(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'review')",
        (req_id,),
    ).lastrowid
    db.execute(
        "INSERT INTO work_items (request_id, task_id, seq, source, status, sampled) VALUES (?,?,?,?,'approved',1)",
        (req_id, task_id, 1, "{}"),
    )
    db.execute(
        "INSERT INTO work_items (request_id, task_id, seq, source, status) VALUES (?,?,?,?,'done')",
        (req_id, task_id, 2, "{}"),
    )
    db.commit()
    acc = finalize_review(db, task_id, seed["a"]["mgr"])
    assert acc == 100.0
    task = db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()
    assert task["status"] == "completed"
    bulk = db.execute("SELECT status FROM work_items WHERE task_id=? AND seq=2", (task_id,)).fetchone()
    assert bulk["status"] == "approved"


def test_hold_item_flows_to_client_review_and_back(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    worker = login("a_wk")
    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/start", data={"csrf_token": tok}, follow_redirects=False)

    items = db.execute("SELECT id, seq FROM work_items WHERE task_id=? ORDER BY seq", (task_id,)).fetchall()
    item1 = items[0]["id"]
    tok = csrf(worker, f"/task/{task_id}/work")
    worker.post(f"/task/{task_id}/work/{item1}", data={
        "action": "hold", "hold_note": "판단이 어려움", "csrf_token": tok,
    }, follow_redirects=False)
    for it in items[1:]:
        tok = csrf(worker, f"/task/{task_id}/work")
        worker.post(f"/task/{task_id}/work/{it['id']}", data={
            "action": "save", "f1": "음식점", "csrf_token": tok,
        }, follow_redirects=False)

    held = db.execute("SELECT status FROM work_items WHERE id=?", (item1,)).fetchone()
    assert held["status"] == "hold"

    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/submit", data={"csrf_token": tok}, follow_redirects=False)
    task_status = db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()["status"]
    assert task_status == "review"

    tok = csrf(ca, f"/client/tasks/{task_id}")
    r = ca.post(f"/client/tasks/{task_id}/items/{item1}/reject", data={
        "note": "다시 확인해주세요", "csrf_token": tok,
    }, follow_redirects=False)
    assert r.status_code == 303
    back = db.execute("SELECT status, hold_note FROM work_items WHERE id=?", (item1,)).fetchone()
    assert back["status"] == "open"
    assert back["hold_note"] == "다시 확인해주세요"

    for it in items[1:]:
        tok = csrf(ca, f"/client/tasks/{task_id}")
        ca.post(f"/client/tasks/{task_id}/items/{it['id']}/approve", data={"csrf_token": tok}, follow_redirects=False)
    tok = csrf(ca, f"/client/tasks/{task_id}")
    ca.post(f"/client/tasks/{task_id}/finish", data={"csrf_token": tok}, follow_redirects=False)
    assert db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()["status"] == "returned"

    r = worker.get(f"/task/{task_id}/work")
    assert "보류 답변: 다시 확인해주세요" in r.text
    tok = csrf(worker, f"/task/{task_id}/work")
    worker.post(f"/task/{task_id}/work/{item1}", data={
        "action": "save", "f1": "소매", "csrf_token": tok,
    }, follow_redirects=False)
    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/resubmit", data={"csrf_token": tok}, follow_redirects=False)
    row = db.execute("SELECT status, sampled, round FROM work_items WHERE id=?", (item1,)).fetchone()
    assert (row["status"], row["sampled"], row["round"]) == ("done", 1, 2)

    tok = csrf(ca, f"/client/tasks/{task_id}")
    ca.post(f"/client/tasks/{task_id}/items/{item1}/approve", data={"csrf_token": tok}, follow_redirects=False)
    tok = csrf(ca, f"/client/tasks/{task_id}")
    ca.post(f"/client/tasks/{task_id}/finish", data={"csrf_token": tok}, follow_redirects=False)
    assert db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()["status"] == "completed"


def test_resubmit_resamples_prior_rejects_fully(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'in_progress')",
        (req_id,),
    ).lastrowid
    db.execute(
        "INSERT INTO work_items (request_id, task_id, seq, source, status, review_note) VALUES (?,?,1,?,?,?)",
        (req_id, task_id, "{}", "done", "다시 확인"),
    )
    for seq in range(2, 22):
        db.execute(
            "INSERT INTO work_items (request_id, task_id, seq, source, status) VALUES (?,?,?,?,?)",
            (req_id, task_id, seq, "{}", "done"),
        )
    db.commit()
    run_sample(db, task_id, 10, resubmit=True)

    prior = db.execute("SELECT sampled FROM work_items WHERE task_id=? AND seq=1", (task_id,)).fetchone()
    assert prior["sampled"] == 1

    fresh_sampled = db.execute(
        "SELECT COUNT(*) AS n FROM work_items WHERE task_id=? AND seq>1 AND sampled=1", (task_id,)
    ).fetchone()["n"]
    assert fresh_sampled == 5


def test_export_csv_seq_matches_source(db, seed, login):
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))

    r = ca.get(f"/client/requests/{req_id}/items/export")
    assert r.status_code == 200
    lines = r.text.lstrip("﻿").strip().splitlines()
    assert lines[0].split(",")[0] == "seq"
    body_rows = [l.split(",") for l in lines[1:]]
    assert [row[0] for row in body_rows] == ["1", "2", "3"]
    assert body_rows[0][1] == "가게1"


def test_request_detail_renders_with_item_spec(db, seed, login):
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    r = ca.get(f"/client/requests/{req_id}")
    assert r.status_code == 200
    assert "설정 이어하기" in r.text

    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))
    r = ca.get(f"/client/requests/{req_id}")
    assert r.status_code == 200
    assert "결과 반출" in r.text


def test_worker_work_screen_and_review_screen_render(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    worker = login("a_wk")
    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/start", data={"csrf_token": tok}, follow_redirects=False)

    r = worker.get(f"/task/{task_id}/work")
    assert r.status_code == 200
    assert "저장하고 다음" in r.text

    items = db.execute("SELECT id FROM work_items WHERE task_id=? ORDER BY seq", (task_id,)).fetchall()
    for it in items:
        tok = csrf(worker, f"/task/{task_id}/work")
        worker.post(f"/task/{task_id}/work/{it['id']}", data={
            "action": "save", "f1": "음식점", "csrf_token": tok,
        }, follow_redirects=False)

    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/submit", data={"csrf_token": tok}, follow_redirects=False)

    r = ca.get(f"/client/tasks/{task_id}")
    assert r.status_code == 200
    assert "검수" in r.text


def test_worker_task_detail_renders_for_item_spec_task(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    worker = login("a_wk")
    r = worker.get(f"/task/{task_id}")
    assert r.status_code == 200

    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/start", data={"csrf_token": tok}, follow_redirects=False)
    r = worker.get(f"/task/{task_id}")
    assert r.status_code == 200
    assert "작업 화면 열기" in r.text


def test_finalize_counts_current_round_only(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'review')",
        (req_id,),
    ).lastrowid
    for seq in range(1, 10):
        db.execute(
            "INSERT INTO work_items (request_id, task_id, seq, source, status, sampled, round) "
            "VALUES (?,?,?,'{}','approved',1,1)", (req_id, task_id, seq),
        )
    db.execute(
        "INSERT INTO work_items (request_id, task_id, seq, source, status, sampled, round, review_note) "
        "VALUES (?,?,10,'{}','rejected',1,2,'오타')", (req_id, task_id),
    )
    db.commit()
    acc = finalize_review(db, task_id, seed["a"]["mgr"])
    assert acc == 0
    assert db.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()["status"] == "returned"


def test_resubmit_full_check_for_returned_hold(db, seed):
    a = seed["a"]["company"]
    req_id = db.execute(
        "INSERT INTO work_requests (company_id, title, task_type, status, item_spec) "
        "VALUES (?, '요청', 'data_entry', 'submitted', '{}')", (a,),
    ).lastrowid
    task_id = db.execute(
        "INSERT INTO tasks (work_request_id, title, task_type, status) VALUES (?, '업무', 'data_entry', 'review')",
        (req_id,),
    ).lastrowid
    _mk_task_items(db, req_id, task_id, 20)
    db.execute("UPDATE work_items SET hold_note='기준 문의' WHERE task_id=? AND seq=7", (task_id,))
    db.commit()
    run_sample(db, task_id, 10, resubmit=True)
    row = db.execute("SELECT sampled, round FROM work_items WHERE task_id=? AND seq=7", (task_id,)).fetchone()
    assert row["sampled"] == 1 and row["round"] == 1


def test_parse_xlsx_cell_types():
    raw = _xlsx_bytes(
        ["상호", "등록일", "코드"],
        [["가게1", date(2026, 3, 5), "007"], ["가게2", date(2026, 3, 6), 3.0]],
    )
    headers, rows = parse_xlsx(raw)
    assert headers == ["상호", "등록일", "코드"]
    assert rows[0]["등록일"] == "2026-03-05"
    assert rows[0]["코드"] == "007"
    assert rows[1]["코드"] == "3"


def test_parse_xlsx_rejects_corrupt_file():
    with pytest.raises(CsvError):
        parse_xlsx(b"PK\x03\x04\x00\x00broken")


def test_parse_item_file_dispatch_by_ext_and_magic():
    raw_csv = _csv_bytes(["상호"], [["가게1"]])
    raw_xlsx = _xlsx_bytes(["상호"], [["가게1"]])
    headers, rows = parse_item_file("items.csv", raw_csv)
    assert headers == ["상호"]
    headers, rows = parse_item_file("items.xlsx", raw_xlsx)
    assert headers == ["상호"]
    with pytest.raises(CsvError):
        parse_item_file("items.xlsx", raw_csv)
    with pytest.raises(CsvError):
        parse_item_file("items.csv", raw_xlsx)
    with pytest.raises(CsvError):
        parse_item_file("items.txt", raw_csv)


def test_create_request_with_xlsx_upload(db, seed, login):
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    raw = _xlsx_bytes(["상호", "주소", "비고"], [["가게1", "서울", "x"], ["가게2", "부산", "y"]])
    r = c.post("/client/requests", data={
        "title": "엑셀업로드", "task_type": "data_entry", "volume": "1",
        "priority": "normal", "output_format": "excel", "data_work": "1",
        "csrf_token": tok,
    }, files={"item_csv": ("items.xlsx", raw, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
       follow_redirects=False)
    assert r.status_code == 303
    req_id = int(r.headers["location"].split("/")[3])
    rows = db.execute("SELECT * FROM work_items WHERE request_id=? ORDER BY seq", (req_id,)).fetchall()
    assert len(rows) == 2
    src = json.loads(rows[0]["source"])
    assert src["상호"] == "가게1"


def test_upload_corrupt_xlsx_rejected(db, seed, login):
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    r = c.post("/client/requests", data={
        "title": "손상파일", "task_type": "data_entry", "volume": "1",
        "priority": "normal", "output_format": "excel", "data_work": "1",
        "csrf_token": tok,
    }, files={"item_csv": ("items.xlsx", b"PK\x03\x04broken",
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, follow_redirects=False)
    assert r.status_code == 400


def test_item_spec_form_prefills_blank_columns(db, seed, login):
    c = login("a_mgr")
    tok = csrf(c, "/client/requests/new")
    raw = _csv_bytes(["상호", "주소", "업종"], [["가게1", "서울", ""], ["가게2", "부산", ""]])
    r = c.post("/client/requests", data={
        "title": "빈열", "task_type": "data_entry", "volume": "1",
        "priority": "normal", "output_format": "excel", "data_work": "1",
        "csrf_token": tok,
    }, files={"item_csv": ("items.csv", raw, "text/csv")}, follow_redirects=False)
    req_id = int(r.headers["location"].split("/")[3])
    html = c.get(f"/client/requests/{req_id}/item-spec").text
    assert 'value="업종"' in html
    show_blank = re.search(r'name="show" value="업종"[^>]*', html).group(0)
    show_filled = re.search(r'name="show" value="상호"[^>]*', html).group(0)
    assert "checked" not in show_blank
    assert "checked" in show_filled


def test_field_type_validation(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    tok = csrf(ca, f"/client/requests/{req_id}/item-spec")
    ca.post(f"/client/requests/{req_id}/item-spec", data={
        "show": ["상호", "주소"],
        "field_label": ["수량", "방문일", "연락처"],
        "field_type": ["number", "date", "phone"],
        "field_options": ["", "", ""],
        "field_required": ["0", "0", "0"],
        "sample_rate": "10",
        "csrf_token": tok,
    }, follow_redirects=False)
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    worker = login("a_wk")
    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/start", data={"csrf_token": tok}, follow_redirects=False)
    item_id = db.execute(
        "SELECT id FROM work_items WHERE task_id=? ORDER BY seq LIMIT 1", (task_id,)
    ).fetchone()["id"]

    tok = csrf(worker, f"/task/{task_id}/work")
    r = worker.post(f"/task/{task_id}/work/{item_id}", data={
        "action": "save", "f1": "abc", "f2": "2026-01-05", "f3": "010-1234-5678", "csrf_token": tok,
    }, follow_redirects=False)
    assert "field_invalid" in r.headers["location"]

    tok = csrf(worker, f"/task/{task_id}/work")
    r = worker.post(f"/task/{task_id}/work/{item_id}", data={
        "action": "save", "f1": "12.5", "f2": "20260105", "f3": "010-1234-5678", "csrf_token": tok,
    }, follow_redirects=False)
    assert "field_invalid" in r.headers["location"]

    tok = csrf(worker, f"/task/{task_id}/work")
    r = worker.post(f"/task/{task_id}/work/{item_id}", data={
        "action": "save", "f1": "12.5", "f2": "2026-01-05", "f3": "0101234", "csrf_token": tok,
    }, follow_redirects=False)
    assert "field_invalid" in r.headers["location"]

    tok = csrf(worker, f"/task/{task_id}/work")
    r = worker.post(f"/task/{task_id}/work/{item_id}", data={
        "action": "save", "f1": "-12.5", "f2": "2026-01-05", "f3": "010-1234-5678", "csrf_token": tok,
    }, follow_redirects=False)
    assert r.status_code == 303
    assert "field_invalid" not in r.headers["location"]
    row = db.execute("SELECT status FROM work_items WHERE id=?", (item_id,)).fetchone()
    assert row["status"] == "done"


def test_validate_value_rules():
    assert validate_value("number", "12.5")
    assert validate_value("number", "-3")
    assert not validate_value("number", "12.5.3")
    assert validate_value("date", "2026-01-05")
    assert not validate_value("date", "20260105")
    assert validate_value("phone", "010-1234-5678")
    assert not validate_value("phone", "0101234")
    assert not validate_value("phone", "010-1234-56789012")
    assert validate_value("text", "아무거나")


def test_export_xlsx_seq_and_columns(db, seed, login):
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    _save_spec(ca, req_id, csrf(ca, f"/client/requests/{req_id}/item-spec"))

    r = ca.get(f"/client/requests/{req_id}/items/export?fmt=xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    data_rows = list(ws.iter_rows(values_only=True))
    assert data_rows[0][0] == "seq"
    assert [row[0] for row in data_rows[1:]] == [1, 2, 3]
    assert data_rows[1][1] == "가게1"


def test_task_work_screen_label_for_linking(db, seed, login):
    a = seed["a"]["company"]
    ca = login("a_mgr")
    tok = csrf(ca, "/client/requests/new")
    req_id = _new_req(ca, tok)
    tok = csrf(ca, f"/client/requests/{req_id}/item-spec")
    ca.post(f"/client/requests/{req_id}/item-spec", data={
        "show": ["상호"],
        "field_label": ["수량"],
        "field_type": ["number"],
        "field_options": [""],
        "field_required": ["0"],
        "sample_rate": "10",
        "csrf_token": tok,
    }, follow_redirects=False)
    db.execute("UPDATE work_requests SET status='submitted' WHERE id=?", (req_id,))
    db.commit()
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/new", data={
        "title": "업무", "requested_volume": "3", "csrf_token": tok,
    }, follow_redirects=False)
    task_id = db.execute("SELECT id FROM tasks WHERE work_request_id=?", (req_id,)).fetchone()["id"]
    emp_a = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    tok = csrf(ca, f"/client/requests/{req_id}")
    ca.post(f"/client/requests/{req_id}/tasks/{task_id}/assign", data={
        "employee_id": str(emp_a), "csrf_token": tok,
    }, follow_redirects=False)

    worker = login("a_wk")
    tok = csrf(worker, f"/task/{task_id}")
    worker.post(f"/task/{task_id}/start", data={"csrf_token": tok}, follow_redirects=False)
    html = worker.get(f"/task/{task_id}/work").text
    assert 'for="wf_f1"' in html
    assert 'id="wf_f1"' in html
    assert "work-big" in html
    assert "work-source-label" in html


def test_assignee_auto_assign_on_submit(db, seed, login):
    a = seed["a"]["company"]
    emp_id = db.execute("SELECT id FROM employees WHERE company_id=?", (a,)).fetchone()["id"]
    c = login("a_mgr")
    raw = _csv_bytes(["상호", "주소"], [["가게1", "서울"], ["가게2", "부산"], ["가게3", "대구"]])
    r = c.post("/client/requests", data={
        "title": "한 명 배정", "task_type": "data_entry", "volume": "1", "priority": "normal",
        "output_format": "excel", "data_work": "1", "assignee": str(emp_id),
        "csrf_token": csrf(c, "/client/requests/new"),
    }, files={"item_csv": ("items.csv", raw, "text/csv")}, follow_redirects=False)
    req_id = int(r.headers["location"].split("/")[3])
    _save_spec(c, req_id, csrf(c, f"/client/requests/{req_id}/item-spec"))
    assert "바로 배정됩니다" in c.get(f"/client/requests/{req_id}").text

    c.post(f"/client/requests/{req_id}/submit", data={
        "csrf_token": csrf(c, f"/client/requests/{req_id}"),
    }, follow_redirects=False)
    t = db.execute("SELECT id, status, assigned_to, requested_volume FROM tasks WHERE work_request_id=?", (req_id,)).fetchall()
    assert len(t) == 1
    assert t[0]["status"] == "assigned" and t[0]["assigned_to"] == emp_id and t[0]["requested_volume"] == 3
    n = db.execute("SELECT COUNT(*) AS n FROM work_items WHERE task_id=?", (t[0]["id"],)).fetchone()["n"]
    assert n == 3
    assert db.execute("SELECT status FROM work_requests WHERE id=?", (req_id,)).fetchone()["status"] == "in_progress"


def test_assignee_other_company_ignored(db, seed, login):
    b = seed["b"]["company"]
    other = db.execute("SELECT id FROM employees WHERE company_id=?", (b,)).fetchone()["id"]
    c = login("a_mgr")
    c.post("/client/requests", data={
        "title": "일반 요청", "task_type": "data_entry", "volume": "2", "priority": "normal",
        "output_format": "excel", "assignee": str(other), "submit_now": "1",
        "csrf_token": csrf(c, "/client/requests/new"),
    }, follow_redirects=False)
    wr = db.execute("SELECT id, assigned_to FROM work_requests WHERE title='일반 요청'").fetchone()
    assert wr["assigned_to"] is None
    assert db.execute("SELECT COUNT(*) AS n FROM tasks WHERE work_request_id=?", (wr["id"],)).fetchone()["n"] == 0
