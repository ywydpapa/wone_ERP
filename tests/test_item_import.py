import io

from openpyxl import Workbook, load_workbook

from tests.conftest import csrf


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_form_renders(login):
    r = login("a_mgr").get("/items/import")
    assert r.status_code == 200


def test_import_add_and_update(db, seed, login):
    a = seed["a"]["company"]
    db.execute("INSERT INTO items (item_code, item_name, company_id) VALUES ('00007','기존',?)", (a,))
    db.commit()
    data = _xlsx([
        ["품목코드", "품목명", "유형", "단위", "단가", "안전재고", "분류"],
        ["00007", "기존수정", "완제품", "EA", 1000, 10, "완제품류"],
        [None, "신규품목", None, None, None, None, None],
    ])
    c = login("a_mgr")
    r = c.post("/items/import", files={"file": ("i.xlsx", data,
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert "신규 1건, 수정 1건" in r.text

    row = db.execute("SELECT item_name, item_type, unit_price FROM items WHERE item_code='00007' AND company_id=?", (a,)).fetchone()
    assert row["item_name"] == "기존수정"
    assert row["item_type"] == "product"
    assert row["unit_price"] == 1000

    new_row = db.execute("SELECT item_type FROM items WHERE item_name='신규품목' AND company_id=?", (a,)).fetchone()
    assert new_row["item_type"] == "material"


def test_import_auto_code_assigned(db, seed, login):
    a = seed["a"]["company"]
    db.execute("INSERT INTO items (item_code, item_name, company_id) VALUES ('00003','기존',?)", (a,))
    db.commit()
    data = _xlsx([["품목코드", "품목명"], [None, "신규품목"]])
    c = login("a_mgr")
    c.post("/items/import", files={"file": ("i.xlsx", data,
           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    row = db.execute("SELECT item_code FROM items WHERE item_name='신규품목' AND company_id=?", (a,)).fetchone()
    assert row["item_code"] == "00004"


def test_import_row_error_rejects_whole_file(db, seed, login):
    a = seed["a"]["company"]
    data = _xlsx([
        ["품목코드", "품목명", "단가"],
        [None, "정상품목", 100],
        [None, "오류품목", "숫자아님"],
    ])
    c = login("a_mgr")
    r = c.post("/items/import", files={"file": ("i.xlsx", data,
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert "반영하지 않" in r.text
    assert not db.execute("SELECT 1 FROM items WHERE item_name='정상품목' AND company_id=?", (a,)).fetchone()


def test_import_missing_name_column_is_error(login):
    c = login("a_mgr")
    data = _xlsx([["품목코드"], ["1"]])
    r = c.post("/items/import", files={"file": ("i.xlsx", data,
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert "품목명" in r.text


def test_import_template_roundtrip(db, seed, login):
    a = seed["a"]["company"]
    db.execute("INSERT INTO items (item_code, item_name, item_type, unit_price, company_id) "
               "VALUES ('00001','왕복품목','product',5000,?)", (a,))
    db.commit()
    c = login("a_mgr")
    r = c.get("/items/import/template")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    head = [cell.value for cell in ws[1]]
    assert head[:2] == ["품목코드", "품목명"]
    body = [cell.value for cell in ws[2]]
    row = dict(zip(head, body))
    assert row["품목코드"] == "00001"
    assert row["품목명"] == "왕복품목"
    assert row["유형"] == "완제품"

    buf = io.BytesIO()
    wb.save(buf)
    r2 = c.post("/items/import", files={"file": ("i.xlsx", buf.getvalue(),
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.status_code == 200
    assert "신규 0건, 수정 1건" in r2.text


def test_import_other_company_item_untouched(db, seed, login):
    b = seed["b"]["company"]
    db.execute("INSERT INTO items (item_code, item_name, company_id) VALUES ('00001','남의품목',?)", (b,))
    db.commit()
    data = _xlsx([["품목코드", "품목명"], ["00001", "내가고친이름"]])
    c = login("a_mgr")
    c.post("/items/import", files={"file": ("i.xlsx", data,
           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    row = db.execute("SELECT item_name FROM items WHERE company_id=? AND item_code='00001'", (b,)).fetchone()
    assert row["item_name"] == "남의품목"
