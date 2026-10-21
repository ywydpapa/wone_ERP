import hashlib
import re
import time
from unittest.mock import patch
from io import BytesIO

import msoffcrypto
from openpyxl import load_workbook

from core.crypto import enc_rno
from core.export import mask_rno
import routers.hr_book as hr_book_mod
from routers.hr_book import calc_disability_roster


def _disabled_emp(db, company_id, name="장애인근로자", resident_no="900101-1234567"):
    return db.execute(
        """INSERT INTO employees (name, company_id, status, hire_date, disability_type, disability_grade,
               resident_no, monthly_work_hours, base_salary)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (name, company_id, "active", "2026-01-01", "지체장애", "심하지않은장애",
         enc_rno(resident_no), 174, 2000000),
    ).lastrowid


def _add_hr_user(db, company_id, username="a_hr"):
    pw = hashlib.sha256(b"1234").hexdigest()
    uid = db.execute(
        "INSERT INTO users (username, password, name, role) VALUES (?,?,?,?)",
        (username, pw, "인사담당", "client"),
    ).lastrowid
    db.execute(
        "INSERT INTO client_users (user_id, company_id, client_role) VALUES (?,?,?)",
        (uid, company_id, "hr"),
    )
    return uid


# 마스킹

def test_mask_rno_normal_and_abnormal():
    assert mask_rno("900101-1234567") == "900101-*******"
    assert mask_rno("9001011234567") == "900101-*******"
    assert mask_rno("900101") == "******-*******"
    assert mask_rno("abcdefg1234567") == "******-*******"
    assert mask_rno("") == "-"
    assert mask_rno(None) == "-"


# 사유 필수

def test_subsidy_excel_requires_reason(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.get("/hr/subsidy/excel?year=2026&month=3")
    assert r.status_code == 400


def test_register_print_requires_reason(db, seed, login):
    a = seed["a"]["company"]
    emp_id = _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.get(f"/register/{emp_id}/print")
    assert r.status_code == 400


# 감사로그

def test_pii_export_writes_audit_log(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.get("/hr/subsidy/excel?year=2026&month=3&reason=장려금 신청 준비")
    assert r.status_code == 200
    row = db.execute(
        "SELECT * FROM audit_logs WHERE action='download' ORDER BY id DESC LIMIT 1"
    ).fetchone()
    assert row is not None
    assert "장려금" in row["summary"]
    assert "장려금 신청 준비" in (row["detail"] or "")


# 워터마크

def test_register_print_has_watermark(db, seed, login):
    a = seed["a"]["company"]
    emp_id = _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.get(f"/register/{emp_id}/print?reason=급여산정 확인")
    assert r.status_code == 200
    assert "pii-watermark" in r.text


# 마스킹 적용 확인 (명부 엑셀)

def test_disability_report_excel_roster_sheet_masked(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a, resident_no="900101-1234567")
    db.commit()
    c = login("a_mgr")
    r = c.get("/hr/disability-report/excel?year=2026&reason=테스트조회")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    ws = wb["장애인근로자명부"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        for v in row:
            if isinstance(v, str):
                assert not re.search(r"\d{6}-?[1-4]\d{6}", v)


def test_roster_export_masked(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a, resident_no="900101-1234567")
    db.commit()
    c = login("a_mgr")
    r = c.get("/roster/export?year=2026&reason=테스트조회")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        for v in row:
            if isinstance(v, str):
                assert not re.search(r"\d{6}-?[1-4]\d{6}", v)


# 원본 제출용 엔드포인트 권한

def test_submission_denied_for_hr_role(db, seed, login):
    a = seed["a"]["company"]
    _add_hr_user(db, a)
    db.commit()
    c = login("a_hr")
    r = c.get("/hr/disability-report/submission?year=2026")
    assert r.status_code == 403


def test_submission_denied_for_platform_staff(login):
    c = login("t_staff")
    r = c.get("/hr/disability-report/submission?year=2026", follow_redirects=False)
    assert r.status_code in (302, 303, 403)


def test_roster_tenant_scoped(db, seed):
    a = seed["a"]["company"]
    b = seed["b"]["company"]
    _disabled_emp(db, a, name="가장애인")
    _disabled_emp(db, b, name="나장애인")
    db.commit()
    roster_a = calc_disability_roster(db, a, year=2026)
    names = [r["name"] for r in roster_a]
    assert "가장애인" in names
    assert "나장애인" not in names


def test_submission_wrong_password_rejected(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.post("/hr/disability-report/submission",
                data={"year": 2026, "reason": "공단 제출", "password": "wrong"})
    assert r.status_code == 200
    assert "비밀번호가 일치하지 않습니다" in r.text
    cnt = db.execute(
        "SELECT COUNT(*) c FROM audit_logs WHERE summary LIKE '%공단제출용%'"
    ).fetchone()["c"]
    assert cnt == 0


def test_submission_issue_download_and_notify(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a, name="가장애인", resident_no="900101-1234567")
    db.commit()
    c = login("a_mgr")
    r = c.post("/hr/disability-report/submission",
                data={"year": 2026, "reason": "공단 제출", "password": "1234"})
    assert r.status_code == 200

    pw_match = re.search(r'user-select-all">([^<]+)</p>', r.text)
    assert pw_match
    file_password = pw_match.group(1).strip()

    token_match = re.search(r"token=([0-9a-f]{32})", r.text)
    assert token_match
    token = token_match.group(1)

    dl = c.get(f"/hr/disability-report/submission/download?token={token}")
    assert dl.status_code == 200

    office_file = msoffcrypto.OfficeFile(BytesIO(dl.content))
    office_file.load_key(password=file_password)
    dec = BytesIO()
    office_file.decrypt(dec)
    dec.seek(0)
    wb = load_workbook(dec)
    ws = wb["장애인근로자명부(원본)"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(row[3] == "가장애인" and row[4] == "900101-1234567" for row in rows)

    # 1회용 토큰 재사용 거부. 첫 다운로드에서 세션 토큰도 함께 폐기되므로 403
    dl2 = c.get(f"/hr/disability-report/submission/download?token={token}")
    assert dl2.status_code in (400, 403)

    # 같은 회사 책임자급 전원 알림
    mgr_id = seed["a"]["mgr"]
    notif = db.execute(
        "SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 1", (mgr_id,)
    ).fetchone()
    assert notif is not None
    assert "장애인 근로자명부" in notif["title"]


def test_submission_token_scoped_to_session(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a)
    db.commit()
    c1 = login("a_mgr")
    r = c1.post("/hr/disability-report/submission",
                 data={"year": 2026, "reason": "공단 제출", "password": "1234"})
    token = re.search(r"token=([0-9a-f]{32})", r.text).group(1)

    c2 = login("b_mgr")
    dl = c2.get(f"/hr/disability-report/submission/download?token={token}")
    assert dl.status_code == 403


def test_submission_token_expires(db, seed, login):
    a = seed["a"]["company"]
    _disabled_emp(db, a)
    db.commit()
    c = login("a_mgr")
    r = c.post("/hr/disability-report/submission",
                data={"year": 2026, "reason": "공단 제출", "password": "1234"})
    token = re.search(r"token=([0-9a-f]{32})", r.text).group(1)
    with patch("routers.hr_book.time.time", return_value=time.time() + hr_book_mod.SUBMISSION_TTL + 1):
        dl = c.get(f"/hr/disability-report/submission/download?token={token}")
    assert dl.status_code == 400
