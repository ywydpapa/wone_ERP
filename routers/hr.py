import os
import uuid
import calendar

from core.tz import now_kst

from fastapi import APIRouter, Depends, Form, Request, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse
from core.deps import get_db, require_staff, templates

CERT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static", "uploads", "certs")
from core.constants import APPEAL_TYPE_LABELS, DISABILITY_TYPES, COMPANY_INFO
from routers._helpers import get_page
from core.capability import derive_tier2, derive_tier3
from core.attendance import (
    get_attendance_records,
    get_attendance_summary,
    get_worker_detail,
    get_worker_month_records,
    build_calendar_data,
    worker_month_summary,
    prev_next_month,
    available_months,
    STATUS_LABELS,
)
from core.export import styled_workbook, excel_response, pdf_response

router = APIRouter(prefix="/hr")


@router.get("/employees", response_class=HTMLResponse)
async def employee_list(
    request: Request,
    q: str = "",
    dept: str = "",
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    page = get_page(request)
    per_page = 10
    base_sql = (
        "FROM employees e "
        "LEFT JOIN client_companies cc ON e.company_id = cc.id "
        "WHERE 1=1"
    )
    params = []
    if q:
        base_sql += " AND e.name LIKE ?"
        params.append(f"%{q}%")
    if dept:
        base_sql += " AND e.dept = ?"
        params.append(dept)
    total = conn.execute("SELECT COUNT(*) " + base_sql, params).fetchone()[0]
    sql = "SELECT e.*, cc.name AS company_name " + base_sql + " ORDER BY e.employee_no LIMIT ? OFFSET ?"
    employees = conn.execute(sql, params + [per_page, (page - 1) * per_page]).fetchall()
    depts = conn.execute(
        "SELECT DISTINCT dept FROM employees ORDER BY dept"
    ).fetchall()

    total_pages = max(1, (total + per_page - 1) // per_page)

    return templates.TemplateResponse(
        request=request,
        name="hr/employee_list.html",
        context={
            "request": request,
            "page_title": "인사관리 - 직원 목록",
            "user_name": user["name"],
            "user_role": user["role"],
            "employees": employees,
            "depts": [r["dept"] for r in depts],
            "q": q,
            "selected_dept": dept,
            "total_count": total,
            "page": page,
            "total_pages": total_pages,
            "base_url": f"/hr/employees?q={q}&dept={dept}&",
        },
    )


@router.get("/employees/new", response_class=HTMLResponse)
async def employee_new(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()
    return templates.TemplateResponse(
        request=request,
        name="hr/employee_form.html",
        context={
            "request": request,
            "page_title": "직원 신규 등록",
            "user_name": user["name"],
            "user_role": user["role"],
            "employee": None,
            "companies": companies,
            "disability_types": DISABILITY_TYPES,
            "error": "",
        },
    )


async def _save_cert_file(upload):
    if not upload or not upload.filename:
        return ""
    os.makedirs(CERT_DIR, exist_ok=True)
    ext = os.path.splitext(upload.filename)[1]
    fname = f"{uuid.uuid4().hex[:12]}{ext}"
    path = os.path.join(CERT_DIR, fname)
    content = await upload.read()
    with open(path, "wb") as f:
        f.write(content)
    return f"uploads/certs/{fname}"


@router.post("/employees", response_class=HTMLResponse)
async def employee_create(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
    name: str = Form(...),
    employee_no: str = Form(""),
    dept: str = Form("경영지원팀"),
    position: str = Form("사원"),
    hire_date: str = Form(""),
    status: str = Form("active"),
    disability_type: str = Form(""),
    disability_grade: str = Form(""),
    emergency_contact: str = Form(""),
    emergency_phone: str = Form(""),
    notes: str = Form(""),
    company_id: str = Form(""),
    resident_no: str = Form(""),
    job_title: str = Form(""),
    monthly_work_hours: int = Form(0),
    employment_insurance: int = Form(1),
    disability_cert_file: UploadFile = File(None),
):
    cert_path = await _save_cert_file(disability_cert_file)
    try:
        cur = conn.execute(
            """INSERT INTO employees
               (name, employee_no, dept, position, hire_date, status,
                disability_type, disability_grade,
                emergency_contact, emergency_phone, notes, company_id,
                resident_no, job_title, monthly_work_hours,
                employment_insurance, disability_cert_file)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (name, employee_no or None, dept, position, hire_date, status,
             disability_type, disability_grade,
             emergency_contact, emergency_phone, notes,
             int(company_id) if company_id else None,
             resident_no, job_title, monthly_work_hours,
             employment_insurance, cert_path),
        )
        conn.commit()
        new_id = cur.lastrowid
    except Exception:
        companies = conn.execute(
            "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
        ).fetchall()
        return templates.TemplateResponse(
            request=request,
            name="hr/employee_form.html",
            context={
                "request": request,
                "page_title": "직원 신규 등록",
                "user_name": user["name"],
                "user_role": user["role"],
                "employee": None,
                "companies": companies,
                "disability_types": DISABILITY_TYPES,
                "error": "사번이 중복되었거나 필수 항목이 누락되었습니다.",
            },
        )
    return RedirectResponse(url=f"/hr/employees/{new_id}", status_code=303)


@router.get("/employees/{emp_id}", response_class=HTMLResponse)
async def employee_detail(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    employee = conn.execute(
        "SELECT e.*, cc.name AS company_name "
        "FROM employees e "
        "LEFT JOIN client_companies cc ON e.company_id = cc.id "
        "WHERE e.id=?",
        (emp_id,),
    ).fetchone()
    if not employee:
        return RedirectResponse(url="/hr/employees", status_code=303)
    return templates.TemplateResponse(
        request=request,
        name="hr/employee_detail.html",
        context={
            "request": request,
            "page_title": f"{employee['name']} - 직원 상세",
            "user_name": user["name"],
            "user_role": user["role"],
            "employee": employee,
        },
    )


@router.get("/employees/{emp_id}/edit", response_class=HTMLResponse)
async def employee_edit(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    employee = conn.execute(
        "SELECT * FROM employees WHERE id=?", (emp_id,)
    ).fetchone()
    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()
    if not employee:
        return RedirectResponse(url="/hr/employees", status_code=303)
    return templates.TemplateResponse(
        request=request,
        name="hr/employee_form.html",
        context={
            "request": request,
            "page_title": f"{employee['name']} - 정보 수정",
            "user_name": user["name"],
            "user_role": user["role"],
            "employee": employee,
            "companies": companies,
            "disability_types": DISABILITY_TYPES,
            "error": "",
        },
    )


@router.post("/employees/{emp_id}/edit", response_class=HTMLResponse)
async def employee_update(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
    name: str = Form(...),
    employee_no: str = Form(""),
    dept: str = Form("경영지원팀"),
    position: str = Form("사원"),
    hire_date: str = Form(""),
    status: str = Form("active"),
    disability_type: str = Form(""),
    disability_grade: str = Form(""),
    emergency_contact: str = Form(""),
    emergency_phone: str = Form(""),
    notes: str = Form(""),
    company_id: str = Form(""),
    resident_no: str = Form(""),
    job_title: str = Form(""),
    monthly_work_hours: int = Form(0),
    employment_insurance: int = Form(1),
    disability_cert_file: UploadFile = File(None),
):
    cert_path = await _save_cert_file(disability_cert_file)
    try:
        sql = """UPDATE employees SET
               name=?, employee_no=?, dept=?, position=?, hire_date=?, status=?,
               disability_type=?, disability_grade=?,
               emergency_contact=?, emergency_phone=?, notes=?, company_id=?,
               resident_no=?, job_title=?, monthly_work_hours=?,
               employment_insurance=?,
               updated_at=datetime('now','localtime')"""
        params = [
            name, employee_no or None, dept, position, hire_date, status,
            disability_type, disability_grade,
            emergency_contact, emergency_phone, notes,
            int(company_id) if company_id else None,
            resident_no, job_title, monthly_work_hours,
            employment_insurance,
        ]
        if cert_path:
            sql += ", disability_cert_file=?"
            params.append(cert_path)
        sql += " WHERE id=?"
        params.append(emp_id)
        conn.execute(sql, params)
        conn.commit()
    except Exception:
        employee = conn.execute(
            "SELECT * FROM employees WHERE id=?", (emp_id,)
        ).fetchone()
        companies = conn.execute(
            "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
        ).fetchall()
        return templates.TemplateResponse(
            request=request,
            name="hr/employee_form.html",
            context={
                "request": request,
                "page_title": "직원 정보 수정",
                "user_name": user["name"],
                "user_role": user["role"],
                "employee": employee,
                "companies": companies,
                "disability_types": DISABILITY_TYPES,
                "error": "사번이 중복되었거나 필수 항목이 누락되었습니다.",
            },
        )
    return RedirectResponse(url=f"/hr/employees/{emp_id}", status_code=303)


PROFILE_COLS = [
    "hand_left", "hand_right", "arm_left", "arm_right", "neck",
    "foot_left", "foot_right", "posture_maintenance",
    "vision", "hearing", "eye_movement", "eyelid_control",
    "speech", "breath_control",
    "reading_level", "sustained_focus", "memory_aid_needed",
    "continuous_work_minutes", "fatigue_pattern", "posture_change_interval",
    "input_overrides", "notes",
]


@router.get("/employees/{emp_id}/profile", response_class=HTMLResponse)
async def capability_profile_form(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    employee = conn.execute(
        "SELECT * FROM employees WHERE id=?", (emp_id,)
    ).fetchone()
    if not employee:
        return RedirectResponse(url="/hr/employees", status_code=303)
    profile = conn.execute(
        "SELECT * FROM capability_profiles WHERE employee_id=? ORDER BY effective_date DESC LIMIT 1",
        (emp_id,),
    ).fetchone()

    tier2, tier3 = None, None
    bw_class, bw_label = "bg-secondary", "저속"
    if profile:
        p = dict(profile)
        tier2 = derive_tier2(p)
        tier3 = derive_tier3(tier2, p)
        bw_map = {"high": ("bg-success", "고속"), "medium": ("bg-warning text-dark", "중속")}
        bw_class, bw_label = bw_map.get(tier2.bandwidth, ("bg-secondary", "저속"))

    return templates.TemplateResponse(
        request=request,
        name="hr/capability_profile.html",
        context={
            "request": request,
            "page_title": f"{employee['name']} - 능력 프로필",
            "user_name": user["name"],
            "user_role": user["role"],
            "employee": employee,
            "profile": dict(profile) if profile else None,
            "tier2": tier2,
            "tier3": tier3,
            "bw_class": bw_class,
            "bw_label": bw_label,
        },
    )


@router.post("/employees/{emp_id}/profile", response_class=HTMLResponse)
async def capability_profile_save(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    form = await request.form()
    employee = conn.execute(
        "SELECT * FROM employees WHERE id=?", (emp_id,)
    ).fetchone()
    if not employee:
        return RedirectResponse(url="/hr/employees", status_code=303)

    effective_date = form.get("effective_date", "")
    values = {}
    for col in PROFILE_COLS:
        val = form.get(col, "")
        if col in ("posture_maintenance", "eye_movement", "eyelid_control",
                    "breath_control", "sustained_focus", "memory_aid_needed"):
            values[col] = int(val) if val not in ("", None) else -1
        elif col in ("continuous_work_minutes", "posture_change_interval"):
            values[col] = int(val) if val not in ("", None) else None
        else:
            values[col] = val

    existing = conn.execute(
        "SELECT id FROM capability_profiles WHERE employee_id=? AND effective_date=?",
        (emp_id, effective_date),
    ).fetchone()

    cols_str = ", ".join(PROFILE_COLS)
    placeholders = ", ".join(["?"] * len(PROFILE_COLS))
    col_vals = [values[c] for c in PROFILE_COLS]

    if existing:
        set_clause = ", ".join(f"{c}=?" for c in PROFILE_COLS)
        conn.execute(
            f"UPDATE capability_profiles SET {set_clause}, updated_at=datetime('now','localtime') WHERE id=?",
            col_vals + [existing["id"]],
        )
    else:
        conn.execute(
            f"INSERT INTO capability_profiles (employee_id, effective_date, {cols_str}) VALUES (?,?,{placeholders})",
            [emp_id, effective_date] + col_vals,
        )
    conn.commit()
    return RedirectResponse(url=f"/hr/employees/{emp_id}/profile", status_code=303)


@router.get("/employees/{emp_id}/profile/result", response_class=HTMLResponse)
async def capability_result(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    employee = conn.execute(
        "SELECT * FROM employees WHERE id=?", (emp_id,)
    ).fetchone()
    if not employee:
        return RedirectResponse(url="/hr/employees", status_code=303)
    profile = conn.execute(
        "SELECT * FROM capability_profiles WHERE employee_id=? ORDER BY effective_date DESC LIMIT 1",
        (emp_id,),
    ).fetchone()

    if not profile:
        return RedirectResponse(url=f"/hr/employees/{emp_id}/profile", status_code=303)

    p = dict(profile)
    tier2 = derive_tier2(p)
    tier3 = derive_tier3(tier2, p)

    return templates.TemplateResponse(
        request=request,
        name="hr/capability_result.html",
        context={
            "request": request,
            "page_title": f"{employee['name']} - 능력 분석 결과",
            "user_name": user["name"],
            "user_role": user["role"],
            "employee": employee,
            "profile": p,
            "tier2": tier2,
            "tier3": tier3,
        },
    )


@router.get("/attendance", response_class=HTMLResponse)
async def hr_attendance(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    status_filter = request.query_params.get("status", "")
    company_filter = request.query_params.get("company", "")

    emp_sql = "SELECT id FROM employees WHERE status = 'active'"
    emp_params = []
    if company_filter:
        emp_sql += " AND company_id = ?"
        emp_params.append(int(company_filter))
    emp_rows = conn.execute(emp_sql, emp_params).fetchall()
    emp_ids = [r["id"] for r in emp_rows]

    records = get_attendance_records(conn, emp_ids, month)
    if status_filter:
        records = [r for r in records if r["status"] == status_filter]
    for r in records:
        r.setdefault("company_name", r.get("company_name") or "-")

    total_workers = len(emp_ids)
    summary = get_attendance_summary(records, total_workers)

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status = 'active' ORDER BY name"
    ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance.html",
        context={
            "request": request,
            "page_title": "근태 관리",
            "user_name": user["name"],
            "user_role": user["role"],
            "records": records,
            "summary": summary,
            "current_month": month,
            "status_filter": status_filter,
            "company_filter": company_filter,
            "companies": companies,
            "months": available_months(),
        },
    )


def _build_sheet_data(conn, year, month, company_id):
    _, days_in_month = calendar.monthrange(year, month)
    ym = f"{year}-{month:02d}"

    emp_sql = "SELECT id, name, employee_no, monthly_work_hours FROM employees WHERE status='active'"
    emp_params = []
    if company_id:
        emp_sql += " AND company_id=?"
        emp_params.append(company_id)
    emp_sql += " ORDER BY employee_no, name"
    employees = conn.execute(emp_sql, emp_params).fetchall()
    emp_ids = [e["id"] for e in employees]

    if not emp_ids:
        return [], days_in_month, []

    ph = ",".join("?" * len(emp_ids))
    att_rows = conn.execute(
        f"""SELECT employee_id, work_date, status, clock_in, clock_out,
                   work_minutes, overtime_minutes, break_minutes, notes
            FROM attendance
            WHERE employee_id IN ({ph}) AND work_date LIKE ?
            ORDER BY work_date""",
        [*emp_ids, f"{ym}%"],
    ).fetchall()

    att_map = {}
    for r in att_rows:
        att_map.setdefault(r["employee_id"], {})[r["work_date"]] = dict(r)

    rows = []
    for emp in employees:
        eid = emp["id"]
        days = att_map.get(eid, {})
        cells = []
        work_days = 0
        late_cnt = 0
        early_cnt = 0
        absent_cnt = 0
        holiday_cnt = 0
        total_mins = 0
        ot_mins = 0
        for d in range(1, days_in_month + 1):
            dt = f"{ym}-{d:02d}"
            rec = days.get(dt)
            if rec:
                cells.append(rec)
                if rec["status"] in ("normal", "late", "early_leave"):
                    work_days += 1
                    total_mins += rec.get("work_minutes") or 0
                    ot_mins += rec.get("overtime_minutes") or 0
                if rec["status"] == "late":
                    late_cnt += 1
                elif rec["status"] == "early_leave":
                    early_cnt += 1
                elif rec["status"] == "absent":
                    absent_cnt += 1
                elif rec["status"] == "holiday":
                    holiday_cnt += 1
            else:
                cells.append(None)
        rows.append({
            "employee": dict(emp),
            "cells": cells,
            "work_days": work_days,
            "late": late_cnt,
            "early_leave": early_cnt,
            "absent": absent_cnt,
            "holiday_days": holiday_cnt,
            "monthly_work_hours": emp["monthly_work_hours"] or 0,
            "total_mins": total_mins,
            "total_hours": total_mins // 60,
            "total_mins_r": total_mins % 60,
            "ot_hours": ot_mins // 60,
            "ot_mins_r": ot_mins % 60,
        })
    return rows, days_in_month, employees


@router.get("/attendance/sheet", response_class=HTMLResponse)
async def attendance_sheet(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    page = int(request.query_params.get("page", 1))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None

    rows, days_in_month, _ = _build_sheet_data(conn, year, month, company_id)

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()

    # 고객사 정보 (문서 헤더용)
    company_name = ""
    company_info = None
    if company_id:
        ci_row = conn.execute(
            "SELECT name, business_no, representative, address FROM client_companies WHERE id=?",
            (company_id,),
        ).fetchone()
        if ci_row:
            company_info = dict(ci_row)
            company_name = ci_row["name"]

    first_weekday = calendar.monthrange(year, month)[0]
    weekdays = [(first_weekday + d) % 7 for d in range(days_in_month)]

    # 배너 집계
    weekday_count = sum(1 for wd in weekdays if wd < 5)
    total_emp = len(rows)
    sum_work = sum(r["work_days"] for r in rows)
    sum_late = sum(r["late"] for r in rows)
    sum_absent = sum(r["absent"] for r in rows)
    expected = weekday_count * total_emp
    banner = {
        "total_emp": total_emp,
        "work": sum_work,
        "late": sum_late,
        "absent": sum_absent,
        "rate": round(sum_work / expected * 100, 1) if expected else 0,
    }

    # 페이지네이션
    per_page = 50
    total_pages = max(1, (total_emp + per_page - 1) // per_page)
    if page < 1:
        page = 1
    if page > total_pages:
        page = total_pages
    rows_page = rows[(page - 1) * per_page: page * per_page]

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance_sheet.html",
        context={
            "request": request,
            "page_title": "출근부",
            "user_name": user["name"],
            "user_role": user["role"],
            "rows": rows_page,
            "all_rows": rows,
            "year": year,
            "month": month,
            "days_in_month": days_in_month,
            "weekdays": weekdays,
            "companies": companies,
            "company_id": company_id or "",
            "company_name": company_name,
            "company_info": company_info,
            "banner": banner,
            "page": page,
            "total_pages": total_pages,
            "months": available_months(),
            "status_labels": STATUS_LABELS,
        },
    )


@router.get("/attendance/sheet/export")
async def attendance_sheet_export(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None
    fmt = request.query_params.get("format", "xlsx")

    rows, days_in_month, _ = _build_sheet_data(conn, year, month, company_id)

    company_name = ""
    if company_id:
        row = conn.execute("SELECT name FROM client_companies WHERE id=?", (company_id,)).fetchone()
        if row:
            company_name = row["name"]

    ym_str = f"{year}년 {month}월"
    filename_base = f"출근부_{year}{month:02d}"
    if company_name:
        filename_base += f"_{company_name}"

    if fmt == "pdf":
        html = templates.get_template("hr/attendance_sheet_pdf.html").render(
            rows=rows, year=year, month=month, days_in_month=days_in_month,
            company_name=company_name, ym_str=ym_str,
            company_info=COMPANY_INFO, status_labels=STATUS_LABELS,
        )
        return pdf_response(html, f"{filename_base}.pdf")

    from openpyxl.styles import Font, Alignment
    from core.export import HEADER_FONT, HEADER_FILL, CELL_FONT, THIN_BORDER
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "출근부"

    title_text = f"출근부 - {ym_str}"
    if company_name:
        title_text += f" ({company_name})"
    sum_cols = ["출근", "지각", "조퇴", "결근", "휴일", "연장", "총근무"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_in_month + 2 + len(sum_cols))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="맑은 고딕", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 3
    headers = ["사번", "이름"] + [str(d) for d in range(1, days_in_month + 1)] + sum_cols
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    status_short = {"normal": "O", "late": "△", "early_leave": "▽", "absent": "X", "holiday": "-"}
    for ri, r in enumerate(rows, header_row + 1):
        emp = r["employee"]
        ws.cell(row=ri, column=1, value=emp.get("employee_no", "")).font = CELL_FONT
        ws.cell(row=ri, column=1).border = THIN_BORDER
        ws.cell(row=ri, column=2, value=emp["name"]).font = CELL_FONT
        ws.cell(row=ri, column=2).border = THIN_BORDER
        for di, cell_data in enumerate(r["cells"], 3):
            cell = ws.cell(row=ri, column=di)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.font = CELL_FONT
            if cell_data:
                cell.value = status_short.get(cell_data["status"], "")
            else:
                cell.value = ""
        base_col = days_in_month + 3
        vals = [
            r["work_days"], r["late"], r["early_leave"], r["absent"],
            r["holiday_days"],
            f"{r['ot_hours']}:{r['ot_mins_r']:02d}",
            f"{r['total_hours']}:{r['total_mins_r']:02d}",
        ]
        for offset, val in enumerate(vals):
            cell = ws.cell(row=ri, column=base_col + offset, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 10
    for ci in range(3, days_in_month + 3):
        col_letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[col_letter].width = 4
    for ci in range(days_in_month + 3, days_in_month + 3 + len(sum_cols)):
        col_letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[col_letter].width = 9

    return excel_response(wb, f"{filename_base}.xlsx")


@router.get("/attendance/appeals", response_class=HTMLResponse)
async def appeal_list(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    rows = conn.execute(
        "SELECT a.*, e.name AS worker_name, cc.name AS company_name "
        "FROM attendance a JOIN employees e ON a.employee_id=e.id "
        "LEFT JOIN client_companies cc ON e.company_id=cc.id "
        "WHERE a.appeal IS NOT NULL "
        "ORDER BY CASE a.appeal WHEN 'pending' THEN 0 ELSE 1 END, a.work_date DESC",
    ).fetchall()
    appeals = []
    for r in rows:
        d = dict(r)
        d["appeal_type_label"] = APPEAL_TYPE_LABELS.get(d.get("appeal_type", ""), d.get("appeal_type", ""))
        appeals.append(d)
    pending_cnt = sum(1 for a in appeals if a["appeal"] == "pending")
    return templates.TemplateResponse(
        request=request, name="hr/appeal_list.html", context={
            "request": request,
            "page_title": "소명 관리",
            "user_name": user["name"],
            "appeals": appeals,
            "pending_cnt": pending_cnt,
        }
    )


@router.post("/attendance/{att_id}/appeal-review")
async def appeal_review(
    request: Request,
    att_id: int,
    action: str = Form(...),
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT id, appeal_type FROM attendance WHERE id=? AND appeal='pending'", (att_id,)
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance/appeals", status_code=303)
    result = "approved" if action == "approve" else "rejected"
    if result == "approved":
        note = f"소명 승인 ({row['appeal_type'] or ''})"
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=?, "
            "status='normal', notes=CASE WHEN notes IS NULL OR notes='' THEN ? "
            "ELSE notes || ' / ' || ? END WHERE id=?",
            (result, user["id"], now_kst().isoformat(), note, note, att_id),
        )
    else:
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=? WHERE id=?",
            (result, user["id"], now_kst().isoformat(), att_id),
        )
    conn.commit()
    return RedirectResponse(url="/hr/attendance/appeals", status_code=303)


@router.post("/attendance/mark-absent")
async def mark_absent(
    request: Request,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    form = await request.form()
    year = int(form.get("year", now_kst().year))
    month = int(form.get("month", now_kst().month))
    company_id = form.get("company_id", "")
    company_id = int(company_id) if company_id else None

    _, days_in_month = calendar.monthrange(year, month)
    ym = f"{year}-{month:02d}"

    emp_sql = "SELECT id, hire_date FROM employees WHERE status='active'"
    emp_params = []
    if company_id:
        emp_sql += " AND company_id=?"
        emp_params.append(company_id)
    employees = conn.execute(emp_sql, emp_params).fetchall()

    inserted = 0
    for emp in employees:
        hire_date = emp["hire_date"] or "2000-01-01"
        for d in range(1, days_in_month + 1):
            dt = f"{ym}-{d:02d}"
            weekday = calendar.weekday(year, month, d)
            if weekday >= 5:
                continue
            if dt < hire_date:
                continue
            exists = conn.execute(
                "SELECT id FROM attendance WHERE employee_id=? AND work_date=?",
                (emp["id"], dt),
            ).fetchone()
            if exists:
                continue
            conn.execute(
                "INSERT INTO attendance (employee_id, work_date, status) VALUES (?, ?, 'absent')",
                (emp["id"], dt),
            )
            inserted += 1
    conn.commit()

    url = f"/hr/attendance/sheet?year={year}&month={month}"
    if company_id:
        url += f"&company_id={company_id}"
    return RedirectResponse(url=url, status_code=303)


@router.get("/attendance/{att_id}/edit", response_class=HTMLResponse)
async def attendance_edit_form(
    request: Request,
    att_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT a.*, e.name AS worker_name FROM attendance a "
        "JOIN employees e ON a.employee_id=e.id WHERE a.id=?", (att_id,)
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    att = dict(row)
    logs = conn.execute(
        "SELECT al.*, u.name AS editor_name FROM attendance_logs al "
        "JOIN users u ON al.changed_by=u.id "
        "WHERE al.attendance_id=? ORDER BY al.created_at DESC", (att_id,)
    ).fetchall()
    return templates.TemplateResponse(
        request=request, name="hr/attendance_edit.html", context={
            "request": request,
            "page_title": "근태 수정",
            "user_name": user["name"],
            "user_role": user["role"],
            "att": att,
            "logs": [dict(l) for l in logs],
            "status_labels": STATUS_LABELS,
        }
    )


@router.post("/attendance/{att_id}/edit")
async def attendance_edit_save(
    request: Request,
    att_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    row = conn.execute("SELECT * FROM attendance WHERE id=?", (att_id,)).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    old = dict(row)
    form = await request.form()
    reason = form.get("reason", "")

    fields = {
        "clock_in": form.get("clock_in", "").strip() or None,
        "clock_out": form.get("clock_out", "").strip() or None,
        "status": form.get("status", old["status"]),
        "overtime_minutes": int(form.get("overtime_minutes") or 0),
        "break_minutes": int(form.get("break_minutes") or 0),
        "notes": form.get("notes", "").strip(),
    }

    # work_minutes 자동계산
    cin, cout = fields["clock_in"], fields["clock_out"]
    if cin and cout:
        cin_h, cin_m = int(cin[:2]), int(cin[3:5])
        out_h, out_m = int(cout[:2]), int(cout[3:5])
        work_mins = max(0, (out_h * 60 + out_m) - (cin_h * 60 + cin_m) - fields["break_minutes"])
        fields["work_minutes"] = work_mins
        # 연장근로 자동계산 (18:00 이후)
        if out_h >= 18 and not form.get("overtime_manual"):
            ot = max(0, (out_h * 60 + out_m) - 18 * 60)
            fields["overtime_minutes"] = ot
    else:
        fields["work_minutes"] = old.get("work_minutes") or 0

    # 변경 이력
    track_fields = ["clock_in", "clock_out", "status", "overtime_minutes", "break_minutes", "notes"]
    for f in track_fields:
        old_val = str(old.get(f) or "")
        new_val = str(fields.get(f) or "")
        if old_val != new_val:
            conn.execute(
                "INSERT INTO attendance_logs (attendance_id, changed_by, field, old_value, new_value, reason) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (att_id, user["id"], f, old_val, new_val, reason),
            )

    conn.execute(
        "UPDATE attendance SET clock_in=?, clock_out=?, status=?, work_minutes=?, "
        "overtime_minutes=?, break_minutes=?, notes=? WHERE id=?",
        (fields["clock_in"], fields["clock_out"], fields["status"],
         fields["work_minutes"], fields["overtime_minutes"],
         fields["break_minutes"], fields["notes"], att_id),
    )
    conn.commit()

    emp_id = old["employee_id"]
    month = old["work_date"][:7]
    return RedirectResponse(url=f"/hr/attendance/{emp_id}?month={month}", status_code=303)


@router.get("/attendance/{employee_id}", response_class=HTMLResponse)
async def hr_attendance_detail(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    worker = get_worker_detail(conn, employee_id)
    if not worker:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    records = get_worker_month_records(conn, employee_id, month)
    summary = worker_month_summary(records)
    cal = build_calendar_data(records, month)
    prev_m, next_m = prev_next_month(month)

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance_detail.html",
        context={
            "request": request,
            "page_title": f"{worker['name']} 근태 상세",
            "user_name": user["name"],
            "user_role": user["role"],
            "worker": worker,
            "records": records,
            "summary": summary,
            "cal": cal,
            "current_month": month,
            "prev_month": prev_m,
            "next_month": next_m,
            "back_url": "/hr/attendance",
        },
    )
