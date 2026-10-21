import io
import json
import os
import zipfile
from datetime import date, datetime
from urllib.parse import quote

from fastapi import APIRouter, Request, Form, Depends, File, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from starlette.responses import JSONResponse

from openpyxl import Workbook

from core.export import styled_sheet, excel_response, csv_response
from core.erp_export import EXPORT_KEY, EXPORT_SCHEMAS, fetch_rows, std_lines, apply_template, csv_bytes

from core.constants import ERP_MODULES
from core.deps import get_db, require_client_with_company, templates
from core.tz import now_kst
from core.export import pdf_response
from core.notification import create_notification, get_notifications, mark_all_read
from routers.platform import TASK_TYPE_LABELS, STATUS_LABELS, REQUEST_STATUS_LABELS
from core.workflow import (
    transition_request, transition_task,
    log_activity, get_activity_logs, get_due_status,
    sync_request_volume, advance_request, REQUEST_TRANSITIONS, REQUEST_ROLE_ACTIONS,
)
from core.data_work import (
    CsvError, parse_item_file, get_item_spec, assign_items, release_items,
    review_cnt, finalize_review, export_rows, PASS_ACCURACY, validate_value, FIELD_TYPES,
    STATUS_LABELS as WI_STATUS_LABELS,
)
from routers._helpers import save_upload

router = APIRouter(prefix="/client")

PRIORITY_LABELS = {
    "low": "낮음",
    "normal": "보통",
    "high": "높음",
    "urgent": "긴급",
}

OUTPUT_FORMAT_LABELS = {
    "excel": "엑셀",
    "json": "JSON",
    "pdf": "PDF",
    "erp": "ERP 연동",
}


REQ_SORTS = {"title": "wr.title", "type": "wr.task_type", "volume": "wr.volume",
             "status": "wr.status", "due": "wr.due_date"}


@router.get("/requests", response_class=HTMLResponse)
async def request_list(
    request: Request,
    sort: str = "",
    dir: str = "",
    export: str = "",
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    status_filter = request.query_params.get("status", "")

    company_id = user["company_id"]
    sql = (
        "SELECT wr.* "
        "FROM work_requests wr "
        "WHERE wr.company_id = ?"
    )
    params = [company_id]
    if status_filter:
        sql += " AND wr.status = ?"
        params.append(status_filter)
    order = REQ_SORTS.get(sort)
    sql += f" ORDER BY {order} {'DESC' if dir == 'desc' else 'ASC'}" if order else " ORDER BY wr.created_at DESC"
    rows = conn.execute(sql, params).fetchall()
    reqs = [dict(r) for r in rows]

    if export == "csv":
        data = [[r["title"], TASK_TYPE_LABELS.get(r["task_type"], r["task_type"]), r["volume"],
                 REQUEST_STATUS_LABELS.get(r["status"], r["status"]), r["due_date"] or "-"] for r in reqs]
        return csv_response(["제목", "업무 유형", "수량", "상태", "마감일"], data, "requests.csv")

    total_count = conn.execute(
        "SELECT COUNT(*) FROM work_requests WHERE company_id = ?", (company_id,)
    ).fetchone()[0]

    return templates.TemplateResponse(
        request=request, name="client/request_list.html", context={
            "request": request,
            "page_title": "업무 요청 목록",
            "user_name": user["name"],
            "requests": reqs,
            "status_filter": status_filter,
            "total_count": total_count,
            "task_type_labels": TASK_TYPE_LABELS,
            "request_status_labels": REQUEST_STATUS_LABELS,
            "sort": sort,
            "dir": dir,
        }
    )


def _active_workers(conn, company_id):
    return [dict(r) for r in conn.execute(
        "SELECT id, name, dept FROM employees WHERE company_id=? AND status='active' ORDER BY name",
        (company_id,),
    ).fetchall()]


def _pick_worker(conn, company_id, emp_id):
    if not emp_id.isdigit():
        return None
    row = conn.execute(
        "SELECT id FROM employees WHERE id=? AND company_id=? AND status='active'",
        (int(emp_id), company_id),
    ).fetchone()
    return row["id"] if row else None


# 요청 단위 담당자 지정 시 제출과 동시에 전체를 세부 업무 1건으로 배정
def _auto_assign(conn, req_id, user):
    wr = conn.execute("SELECT * FROM work_requests WHERE id=?", (req_id,)).fetchone()
    if not wr["assigned_to"] or wr["status"] != "submitted":
        return
    if conn.execute("SELECT 1 FROM tasks WHERE work_request_id=? LIMIT 1", (req_id,)).fetchone():
        return
    emp = conn.execute(
        "SELECT id, user_id FROM employees WHERE id=? AND company_id=? AND status='active'",
        (wr["assigned_to"], wr["company_id"]),
    ).fetchone()
    if not emp:
        return
    vol = wr["volume"] or 1
    if wr["item_spec"]:
        vol = conn.execute(
            "SELECT COUNT(*) FROM work_items WHERE request_id=? AND task_id IS NULL AND status='open'",
            (req_id,),
        ).fetchone()[0]
    tid = conn.execute(
        "INSERT INTO tasks "
        "(work_request_id, title, description, task_type, status, priority, due_date, assigned_to, assigned_by, requested_volume) "
        "VALUES (?, ?, ?, ?, 'assigned', ?, ?, ?, ?, ?)",
        (req_id, wr["title"], wr["description"], wr["task_type"], wr["priority"], wr["due_date"],
         emp["id"], user["id"], vol),
    ).lastrowid
    if wr["item_spec"]:
        assign_items(conn, req_id, tid, vol)
    log_activity(conn, 'work_request', req_id, user["id"], 'task_assigned', f"태스크 #{tid} 배정")
    create_notification(conn, emp["user_id"], f'업무 배정: {wr["title"]}', link=f'/task/{tid}')
    sync_request_volume(conn, req_id)
    advance_request(conn, req_id, 'in_progress', user["id"])
    conn.commit()


@router.get("/requests/new", response_class=HTMLResponse)
async def request_new(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    return templates.TemplateResponse(
        request=request, name="client/request_form.html", context={
            "request": request,
            "page_title": "새 업무 요청",
            "user_name": user["name"],
            "task_type_labels": TASK_TYPE_LABELS,
            "priority_labels": PRIORITY_LABELS,
            "output_format_labels": OUTPUT_FORMAT_LABELS,
            "workers": _active_workers(conn, user["company_id"]),
        }
    )


@router.post("/requests")
async def request_create(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
    title: str = Form(...),
    description: str = Form(""),
    task_type: str = Form(...),
    volume: int = Form(1),
    priority: str = Form("normal"),
    due_date: str = Form(""),
    output_format: str = Form("excel"),
    submit_now: str = Form(""),
    attachment: UploadFile = File(None),
    data_work: str = Form(""),
    item_csv: UploadFile = File(None),
    assignee: str = Form(""),
):
    company_id = user["company_id"]

    headers, rows = [], []
    if data_work == "1":
        if not item_csv or not item_csv.filename:
            return JSONResponse({"error": "건 단위 작업은 CSV 또는 엑셀 파일이 필요합니다."}, status_code=400)
        raw = await item_csv.read()
        try:
            headers, rows = parse_item_file(item_csv.filename, raw)
        except CsvError as e:
            return JSONResponse({"error": str(e)}, status_code=400)
        volume = len(rows)

    saved = None
    if attachment and attachment.filename:
        saved = await save_upload(attachment, upload_dir="var/uploads/work_requests")
        if isinstance(saved, JSONResponse):
            return saved
    conn.execute(
        "INSERT INTO work_requests "
        "(company_id, requested_by, title, description, task_type, volume, priority, status, due_date, "
        "output_format, total_volume, assigned_to) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, 'draft', ?, ?, ?, ?)",
        (company_id, user["id"], title, description, task_type, volume, priority,
         due_date or None, output_format, volume if rows else 0, _pick_worker(conn, company_id, assignee)),
    )
    new_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    log_activity(conn, 'work_request', new_id, user["id"], 'created', title)

    if rows:
        # 표시 열 선택 전 임시 보관. item-spec 저장 시 show 열 외 삭제
        for seq, row in enumerate(rows, 1):
            conn.execute(
                "INSERT INTO work_items (request_id, seq, source) VALUES (?, ?, ?)",
                (new_id, seq, json.dumps(row, ensure_ascii=False)),
            )

    if saved:
        original = saved.split("|")[0] if "|" in saved else saved
        conn.execute(
            "INSERT INTO work_request_attachments (request_id, user_id, file_name, file_path) VALUES (?,?,?,?)",
            (new_id, user["id"], original, saved),
        )
    conn.commit()

    if rows:
        return RedirectResponse(url=f"/client/requests/{new_id}/item-spec", status_code=303)

    if submit_now == "1":
        transition_request(conn, new_id, 'submit', user["id"], 'client')
        _auto_assign(conn, new_id, user)

    return RedirectResponse(url=f"/client/requests/{new_id}", status_code=303)


@router.get("/requests/{req_id}/item-spec", response_class=HTMLResponse)
async def item_spec_form(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    work_req = conn.execute(
        "SELECT * FROM work_requests WHERE id=? AND company_id=? AND status='draft'",
        (req_id, user["company_id"]),
    ).fetchone()
    if not work_req or work_req["item_spec"]:
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)
    item_rows = conn.execute(
        "SELECT source FROM work_items WHERE request_id=? ORDER BY seq", (req_id,)
    ).fetchall()
    if not item_rows:
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)
    sources = [json.loads(r["source"]) for r in item_rows]
    headers = list(sources[0].keys())
    row_count = len(sources)
    # 모든 행에서 비어 있는 열은 근로자 입력이 필요한 열로 보고 입력 항목으로 미리 채움
    blank_cols = [h for h in headers if all(not (src.get(h) or "").strip() for src in sources)]
    field_rows = [{"label": c, "type": "text"} for c in blank_cols] or [{"label": "", "type": "text"}]
    return templates.TemplateResponse(
        request=request, name="client/item_spec_form.html", context={
            "request": request, "page_title": "건 단위 작업 설정", "user_name": user["name"],
            "work_req": dict(work_req), "headers": headers, "row_count": row_count,
            "blank_cols": blank_cols, "field_rows": field_rows,
        }
    )


@router.post("/requests/{req_id}/item-spec")
async def item_spec_save(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    work_req = conn.execute(
        "SELECT * FROM work_requests WHERE id=? AND company_id=? AND status='draft'",
        (req_id, user["company_id"]),
    ).fetchone()
    if not work_req or work_req["item_spec"]:
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)

    form = await request.form()
    show = form.getlist("show")
    labels = form.getlist("field_label")
    types = form.getlist("field_type")
    options = form.getlist("field_options")
    requireds = form.getlist("field_required")
    try:
        sample_rate = int(form.get("sample_rate") or 10)
    except ValueError:
        sample_rate = 10

    if not show:
        return RedirectResponse(url=f"/client/requests/{req_id}/item-spec", status_code=303)

    fields = []
    for i, label in enumerate(labels):
        label = label.strip()
        if not label:
            continue
        ftype = types[i] if i < len(types) else "text"
        opts = [o.strip() for o in (options[i] if i < len(options) else "").split(",") if o.strip()]
        fields.append({
            "key": f"f{i + 1}",
            "label": label,
            "type": ftype if ftype in FIELD_TYPES else "text",
            "options": opts,
            "required": (requireds[i] if i < len(requireds) else "0") == "1",
        })
    if not fields:
        return RedirectResponse(url=f"/client/requests/{req_id}/item-spec", status_code=303)

    item_spec = {"show": show, "fields": fields, "sample_rate": sample_rate}

    items = conn.execute("SELECT id, source FROM work_items WHERE request_id=?", (req_id,)).fetchall()
    for it in items:
        src = json.loads(it["source"])
        trimmed = {c: src.get(c, "") for c in show}
        conn.execute("UPDATE work_items SET source=? WHERE id=?", (json.dumps(trimmed, ensure_ascii=False), it["id"]))

    conn.execute("UPDATE work_requests SET item_spec=? WHERE id=?", (json.dumps(item_spec, ensure_ascii=False), req_id))
    conn.commit()
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=item_spec_saved", status_code=303)


@router.get("/requests/{req_id}/edit", response_class=HTMLResponse)
async def request_edit(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT * FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, company_id),
    ).fetchone()
    if not row or row["status"] not in ("draft",):
        return RedirectResponse(url="/client/requests", status_code=303)
    work_req = dict(row)

    attachments = conn.execute(
        "SELECT * FROM work_request_attachments WHERE request_id=? ORDER BY created_at",
        (req_id,),
    ).fetchall()

    return templates.TemplateResponse(
        request=request, name="client/request_edit.html", context={
            "request": request,
            "page_title": "요청 수정",
            "user_name": user["name"],
            "task_type_labels": TASK_TYPE_LABELS,
            "priority_labels": PRIORITY_LABELS,
            "output_format_labels": OUTPUT_FORMAT_LABELS,
            "work_req": work_req,
            "attachments": [dict(a) for a in attachments],
            "workers": _active_workers(conn, company_id),
        }
    )


@router.post("/requests/{req_id}/edit")
async def request_update(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
    title: str = Form(...),
    description: str = Form(""),
    task_type: str = Form(...),
    volume: int = Form(1),
    priority: str = Form("normal"),
    due_date: str = Form(""),
    output_format: str = Form("excel"),
    attachment: UploadFile = File(None),
    assignee: str = Form(""),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT status FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, company_id),
    ).fetchone()
    if not row or row["status"] not in ("draft",):
        return JSONResponse({"error": "이미 제출된 요청은 수정할 수 없습니다."}, status_code=400)
    saved = None
    if attachment and attachment.filename:
        saved = await save_upload(attachment, upload_dir="var/uploads/work_requests")
        if isinstance(saved, JSONResponse):
            return saved
    conn.execute(
        "UPDATE work_requests "
        "SET title=?, description=?, task_type=?, volume=?, priority=?, due_date=?, "
        "output_format=?, assigned_to=?, updated_at=datetime('now','localtime') "
        "WHERE id=?",
        (title, description, task_type, volume, priority, due_date or None,
         output_format, _pick_worker(conn, company_id, assignee), req_id),
    )

    if saved:
        original = saved.split("|")[0] if "|" in saved else saved
        conn.execute(
            "INSERT INTO work_request_attachments (request_id, user_id, file_name, file_path) VALUES (?,?,?,?)",
            (req_id, user["id"], original, saved),
        )

    conn.commit()
    return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)


@router.post("/requests/{req_id}/attachments/{att_id}/delete")
async def request_attachment_delete(
    request: Request,
    req_id: int,
    att_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT wr.status FROM work_requests wr "
        "WHERE wr.id=? AND wr.company_id=?",
        (req_id, user["company_id"]),
    ).fetchone()
    if not row or row["status"] != "draft":
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)
    conn.execute(
        "DELETE FROM work_request_attachments WHERE id=? AND request_id=?",
        (att_id, req_id),
    )
    conn.commit()
    return RedirectResponse(url=f"/client/requests/{req_id}/edit", status_code=303)


@router.post("/requests/{req_id}/submit")
async def request_submit(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT id, item_spec FROM work_requests WHERE id=? AND company_id=?",
        (req_id, company_id),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    if not row["item_spec"] and conn.execute(
            "SELECT 1 FROM work_items WHERE request_id=? LIMIT 1", (req_id,)).fetchone():
        return RedirectResponse(url=f"/client/requests/{req_id}/item-spec", status_code=303)
    transition_request(conn, req_id, 'submit', user["id"], 'client')
    _auto_assign(conn, req_id, user)
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=submitted", status_code=303)


@router.post("/requests/{req_id}/resubmit")
async def request_resubmit(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT id FROM work_requests WHERE id=? AND company_id=?",
        (req_id, company_id),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    transition_request(conn, req_id, 'resubmit', user["id"], 'client')
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=resubmitted", status_code=303)


@router.post("/requests/{req_id}/close")
async def request_close(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT id FROM work_requests WHERE id=? AND company_id=?",
        (req_id, company_id),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    transition_request(conn, req_id, 'close', user["id"], 'client')
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=closed", status_code=303)


@router.post("/requests/{req_id}/cancel")
async def request_cancel(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT status, item_spec FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, company_id),
    ).fetchone()
    if row and row["status"] in ("draft", "submitted"):
        transition_request(conn, req_id, 'cancel', user["id"], 'client')
        # 표시 열 미확정 원본 전체 행 즉시 삭제
        if not row["item_spec"]:
            conn.execute("DELETE FROM work_items WHERE request_id=?", (req_id,))
            conn.commit()

    return RedirectResponse(url="/client/requests?msg=cancelled", status_code=303)


@router.post("/requests/{req_id}/comments")
async def request_add_comment(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
    content: str = Form(...),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT id FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, company_id),
    ).fetchone()
    if row:
        conn.execute(
            "INSERT INTO work_request_comments (request_id, user_id, author, content) VALUES (?, ?, ?, ?)",
            (req_id, user["id"], user["name"], content),
        )
        log_activity(conn, 'work_request', req_id, user["id"], 'commented', content[:50])
        staff_ids = conn.execute("SELECT id FROM users WHERE role='platform_staff'").fetchall()
        for s in staff_ids:
            create_notification(conn, s["id"], f'댓글: {content[:30]}', link=f'/platform/work-requests/{req_id}')

    return RedirectResponse(url=f"/client/requests/{req_id}?msg=comment", status_code=303)


@router.get("/requests/{req_id}/export")
async def request_export(
    req_id: int,
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    wr = conn.execute(
        "SELECT id FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, user["company_id"])).fetchone()
    if not wr:
        return RedirectResponse("/client/requests", status_code=303)

    tasks = conn.execute(
        "SELECT id, erp_module FROM tasks WHERE work_request_id = ? AND status = 'completed'",
        (req_id,)).fetchall()
    groups = {}
    for t in tasks:
        key = EXPORT_KEY.get(t["erp_module"] or "")
        if key:
            groups.setdefault(key, []).append(t["id"])

    datasets = []
    for key, ids in groups.items():
        rows = fetch_rows(conn, key, ids)
        if not rows:
            continue
        tpl = conn.execute(
            "SELECT * FROM export_templates WHERE company_id = ? AND module = ?",
            (user["company_id"], key)).fetchone()
        if tpl:
            headers, lines = apply_template(key, tpl, rows)
            datasets.append((key, headers, lines, tpl["file_format"], tpl["id"]))
        else:
            headers, lines = std_lines(key, rows)
            datasets.append((key, headers, lines, "xlsx", None))

    if not datasets:
        return RedirectResponse(f"/client/requests/{req_id}?msg=no_export", status_code=303)

    total = sum(len(d[2]) for d in datasets)
    tpl_id = datasets[0][4] if len(datasets) == 1 else None
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    if all(d[3] == "xlsx" for d in datasets):
        fname = f"export_req{req_id}_{stamp}.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        for key, headers, lines, _, _ in datasets:
            styled_sheet(wb, EXPORT_SCHEMAS[key]["label"], headers, lines)
        resp = excel_response(wb, fname)
    elif len(datasets) == 1:
        key, headers, lines, _, _ = datasets[0]
        fname = f"export_req{req_id}_{stamp}.csv"
        resp = Response(content=csv_bytes(headers, lines), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
    else:
        fname = f"export_req{req_id}_{stamp}.zip"
        mem = io.BytesIO()
        with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as z:
            for key, headers, lines, fmt, _ in datasets:
                label = EXPORT_SCHEMAS[key]["label"]
                if fmt == "csv":
                    z.writestr(f"{label}.csv", csv_bytes(headers, lines))
                else:
                    b = io.BytesIO()
                    swb = Workbook()
                    swb.remove(swb.active)
                    styled_sheet(swb, label, headers, lines)
                    swb.save(b)
                    z.writestr(f"{label}.xlsx", b.getvalue())
        mem.seek(0)
        resp = Response(content=mem.getvalue(), media_type="application/zip",
                        headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    conn.execute(
        "INSERT INTO export_logs (company_id, user_id, work_request_id, kind, template_id, row_count, file_name) "
        "VALUES (?, ?, ?, 'data', ?, ?, ?)",
        (user["company_id"], user["id"], req_id, tpl_id, total, fname))
    conn.commit()
    return resp


@router.get("/requests/{req_id}/items/export")
async def request_items_export(
    req_id: int,
    fmt: str = "csv",
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    work_req = conn.execute(
        "SELECT * FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, user["company_id"])).fetchone()
    item_spec = get_item_spec(work_req)
    if not work_req or not item_spec:
        return RedirectResponse("/client/requests", status_code=303)

    headers, rows = export_rows(conn, req_id, item_spec)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    if fmt == "xlsx":
        fname = f"items_req{req_id}_{stamp}.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        styled_sheet(wb, "결과", headers, rows)
        resp = excel_response(wb, fname)
    else:
        fname = f"items_req{req_id}_{stamp}.csv"
        resp = csv_response(headers, rows, fname)
    conn.execute(
        "INSERT INTO export_logs (company_id, user_id, work_request_id, kind, row_count, file_name) "
        "VALUES (?, ?, ?, 'data', ?, ?)",
        (user["company_id"], user["id"], req_id, len(rows), fname))
    conn.commit()
    return resp


@router.get("/requests/{req_id}/deliverables.zip")
async def request_deliverables_zip(
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    wr = conn.execute(
        "SELECT id FROM work_requests WHERE id = ? AND company_id = ?",
        (req_id, user["company_id"])).fetchone()
    if not wr:
        return RedirectResponse("/client/requests", status_code=303)
    files = conn.execute(
        "SELECT d.file_name, d.file_path FROM task_deliverables d "
        "JOIN tasks t ON t.id = d.task_id "
        "WHERE t.work_request_id = ? AND t.status = 'completed'", (req_id,)).fetchall()
    if not files:
        return RedirectResponse(f"/client/requests/{req_id}?msg=no_export", status_code=303)
    mem = io.BytesIO()
    cnt = 0
    with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            stored = f["file_path"].split("|")[-1]
            path = os.path.join("var", "uploads", "deliverables", os.path.basename(stored))
            if os.path.exists(path):
                z.write(path, os.path.basename(f["file_name"] or stored))
                cnt += 1
    mem.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    fname = f"deliverables_req{req_id}_{stamp}.zip"
    conn.execute(
        "INSERT INTO export_logs (company_id, user_id, work_request_id, kind, row_count, file_name) "
        "VALUES (?, ?, ?, 'files', ?, ?)",
        (user["company_id"], user["id"], req_id, cnt, fname))
    conn.commit()
    return Response(content=mem.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/requests/{req_id}", response_class=HTMLResponse)
async def request_detail(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    work_req = conn.execute(
        "SELECT wr.*, cc.name as company_name "
        "FROM work_requests wr "
        "JOIN client_companies cc ON wr.company_id = cc.id "
        "WHERE wr.id = ? AND wr.company_id = ?",
        (req_id, company_id),
    ).fetchone()
    if not work_req:
        return RedirectResponse(url="/", status_code=303)
    work_req = dict(work_req)
    work_req["due_status"] = get_due_status(work_req.get("due_date"))

    tasks = conn.execute(
        "SELECT t.*, e.name as worker_name "
        "FROM tasks t "
        "LEFT JOIN employees e ON t.assigned_to = e.id "
        "WHERE t.work_request_id = ? "
        "ORDER BY t.id",
        (req_id,),
    ).fetchall()
    tasks = [dict(t) for t in tasks]

    comments = conn.execute(
        "SELECT * FROM work_request_comments WHERE request_id = ? AND is_internal = 0 ORDER BY created_at",
        (req_id,),
    ).fetchall()
    comments = [dict(c) for c in comments]

    activities = get_activity_logs(conn, 'work_request', req_id)
    activities = [dict(a) for a in activities]

    attachments = conn.execute(
        "SELECT * FROM work_request_attachments WHERE request_id=? ORDER BY created_at",
        (req_id,),
    ).fetchall()
    attachments = [dict(a) for a in attachments]
    for a in attachments:
        fp = a.get("file_path", "")
        a["file_url"] = fp.split("|")[1] if "|" in fp else fp

    workers = conn.execute(
        "SELECT e.id, e.name, e.dept, "
        "(SELECT COUNT(*) FROM tasks t2 WHERE t2.assigned_to=e.id "
        " AND t2.status NOT IN ('completed','cancelled')) AS task_count "
        "FROM employees e WHERE e.company_id=? AND e.status='active' ORDER BY e.name",
        (company_id,),
    ).fetchall()

    allowed_actions = []
    for action in REQUEST_ROLE_ACTIONS.get('client', set()):
        if (work_req['status'], action) in REQUEST_TRANSITIONS:
            allowed_actions.append(action)

    approved_cnt = conn.execute(
        "SELECT COUNT(*) FROM tasks WHERE work_request_id = ? AND status = 'completed'",
        (req_id,)).fetchone()[0]
    has_deliverables = bool(conn.execute(
        "SELECT 1 FROM task_deliverables d JOIN tasks t ON t.id = d.task_id "
        "WHERE t.work_request_id = ? AND t.status = 'completed' LIMIT 1",
        (req_id,)).fetchone())
    exports = [dict(r) for r in conn.execute(
        "SELECT e.*, u.name AS user_name FROM export_logs e JOIN users u ON u.id = e.user_id "
        "WHERE e.work_request_id = ? ORDER BY e.created_at DESC", (req_id,)).fetchall()]

    item_spec = get_item_spec(work_req)
    unassigned_cnt = 0
    item_spec_pending = False
    if item_spec:
        unassigned_cnt = conn.execute(
            "SELECT COUNT(*) FROM work_items WHERE request_id=? AND task_id IS NULL AND status='open'",
            (req_id,)).fetchone()[0]
    else:
        item_spec_pending = bool(conn.execute(
            "SELECT 1 FROM work_items WHERE request_id=? LIMIT 1", (req_id,)).fetchone())

    return templates.TemplateResponse(
        request=request, name="client/request_detail.html", context={
            "request": request,
            "page_title": work_req["title"],
            "user_name": user["name"],
            "work_req": work_req,
            "tasks": tasks,
            "comments": comments,
            "activities": activities,
            "attachments": attachments,
            "allowed_actions": allowed_actions,
            "task_type_labels": TASK_TYPE_LABELS,
            "status_labels": STATUS_LABELS,
            "request_status_labels": REQUEST_STATUS_LABELS,
            "workers": [dict(w) for w in workers],
            "assignee": next((w["name"] for w in workers if w["id"] == work_req["assigned_to"]), ""),
            "erp_modules": ERP_MODULES,
            "priority_labels": PRIORITY_LABELS,
            "approved_cnt": approved_cnt,
            "has_deliverables": has_deliverables,
            "exports": exports,
            "item_spec": item_spec,
            "unassigned_cnt": unassigned_cnt,
            "item_spec_pending": item_spec_pending,
            "msg": request.query_params.get("msg", ""),
        }
    )


@router.post("/requests/{req_id}/tasks/new")
async def client_task_create(
    request: Request,
    req_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
    title: str = Form(...),
    description: str = Form(""),
    due_date: str = Form(""),
    requested_volume: int = Form(0),
    erp_module: str = Form(""),
    priority: str = Form(""),
):
    work_req = conn.execute(
        "SELECT * FROM work_requests WHERE id=? AND company_id=?",
        (req_id, user["company_id"]),
    ).fetchone()
    if not work_req or work_req["status"] in ("draft", "closed", "cancelled"):
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)
    if work_req["item_spec"]:
        # 건 단위 작업은 완료 수량을 건 기준으로 산정. ERP 문서 건수 동기화 제외
        erp_module = ""
        left = conn.execute(
            "SELECT COUNT(*) FROM work_items WHERE request_id=? AND task_id IS NULL AND status='open'",
            (req_id,),
        ).fetchone()[0]
        requested_volume = min(requested_volume, left)
        if requested_volume <= 0:
            return RedirectResponse(url=f"/client/requests/{req_id}?msg=no_items_left", status_code=303)

    conn.execute(
        "INSERT INTO tasks "
        "(work_request_id, title, description, task_type, status, priority, due_date, assigned_by, requested_volume, erp_module) "
        "VALUES (?, ?, ?, ?, 'pending', ?, ?, ?, ?, ?)",
        (req_id, title, description, work_req["task_type"], priority or work_req["priority"],
         (due_date or work_req["due_date"]) or None, user["id"], requested_volume, erp_module),
    )
    if work_req["item_spec"]:
        tid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        assign_items(conn, req_id, tid, requested_volume)
    log_activity(conn, 'work_request', req_id, user["id"], 'task_created', title)
    sync_request_volume(conn, req_id)
    conn.commit()
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=task_created", status_code=303)


@router.post("/requests/{req_id}/tasks/{tid}/assign")
async def client_task_assign(
    request: Request,
    req_id: int,
    tid: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
    employee_id: str = Form(""),
):
    company_id = user["company_id"]
    task = conn.execute(
        "SELECT t.* FROM tasks t JOIN work_requests wr ON t.work_request_id=wr.id "
        "WHERE t.id=? AND t.work_request_id=? AND wr.company_id=?",
        (tid, req_id, company_id),
    ).fetchone()
    if not task or task["status"] not in ("pending", "assigned"):
        return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)

    if employee_id:
        emp = conn.execute(
            "SELECT id, user_id FROM employees WHERE id=? AND company_id=? AND status='active'",
            (int(employee_id), company_id),
        ).fetchone()
        if not emp:
            return RedirectResponse(url=f"/client/requests/{req_id}", status_code=303)
        conn.execute(
            "UPDATE tasks SET assigned_to=?, assigned_by=?, status='assigned', "
            "updated_at=datetime('now','localtime') WHERE id=?",
            (emp["id"], user["id"], tid),
        )
        log_activity(conn, 'work_request', req_id, user["id"], 'task_assigned', f"태스크 #{tid} 배정")
        create_notification(conn, emp["user_id"], f'업무 배정: {task["title"]}', link=f'/task/{tid}')
        advance_request(conn, req_id, 'in_progress', user["id"])
        conn.commit()
        return RedirectResponse(url=f"/client/requests/{req_id}?msg=assigned", status_code=303)
    else:
        conn.execute(
            "UPDATE tasks SET assigned_to=NULL, status='pending', "
            "updated_at=datetime('now','localtime') WHERE id=?",
            (tid,),
        )
        log_activity(conn, 'work_request', req_id, user["id"], 'task_unassigned', f"태스크 #{tid} 배정 해제")

    conn.commit()
    return RedirectResponse(url=f"/client/requests/{req_id}?msg=unassigned", status_code=303)


def _get_company_task(conn, task_id, company_id):
    return conn.execute(
        "SELECT t.*, wr.company_id, wr.title AS request_title, e.name AS worker_name "
        "FROM tasks t "
        "JOIN work_requests wr ON t.work_request_id = wr.id "
        "LEFT JOIN employees e ON t.assigned_to = e.id "
        "WHERE t.id = ? AND wr.company_id = ?",
        (task_id, company_id),
    ).fetchone()


@router.get("/tasks/{task_id}", response_class=HTMLResponse)
async def client_task_detail(
    request: Request,
    task_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    task = dict(row)

    deliverables = conn.execute(
        "SELECT * FROM task_deliverables WHERE task_id=? ORDER BY created_at DESC",
        (task_id,),
    ).fetchall()
    deliverable_list = []
    for d in deliverables:
        dd = dict(d)
        parts = dd["file_path"].split("|", 1)
        dd["stored_name"] = parts[1] if len(parts) == 2 else parts[0]
        deliverable_list.append(dd)

    comments = conn.execute(
        "SELECT tc.*, u.role as user_role "
        "FROM task_comments tc JOIN users u ON tc.user_id = u.id "
        "WHERE tc.task_id=? ORDER BY tc.created_at ASC",
        (task_id,),
    ).fetchall()

    work_req = conn.execute("SELECT * FROM work_requests WHERE id=?", (task["work_request_id"],)).fetchone()
    item_spec = get_item_spec(work_req)
    review_items = []
    item_progress = None
    round_stat = None
    if item_spec:
        total = conn.execute("SELECT COUNT(*) FROM work_items WHERE task_id=?", (task_id,)).fetchone()[0]
        done = conn.execute(
            "SELECT COUNT(*) FROM work_items WHERE task_id=? AND status IN ('done','approved','rejected')",
            (task_id,),
        ).fetchone()[0]
        item_progress = {"total": total, "done": done}
        if task["status"] == "review":
            rows = conn.execute(
                "SELECT * FROM work_items WHERE task_id=? AND (status='hold' OR (status='done' AND sampled=1)) "
                "ORDER BY seq",
                (task_id,),
            ).fetchall()
            for r in rows:
                it = dict(r)
                src = json.loads(it["source"]) if it["source"] else {}
                res = json.loads(it["result"]) if it["result"] else {}
                it["src_line"] = " / ".join(str(src.get(c, "")) for c in item_spec["show"])
                it["res_text"] = {}
                for f in item_spec["fields"]:
                    v = res.get(f["key"], "")
                    it["res_text"][f["key"]] = ", ".join(v) if isinstance(v, list) else v
                review_items.append(it)
            stat = conn.execute(
                "SELECT COUNT(*), SUM(status='approved'), SUM(status='rejected') FROM work_items "
                "WHERE task_id=? AND sampled=1 AND status IN ('approved','rejected') "
                "AND round = (SELECT MAX(round) FROM work_items WHERE task_id=?)",
                (task_id, task_id),
            ).fetchone()
            back = conn.execute(
                "SELECT COUNT(*) FROM work_items WHERE task_id=? AND status='open'", (task_id,)
            ).fetchone()[0]
            if stat[0]:
                acc = round(stat[1] / stat[0] * 100)
                round_stat = {"total": stat[0], "ok": stat[1], "bad": stat[2], "back": back, "acc": acc,
                              "pass": stat[1] / stat[0] * 100 >= PASS_ACCURACY and not back}

    return templates.TemplateResponse(
        request=request, name="client/task_detail.html", context={
            "request": request,
            "page_title": task["title"],
            "user_name": user["name"],
            "task": task,
            "deliverables": deliverable_list,
            "comments": [dict(c) for c in comments],
            "status_labels": STATUS_LABELS,
            "item_spec": item_spec,
            "item_progress": item_progress,
            "review_items": review_items,
            "pending_review": review_cnt(conn, task_id) if item_spec else 0,
            "round_stat": round_stat,
            "pass_acc": PASS_ACCURACY,
            "wi_status_labels": WI_STATUS_LABELS,
        }
    )


@router.post("/tasks/{task_id}/approve")
async def client_task_approve(
    request: Request,
    task_id: int,
    review_notes: str = Form(""),
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row or row["status"] != "review":
        return RedirectResponse(url="/client/requests", status_code=303)
    work_req = conn.execute("SELECT item_spec FROM work_requests WHERE id=?", (row["work_request_id"],)).fetchone()
    if get_item_spec(work_req):
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)
    transition_task(conn, task_id, 'approve', user["id"], 'client', review_notes)
    sync_request_volume(conn, row["work_request_id"])
    left = conn.execute(
        "SELECT COUNT(*) FROM tasks WHERE work_request_id=? AND status NOT IN ('completed','cancelled')",
        (row["work_request_id"],),
    ).fetchone()[0]
    if left == 0:
        advance_request(conn, row["work_request_id"], 'completed', user["id"])
    conn.commit()
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=approved", status_code=303)


@router.post("/tasks/{task_id}/reject")
async def client_task_reject(
    request: Request,
    task_id: int,
    review_notes: str = Form(...),
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row or row["status"] != "review":
        return RedirectResponse(url="/client/requests", status_code=303)
    work_req = conn.execute("SELECT item_spec FROM work_requests WHERE id=?", (row["work_request_id"],)).fetchone()
    if get_item_spec(work_req):
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)
    transition_task(conn, task_id, 'reject', user["id"], 'client', review_notes)
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=rejected", status_code=303)


@router.post("/tasks/{task_id}/items/{item_id}/approve")
async def client_item_approve(
    request: Request,
    task_id: int,
    item_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row or row["status"] != "review":
        return RedirectResponse(url="/client/requests", status_code=303)
    item = conn.execute("SELECT * FROM work_items WHERE id=? AND task_id=?", (item_id, task_id)).fetchone()
    if not item:
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)

    now = now_kst().strftime('%Y-%m-%d %H:%M:%S')
    if item["status"] == "hold":
        work_req = conn.execute("SELECT item_spec FROM work_requests WHERE id=?", (row["work_request_id"],)).fetchone()
        spec = get_item_spec(work_req)
        form = await request.form()
        result = {}
        for f in spec["fields"]:
            v = form.getlist(f["key"]) if f["type"] == "tags" else form.get(f["key"], "")
            if f.get("required") and not v:
                return RedirectResponse(url=f"/client/tasks/{task_id}?msg=field_required", status_code=303)
            if f["type"] in ("number", "date", "phone") and v and not validate_value(f["type"], v):
                return RedirectResponse(
                    url=f"/client/tasks/{task_id}?msg=field_invalid&label={quote(f['label'])}", status_code=303)
            result[f["key"]] = v
        conn.execute(
            "UPDATE work_items SET status='approved', result=?, reviewed_by=?, reviewed_at=? WHERE id=?",
            (json.dumps(result, ensure_ascii=False), user["id"], now, item_id),
        )
    elif item["status"] == "done":
        conn.execute(
            "UPDATE work_items SET status='approved', reviewed_by=?, reviewed_at=? WHERE id=?",
            (user["id"], now, item_id),
        )
    conn.commit()
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=item_approved", status_code=303)


@router.post("/tasks/{task_id}/items/{item_id}/reject")
async def client_item_reject(
    request: Request,
    task_id: int,
    item_id: int,
    note: str = Form(...),
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row or row["status"] != "review":
        return RedirectResponse(url="/client/requests", status_code=303)
    item = conn.execute("SELECT * FROM work_items WHERE id=? AND task_id=?", (item_id, task_id)).fetchone()
    if not item:
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)

    now = now_kst().strftime('%Y-%m-%d %H:%M:%S')
    if item["status"] == "hold":
        conn.execute(
            "UPDATE work_items SET status='open', hold_note=?, reviewed_by=?, reviewed_at=? WHERE id=?",
            (note, user["id"], now, item_id),
        )
    elif item["status"] == "done":
        conn.execute(
            "UPDATE work_items SET status='rejected', review_note=?, reviewed_by=?, reviewed_at=? WHERE id=?",
            (note, user["id"], now, item_id),
        )
    conn.commit()
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=item_rejected", status_code=303)


@router.post("/tasks/{task_id}/finish")
async def client_task_finish(
    request: Request,
    task_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row or row["status"] != "review":
        return RedirectResponse(url="/client/requests", status_code=303)
    work_req = conn.execute("SELECT item_spec FROM work_requests WHERE id=?", (row["work_request_id"],)).fetchone()
    if not get_item_spec(work_req):
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)
    if review_cnt(conn, task_id) > 0:
        return RedirectResponse(url=f"/client/tasks/{task_id}", status_code=303)

    finalize_review(conn, task_id, user["id"])
    sync_request_volume(conn, row["work_request_id"])
    left = conn.execute(
        "SELECT COUNT(*) FROM tasks WHERE work_request_id=? AND status NOT IN ('completed','cancelled')",
        (row["work_request_id"],),
    ).fetchone()[0]
    if left == 0:
        advance_request(conn, row["work_request_id"], 'completed', user["id"])
    conn.commit()
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=finished", status_code=303)


@router.post("/tasks/{task_id}/cancel")
async def client_task_cancel(
    request: Request,
    task_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    transition_task(conn, task_id, 'cancel', user["id"], 'client')
    release_items(conn, task_id)
    sync_request_volume(conn, row["work_request_id"])
    log_activity(conn, 'work_request', row["work_request_id"], user["id"], 'task_cancelled', f"태스크 #{task_id} 취소")
    conn.commit()
    return RedirectResponse(url=f"/client/requests/{row['work_request_id']}?msg=task_cancelled", status_code=303)


@router.post("/tasks/{task_id}/comments")
async def client_task_comment(
    request: Request,
    task_id: int,
    content: str = Form(...),
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    row = _get_company_task(conn, task_id, user["company_id"])
    if not row:
        return RedirectResponse(url="/client/requests", status_code=303)
    conn.execute(
        "INSERT INTO task_comments (task_id, user_id, author, content) VALUES (?,?,?,?)",
        (task_id, user["id"], user["name"], content),
    )
    if row["assigned_to"]:
        worker_uid = conn.execute(
            "SELECT user_id FROM employees WHERE id=?", (row["assigned_to"],)
        ).fetchone()
        if worker_uid:
            create_notification(conn, worker_uid["user_id"], f'댓글: {content[:30]}', link=f'/task/{task_id}')
    conn.commit()
    return RedirectResponse(url=f"/client/tasks/{task_id}?msg=comment", status_code=303)


@router.get("/notifications", response_class=HTMLResponse)
async def client_notifications(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    notifs = get_notifications(conn, user["id"], limit=50)
    mark_all_read(conn, user["id"])
    return templates.TemplateResponse(
        request=request,
        name="client/notifications.html",
        context={
            "request": request,
            "page_title": "알림",
            "user_name": user["name"],
            "notifications": [dict(n) for n in notifs],
        },
    )


@router.get("/monthly-report", response_class=HTMLResponse)
async def monthly_report_list(
    request: Request,
    user: dict = Depends(require_client_with_company("hr")),
    conn=Depends(get_db),
):
    company = conn.execute(
        "SELECT name FROM client_companies WHERE id=?", (user["company_id"],)
    ).fetchone()

    reports = conn.execute(
        "SELECT * FROM monthly_reports "
        "WHERE company_id=? AND status='published' "
        "ORDER BY year DESC, month DESC",
        (user["company_id"],),
    ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="client/monthly_report.html",
        context={
            "request": request,
            "page_title": "월간 리포트",
            "user_name": user["name"],
            "company_name": company["name"] if company else "",
            "reports": [dict(r) for r in reports],
        },
    )


@router.get("/monthly-report/{year}/{month}", response_class=HTMLResponse)
async def monthly_report_detail(
    request: Request,
    year: int,
    month: int,
    user: dict = Depends(require_client_with_company("hr")),
    conn=Depends(get_db),
):
    report = conn.execute(
        "SELECT mr.*, cc.name AS company_name "
        "FROM monthly_reports mr "
        "JOIN client_companies cc ON mr.company_id=cc.id "
        "WHERE mr.company_id=? AND mr.year=? AND mr.month=? AND mr.status='published'",
        (user["company_id"], year, month),
    ).fetchone()
    if not report:
        return RedirectResponse(url="/client/monthly-report", status_code=303)

    data = json.loads(report["data"])

    return templates.TemplateResponse(
        request=request,
        name="client/monthly_report_detail.html",
        context={
            "request": request,
            "page_title": f"{year}년 {month}월 리포트",
            "user_name": user["name"],
            "company_name": report["company_name"],
            "report": dict(report),
            "data": data,
            "task_type_labels": TASK_TYPE_LABELS,
        },
    )


@router.get("/monthly-report/{year}/{month}/pdf")
async def monthly_report_pdf(
    request: Request,
    year: int,
    month: int,
    user: dict = Depends(require_client_with_company("hr")),
    conn=Depends(get_db),
):
    report = conn.execute(
        "SELECT mr.*, cc.name AS company_name "
        "FROM monthly_reports mr "
        "JOIN client_companies cc ON mr.company_id=cc.id "
        "WHERE mr.company_id=? AND mr.year=? AND mr.month=? AND mr.status='published'",
        (user["company_id"], year, month),
    ).fetchone()
    if not report:
        return RedirectResponse(url="/client/monthly-report", status_code=303)

    data = json.loads(report["data"])
    html = templates.get_template("report/monthly_report_pdf.html").render(
        report=dict(report), data=data,
        task_type_labels=TASK_TYPE_LABELS,
    )
    filename = f"report_{report['company_name']}_{year}_{month:02d}.pdf"
    return pdf_response(html, filename)
