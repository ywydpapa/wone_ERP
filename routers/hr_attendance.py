import calendar
import json

from core.tz import now_kst

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse
from core.deps import get_db, require_login, templates
from core.constants import APPEAL_TYPE_LABELS, COMPANY_INFO
from core.attendance import (
    get_attendance_records,
    get_attendance_summary,
    get_worker_detail,
    get_worker_month_records,
    get_worker_month_leaves,
    merge_records_leaves,
    build_calendar_data,
    worker_month_summary,
    prev_next_month,
    available_months,
    is_month_closed,
    get_closing_info,
    STATUS_LABELS,
)
from core.export import styled_workbook, excel_response, pdf_response

router = APIRouter()


def _build_sheet_data(conn, year, month, company_id):
    _, days_in_month = calendar.monthrange(year, month)
    ym = f"{year}-{month:02d}"

    emp_sql = "SELECT id, name, employee_no, monthly_work_hours, hire_date FROM employees WHERE status='active'"
    emp_params = []
    if company_id:
        emp_sql += " AND company_id=?"
        emp_params.append(company_id)
    emp_sql += " ORDER BY employee_no, name"
    employees = conn.execute(emp_sql, emp_params).fetchall()
    emp_ids = [e["id"] for e in employees]

    if not emp_ids:
        return [], days_in_month, []

    ph = ",".join("?" * len(emp_ids))
    att_rows = conn.execute(
        f"""SELECT employee_id, work_date, status, clock_in, clock_out,
                   work_minutes, overtime_minutes, break_minutes, notes
            FROM attendance
            WHERE employee_id IN ({ph}) AND work_date LIKE ?
            ORDER BY work_date""",
        [*emp_ids, f"{ym}%"],
    ).fetchall()

    att_map = {}
    for r in att_rows:
        att_map.setdefault(r["employee_id"], {})[r["work_date"]] = dict(r)

    rows = []
    for emp in employees:
        eid = emp["id"]
        days = att_map.get(eid, {})
        cells = []
        work_days = 0
        late_cnt = 0
        early_cnt = 0
        absent_cnt = 0
        holiday_cnt = 0
        total_mins = 0
        ot_mins = 0
        for d in range(1, days_in_month + 1):
            dt = f"{ym}-{d:02d}"
            rec = days.get(dt)
            if rec:
                cells.append(rec)
                if rec["status"] in ("normal", "late", "early_leave"):
                    work_days += 1
                    total_mins += rec.get("work_minutes") or 0
                    ot_mins += rec.get("overtime_minutes") or 0
                if rec["status"] == "late":
                    late_cnt += 1
                elif rec["status"] == "early_leave":
                    early_cnt += 1
                elif rec["status"] == "absent":
                    absent_cnt += 1
                elif rec["status"] == "holiday":
                    holiday_cnt += 1
            else:
                cells.append(None)
        rows.append({
            "employee": dict(emp),
            "cells": cells,
            "work_days": work_days,
            "late": late_cnt,
            "early_leave": early_cnt,
            "absent": absent_cnt,
            "holiday_days": holiday_cnt,
            "monthly_work_hours": emp["monthly_work_hours"] or 0,
            "total_mins": total_mins,
            "total_hours": total_mins // 60,
            "total_mins_r": total_mins % 60,
            "ot_hours": ot_mins // 60,
            "ot_mins_r": ot_mins % 60,
        })
    return rows, days_in_month, employees


@router.get("/attendance", response_class=HTMLResponse)
async def hr_attendance(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    status_filter = request.query_params.get("status", "")
    company_filter = request.query_params.get("company", "")

    emp_sql = "SELECT id FROM employees WHERE status = 'active'"
    emp_params = []
    if company_filter:
        emp_sql += " AND company_id = ?"
        emp_params.append(int(company_filter))
    emp_rows = conn.execute(emp_sql, emp_params).fetchall()
    emp_ids = [r["id"] for r in emp_rows]

    records = get_attendance_records(conn, emp_ids, month)
    summary = get_attendance_summary(records, len(emp_ids))
    if status_filter:
        statuses = status_filter.split(",")
        records = [r for r in records if r["status"] in statuses]
    for r in records:
        r.setdefault("company_name", r.get("company_name") or "-")

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status = 'active' ORDER BY name"
    ).fetchall()

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance.html",
        context={
            "request": request,
            "page_title": "근태 관리",
            "user_name": user["name"],
            "user_role": user["role"],
            "records": records,
            "summary": summary,
            "current_month": month,
            "status_filter": status_filter,
            "company_filter": company_filter,
            "companies": companies,
            "months": available_months(),
        },
    )


@router.get("/attendance/sheet", response_class=HTMLResponse)
async def attendance_sheet(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    page = int(request.query_params.get("page", 1))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None

    rows, days_in_month, _ = _build_sheet_data(conn, year, month, company_id)

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()

    # 고객사 정보 (문서 헤더용)
    company_name = ""
    company_info = None
    if company_id:
        ci_row = conn.execute(
            "SELECT name, business_no, representative, address FROM client_companies WHERE id=?",
            (company_id,),
        ).fetchone()
        if ci_row:
            company_info = dict(ci_row)
            company_name = ci_row["name"]

    first_weekday = calendar.monthrange(year, month)[0]
    weekdays = [(first_weekday + d) % 7 for d in range(days_in_month)]

    # 배너 집계
    weekday_count = sum(1 for wd in weekdays if wd < 5)
    total_emp = len(rows)
    sum_work = sum(r["work_days"] for r in rows)
    sum_late = sum(r["late"] for r in rows)
    sum_absent = sum(r["absent"] for r in rows)
    expected = weekday_count * total_emp
    banner = {
        "total_emp": total_emp,
        "work": sum_work,
        "late": sum_late,
        "absent": sum_absent,
        "rate": round(sum_work / expected * 100, 1) if expected else 0,
    }

    # 미처리 집계 (근태마감용)
    today_str = now.strftime("%Y-%m-%d")
    unrecorded = []
    for r in rows:
        emp = r["employee"]
        hire = emp.get("hire_date") or "2000-01-01"
        cnt = 0
        for d in range(days_in_month):
            if weekdays[d] >= 5:
                continue
            dt = f"{year}-{month:02d}-{d+1:02d}"
            if dt < hire or dt > today_str:
                continue
            if r["cells"][d] is None:
                cnt += 1
        if cnt:
            unrecorded.append({"id": emp["id"], "name": emp["name"], "no": emp["employee_no"], "count": cnt})
    unrecorded_total = sum(u["count"] for u in unrecorded)

    # 마감 상태
    closing_info = get_closing_info(conn, year, month, company_id) if company_id else None
    is_closed = closing_info and closing_info["status"] == "closed"
    if closing_info and closing_info.get("closed_at"):
        closing_info["closed_date"] = closing_info["closed_at"][:10]

    # 미처리 소명
    pending_appeals = 0
    if company_id:
        pending_appeals = conn.execute(
            "SELECT COUNT(*) as cnt FROM attendance a JOIN employees e ON a.employee_id=e.id "
            "WHERE a.appeal='pending' AND e.company_id=? AND a.work_date LIKE ?",
            (company_id, f"{year}-{month:02d}%"),
        ).fetchone()["cnt"]

    # 페이지네이션
    per_page = 50
    total_pages = max(1, (total_emp + per_page - 1) // per_page)
    if page < 1:
        page = 1
    if page > total_pages:
        page = total_pages
    rows_page = rows[(page - 1) * per_page: page * per_page]

    sheet_totals = {
        "work": sum(r["work_days"] for r in rows),
        "late": sum(r["late"] for r in rows),
        "early": sum(r["early_leave"] for r in rows),
        "absent": sum(r["absent"] for r in rows),
        "holiday": sum(r["holiday_days"] for r in rows),
        "ot_m": sum(r["ot_hours"] * 60 + r["ot_mins_r"] for r in rows),
        "total_m": sum(r["total_mins"] for r in rows),
    }

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance_sheet.html",
        context={
            "request": request,
            "page_title": "출근부",
            "user_name": user["name"],
            "user_role": user["role"],
            "rows": rows_page,
            "all_rows": rows,
            "sheet_totals": sheet_totals,
            "year": year,
            "month": month,
            "days_in_month": days_in_month,
            "weekdays": weekdays,
            "companies": companies,
            "company_id": company_id or "",
            "company_name": company_name,
            "company_info": company_info,
            "banner": banner,
            "page": page,
            "total_pages": total_pages,
            "months": available_months(),
            "status_labels": STATUS_LABELS,
            "unrecorded": unrecorded,
            "unrecorded_total": unrecorded_total,
            "is_closed": is_closed,
            "closing_info": closing_info,
            "pending_appeals": pending_appeals,
        },
    )


@router.get("/attendance/sheet/export")
async def attendance_sheet_export(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None
    fmt = request.query_params.get("format", "xlsx")

    rows, days_in_month, _ = _build_sheet_data(conn, year, month, company_id)

    company_name = ""
    if company_id:
        row = conn.execute("SELECT name FROM client_companies WHERE id=?", (company_id,)).fetchone()
        if row:
            company_name = row["name"]

    ym_str = f"{year}년 {month}월"
    filename_base = f"출근부_{year}{month:02d}"
    if company_name:
        filename_base += f"_{company_name}"

    if fmt == "pdf":
        html = templates.get_template("hr/attendance_sheet_pdf.html").render(
            rows=rows, year=year, month=month, days_in_month=days_in_month,
            company_name=company_name, ym_str=ym_str,
            company_info=COMPANY_INFO, status_labels=STATUS_LABELS,
        )
        return pdf_response(html, f"{filename_base}.pdf")

    from openpyxl.styles import Font, Alignment
    from core.export import HEADER_FONT, HEADER_FILL, CELL_FONT, THIN_BORDER
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "출근부"

    title_text = f"출근부 - {ym_str}"
    if company_name:
        title_text += f" ({company_name})"
    sum_cols = ["출근", "지각", "조퇴", "결근", "휴일", "연장", "총근무"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_in_month + 2 + len(sum_cols))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="맑은 고딕", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 3
    headers = ["사번", "이름"] + [str(d) for d in range(1, days_in_month + 1)] + sum_cols
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    status_short = {"normal": "O", "late": "△", "early_leave": "▽", "absent": "X", "holiday": "-"}
    for ri, r in enumerate(rows, header_row + 1):
        emp = r["employee"]
        ws.cell(row=ri, column=1, value=emp.get("employee_no", "")).font = CELL_FONT
        ws.cell(row=ri, column=1).border = THIN_BORDER
        ws.cell(row=ri, column=2, value=emp["name"]).font = CELL_FONT
        ws.cell(row=ri, column=2).border = THIN_BORDER
        for di, cell_data in enumerate(r["cells"], 3):
            cell = ws.cell(row=ri, column=di)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.font = CELL_FONT
            if cell_data:
                cell.value = status_short.get(cell_data["status"], "")
            else:
                cell.value = ""
        base_col = days_in_month + 3
        vals = [
            r["work_days"], r["late"], r["early_leave"], r["absent"],
            r["holiday_days"],
            f"{r['ot_hours']}:{r['ot_mins_r']:02d}",
            f"{r['total_hours']}:{r['total_mins_r']:02d}",
        ]
        for offset, val in enumerate(vals):
            cell = ws.cell(row=ri, column=base_col + offset, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 10
    for ci in range(3, days_in_month + 3):
        col_letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[col_letter].width = 4
    for ci in range(days_in_month + 3, days_in_month + 3 + len(sum_cols)):
        col_letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[col_letter].width = 9

    return excel_response(wb, f"{filename_base}.xlsx")


@router.get("/attendance/appeals", response_class=HTMLResponse)
async def appeal_list(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    rows = conn.execute(
        "SELECT a.*, e.name AS worker_name, cc.name AS company_name "
        "FROM attendance a JOIN employees e ON a.employee_id=e.id "
        "LEFT JOIN client_companies cc ON e.company_id=cc.id "
        "WHERE a.appeal IS NOT NULL "
        "ORDER BY CASE a.appeal WHEN 'pending' THEN 0 ELSE 1 END, a.work_date DESC",
    ).fetchall()
    appeals = []
    for r in rows:
        d = dict(r)
        d["appeal_type_label"] = APPEAL_TYPE_LABELS.get(d.get("appeal_type", ""), d.get("appeal_type", ""))
        appeals.append(d)
    pending_cnt = sum(1 for a in appeals if a["appeal"] == "pending")
    return templates.TemplateResponse(
        request=request, name="hr/appeal_list.html", context={
            "request": request,
            "page_title": "소명 관리",
            "user_name": user["name"],
            "appeals": appeals,
            "pending_cnt": pending_cnt,
        }
    )


@router.get("/attendance/appeals/{att_id}", response_class=HTMLResponse)
async def appeal_detail(
    request: Request,
    att_id: int,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT a.*, e.name AS worker_name, cc.name AS company_name, "
        "u.name AS reviewer_name "
        "FROM attendance a JOIN employees e ON a.employee_id=e.id "
        "LEFT JOIN client_companies cc ON e.company_id=cc.id "
        "LEFT JOIN users u ON a.appeal_reviewed_by=u.id "
        "WHERE a.id=? AND a.appeal IS NOT NULL",
        (att_id,),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance/appeals", status_code=303)
    a = dict(row)
    a["appeal_type_label"] = APPEAL_TYPE_LABELS.get(a.get("appeal_type", ""), a.get("appeal_type", ""))
    return templates.TemplateResponse(
        request=request, name="hr/appeal_detail.html", context={
            "request": request,
            "page_title": "소명 상세",
            "user_name": user["name"],
            "a": a,
        }
    )


@router.post("/attendance/{att_id}/appeal-review")
async def appeal_review(
    request: Request,
    att_id: int,
    action: str = Form(...),
    reject_reason: str = Form(""),
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT id, appeal_type FROM attendance WHERE id=? AND appeal='pending'", (att_id,)
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance/appeals", status_code=303)
    result = "approved" if action == "approve" else "rejected"
    if result == "approved":
        note = f"소명 승인 ({row['appeal_type'] or ''})"
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=?, "
            "status='normal', notes=CASE WHEN notes IS NULL OR notes='' THEN ? "
            "ELSE notes || ' / ' || ? END WHERE id=?",
            (result, user["id"], now_kst().isoformat(), note, note, att_id),
        )
    else:
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=?, "
            "appeal_reject_reason=? WHERE id=?",
            (result, user["id"], now_kst().isoformat(), reject_reason.strip(), att_id),
        )
    conn.commit()
    return RedirectResponse(url="/hr/attendance/appeals", status_code=303)


@router.post("/attendance/close")
async def close_attendance(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    if not user.get("is_hr_admin"):
        return RedirectResponse(url="/hr/attendance/sheet", status_code=303)

    form = await request.form()
    year = int(form.get("year", now_kst().year))
    month = int(form.get("month", now_kst().month))
    company_id = form.get("company_id", "")
    company_id = int(company_id) if company_id else None

    url = f"/hr/attendance/sheet?year={year}&month={month}"
    if company_id:
        url += f"&company_id={company_id}"

    if not company_id:
        return RedirectResponse(url=url, status_code=303)

    if is_month_closed(conn, year, month, company_id):
        return RedirectResponse(url=url, status_code=303)

    selected = form.getlist("emp_ids")

    _, days_in_month = calendar.monthrange(year, month)
    ym = f"{year}-{month:02d}"
    today_str = now_kst().strftime("%Y-%m-%d")

    emp_ids = [int(eid) for eid in selected] if selected else []
    absent_count = 0

    if emp_ids:
        ph = ",".join("?" * len(emp_ids))
        employees = conn.execute(
            f"SELECT id, hire_date FROM employees WHERE id IN ({ph})", emp_ids
        ).fetchall()
    else:
        employees = []

    for emp in employees:
        hire_date = emp["hire_date"] or "2000-01-01"
        for d in range(1, days_in_month + 1):
            dt = f"{ym}-{d:02d}"
            if calendar.weekday(year, month, d) >= 5:
                continue
            if dt < hire_date or dt > today_str:
                continue
            exists = conn.execute(
                "SELECT id FROM attendance WHERE employee_id=? AND work_date=?",
                (emp["id"], dt),
            ).fetchone()
            if exists:
                continue
            conn.execute(
                "INSERT INTO attendance (employee_id, work_date, status, notes) VALUES (?, ?, 'absent', '근태마감')",
                (emp["id"], dt),
            )
            absent_count += 1

    conn.execute(
        "INSERT OR REPLACE INTO attendance_closings "
        "(year, month, company_id, status, closed_by, closed_at, absent_count) "
        "VALUES (?, ?, ?, 'closed', ?, ?, ?)",
        (year, month, company_id, user["id"], now_kst().isoformat(), absent_count),
    )
    closing_id = conn.execute(
        "SELECT id FROM attendance_closings WHERE year=? AND month=? AND company_id=?",
        (year, month, company_id),
    ).fetchone()["id"]

    conn.execute(
        "INSERT INTO attendance_closing_logs (closing_id, action, performed_by, detail) "
        "VALUES (?, 'closed', ?, ?)",
        (closing_id, user["id"], json.dumps({"absent_count": absent_count, "emp_ids": emp_ids})),
    )
    conn.commit()

    return RedirectResponse(url=url, status_code=303)


@router.post("/attendance/reopen")
async def reopen_attendance(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    if user["role"] != "admin":
        return RedirectResponse(url="/hr/attendance/sheet", status_code=303)

    form = await request.form()
    year = int(form.get("year", now_kst().year))
    month = int(form.get("month", now_kst().month))
    company_id = form.get("company_id", "")
    company_id = int(company_id) if company_id else None
    reason = form.get("reopen_reason", "").strip()

    url = f"/hr/attendance/sheet?year={year}&month={month}"
    if company_id:
        url += f"&company_id={company_id}"

    if not company_id or not reason:
        return RedirectResponse(url=url, status_code=303)

    closing = conn.execute(
        "SELECT id FROM attendance_closings WHERE year=? AND month=? AND company_id=? AND status='closed'",
        (year, month, company_id),
    ).fetchone()
    if not closing:
        return RedirectResponse(url=url, status_code=303)

    conn.execute(
        "UPDATE attendance_closings SET status='open', reopened_by=?, reopened_at=?, reopen_reason=? WHERE id=?",
        (user["id"], now_kst().isoformat(), reason, closing["id"]),
    )
    conn.execute(
        "INSERT INTO attendance_closing_logs (closing_id, action, performed_by, reason) "
        "VALUES (?, 'reopened', ?, ?)",
        (closing["id"], user["id"], reason),
    )
    conn.commit()

    return RedirectResponse(url=url, status_code=303)


@router.get("/attendance/{att_id}/edit", response_class=HTMLResponse)
async def attendance_edit_form(
    request: Request,
    att_id: int,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    row = conn.execute(
        "SELECT a.*, e.name AS worker_name FROM attendance a "
        "JOIN employees e ON a.employee_id=e.id WHERE a.id=?", (att_id,)
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    att = dict(row)

    # 마감 여부
    wd = att["work_date"]
    y, m = int(wd[:4]), int(wd[5:7])
    cid_row = conn.execute("SELECT company_id FROM employees WHERE id=?", (att["employee_id"],)).fetchone()
    edit_locked = False
    if cid_row and cid_row["company_id"] and user["role"] != "admin":
        edit_locked = is_month_closed(conn, y, m, cid_row["company_id"])

    logs = conn.execute(
        "SELECT al.*, u.name AS editor_name FROM attendance_logs al "
        "JOIN users u ON al.changed_by=u.id "
        "WHERE al.attendance_id=? ORDER BY al.created_at DESC", (att_id,)
    ).fetchall()
    return templates.TemplateResponse(
        request=request, name="hr/attendance_edit.html", context={
            "request": request,
            "page_title": "근태 수정",
            "user_name": user["name"],
            "user_role": user["role"],
            "att": att,
            "logs": [dict(l) for l in logs],
            "status_labels": STATUS_LABELS,
            "edit_locked": edit_locked,
        }
    )


@router.post("/attendance/{att_id}/edit")
async def attendance_edit_save(
    request: Request,
    att_id: int,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    row = conn.execute("SELECT * FROM attendance WHERE id=?", (att_id,)).fetchone()
    if not row:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    old = dict(row)

    # 마감 잠금
    if user["role"] != "admin":
        wd = old["work_date"]
        y, m = int(wd[:4]), int(wd[5:7])
        cid = conn.execute(
            "SELECT company_id FROM employees WHERE id=?", (old["employee_id"],)
        ).fetchone()
        if cid and is_month_closed(conn, y, m, cid["company_id"]):
            return RedirectResponse(
                url=f"/hr/attendance/{old['employee_id']}?month={wd[:7]}", status_code=303
            )

    form = await request.form()
    reason = form.get("reason", "")

    fields = {
        "clock_in": form.get("clock_in", "").strip() or None,
        "clock_out": form.get("clock_out", "").strip() or None,
        "status": form.get("status", old["status"]),
        "overtime_minutes": int(form.get("overtime_minutes") or 0),
        "break_minutes": int(form.get("break_minutes") or 0),
        "notes": form.get("notes", "").strip(),
    }

    # work_minutes 자동계산
    cin, cout = fields["clock_in"], fields["clock_out"]
    if cin and cout:
        cin_h, cin_m = int(cin[:2]), int(cin[3:5])
        out_h, out_m = int(cout[:2]), int(cout[3:5])
        work_mins = max(0, (out_h * 60 + out_m) - (cin_h * 60 + cin_m) - fields["break_minutes"])
        fields["work_minutes"] = work_mins
        # 연장근로 자동계산 (18:00 이후)
        if out_h >= 18 and not form.get("overtime_manual"):
            ot = max(0, (out_h * 60 + out_m) - 18 * 60)
            fields["overtime_minutes"] = ot
    else:
        fields["work_minutes"] = old.get("work_minutes") or 0

    # 변경 이력
    track_fields = ["clock_in", "clock_out", "status", "overtime_minutes", "break_minutes", "notes"]
    for f in track_fields:
        old_val = str(old.get(f) or "")
        new_val = str(fields.get(f) or "")
        if old_val != new_val:
            conn.execute(
                "INSERT INTO attendance_logs (attendance_id, changed_by, field, old_value, new_value, reason) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (att_id, user["id"], f, old_val, new_val, reason),
            )

    conn.execute(
        "UPDATE attendance SET clock_in=?, clock_out=?, status=?, work_minutes=?, "
        "overtime_minutes=?, break_minutes=?, notes=? WHERE id=?",
        (fields["clock_in"], fields["clock_out"], fields["status"],
         fields["work_minutes"], fields["overtime_minutes"],
         fields["break_minutes"], fields["notes"], att_id),
    )
    conn.commit()

    emp_id = old["employee_id"]
    month = old["work_date"][:7]
    return RedirectResponse(url=f"/hr/attendance/{emp_id}?month={month}", status_code=303)


@router.get("/attendance/daily", response_class=HTMLResponse)
async def attendance_daily(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    today = now_kst().strftime("%Y-%m-%d")
    target_date = request.query_params.get("date", today)
    company_filter = request.query_params.get("company", "")

    emp_sql = "SELECT id, name, employee_no, dept, company_id FROM employees WHERE status='active'"
    emp_params = []
    if company_filter:
        emp_sql += " AND company_id=?"
        emp_params.append(int(company_filter))
    emp_sql += " ORDER BY name"
    workers = [dict(r) for r in conn.execute(emp_sql, emp_params).fetchall()]

    # 해당일 기존 기록
    existing = {}
    if workers:
        ph = ",".join("?" * len(workers))
        rows = conn.execute(
            f"SELECT * FROM attendance WHERE work_date=? AND employee_id IN ({ph})",
            [target_date] + [w["id"] for w in workers],
        ).fetchall()
        for r in rows:
            existing[r["employee_id"]] = dict(r)

    for w in workers:
        w["att"] = existing.get(w["id"])

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()

    # 마감 잠금
    daily_locked = False
    if company_filter and user["role"] != "admin":
        y, m = int(target_date[:4]), int(target_date[5:7])
        daily_locked = is_month_closed(conn, y, m, int(company_filter))

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance_daily.html",
        context={
            "request": request,
            "page_title": "일일 근태 입력",
            "user_name": user["name"],
            "user_role": user["role"],
            "workers": workers,
            "target_date": target_date,
            "today": today,
            "company_filter": company_filter,
            "companies": companies,
            "status_labels": STATUS_LABELS,
            "daily_locked": daily_locked,
        },
    )


@router.post("/attendance/daily")
async def attendance_daily_save(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    form = await request.form()
    target_date = form.get("target_date", now_kst().strftime("%Y-%m-%d"))
    company_filter = form.get("company_filter", "")
    emp_ids = form.getlist("emp_ids")

    # 마감 잠금 (고객사 필터 있을 때)
    if user["role"] != "admin" and company_filter:
        y, m = int(target_date[:4]), int(target_date[5:7])
        if is_month_closed(conn, y, m, int(company_filter)):
            url = f"/hr/attendance/daily?date={target_date}&company={company_filter}"
            return RedirectResponse(url=url, status_code=303)

    for eid in emp_ids:
        eid = int(eid)
        clock_in = form.get(f"clock_in_{eid}", "").strip() or None
        clock_out = form.get(f"clock_out_{eid}", "").strip() or None
        status = form.get(f"status_{eid}", "normal")
        notes = form.get(f"notes_{eid}", "").strip()

        # work_minutes 계산
        work_mins = 0
        ot_mins = 0
        if clock_in and clock_out:
            cin_h, cin_m = int(clock_in[:2]), int(clock_in[3:5])
            out_h, out_m = int(clock_out[:2]), int(clock_out[3:5])
            work_mins = max(0, (out_h * 60 + out_m) - (cin_h * 60 + cin_m) - 60)
            if out_h >= 18:
                ot_mins = max(0, (out_h * 60 + out_m) - 18 * 60)

        existing = conn.execute(
            "SELECT id FROM attendance WHERE employee_id=? AND work_date=?",
            (eid, target_date),
        ).fetchone()

        if existing:
            conn.execute(
                "UPDATE attendance SET clock_in=?, clock_out=?, status=?, "
                "work_minutes=?, overtime_minutes=?, notes=? WHERE id=?",
                (clock_in, clock_out, status, work_mins, ot_mins, notes, existing["id"]),
            )
        else:
            conn.execute(
                "INSERT INTO attendance (employee_id, work_date, clock_in, clock_out, "
                "status, work_minutes, overtime_minutes, notes) VALUES (?,?,?,?,?,?,?,?)",
                (eid, target_date, clock_in, clock_out, status, work_mins, ot_mins, notes),
            )

    conn.commit()
    url = f"/hr/attendance/daily?date={target_date}"
    if company_filter:
        url += f"&company={company_filter}"
    return RedirectResponse(url=url, status_code=303)


@router.get("/attendance/{employee_id}", response_class=HTMLResponse)
async def hr_attendance_detail(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    worker = get_worker_detail(conn, employee_id)
    if not worker:
        return RedirectResponse(url="/hr/attendance", status_code=303)
    records = get_worker_month_records(conn, employee_id, month)
    leaves = get_worker_month_leaves(conn, employee_id, month)
    summary = worker_month_summary(records, leaves)
    cal = build_calendar_data(records, month, leaves)
    merged = merge_records_leaves(records, leaves)
    prev_m, next_m = prev_next_month(month)

    return templates.TemplateResponse(
        request=request,
        name="hr/attendance_detail.html",
        context={
            "request": request,
            "page_title": f"{worker['name']} 근태 상세",
            "user_name": user["name"],
            "user_role": user["role"],
            "worker": worker,
            "records": merged,
            "summary": summary,
            "cal": cal,
            "current_month": month,
            "prev_month": prev_m,
            "next_month": next_m,
            "back_url": "/hr/attendance",
        },
    )
