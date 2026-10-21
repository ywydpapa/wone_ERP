import calendar
import csv
import io
from datetime import date, datetime
from urllib.parse import quote, urlencode

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from openpyxl import load_workbook
from starlette.responses import JSONResponse, StreamingResponse
from core.auth import verify_pw
from core.deps import get_db, require_login, templates
from core.tz import today_kst
from core import audit, books, credit, invoices, partner
from core.books import require_book
from core.export import esc_cell
from routers.erp_master import norm_biz_no

router = APIRouter()


def _create_payroll_slip(conn, book, payslip, employee_name, uid):
    year = payslip["pay_year"]
    month = payslip["pay_month"]
    gross = payslip["gross_pay"]
    net = payslip["net_pay"]
    income_tax = payslip["income_tax"] + payslip["resident_tax"]
    pension = payslip["national_pension"]
    health = payslip["health_insurance"] + (payslip.get("long_care") or 0)
    employment = payslip["employment_insurance"]

    emp_id = payslip.get("employee_id")
    emp = conn.execute("SELECT dept FROM employees WHERE id=?", (emp_id,)).fetchone() if emp_id else None
    dept = emp["dept"] or "" if emp else ""
    if dept and not conn.execute("SELECT 1 FROM depts WHERE company_id=? AND name=?", (book, dept)).fetchone():
        dept = ""

    summary = f"{employee_name} {month}월 급여"
    lines = [
        {"code": "801", "debit": gross, "credit": 0, "summary": summary, "dept": dept},
        {"code": "254", "debit": 0, "credit": net, "summary": summary},
    ]
    if income_tax > 0:
        lines.append({"code": "261", "debit": 0, "credit": income_tax, "summary": summary})
    if pension > 0:
        lines.append({"code": "262", "debit": 0, "credit": pension, "summary": summary})
    if health > 0:
        lines.append({"code": "263", "debit": 0, "credit": health, "summary": summary})
    if employment > 0:
        lines.append({"code": "264", "debit": 0, "credit": employment, "summary": summary})

    # 사측 부담분
    er_pension = payslip.get("employer_pension") or 0
    er_health = payslip.get("employer_health") or 0
    er_ei = (payslip.get("employer_employment") or 0) + (payslip.get("employer_accident") or 0)
    er_summary = f"{summary} 사측부담"
    if er_pension > 0:
        lines.append({"code": "815", "debit": er_pension, "credit": 0, "summary": er_summary, "dept": dept})
    if er_health > 0:
        lines.append({"code": "804", "debit": er_health, "credit": 0, "summary": er_summary, "dept": dept})
    if er_ei > 0:
        lines.append({"code": "818", "debit": er_ei, "credit": 0, "summary": er_summary, "dept": dept})
    er_total = er_pension + er_health + er_ei
    if er_total > 0:
        lines.append({"code": "256", "debit": 0, "credit": er_total, "summary": er_summary})

    title = f"급여 전표 ({employee_name} {year}.{month:02d})"
    slip_date = payslip.get("pay_date") or f"{year}-{month:02d}-25"
    content = f"{employee_name} {year}년 {month}월 급여 자동 전표"
    payslip_id = payslip["id"]
    doc_id = books.post(conn, book, uid, slip_date, lines, title, "payroll", payslip_id, content=content)
    conn.execute("UPDATE erp_docs SET payslip_id=? WHERE id=?", (payslip_id, doc_id))
    return doc_id


def _ledger_partners(conn, book):
    return [r[0] for r in conn.execute(
        """SELECT DISTINCT sl.partner FROM slip_lines sl JOIN erp_docs ed ON sl.doc_id=ed.id
           WHERE ed.company_id=? AND sl.partner != '' ORDER BY sl.partner""", (book,)
    ).fetchall()]


SUMS = ("supply_amt", "vat", "total")


# 더존 매입매출장 기준. 일반전표는 제외, 세금계산서/계산서/카드/현금영수증만 집계
def _vat_book(conn, book, kind, date_from, date_to, partner_id, code):
    sql = """SELECT t.*, p.name AS partner_name, p.biz_no, d.slip_no,
                    (SELECT name FROM tax_invoice_items i WHERE i.invoice_id=t.id ORDER BY i.line_no LIMIT 1) AS item,
                    (SELECT COUNT(*) FROM tax_invoice_items i WHERE i.invoice_id=t.id) AS item_cnt
             FROM tax_invoices t LEFT JOIN partners p ON p.id=t.partner_id LEFT JOIN erp_docs d ON d.id=t.doc_id
             WHERE t.company_id=? AND t.kind=? AND t.issue_date BETWEEN ? AND ?"""
    args = [book, kind, date_from, date_to]
    if partner_id:
        sql += " AND t.partner_id=?"
        args.append(partner_id)
    rows = [dict(r) for r in conn.execute(sql + " ORDER BY t.issue_date, t.id", args)]
    for r in rows:
        r["vat_code"] = invoices.vat_code(r)
        r["item"] = f"{r['item']} 외 {r['item_cnt'] - 1}건" if r["item_cnt"] > 1 else (r["item"] or "")
    if code:
        rows = [r for r in rows if r["vat_code"].split(" ")[0] == code]

    types = {}
    for r in rows:
        t = types.setdefault(r["vat_code"] or "-", {"cnt": 0, **dict.fromkeys(SUMS, 0)})
        t["cnt"] += 1
        for k in SUMS:
            t[k] += r[k]

    # 월계/누계는 수정, 취소분 음수 금액 포함 기준
    months = []
    acc = dict.fromkeys(SUMS, 0)
    for r in rows:
        ym = r["issue_date"][:7]
        if not months or months[-1]["ym"] != ym:
            months.append({"ym": ym, "mon": int(ym[5:]), "rows": [], "sum": dict.fromkeys(SUMS, 0)})
        m = months[-1]
        m["rows"].append(r)
        for k in SUMS:
            m["sum"][k] += r[k]
            acc[k] += r[k]
        m["acc"] = dict(acc)
    return rows, months, sorted(types.items()), acc


def _vat_book_page(request, u, conn, side, date_from, date_to, partner_id, code):
    book = u["book"]
    kind = "매입" if side == "purchase" else "매출"
    date_from, date_to = _period(date_from, date_to)
    rows, months, types, tot = _vat_book(conn, book, kind, date_from, date_to, partner_id, code)
    partner = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=?",
                           (partner_id, book)).fetchone() if partner_id else None
    qs = urlencode({k: v for k, v in {"date_from": date_from, "date_to": date_to,
                                      "partner_id": partner["id"] if partner else "", "code": code}.items() if v})
    codes = [(c, v[0]) for c, v in invoices.VAT_TYPES.items() if v[1] == kind]
    return templates.TemplateResponse(request=request, name="erp/ledger_vat.html", context={
        "request": request, "page_title": f"{kind}장", "user_name": u["name"],
        "kind": kind, "side": side, "months": months, "types": types, "tot": tot, "cnt": len(rows),
        "codes": codes, "code": code, "partner": partner, "qs": qs,
        "date_from": date_from, "date_to": date_to, "me": books.company_info(conn, book),
    })


def _vat_book_csv(u, conn, side, date_from, date_to, partner_id, code):
    kind = "매입" if side == "purchase" else "매출"
    date_from, date_to = _period(date_from, date_to)
    rows, _, _, tot = _vat_book(conn, u["book"], kind, date_from, date_to, partner_id, code)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["작성일자", "유형", "거래처", "사업자번호", "품목", "공급가액", "부가세", "합계", "전자", "상태", "전표번호"])
    for r in rows:
        writer.writerow([esc_cell(v) for v in [
            r["issue_date"], r["vat_code"], r["partner_name"] or "", r["biz_no"] or "", r["item"],
            r["supply_amt"], r["vat"], r["total"], "전자" if r["is_electronic"] else "종이", r["status"],
            r["slip_no"] or "",
        ]])
    writer.writerow(["", "", "", "", "합계", tot["supply_amt"], tot["vat"], tot["total"], "", "", ""])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f"attachment; filename={side}_ledger.csv"},
    )


@router.get("/ledger/purchase", response_class=HTMLResponse)
async def purchase_ledger(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                          date_from: str = "", date_to: str = "", partner_id: str = "", code: str = ""):
    return _vat_book_page(request, u, conn, "purchase", date_from, date_to, partner_id, code)


@router.get("/ledger/purchase/excel")
async def purchase_ledger_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                                date_from: str = "", date_to: str = "", partner_id: str = "", code: str = ""):
    return _vat_book_csv(u, conn, "purchase", date_from, date_to, partner_id, code)


@router.get("/ledger/sales", response_class=HTMLResponse)
async def sales_ledger(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                       date_from: str = "", date_to: str = "", partner_id: str = "", code: str = ""):
    return _vat_book_page(request, u, conn, "sales", date_from, date_to, partner_id, code)


@router.get("/ledger/sales/excel")
async def sales_ledger_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                             date_from: str = "", date_to: str = "", partner_id: str = "", code: str = ""):
    return _vat_book_csv(u, conn, "sales", date_from, date_to, partner_id, code)


def _ar_ap_data(conn, book, partner, date_from, date_to):
    rows = []
    opening_balance = 0
    period_debit = 0
    period_credit = 0
    summary_list = []
    summary_total = 0

    if partner:
        pid_row = conn.execute("SELECT id FROM partners WHERE company_id=? AND name=?", (book, partner)).fetchone()
        pid = pid_row["id"] if pid_row else None
        match = "(sl.partner_id = ? OR (sl.partner_id IS NULL AND sl.partner = ?))"

        if date_from:
            ob = conn.execute(f"""
                SELECT COALESCE(SUM(CASE WHEN sl.account_code='110' THEN sl.debit - sl.credit ELSE 0 END),0)
                     + COALESCE(SUM(CASE WHEN sl.account_code='251' THEN sl.credit - sl.debit ELSE 0 END),0)
                FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
                WHERE d.company_id=? AND d.status='done' AND {match} AND sl.account_code IN ('110','251') AND d.slip_date < ?
            """, [book, pid, partner, date_from]).fetchone()
            opening_balance = ob[0] if ob else 0

        conds = ["d.company_id=?", "d.status='done'", match, "sl.account_code IN ('110','251')"]
        params = [book, pid, partner]
        if date_from:
            conds.append("d.slip_date >= ?")
            params.append(date_from)
        if date_to:
            conds.append("d.slip_date <= ?")
            params.append(date_to)
        where = " AND ".join(conds)
        rows = conn.execute(f"""
            SELECT d.id as doc_id, d.slip_date, d.doc_number,
                   sl.account_code, sl.summary, sl.debit, sl.credit
            FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
            WHERE {where}
            ORDER BY d.slip_date, d.id, sl.line_no
        """, params).fetchall()

        bal = opening_balance
        rows = [dict(r) for r in rows]
        for r in rows:
            if r["account_code"] == '110':
                bal += r["debit"] - r["credit"]
            else:
                bal += r["credit"] - r["debit"]
            r["bal"] = bal
        period_debit = sum(r["debit"] for r in rows)
        period_credit = sum(r["credit"] for r in rows)
    else:
        conds = ["d.company_id=?", "d.status='done'", "sl.partner != ''", "sl.account_code IN ('110','251')"]
        params = [book]
        if date_from:
            conds.append("d.slip_date >= ?")
            params.append(date_from)
        if date_to:
            conds.append("d.slip_date <= ?")
            params.append(date_to)
        where = " AND ".join(conds)
        sum_rows = conn.execute(f"""
            SELECT sl.partner,
                   SUM(CASE WHEN sl.account_code='110' THEN sl.debit ELSE 0 END) as ar_debit,
                   SUM(CASE WHEN sl.account_code='110' THEN sl.credit ELSE 0 END) as ar_credit,
                   SUM(CASE WHEN sl.account_code='251' THEN sl.credit ELSE 0 END) as ap_credit,
                   SUM(CASE WHEN sl.account_code='251' THEN sl.debit ELSE 0 END) as ap_debit
            FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
            WHERE {where}
            GROUP BY COALESCE(sl.partner_id, 'n:' || sl.partner)
            ORDER BY sl.partner
        """, params).fetchall()
        for r in sum_rows:
            ar_bal = r["ar_debit"] - r["ar_credit"]
            ap_bal = r["ap_credit"] - r["ap_debit"]
            net = ar_bal - ap_bal
            summary_list.append({
                "partner": r["partner"],
                "ar_bal": ar_bal,
                "ap_bal": ap_bal,
                "net": net,
            })
            summary_total += net

    return rows, opening_balance, period_debit, period_credit, summary_list, summary_total


@router.get("/ledger/ar-ap", response_class=HTMLResponse)
async def ar_ap(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                partner: str = "", date_from: str = "", date_to: str = ""):
    book = u["book"]
    rows, opening_balance, period_debit, period_credit, summary_list, summary_total = _ar_ap_data(
        conn, book, partner, date_from, date_to)
    partner_id = None
    if partner:
        prow = conn.execute("SELECT id FROM partners WHERE company_id=? AND name=?", (book, partner)).fetchone()
        partner_id = prow["id"] if prow else None

    return templates.TemplateResponse(request=request, name="erp/ledger_ar_ap.html", context={
        "request": request, "page_title": "거래처원장",
        "user_name": u["name"],
        "partners": _ledger_partners(conn, book),
        "partner": partner,
        "partner_id": partner_id,
        "date_from": date_from,
        "date_to": date_to,
        "rows": rows,
        "opening_balance": opening_balance,
        "period_debit": period_debit,
        "period_credit": period_credit,
        "summary_list": summary_list,
        "summary_total": summary_total,
    })


@router.get("/ledger/ar-ap/excel")
async def ar_ap_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                      partner: str = "", date_from: str = "", date_to: str = ""):
    rows, opening_balance, period_debit, period_credit, summary_list, summary_total = _ar_ap_data(
        conn, u["book"], partner, date_from, date_to)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    if partner:
        writer.writerow(["일자", "전표번호", "계정", "적요", "차변", "대변", "잔액"])
        if date_from:
            writer.writerow([date_from, "", "이월", "", "", "", opening_balance])
        for r in rows:
            writer.writerow([esc_cell(v) for v in [
                r["slip_date"], r["doc_number"], "채권" if r["account_code"] == "110" else "채무",
                r["summary"] or "", r["debit"] or "", r["credit"] or "", r["bal"],
            ]])
        writer.writerow(["", "", "", "합계", period_debit, period_credit, rows[-1]["bal"] if rows else opening_balance])
    else:
        writer.writerow(["거래처", "채권잔액", "채무잔액", "순잔액"])
        for r in summary_list:
            writer.writerow([esc_cell(v) for v in [r["partner"], r["ar_bal"], r["ap_bal"], r["net"]]])
        writer.writerow(["합계", sum(r["ar_bal"] for r in summary_list), sum(r["ap_bal"] for r in summary_list), summary_total])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=ar_ap_ledger.csv"},
    )


@router.get("/ledger/pnl", response_class=HTMLResponse)
async def pnl_summary(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    year = request.query_params.get("year", "")
    if not year:
        year = today_kst().year
    else:
        year = int(year)

    revenue_rows = conn.execute(
        """SELECT SUBSTR(ed.slip_date,6,2) as mon, SUM(sl.credit) as amt
           FROM slip_lines sl JOIN erp_docs ed ON sl.doc_id=ed.id
           WHERE ed.company_id=? AND ed.status='done' AND sl.account_code BETWEEN '401' AND '403' AND sl.credit > 0
             AND ed.slip_date LIKE ?
           GROUP BY SUBSTR(ed.slip_date,6,2)
           ORDER BY mon""",
        (book, f"{year}-%"),
    ).fetchall()
    revenue = {r["mon"]: r["amt"] for r in revenue_rows}

    expense_rows = conn.execute(
        """SELECT SUBSTR(ed.slip_date,6,2) as mon, SUM(sl.debit) as amt
           FROM slip_lines sl JOIN erp_docs ed ON sl.doc_id=ed.id
           WHERE ed.company_id=? AND ed.status='done' AND (sl.account_code BETWEEN '451' AND '503'
                  OR sl.account_code BETWEEN '801' AND '899'
                  OR sl.account_code BETWEEN '951' AND '990')
             AND sl.debit > 0
             AND ed.slip_date LIKE ?
           GROUP BY SUBSTR(ed.slip_date,6,2)
           ORDER BY mon""",
        (book, f"{year}-%"),
    ).fetchall()
    expense = {r["mon"]: r["amt"] for r in expense_rows}

    months = []
    total_rev = 0
    total_exp = 0
    for m in range(1, 13):
        mk = f"{m:02d}"
        rev = revenue.get(mk, 0)
        exp = expense.get(mk, 0)
        total_rev += rev
        total_exp += exp
        months.append({"month": m, "revenue": rev, "expense": exp, "profit": rev - exp})

    return templates.TemplateResponse(request=request, name="erp/ledger_pnl.html", context={
        "request": request, "page_title": "월별 손익 요약",
        "user_name": u["name"],
        "year": year, "months": months,
        "total_revenue": total_rev, "total_expense": total_exp,
        "total_profit": total_rev - total_exp,
    })


def _general_ledger_data(conn, book, account_code, date_from, date_to):
    rows = []
    opening_debit = 0
    opening_credit = 0

    if account_code:
        conditions = ["d.company_id=?", "d.status='done'", "sl.account_code = ?"]
        params = [book, account_code]

        if date_from:
            opening = conn.execute("""
                SELECT COALESCE(SUM(sl.debit),0), COALESCE(SUM(sl.credit),0)
                FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
                WHERE d.company_id=? AND d.status='done' AND sl.account_code=? AND d.slip_date < ?
            """, [book, account_code, date_from]).fetchone()
            opening_debit, opening_credit = opening[0], opening[1]
            conditions.append("d.slip_date >= ?")
            params.append(date_from)

        if date_to:
            conditions.append("d.slip_date <= ?")
            params.append(date_to)

        where = " AND ".join(conditions)
        rows = conn.execute(f"""
            SELECT d.id as doc_id, d.slip_date, d.doc_number, d.title,
                   sl.summary, sl.partner, sl.debit, sl.credit
            FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
            WHERE {where}
            ORDER BY d.slip_date, d.id, sl.line_no
        """, params).fetchall()

    acct = None
    if account_code:
        acct = conn.execute(
            "SELECT code, name, is_debit FROM accounts WHERE company_id=? AND code=?", [book, account_code]
        ).fetchone()

    display_rows = []
    if acct:
        is_debit = acct["is_debit"]
        bal = (opening_debit - opening_credit) if is_debit else (opening_credit - opening_debit)
        for r in rows:
            r = dict(r)
            bal = bal + r["debit"] - r["credit"] if is_debit else bal + r["credit"] - r["debit"]
            r["bal"] = bal
            display_rows.append(r)
    else:
        display_rows = [dict(r) for r in rows]

    period_debit = sum(r["debit"] for r in display_rows)
    period_credit = sum(r["credit"] for r in display_rows)
    return acct, display_rows, opening_debit, opening_credit, period_debit, period_credit


@router.get("/ledger/general", response_class=HTMLResponse)
async def general_ledger(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                         account_code: str = "", date_from: str = "", date_to: str = ""):
    book = u["book"]
    accounts = conn.execute(
        "SELECT code, name, category, is_debit FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code",
        (book,),
    ).fetchall()
    acct_groups = []
    for a in accounts:
        if not acct_groups or acct_groups[-1][0] != a["category"]:
            acct_groups.append((a["category"], []))
        acct_groups[-1][1].append(a)

    acct, display_rows, opening_debit, opening_credit, period_debit, period_credit = _general_ledger_data(
        conn, book, account_code, date_from, date_to)

    return templates.TemplateResponse(request=request, name="erp/general_ledger.html", context={
        "request": request, "page_title": "총계정원장",
        "user_name": u["name"],
        "acct_groups": acct_groups,
        "rows": display_rows,
        "acct": acct,
        "account_code": account_code,
        "date_from": date_from,
        "date_to": date_to,
        "opening_debit": opening_debit,
        "opening_credit": opening_credit,
        "period_debit": period_debit,
        "period_credit": period_credit,
    })


@router.get("/ledger/general/excel")
async def general_ledger_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                               account_code: str = "", date_from: str = "", date_to: str = ""):
    acct, rows, opening_debit, opening_credit, period_debit, period_credit = _general_ledger_data(
        conn, u["book"], account_code, date_from, date_to)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["일자", "전표번호", "적요", "거래처", "차변", "대변", "잔액"])
    if acct and date_from:
        ob = opening_debit - opening_credit if acct["is_debit"] else opening_credit - opening_debit
        writer.writerow([date_from, "", "이월", "", "", "", ob])
    for r in rows:
        writer.writerow([esc_cell(v) for v in [
            r["slip_date"], r["doc_number"], r["summary"] or "", r["partner"] or "",
            r["debit"] or "", r["credit"] or "", r.get("bal", ""),
        ]])
    writer.writerow(["", "", "", "합계", period_debit, period_credit, ""])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=general_ledger.csv"},
    )


def _trial_balance_data(conn, book, date_from, date_to):
    today = today_kst()
    if not date_from:
        date_from = today.replace(day=1).isoformat()
    if not date_to:
        last_day = calendar.monthrange(today.year, today.month)[1]
        date_to = today.replace(day=last_day).isoformat()

    query = """
        SELECT a.code, a.name, a.category, a.is_debit,
               COALESCE(SUM(CASE WHEN d.slip_date < ? THEN sl.debit ELSE 0 END),0) as prev_debit,
               COALESCE(SUM(CASE WHEN d.slip_date < ? THEN sl.credit ELSE 0 END),0) as prev_credit,
               COALESCE(SUM(CASE WHEN d.slip_date BETWEEN ? AND ? THEN sl.debit ELSE 0 END),0) as period_debit,
               COALESCE(SUM(CASE WHEN d.slip_date BETWEEN ? AND ? THEN sl.credit ELSE 0 END),0) as period_credit
        FROM accounts a
        LEFT JOIN slip_lines sl ON sl.account_code = a.code
        LEFT JOIN erp_docs d ON sl.doc_id = d.id AND d.company_id = ? AND d.status = 'done'
        WHERE a.is_active = 1 AND a.company_id = ?
        GROUP BY a.code, a.name, a.category, a.is_debit
        HAVING prev_debit > 0 OR prev_credit > 0 OR period_debit > 0 OR period_credit > 0
        ORDER BY a.code
    """
    raw = conn.execute(query, [date_from, date_from, date_from, date_to, date_from, date_to, book, book]).fetchall()

    result = []
    for r in raw:
        if r["is_debit"]:
            prev_bal = r["prev_debit"] - r["prev_credit"]
            closing_bal = prev_bal + r["period_debit"] - r["period_credit"]
        else:
            prev_bal = r["prev_credit"] - r["prev_debit"]
            closing_bal = prev_bal + r["period_credit"] - r["period_debit"]
        debit_bal = closing_bal if (r["is_debit"] and closing_bal >= 0) or (not r["is_debit"] and closing_bal < 0) else 0
        credit_bal = abs(closing_bal) - debit_bal if closing_bal else 0
        result.append({
            "code": r["code"], "name": r["name"], "category": r["category"],
            "prev_bal": prev_bal,
            "period_debit": r["period_debit"], "period_credit": r["period_credit"],
            "closing_bal": closing_bal,
            "debit_bal": debit_bal, "credit_bal": credit_bal,
        })

    grouped = []
    cur_cat = ""
    sub = {"prev": 0, "debit": 0, "credit": 0, "debit_bal": 0, "credit_bal": 0}
    for r in result:
        if r["category"] != cur_cat:
            if cur_cat:
                grouped.append({"_type": "subtotal", "category": cur_cat, **sub})
                sub = {"prev": 0, "debit": 0, "credit": 0, "debit_bal": 0, "credit_bal": 0}
            grouped.append({"_type": "header", "category": r["category"]})
            cur_cat = r["category"]
        sub["prev"] += r["prev_bal"]
        sub["debit"] += r["period_debit"]
        sub["credit"] += r["period_credit"]
        sub["debit_bal"] += r["debit_bal"]
        sub["credit_bal"] += r["credit_bal"]
        r["_type"] = "row"
        grouped.append(r)
    if cur_cat:
        grouped.append({"_type": "subtotal", "category": cur_cat, **sub})

    totals = {
        "prev": sum(r["prev_bal"] for r in result),
        "debit": sum(r["period_debit"] for r in result),
        "credit": sum(r["period_credit"] for r in result),
        "debit_bal": sum(r["debit_bal"] for r in result),
        "credit_bal": sum(r["credit_bal"] for r in result),
    }
    return date_from, date_to, result, grouped, totals


@router.get("/ledger/trial-balance", response_class=HTMLResponse)
async def trial_balance(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                        date_from: str = "", date_to: str = ""):
    date_from, date_to, result, grouped, totals = _trial_balance_data(conn, u["book"], date_from, date_to)

    return templates.TemplateResponse(request=request, name="erp/trial_balance.html", context={
        "request": request, "page_title": "합계잔액시산표",
        "user_name": u["name"],
        "rows": result,
        "grouped": grouped,
        "totals": totals,
        "date_from": date_from,
        "date_to": date_to,
    })


@router.get("/ledger/trial-balance/excel")
async def trial_balance_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                              date_from: str = "", date_to: str = ""):
    date_from, date_to, result, grouped, totals = _trial_balance_data(conn, u["book"], date_from, date_to)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["계정코드", "계정명", "이월잔액", "당기차변", "당기대변", "차변잔액", "대변잔액"])
    for r in result:
        writer.writerow([esc_cell(v) for v in [
            r["code"], r["name"], r["prev_bal"], r["period_debit"], r["period_credit"],
            r["debit_bal"], r["credit_bal"],
        ]])
    writer.writerow(["합계", "", totals["prev"], totals["debit"], totals["credit"], totals["debit_bal"], totals["credit_bal"]])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=trial_balance.csv"},
    )


def _period(date_from, date_to):
    today = today_kst()
    q = (today.month - 1) // 3 * 3 + 1
    last = calendar.monthrange(today.year, q + 2)[1]
    def_from, def_to = date(today.year, q, 1).isoformat(), date(today.year, q + 2, last).isoformat()
    try:
        if date_from:
            date.fromisoformat(date_from)
        if date_to:
            date.fromisoformat(date_to)
    except ValueError:
        return def_from, def_to
    return date_from or def_from, date_to or def_to


def _get_inv(conn, inv_id, book):
    return conn.execute(
        """SELECT t.*, p.name AS partner_name, p.biz_no, p.representative, p.address, p.biz_type, p.biz_item,
                  c.name AS contact_name
           FROM tax_invoices t LEFT JOIN partners p ON p.id=t.partner_id
           LEFT JOIN partner_contacts c ON c.id=t.contact_id
           WHERE t.id=? AND t.company_id=?""",
        (inv_id, book),
    ).fetchone()


INV_SORTS = {"date": "t.issue_date", "partner": "p.name", "supply": "t.supply_amt", "vat": "t.vat",
             "total": "t.total", "status": "t.status", "slip": "d.slip_no"}


def _invoice_filter_rows(conn, book, type_, date_from, date_to, partner_id, status, sort="", dir=""):
    kind = {"sales": "매출", "purchase": "매입"}.get(type_)
    sql = """SELECT t.*, p.name AS partner_name, d.slip_no FROM tax_invoices t
             LEFT JOIN partners p ON p.id=t.partner_id LEFT JOIN erp_docs d ON d.id=t.doc_id
             WHERE t.company_id=? AND t.issue_date BETWEEN ? AND ?"""
    args = [book, date_from, date_to]
    if partner_id:
        sql += " AND t.partner_id=?"
        args.append(partner_id)
    order = INV_SORTS.get(sort)
    sql += f" ORDER BY {order} {'DESC' if dir == 'desc' else 'ASC'}, t.id DESC" if order else " ORDER BY t.issue_date DESC, t.id DESC"
    base = [dict(r) for r in conn.execute(sql, args)]
    for r in base:
        r["reason_label"] = invoices.AMEND_REASONS.get(r["amend_reason"], "")
        r["vat_code"] = invoices.vat_code(r)
    scope = [r for r in base if not kind or r["kind"] == kind]
    rows = [r for r in scope if not status or r["status"] == status]
    return base, scope, rows


@router.get("/tax-invoices", response_class=HTMLResponse)
async def tax_invoice_list(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                           type: str = "", date_from: str = "", date_to: str = "",
                           partner_id: str = "", status: str = "", sort: str = "", dir: str = ""):
    date_from, date_to = _period(date_from, date_to)
    base, scope, rows = _invoice_filter_rows(conn, u["book"], type, date_from, date_to, partner_id, status, sort, dir)

    sums = {k: sum(r["total"] for r in base if r["kind"] == k) for k in ("매출", "매입")}
    cnts = {k: sum(1 for r in base if r["kind"] == k) for k in ("매출", "매입")}
    status_cnts = {s: sum(1 for r in scope if r["status"] == s) for s in ("정상", "수정", "취소")}
    status_sums = {s: sum(r["total"] for r in scope if r["status"] == s) for s in ("정상", "수정", "취소")}
    partner = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=?",
                           (partner_id, u["book"])).fetchone() if partner_id else None
    qs = {"date_from": date_from, "date_to": date_to, "partner_id": partner["id"] if partner else ""}
    qs_type = urlencode({k: v for k, v in {**qs, "status": status}.items() if v})
    qs_status = urlencode({k: v for k, v in {**qs, "type": type}.items() if v})
    tot = {k: sum(r[k] for r in rows) for k in ("supply_amt", "vat", "total")}

    page = request.query_params.get("page", "1")
    page = max(1, int(page)) if page.isdigit() else 1
    per_page = 30
    total_pages = max(1, (len(rows) + per_page - 1) // per_page)
    page_rows = rows[(page - 1) * per_page: page * per_page]
    filter_qs = urlencode({k: v for k, v in {**qs, "type": type, "status": status, "sort": sort, "dir": dir}.items() if v})

    return templates.TemplateResponse(request=request, name="erp/tax_invoice_list.html", context={
        "request": request, "page_title": "세금계산서 관리", "user_name": u["name"],
        "rows": page_rows, "sums": sums, "cnts": cnts, "inv_type": type,
        "status_cnts": status_cnts, "status_sums": status_sums, "status": status,
        "date_from": date_from, "date_to": date_to, "partner": partner,
        "qs_type": qs_type, "qs_status": qs_status, "tot": tot, "total_count": len(rows),
        "page": page, "total_pages": total_pages, "base_url": "/tax-invoices?" + filter_qs, "filter_qs": filter_qs,
        "sort": sort, "dir": dir,
    })


@router.get("/tax-invoices/excel")
async def tax_invoice_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                            type: str = "", date_from: str = "", date_to: str = "",
                            partner_id: str = "", status: str = "", sort: str = "", dir: str = ""):
    date_from, date_to = _period(date_from, date_to)
    _, _, rows = _invoice_filter_rows(conn, u["book"], type, date_from, date_to, partner_id, status, sort, dir)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["구분", "작성일자", "거래처", "증빙", "과세유형", "공급가액", "세액", "합계", "상태", "승인번호", "전표번호"])
    for r in rows:
        writer.writerow([esc_cell(v) for v in [
            r["kind"], r["issue_date"], r["partner_name"] or "", r["evidence"], r["tax_type"],
            r["supply_amt"], r["vat"], r["total"], r["status"], r["approval_no"] or "", r["slip_no"] or "",
        ]])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=tax_invoices.csv"},
    )


PAY_CODES = {"매출": ("110", "101", "103"), "매입": ("251", "254", "101", "103")}


def _inv_accts(conn, book):
    rows = conn.execute("SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code", (book,)).fetchall()
    sales = [r for r in rows if "401" <= r["code"] <= "450"]
    buys = [r for r in rows if books.acct_group(r["code"]) in ("재고자산", "유형자산", "매출원가", "판매관리비")]
    return sales, buys


@router.get("/tax-invoices/new", response_class=HTMLResponse)
async def tax_invoice_new(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    return templates.TemplateResponse(request=request, name="erp/tax_invoice_form.html", context={
        "request": request, "page_title": "세금계산서 입력", "user_name": u["name"],
        "me": books.company_info(conn, u["book"]), "today": today_kst().isoformat(), "error": "",
        "vat_types": invoices.VAT_TYPES,
    })


# 전표 + 세금계산서 저장. 커밋은 호출부에서 처리
def save_invoice(conn, book, u, inv, items, acct, pay, contact_id=None):
    partner_row = conn.execute("SELECT name FROM partners WHERE id=? AND company_id=?",
                               (inv["partner_id"], book)).fetchone()
    summary = f"{inv['kind']} {inv['evidence']} {partner_row['name']}"
    doc_id = books.post(conn, book, u["id"], inv["issue_date"], invoices.inv_lines(inv, acct, pay, summary),
                        f"{inv['kind']} {inv['evidence']} ({partner_row['name']})", "tax", None)
    inv_id = invoices.add(conn, book, u["id"], {**inv, "doc_id": doc_id}, items)
    conn.execute("UPDATE tax_invoices SET contact_id=? WHERE id=? AND company_id=?", (contact_id, inv_id, book))
    conn.execute("UPDATE erp_docs SET source_id=? WHERE id=?", (inv_id, doc_id))
    return inv_id


@router.post("/tax-invoices")
async def tax_invoice_create(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    form = await request.form()
    book = u["book"]
    kind, tax_type, evidence = form.get("kind", "매출"), form.get("tax_type", "과세"), form.get("evidence", "세금계산서")
    issue_date = form.get("issue_date", "") or today_kst().isoformat()
    acct, pay = form.get("acct", ""), form.get("pay", "")
    receipt = form.get("receipt_type", "청구")
    save_mode = form.get("save_mode", "detail")

    try:
        partner_id = int(form.get("partner_id") or 0)
    except ValueError:
        partner_id = 0
    partner_row = conn.execute(
        "SELECT id, name FROM partners WHERE id=? AND company_id=?", (partner_id, book)
    ).fetchone() if partner_id else None

    line_nos = [x for x in form.get("line_count", "").split(",") if x.strip()] or ["1"]

    def fail(msg):
        return JSONResponse({"error": msg}, status_code=400)

    if kind not in PAY_CODES or tax_type not in ("과세", "영세", "면세") or evidence not in invoices.EVIDENCES:
        return fail("구분을 확인해 주세요.")
    if evidence == "기타" and (kind == "매입" or tax_type == "면세"):
        return fail("증빙 없는 거래는 매출 과세(건별), 영세(수출)만 입력할 수 있습니다.")
    if evidence == "계산서" and tax_type != "면세":
        return fail("계산서는 면세 거래만 입력할 수 있습니다.")
    if evidence == "세금계산서" and tax_type == "면세":
        return fail("면세 거래는 계산서로 입력해 주세요.")
    if receipt not in ("영수", "청구"):
        return fail("영수/청구 구분을 확인해 주세요.")

    def num(key):
        return int(str(form.get(key) or 0).replace(",", ""))

    try:
        date.fromisoformat(issue_date)
        items = []
        for n in line_nos:
            name, supply = form.get(f"item_name_{n}", "").strip(), num(f"item_supply_{n}")
            if not name and not supply:
                continue
            mm, dd = form.get(f"item_month_{n}", "").strip(), form.get(f"item_day_{n}", "").strip()
            item_date = f"{issue_date[:4]}-{int(mm):02d}-{int(dd):02d}" if mm and dd else issue_date
            date.fromisoformat(item_date)
            items.append({"item_date": item_date, "name": name, "spec": form.get(f"item_spec_{n}", "").strip(),
                          "qty": num(f"item_qty_{n}"), "unit_price": num(f"item_price_{n}"), "supply_amt": supply,
                          "vat": num(f"item_vat_{n}") if tax_type == "과세" else 0,
                          "note": form.get(f"item_note_{n}", "").strip()})
    except ValueError:
        return fail("날짜나 금액 형식이 잘못되었습니다.")
    if any(it["qty"] < 0 or it["unit_price"] < 0 for it in items):
        return fail("품목 수량/단가를 확인해 주세요.")

    sales, buys = _inv_accts(conn, book)
    if acct not in [r["code"] for r in (sales if kind == "매출" else buys)] or pay not in PAY_CODES[kind]:
        return fail("계정을 확인해 주세요.")

    supply = sum(it["supply_amt"] for it in items)
    vat = sum(it["vat"] for it in items)
    if supply <= 0:
        return fail("금액을 입력해 주세요.")

    if not partner_row:
        return fail("거래처를 선택해 주세요.")
    over = partner.guard(conn, book, partner_id, supply + vat, issue_date, "세금계산서", ar=kind == "매출")
    if over and (over["block"] or not form.get("credit_ok")):
        return partner.credit_reply(over)

    contact_id = partner.own_contact(conn, book, partner_id, form.get("contact_id"))

    deduct = 1 if kind == "매출" else int(form.get("deduct", "1") == "1")
    inv = {"kind": kind, "tax_type": tax_type, "evidence": evidence, "issue_date": issue_date,
           "partner_id": partner_id or None, "supply_amt": supply, "vat": vat,
           "approval_no": form.get("approval_no", "").strip(), "is_electronic": int(form.get("is_electronic") == "1"),
           "deduct": deduct, "source_type": "manual", "receipt_type": receipt, "remark": form.get("remark", "").strip()}
    try:
        inv_id = save_invoice(conn, book, u, inv, items, acct, pay, contact_id)
    except ValueError as e:
        conn.rollback()
        return fail(str(e))
    audit.log(conn, u["id"], "invoice_create", "invoice", inv_id, f"{partner_row['name']} {supply + vat:,}",
              company_id=book)
    conn.commit()
    if save_mode == "new":
        return RedirectResponse("/tax-invoices/new", status_code=303)
    return RedirectResponse(f"/tax-invoices/{inv_id}", status_code=303)


def _cell(row, i):
    return row[i] if i is not None and i < len(row) else None


# 홈택스 조회결과 엑셀 열 위치 파악. 상호는 공급자/공급받는자 두 번 나옴
def _hometax_cols(head):
    col = {h: head.index(h) for h in
           ("작성일자", "승인번호", "공급자사업자등록번호", "공급받는자사업자등록번호", "합계금액", "공급가액", "세액")
           if h in head}
    col["분류"] = next((i for i, h in enumerate(head) if "분류" in h), None)
    col["종류"] = next((i for i, h in enumerate(head) if "종류" in h), None)
    seller, buyer = col.get("공급자사업자등록번호"), col.get("공급받는자사업자등록번호")
    col["공급자상호"] = next((i for i, h in enumerate(head) if h == "상호" and seller is not None and i > seller), None)
    col["공급받는자상호"] = next((i for i, h in enumerate(head) if h == "상호" and buyer is not None and i > buyer), None)
    return col


# 홈택스 작성일자는 8자리 숫자 또는 날짜셀
def _hometax_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    d = "".join(c for c in s if c.isdigit())
    return f"{d[:4]}-{d[4:6]}-{d[6:8]}" if len(d) == 8 else s


def _num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    return int(str(v).replace(",", "").strip() or 0)


IMPORT_NEED_COLS = ("작성일자", "승인번호", "공급자사업자등록번호", "공급받는자사업자등록번호",
                    "합계금액", "공급가액", "세액", "공급자상호", "공급받는자상호")


@router.get("/tax-invoices/import", response_class=HTMLResponse)
async def tax_invoice_import_form(request: Request, u: dict = Depends(require_book)):
    return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context={
        "request": request, "page_title": "홈택스 세금계산서 불러오기", "user_name": u["name"],
        "kind": "매출", "rows": None, "errors": [], "result": None,
    })


@router.post("/tax-invoices/import", response_class=HTMLResponse)
async def tax_invoice_import_preview(request: Request, file: UploadFile = File(...), kind: str = Form("매출"),
                                     u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    ctx = {"request": request, "page_title": "홈택스 세금계산서 불러오기", "user_name": u["name"],
           "kind": kind, "rows": None, "errors": [], "result": None}
    if kind not in PAY_CODES:
        ctx["errors"] = ["구분을 확인해 주세요."]
        return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)
    try:
        ws = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True).active
        sheet = list(ws.iter_rows(values_only=True))
    except Exception:
        ctx["errors"] = ["엑셀(xlsx) 파일을 읽을 수 없습니다."]
        return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)

    head_no = next((i for i, row in enumerate(sheet[:20])
                    if any(str(c).strip() == "승인번호" for c in row if c is not None)), None)
    if head_no is None:
        ctx["errors"] = ["승인번호 열을 찾을 수 없습니다. 홈택스 조회결과 파일인지 확인해 주세요."]
        return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)

    head = [str(c).strip() if c is not None else "" for c in sheet[head_no]]
    col = _hometax_cols(head)
    if any(col.get(k) is None for k in IMPORT_NEED_COLS):
        ctx["errors"] = ["필수 열이 없습니다. 홈택스 조회결과 파일인지 확인해 주세요."]
        return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)

    # 매출이면 상대방 = 공급받는자, 매입이면 공급자
    other = "공급자" if kind == "매입" else "공급받는자"
    biz_col, name_col = col[f"{other}사업자등록번호"], col[f"{other}상호"]
    sales, buys = _inv_accts(conn, book)
    acct_opts = buys if kind == "매입" else sales
    dup_set = {r["approval_no"] for r in conn.execute(
        "SELECT approval_no FROM tax_invoices WHERE company_id=? AND approval_no!=''", (book,))}

    rows = []
    for raw in sheet[head_no + 1:]:
        appno = str(_cell(raw, col["승인번호"]) or "").strip()
        if not appno:
            continue
        biz_raw = str(_cell(raw, biz_col) or "").strip()
        biz_no = norm_biz_no(biz_raw) or ""
        name = str(_cell(raw, name_col) or "").strip()
        supply, vat = _num(_cell(raw, col["공급가액"])), _num(_cell(raw, col["세액"]))
        classify = str(_cell(raw, col["분류"]) or "").strip()
        typ = str(_cell(raw, col["종류"]) or "").strip()
        if vat == 0 and "계산서" in classify + typ:
            tax_type = "면세"
        elif "영세" in classify + typ:
            tax_type = "영세"
        else:
            tax_type = "과세"

        partner_row = conn.execute("SELECT id, name FROM partners WHERE company_id=? AND biz_no=?",
                                   (book, biz_no)).fetchone() if biz_no else None
        acct_default = "120"
        if kind == "매입" and partner_row:
            last = conn.execute(
                """SELECT sl.account_code FROM tax_invoices t JOIN slip_lines sl ON sl.doc_id=t.doc_id
                   WHERE t.company_id=? AND t.partner_id=? AND t.kind='매입' AND sl.debit>0 AND sl.account_code!='135'
                   ORDER BY t.id DESC, sl.line_no LIMIT 1""",
                (book, partner_row["id"]),
            ).fetchone()
            if last:
                acct_default = last["account_code"]

        rows.append({
            "no": len(rows) + 1, "approval_no": appno, "issue_date": _hometax_date(_cell(raw, col["작성일자"])),
            "name": partner_row["name"] if partner_row else name, "biz_no": biz_no, "supply": supply, "vat": vat,
            "total": supply + vat, "tax_type": tax_type,
            "partner_id": partner_row["id"] if partner_row else "", "create": not partner_row,
            "acct": acct_default, "dup": appno in dup_set,
        })
    if not rows:
        ctx["errors"] = ["불러올 행이 없습니다."]
        return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)

    ctx["rows"], ctx["acct_opts"] = rows, acct_opts
    return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context=ctx)


# 생성 체크 시 상호+사업자번호로 신규 거래처 생성. 확정 직전 사업자번호로 재조회해 중복 생성 방지
def _import_partner(conn, book, u, kind, name, biz_no):
    if biz_no:
        existing = conn.execute("SELECT id FROM partners WHERE company_id=? AND biz_no=?", (book, biz_no)).fetchone()
        if existing:
            return existing["id"]
    last = conn.execute(
        "SELECT code FROM partners WHERE company_id=? AND code GLOB '[0-9][0-9][0-9][0-9][0-9]' ORDER BY code DESC LIMIT 1",
        (book,),
    ).fetchone()
    code = f"{(int(last['code']) + 1) if last else 1:05d}"
    ptype = "customer" if kind == "매출" else "supplier"
    return conn.execute(
        "INSERT INTO partners (company_id, code, name, partner_type, biz_no, created_by) VALUES (?,?,?,?,?,?)",
        (book, code, name or "거래처", ptype, biz_no, u["id"]),
    ).lastrowid


@router.post("/tax-invoices/import/confirm")
async def tax_invoice_import_confirm(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    form = await request.form()
    book = u["book"]
    kind = form.get("kind", "매출")
    nos = [x for x in form.get("row_count", "").split(",") if x.strip()]

    def fail(msg):
        return JSONResponse({"error": msg}, status_code=400)

    if kind not in PAY_CODES:
        return fail("구분을 확인해 주세요.")

    pay = "110" if kind == "매출" else "251"
    sales, buys = _inv_accts(conn, book)
    acct_codes = [r["code"] for r in (sales if kind == "매출" else buys)]

    added = skipped = dup = 0
    try:
        for n in nos:
            approval_no = form.get(f"row_{n}_approval_no", "").strip()
            if not approval_no:
                continue
            if conn.execute("SELECT 1 FROM tax_invoices WHERE company_id=? AND approval_no=?",
                            (book, approval_no)).fetchone():
                dup += 1
                continue

            issue_date = form.get(f"row_{n}_issue_date", "").strip()
            tax_type = form.get(f"row_{n}_tax_type", "과세")
            if tax_type not in ("과세", "영세", "면세"):
                raise ValueError(f"{n}행: 과세구분을 확인해 주세요.")
            try:
                date.fromisoformat(issue_date)
                supply = int(form.get(f"row_{n}_supply", "0").replace(",", ""))
                vat = int(form.get(f"row_{n}_vat", "0").replace(",", ""))
            except ValueError:
                raise ValueError(f"{n}행: 날짜나 금액 형식이 잘못되었습니다.")

            partner_id = form.get(f"row_{n}_partner_id", "").strip()
            partner_id = int(partner_id) if partner_id.isdigit() else None
            if not partner_id:
                if form.get(f"row_{n}_create") != "1":
                    skipped += 1
                    continue
                name = form.get(f"row_{n}_name", "").strip()
                biz_no = form.get(f"row_{n}_biz_no", "").strip()
                partner_id = _import_partner(conn, book, u, kind, name, biz_no)

            acct = form.get(f"row_{n}_acct", "120") if kind == "매입" else "401"
            if acct not in acct_codes:
                raise ValueError(f"{n}행: 계정을 확인해 주세요.")

            evidence = "계산서" if tax_type == "면세" else "세금계산서"
            inv = {"kind": kind, "tax_type": tax_type, "evidence": evidence, "issue_date": issue_date,
                   "partner_id": partner_id, "supply_amt": supply, "vat": vat, "approval_no": approval_no,
                   "is_electronic": 1, "deduct": 1, "source_type": "hometax", "receipt_type": "청구", "remark": ""}
            items = [{"name": "홈택스 불러오기", "supply_amt": supply, "vat": vat}]
            inv_id = save_invoice(conn, book, u, inv, items, acct, pay)
            partner_name = conn.execute("SELECT name FROM partners WHERE id=?", (partner_id,)).fetchone()["name"]
            audit.log(conn, u["id"], "invoice_import", "invoice", inv_id,
                      f"{partner_name} {supply + vat:,}", company_id=book)
            added += 1
    except ValueError as e:
        conn.rollback()
        return fail(str(e))

    conn.commit()
    return templates.TemplateResponse(request=request, name="erp/tax_invoice_import.html", context={
        "request": request, "page_title": "홈택스 세금계산서 불러오기", "user_name": u["name"],
        "kind": kind, "rows": None, "errors": [],
        "result": {"added": added, "skipped": skipped, "dup": dup},
    })


@router.post("/tax-invoices/{inv_id}/amend")
async def tax_invoice_amend(inv_id: int, issue_date: str = Form(...), supply: str = Form(...),
                            reason: str = Form(""), u: dict = Depends(require_book), conn=Depends(get_db)):
    orig = _get_inv(conn, inv_id, u["book"])
    if not orig:
        return HTMLResponse("<h2>세금계산서를 찾을 수 없습니다</h2>", status_code=404)
    try:
        date.fromisoformat(issue_date)
        supply = int(supply)
        reason_code = int(reason) if reason else None
    except ValueError:
        return RedirectResponse(f"/tax-invoices/{inv_id}?error={quote('날짜나 금액 형식이 잘못되었습니다.')}", status_code=303)
    if supply < 0:
        return RedirectResponse(f"/tax-invoices/{inv_id}?error={quote('금액을 확인해 주세요.')}", status_code=303)
    try:
        new_id = invoices.amend(conn, u["book"], u["id"], inv_id, issue_date, supply, reason_code)
    except ValueError as e:
        conn.rollback()
        return RedirectResponse(f"/tax-invoices/{inv_id}?error={quote(str(e))}", status_code=303)
    # 사유별로 amend 내부에서 일자/금액 보정됨. 저장된 수정분 기준 기록
    amd = conn.execute("SELECT issue_date, supply_amt FROM tax_invoices WHERE id=?", (new_id,)).fetchone()
    audit.log(conn, u["id"], "invoice_amend", "invoice", new_id,
              f"{orig['partner_name']} 수정 {amd['supply_amt']:+,}",
              detail=audit.diff(dict(orig), {"issue_date": amd["issue_date"]}), company_id=u["book"])
    conn.commit()
    return RedirectResponse(f"/tax-invoices/{new_id}", status_code=303)


@router.post("/tax-invoices/{inv_id}/approval")
async def tax_invoice_approval(inv_id: int, approval_no: str = Form(""),
                               u: dict = Depends(require_book), conn=Depends(get_db)):
    if not _get_inv(conn, inv_id, u["book"]):
        return HTMLResponse("<h2>세금계산서를 찾을 수 없습니다</h2>", status_code=404)
    conn.execute("UPDATE tax_invoices SET approval_no=? WHERE id=? AND company_id=?",
                 (approval_no.strip(), inv_id, u["book"]))
    conn.commit()
    return RedirectResponse(f"/tax-invoices/{inv_id}", status_code=303)


def _parties(conn, book, inv):
    me = books.company_info(conn, book)
    partner = {k: inv[k] or "" for k in ("biz_no", "representative", "address", "biz_type", "biz_item")}
    partner["name"] = inv["partner_name"] or ""
    return (me, partner) if inv["kind"] == "매출" else (partner, me)


@router.get("/tax-invoices/{inv_id}", response_class=HTMLResponse)
async def tax_invoice_detail(request: Request, inv_id: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    inv = _get_inv(conn, inv_id, u["book"])
    if not inv:
        return HTMLResponse("<h2>세금계산서를 찾을 수 없습니다</h2>", status_code=404)
    items = [dict(r) for r in conn.execute(
        "SELECT * FROM tax_invoice_items WHERE invoice_id=? ORDER BY line_no", (inv_id,)
    ).fetchall()]
    for it in items:
        it["mm"] = it["item_date"][5:7] if it["item_date"] else ""
        it["dd"] = it["item_date"][8:10] if it["item_date"] else ""
    root = inv["orig_id"] or inv["id"]
    related = [dict(r) for r in conn.execute(
        "SELECT id, issue_date, status, supply_amt, vat, total, amend_reason FROM tax_invoices WHERE (id=? OR orig_id=?) AND id!=? ORDER BY id",
        (root, root, inv_id),
    ).fetchall()]
    for r in related:
        r["reason_label"] = invoices.AMEND_REASONS.get(r["amend_reason"], "")
    reason_label = invoices.AMEND_REASONS.get(inv["amend_reason"], "")
    slip_no = conn.execute("SELECT slip_no FROM erp_docs WHERE id=?", (inv["doc_id"],)).fetchone() if inv["doc_id"] else None
    supplier, buyer = _parties(conn, u["book"], inv)
    return templates.TemplateResponse(request=request, name="erp/tax_invoice_detail.html", context={
        "request": request, "page_title": "세금계산서", "user_name": u["name"],
        "inv": inv, "items": items, "related": related, "slip_no": slip_no["slip_no"] if slip_no else "",
        "supplier": supplier, "buyer": buyer, "reason_label": reason_label,
        "amend_reasons": [(c, label, c in invoices.FIX_DATE, c in invoices.FULL_CANCEL)
                          for c, label in invoices.AMEND_REASONS.items()],
        "error": request.query_params.get("error", ""),
    })


@router.get("/tax-invoices/{inv_id}/print", response_class=HTMLResponse)
async def tax_invoice_print(request: Request, inv_id: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    inv = _get_inv(conn, inv_id, u["book"])
    if not inv:
        return HTMLResponse("<h2>세금계산서를 찾을 수 없습니다</h2>", status_code=404)
    items = [dict(r) for r in conn.execute(
        "SELECT * FROM tax_invoice_items WHERE invoice_id=? ORDER BY line_no", (inv_id,)
    ).fetchall()]
    for it in items:
        it["mm"] = it["item_date"][5:7] if it["item_date"] else ""
        it["dd"] = it["item_date"][8:10] if it["item_date"] else ""
    supplier, buyer = _parties(conn, u["book"], inv)
    title = "계산서" if inv["evidence"] == "계산서" else "세금계산서"
    reason_label = invoices.AMEND_REASONS.get(inv["amend_reason"], "")
    return templates.TemplateResponse(request=request, name="erp/tax_invoice.html", context={
        "request": request, "inv": inv, "items": items, "supplier": supplier, "buyer": buyer, "title": title,
        "mm": inv["issue_date"][5:7], "dd": inv["issue_date"][8:10], "reason_label": reason_label,
    })


@router.get("/tax/purchase-summary", response_class=HTMLResponse)
async def tax_purchase_summary(date_from: str = "", date_to: str = ""):
    qs = f"date_from={date_from}&date_to={date_to}" if date_from or date_to else ""
    return RedirectResponse(f"/tax/sales-summary?{qs}#purchase", status_code=307)


@router.get("/tax/sales-summary", response_class=HTMLResponse)
async def tax_sales_summary(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                            date_from: str = "", date_to: str = ""):
    date_from, date_to = _period(date_from, date_to)
    sales_rows, _ = invoices.partner_totals(conn, u["book"], "매출", date_from, date_to)
    buy_rows, _ = invoices.partner_totals(conn, u["book"], "매입", date_from, date_to)
    return templates.TemplateResponse(request=request, name="erp/tax_sales_summary.html", context={
        "request": request, "page_title": "거래처별 세금계산서 합계표", "user_name": u["name"],
        "sales_rows": sales_rows, "sales_summary": invoices.summary_breakdown(sales_rows),
        "buy_rows": buy_rows, "buy_summary": invoices.summary_breakdown(buy_rows),
        "date_from": date_from, "date_to": date_to,
    })


@router.get("/tax/vat-return", response_class=HTMLResponse)
async def tax_vat_return(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                         date_from: str = "", date_to: str = ""):
    date_from, date_to = _period(date_from, date_to)
    year = int(date_from[:4])
    quick_periods = [
        ("1기 예정 (1~3월)", date(year, 1, 1).isoformat(), date(year, 3, 31).isoformat()),
        ("1기 확정 (4~6월)", date(year, 4, 1).isoformat(), date(year, 6, 30).isoformat()),
        ("2기 예정 (7~9월)", date(year, 7, 1).isoformat(), date(year, 9, 30).isoformat()),
        ("2기 확정 (10~12월)", date(year, 10, 1).isoformat(), date(year, 12, 31).isoformat()),
    ]
    return templates.TemplateResponse(request=request, name="erp/tax_vat_return.html", context={
        "request": request, "page_title": "부가세 신고서", "user_name": u["name"],
        "date_from": date_from, "date_to": date_to, "r": invoices.vat_return(conn, u["book"], date_from, date_to),
        "quick_periods": quick_periods,
    })


@router.get("/fi/closing", response_class=HTMLResponse)
async def fi_closing_list(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    error = request.query_params.get("error", "")
    book = u["book"]

    earliest = conn.execute(
        "SELECT MIN(slip_date) as d FROM erp_docs WHERE company_id=? AND slip_date IS NOT NULL AND slip_date != ''",
        [book],
    ).fetchone()

    months = []
    if earliest and earliest["d"]:
        y, m = int(earliest["d"][:4]), int(earliest["d"][5:7])
        t = today_kst()
        end_y, end_m = t.year, t.month
        while (y, m) <= (end_y, end_m):
            closing = conn.execute(
                "SELECT * FROM fi_closings WHERE company_id=? AND year=? AND month=?", [book, y, m]
            ).fetchone()

            slip_count = conn.execute(
                "SELECT COUNT(*) as cnt FROM erp_docs WHERE company_id=? AND slip_date BETWEEN ? AND ?",
                [book, f"{y}-{m:02d}-01", f"{y}-{m:02d}-{calendar.monthrange(y, m)[1]}"]
            ).fetchone()["cnt"]

            closer_name = ""
            reopener_name = ""
            if closing:
                if closing["closed_by"]:
                    cu = conn.execute("SELECT name FROM users WHERE id=?", [closing["closed_by"]]).fetchone()
                    closer_name = cu["name"] if cu else ""
                if closing["reopened_by"]:
                    ru = conn.execute("SELECT name FROM users WHERE id=?", [closing["reopened_by"]]).fetchone()
                    reopener_name = ru["name"] if ru else ""

            months.append({
                "year": y, "month": m,
                "status": closing["status"] if closing else "open",
                "closing_id": closing["id"] if closing else None,
                "closed_at": (closing["closed_at"][:16] if closing and closing["closed_at"] else ""),
                "closer_name": closer_name,
                "reopened_at": (closing["reopened_at"][:10] if closing and closing["reopened_at"] else ""),
                "reopener_name": reopener_name,
                "reopen_reason": closing["reopen_reason"] if closing else "",
                "slip_count": slip_count,
            })
            m += 1
            if m > 12:
                m = 1
                y += 1

    months.reverse()

    return templates.TemplateResponse(request=request, name="erp/fi_closing.html", context={
        "request": request, "page_title": "월마감", "user_name": u["name"],
        "months": months, "error": error,
    })


@router.post("/fi/closing/close")
async def fi_close(request: Request, conn=Depends(get_db), u: dict = Depends(require_book),
                   year: int = Form(...), month: int = Form(...)):
    if u["role"] not in ("client", "platform_staff"):
        return HTMLResponse("<h2>월마감 권한이 없습니다</h2>", status_code=403)

    book = u["book"]
    last_day = calendar.monthrange(year, month)[1]

    check = conn.execute("""
        SELECT COALESCE(SUM(sl.debit),0) as total_d, COALESCE(SUM(sl.credit),0) as total_c
        FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
        WHERE d.company_id=? AND d.status='done' AND d.slip_date BETWEEN ? AND ?
    """, [book, f"{year}-{month:02d}-01", f"{year}-{month:02d}-{last_day}"]).fetchone()
    if check["total_d"] != check["total_c"]:
        return RedirectResponse(f"/fi/closing?error=imbalance", status_code=303)

    conn.execute("""
        INSERT INTO fi_closings (company_id, year, month, status, closed_by, closed_at)
        VALUES (?, ?, ?, 'closed', ?, datetime('now','localtime'))
        ON CONFLICT(company_id, year, month) DO UPDATE SET
            status='closed', closed_by=excluded.closed_by, closed_at=excluded.closed_at,
            reopened_by=NULL, reopened_at=NULL, reopen_reason=NULL
    """, [book, year, month, u["id"]])
    closing_id = conn.execute(
        "SELECT id FROM fi_closings WHERE company_id=? AND year=? AND month=?", [book, year, month]
    ).fetchone()[0]
    audit.log(conn, u["id"], "close", "closing", closing_id, f"{year}년 {month}월 마감", company_id=book)
    conn.commit()
    return RedirectResponse("/fi/closing", status_code=303)


@router.post("/fi/closing/reopen")
async def fi_reopen(request: Request, conn=Depends(get_db), u: dict = Depends(require_book),
                    closing_id: int = Form(...), reason: str = Form(...), password: str = Form(...)):
    if u["role"] not in ("client", "platform_staff"):
        return HTMLResponse("<h2>월마감 권한이 없습니다</h2>", status_code=403)

    row = conn.execute("SELECT password FROM users WHERE id=?", (u["id"],)).fetchone()
    if not row or not verify_pw(password, row["password"]):
        return RedirectResponse("/fi/closing?error=wrong_password", status_code=303)
    c = conn.execute("SELECT year FROM fi_closings WHERE id=? AND company_id=?", [closing_id, u["book"]]).fetchone()
    if c and conn.execute(
        "SELECT 1 FROM fiscal_closings WHERE company_id=? AND year=? AND status='closed'", [u["book"], c["year"]]
    ).fetchone():
        return RedirectResponse("/fi/closing?error=fiscal_closed", status_code=303)

    conn.execute("""
        UPDATE fi_closings SET status='open', reopened_by=?, reopened_at=datetime('now','localtime'),
               reopen_reason=?
        WHERE id=? AND company_id=?
    """, [u["id"], reason, closing_id, u["book"]])
    audit.log(conn, u["id"], "reopen", "closing", closing_id, f"마감 취소: {reason}", company_id=u["book"])
    conn.commit()
    return RedirectResponse("/fi/closing", status_code=303)


@router.get("/fi/accounts", response_class=HTMLResponse)
async def fi_accounts(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    rows = conn.execute("SELECT * FROM accounts WHERE company_id=? ORDER BY code", (book,)).fetchall()

    grouped = []
    cur_grp = None
    for r in rows:
        grp = books.acct_group(r["code"])
        if grp != cur_grp:
            grouped.append({"_type": "header", "group": grp})
            cur_grp = grp
        grouped.append({"_type": "row", **dict(r)})

    return templates.TemplateResponse(request=request, name="erp/fi_accounts.html", context={
        "request": request, "page_title": "계정과목",
        "user_name": u["name"],
        "grouped": grouped,
        "can_write": u["role"] in ("client", "platform_staff"),
    })


@router.post("/fi/accounts")
async def fi_account_add(request: Request, conn=Depends(get_db), u: dict = Depends(require_book),
                         code: str = Form(...), name: str = Form(...), category: str = Form(...)):
    if u["role"] not in ("client", "platform_staff"):
        return JSONResponse({"error": "계정과목 등록 권한이 없습니다."}, status_code=403)

    if not (len(code) == 3 and code.isdigit()):
        return JSONResponse({"error": "계정코드는 3자리 숫자여야 합니다."}, status_code=400)

    book = u["book"]
    if conn.execute("SELECT 1 FROM accounts WHERE company_id=? AND code=?", (book, code)).fetchone():
        return JSONResponse({"error": "이미 존재하는 계정코드입니다."}, status_code=400)

    cats = [(101, 250, "자산"), (251, 330, "부채"), (331, 400, "자본"), (401, 450, "수익"),
            (451, 900, "비용"), (901, 950, "수익"), (951, 999, "비용")]
    cat = next((c for lo, hi, c in cats if lo <= int(code) <= hi), None)
    if cat != category:
        return JSONResponse({"error": "계정코드와 계정구분이 맞지 않습니다."}, status_code=400)

    is_debit = 1 if category in ("자산", "비용") else 0
    conn.execute(
        "INSERT INTO accounts (company_id, code, name, category, is_debit, is_active) VALUES (?,?,?,?,?,1)",
        (book, code, name, category, is_debit),
    )
    conn.commit()
    return RedirectResponse("/fi/accounts", status_code=303)


@router.post("/fi/accounts/{code}")
async def fi_account_edit(request: Request, code: str, conn=Depends(get_db), u: dict = Depends(require_book),
                          name: str = Form(...), is_active: int = Form(...),
                          req_partner: int = Form(...), req_dept: int = Form(...)):
    if u["role"] not in ("client", "platform_staff"):
        return JSONResponse({"error": "계정과목 수정 권한이 없습니다."}, status_code=403)

    book = u["book"]
    if not conn.execute("SELECT 1 FROM accounts WHERE company_id=? AND code=?", (book, code)).fetchone():
        return JSONResponse({"error": "계정을 찾을 수 없습니다."}, status_code=404)

    conn.execute(
        "UPDATE accounts SET name=?, is_active=?, req_partner=?, req_dept=? WHERE company_id=? AND code=?",
        (name, is_active, req_partner, req_dept, book, code),
    )
    conn.commit()
    return RedirectResponse("/fi/accounts", status_code=303)


def _income_statement_data(conn, book, date_from, date_to):
    today = today_kst()
    if not date_from:
        date_from = f"{today.year}-01-01"
    if not date_to:
        date_to = today.isoformat()

    rows = conn.execute(
        """SELECT a.code, a.name, a.category, a.is_debit,
                  COALESCE(SUM(CASE WHEN ed.id IS NOT NULL THEN sl.debit ELSE 0 END),0) as d,
                  COALESCE(SUM(CASE WHEN ed.id IS NOT NULL THEN sl.credit ELSE 0 END),0) as c
           FROM accounts a
           LEFT JOIN slip_lines sl ON sl.account_code = a.code
           LEFT JOIN erp_docs ed ON sl.doc_id = ed.id AND ed.company_id = ? AND ed.status = 'done'
                AND ed.slip_date BETWEEN ? AND ?
           WHERE a.company_id = ? AND a.category IN ('수익','비용')
           GROUP BY a.code, a.name, a.category, a.is_debit""",
        (book, date_from, date_to, book),
    ).fetchall()

    sections = {"매출": [], "매출원가": [], "판매관리비": [], "영업외수익": [], "영업외비용": []}
    for r in rows:
        bal = (r["c"] - r["d"]) if not r["is_debit"] else (r["d"] - r["c"])
        if bal == 0:
            continue
        grp = books.acct_group(r["code"])
        if grp in sections:
            sections[grp].append({"code": r["code"], "name": r["name"], "balance": bal})

    def total(key):
        return sum(x["balance"] for x in sections[key])

    revenue = total("매출")
    cogs = total("매출원가")
    gross_profit = revenue - cogs
    sga = total("판매관리비")
    operating_income = gross_profit - sga
    other_income = total("영업외수익")
    other_expense = total("영업외비용")
    net_income = operating_income + other_income - other_expense

    return date_from, date_to, sections, {
        "revenue": revenue, "cogs": cogs, "gross_profit": gross_profit,
        "sga": sga, "operating_income": operating_income,
        "other_income": other_income, "other_expense": other_expense,
        "net_income": net_income,
    }


@router.get("/report/income-statement", response_class=HTMLResponse)
async def income_statement(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                           date_from: str = "", date_to: str = ""):
    date_from, date_to, sections, totals = _income_statement_data(conn, u["book"], date_from, date_to)

    return templates.TemplateResponse(request=request, name="erp/income_statement.html", context={
        "request": request, "page_title": "손익계산서",
        "user_name": u["name"],
        "date_from": date_from, "date_to": date_to,
        "sections": sections,
        **totals,
    })


@router.get("/report/income-statement/excel")
async def income_statement_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                                 date_from: str = "", date_to: str = ""):
    date_from, date_to, sections, totals = _income_statement_data(conn, u["book"], date_from, date_to)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["구분", "계정코드", "계정명", "금액"])
    labels = {"매출": "매출액", "매출원가": "매출원가", "판매관리비": "판매비와관리비",
              "영업외수익": "영업외수익", "영업외비용": "영업외비용"}
    for grp, label in labels.items():
        for a in sections[grp]:
            writer.writerow([esc_cell(v) for v in [label, a["code"], a["name"], a["balance"]]])
    writer.writerow(["매출총이익", "", "", totals["gross_profit"]])
    writer.writerow(["영업이익", "", "", totals["operating_income"]])
    writer.writerow(["당기순이익", "", "", totals["net_income"]])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=income_statement.csv"},
    )


def _balance_sheet_data(conn, book, as_of):
    if not as_of:
        as_of = today_kst().isoformat()

    rows = conn.execute("""
        SELECT a.code, a.name, a.category, a.is_debit,
               COALESCE(SUM(CASE WHEN d.id IS NOT NULL THEN sl.debit ELSE 0 END),0) as total_debit,
               COALESCE(SUM(CASE WHEN d.id IS NOT NULL THEN sl.credit ELSE 0 END),0) as total_credit
        FROM accounts a
        LEFT JOIN slip_lines sl ON sl.account_code = a.code
        LEFT JOIN erp_docs d ON sl.doc_id = d.id AND d.slip_date <= ? AND d.company_id = ? AND d.status = 'done'
        WHERE a.is_active = 1 AND a.company_id = ? AND a.category NOT IN ('수익','비용')
        GROUP BY a.code, a.name, a.category, a.is_debit
        HAVING total_debit > 0 OR total_credit > 0
        ORDER BY a.code
    """, [as_of, book, book]).fetchall()

    assets, liabilities, equity = [], [], []
    asset_total = liability_total = equity_total = 0

    for r in rows:
        bal = r["total_debit"] - r["total_credit"] if r["is_debit"] else r["total_credit"] - r["total_debit"]
        if bal == 0:
            continue
        item = {"code": r["code"], "name": r["name"], "balance": bal}
        cat = r["category"]
        if cat == "자산":
            assets.append(item)
            asset_total += bal
        elif cat == "부채":
            liabilities.append(item)
            liability_total += bal
        elif cat == "자본":
            equity.append(item)
            equity_total += bal

    year_start = f"{as_of[:4]}-01-01"
    pl = conn.execute("""
        SELECT a.category, a.is_debit,
               COALESCE(SUM(CASE WHEN d.slip_date < ? THEN sl.debit ELSE 0 END),0) as pre_d,
               COALESCE(SUM(CASE WHEN d.slip_date < ? THEN sl.credit ELSE 0 END),0) as pre_c,
               COALESCE(SUM(CASE WHEN d.slip_date >= ? THEN sl.debit ELSE 0 END),0) as fy_d,
               COALESCE(SUM(CASE WHEN d.slip_date >= ? THEN sl.credit ELSE 0 END),0) as fy_c
        FROM accounts a
        LEFT JOIN slip_lines sl ON sl.account_code = a.code
        LEFT JOIN erp_docs d ON sl.doc_id = d.id AND d.slip_date <= ? AND d.company_id = ? AND d.status = 'done'
        WHERE a.company_id = ? AND a.category IN ('수익','비용')
        GROUP BY a.category, a.is_debit
    """, [year_start, year_start, year_start, year_start, as_of, book, book]).fetchall()

    pre_net = fy_net = 0
    for r in pl:
        pre_bal = r["pre_d"] - r["pre_c"] if r["is_debit"] else r["pre_c"] - r["pre_d"]
        fy_bal = r["fy_d"] - r["fy_c"] if r["is_debit"] else r["fy_c"] - r["fy_d"]
        sign = -1 if r["category"] == "비용" else 1
        pre_net += sign * pre_bal
        fy_net += sign * fy_bal

    if pre_net != 0:
        carried = next((it for it in equity if it["code"] in ("375", "377")), None)
        if carried:
            carried["balance"] += pre_net
        else:
            equity.append({"code": "375", "name": "이월이익잉여금", "balance": pre_net})
        equity_total += pre_net

    if fy_net != 0:
        equity.append({"code": "-", "name": "당기순이익", "balance": fy_net})
        equity_total += fy_net

    return as_of, assets, liabilities, equity, asset_total, liability_total, equity_total


@router.get("/report/balance-sheet", response_class=HTMLResponse)
async def balance_sheet(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                        as_of: str = ""):
    as_of, assets, liabilities, equity, asset_total, liability_total, equity_total = _balance_sheet_data(
        conn, u["book"], as_of)

    return templates.TemplateResponse(request=request, name="erp/balance_sheet.html", context={
        "request": request, "page_title": "재무상태표",
        "user_name": u["name"],
        "as_of": as_of,
        "assets": assets, "liabilities": liabilities, "equity": equity,
        "asset_total": asset_total, "liability_total": liability_total,
        "equity_total": equity_total,
        "is_balanced": asset_total == liability_total + equity_total,
    })


@router.get("/report/balance-sheet/excel")
async def balance_sheet_excel(u: dict = Depends(require_book), conn=Depends(get_db), as_of: str = ""):
    as_of, assets, liabilities, equity, asset_total, liability_total, equity_total = _balance_sheet_data(
        conn, u["book"], as_of)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["구분", "계정코드", "계정명", "금액"])
    for label, items, total in (("자산", assets, asset_total), ("부채", liabilities, liability_total),
                                 ("자본", equity, equity_total)):
        for item in items:
            writer.writerow([esc_cell(v) for v in [label, item["code"], item["name"], item["balance"]]])
        writer.writerow([f"{label} 합계", "", "", total])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=balance_sheet.csv"},
    )


def _period_movements(conn, book, date_from, date_to):
    rows = conn.execute(
        """SELECT a.code, a.category,
                  COALESCE(SUM(CASE WHEN ed.id IS NOT NULL THEN sl.debit ELSE 0 END),0) as d,
                  COALESCE(SUM(CASE WHEN ed.id IS NOT NULL THEN sl.credit ELSE 0 END),0) as c
           FROM accounts a
           LEFT JOIN slip_lines sl ON sl.account_code = a.code
           LEFT JOIN erp_docs ed ON sl.doc_id = ed.id AND ed.company_id = ? AND ed.status = 'done'
                AND ed.slip_date BETWEEN ? AND ?
           WHERE a.company_id = ?
           GROUP BY a.code, a.category""",
        (book, date_from, date_to, book),
    ).fetchall()
    return {r["code"]: {"category": r["category"], "d": r["d"], "c": r["c"]} for r in rows}


def _cash_flow_data(conn, book, date_from, date_to):
    today = today_kst()
    if not date_from:
        date_from = today.replace(day=1).isoformat()
    if not date_to:
        last_day = calendar.monthrange(today.year, today.month)[1]
        date_to = today.replace(day=last_day).isoformat()

    movements = _period_movements(conn, book, date_from, date_to)

    net_income = dep_cf = retire_cf = 0
    ar_cf = ap_cf = inv_asset_cf = other_asset_cf = other_liab_cf = other_op_cf = 0
    stinv_cf = fa_buy = fa_sell = 0
    borrow_in = borrow_out = ltborrow_in = ltborrow_out = 0
    capital_cf = other_equity_cf = 0

    for code, m in movements.items():
        if code in ("101", "102", "103"):
            continue
        d, c, cat = m["d"], m["c"], m["category"]
        cf = c - d
        if d == 0 and c == 0:
            continue
        n = int(code)
        if cat in ("수익", "비용") or code == "400":
            net_income += cf
        elif 207 <= n <= 211:
            dep_cf += cf
        elif code == "292":
            retire_cf += cf
        elif code == "110":
            ar_cf += cf
        elif 120 <= n <= 129:
            inv_asset_cf += cf
        elif 110 <= n <= 199:
            other_asset_cf += cf
        elif code == "251":
            ap_cf += cf
        elif (251 <= n <= 258) or (261 <= n <= 264):
            other_liab_cf += cf
        elif code == "108":
            stinv_cf += cf
        elif 201 <= n <= 206:
            fa_buy += d
            fa_sell += c
        elif code == "259":
            borrow_in += c
            borrow_out += d
        elif code == "291":
            ltborrow_in += c
            ltborrow_out += d
        elif code == "331":
            capital_cf += cf
        elif code in ("341", "375", "377"):
            other_equity_cf += cf
        else:
            other_op_cf += cf

    op_lines = [
        {"label": "당기순이익", "amt": net_income},
        {"label": "감가상각비", "amt": dep_cf},
        {"label": "퇴직급여충당부채의 증가(감소)", "amt": retire_cf},
        {"label": "매출채권의 감소(증가)", "amt": ar_cf},
        {"label": "재고자산의 감소(증가)", "amt": inv_asset_cf},
        {"label": "기타 유동자산의 감소(증가)", "amt": other_asset_cf},
        {"label": "매입채무의 증가(감소)", "amt": ap_cf},
        {"label": "기타 유동부채의 증가(감소)", "amt": other_liab_cf},
    ]
    if other_op_cf:
        op_lines.append({"label": "기타 영업활동", "amt": other_op_cf})
    op_cf = sum(l["amt"] for l in op_lines)

    inv_lines = [
        {"label": "단기금융상품의 증가(감소)", "amt": stinv_cf},
        {"label": "유형자산의 취득", "amt": -fa_buy},
        {"label": "유형자산의 처분", "amt": fa_sell},
    ]
    inv_cf = sum(l["amt"] for l in inv_lines)

    fin_lines = [
        {"label": "단기차입금의 증가", "amt": borrow_in},
        {"label": "단기차입금의 상환", "amt": -borrow_out},
        {"label": "장기차입금의 증가", "amt": ltborrow_in},
        {"label": "장기차입금의 상환", "amt": -ltborrow_out},
        {"label": "자본금의 증가(감소)", "amt": capital_cf},
        {"label": "기타자본잉여금의 증감", "amt": other_equity_cf},
    ]
    fin_cf = sum(l["amt"] for l in fin_lines)

    net_chg = op_cf + inv_cf + fin_cf

    cash_row = conn.execute(
        """SELECT COALESCE(SUM(sl.debit - sl.credit), 0) AS v
           FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
           WHERE d.company_id=? AND d.status='done' AND CAST(sl.account_code AS INTEGER) BETWEEN 101 AND 103
           AND d.slip_date < ?""",
        (book, date_from),
    ).fetchone()
    cash_begin = cash_row["v"]
    cash_end = cash_begin + net_chg

    return (date_from, date_to, op_lines, op_cf, inv_lines, inv_cf, fin_lines, fin_cf,
            cash_begin, net_chg, cash_end)


@router.get("/report/cash-flow", response_class=HTMLResponse)
async def cash_flow_report(request: Request, u: dict = Depends(require_book), conn=Depends(get_db),
                           date_from: str = "", date_to: str = ""):
    (date_from, date_to, op_lines, op_cf, inv_lines, inv_cf, fin_lines, fin_cf,
     cash_begin, net_chg, cash_end) = _cash_flow_data(conn, u["book"], date_from, date_to)

    return templates.TemplateResponse(request=request, name="erp/cash_flow.html", context={
        "request": request, "page_title": "현금흐름표", "user_name": u["name"],
        "date_from": date_from, "date_to": date_to,
        "op_lines": op_lines, "op_cf": op_cf,
        "inv_lines": inv_lines, "inv_cf": inv_cf,
        "fin_lines": fin_lines, "fin_cf": fin_cf,
        "cash_begin": cash_begin, "net_chg": net_chg, "cash_end": cash_end,
    })


@router.get("/report/cash-flow/excel")
async def cash_flow_excel(u: dict = Depends(require_book), conn=Depends(get_db),
                          date_from: str = "", date_to: str = ""):
    (date_from, date_to, op_lines, op_cf, inv_lines, inv_cf, fin_lines, fin_cf,
     cash_begin, net_chg, cash_end) = _cash_flow_data(conn, u["book"], date_from, date_to)

    buf = io.StringIO()
    buf.write('﻿')
    writer = csv.writer(buf)
    writer.writerow(["구분", "항목", "금액"])
    for label, lines, subtotal in (("영업활동", op_lines, op_cf), ("투자활동", inv_lines, inv_cf),
                                    ("재무활동", fin_lines, fin_cf)):
        for l in lines:
            writer.writerow([esc_cell(v) for v in [label, l["label"], l["amt"]]])
        writer.writerow([f"{label} 소계", "", subtotal])
    writer.writerow(["현금 증감", "", net_chg])
    writer.writerow(["기초 현금", "", cash_begin])
    writer.writerow(["기말 현금", "", cash_end])

    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=cash_flow.csv"},
    )


@router.get("/ledger/aging", response_class=HTMLResponse)
async def aging_report(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    view = request.query_params.get("view", "ar")
    basis = request.query_params.get("basis", "date")
    acct_code = "110" if view == "ar" else "251"
    as_of = today_kst()

    groups = []
    for pid, info in credit.open_items(conn, book, acct_code).items():
        prow = conn.execute("SELECT name, pay_terms, pay_days FROM partners WHERE id=? AND company_id=?",
                            (pid, book)).fetchone()
        if prow:
            groups.append((dict(prow, id=pid), info))
    # 거래처 마스터와 연결되지 않은 과거 전표는 이름 기준 표시. 결제조건 없이 전표일을 만기일로 처리
    for name, info in credit.open_items(conn, book, acct_code, orphan=True).items():
        groups.append(({"name": name, "pay_terms": "", "pay_days": 0, "id": None}, info))

    buckets = []
    for prow, info in groups:
        if basis == "due":
            b = {"name": prow["name"], "pid": prow["id"], "notdue": 0, "d30": 0, "d60": 0, "d90": 0, "over90": 0, "total": 0}
            for it in info["items"]:
                if it["remain"] <= 0:
                    continue
                due = partner.due_date(it["date"], prow["pay_terms"], prow["pay_days"])
                age = (as_of - date.fromisoformat(due)).days
                if age <= 0:
                    b["notdue"] += it["remain"]
                elif age <= 30:
                    b["d30"] += it["remain"]
                elif age <= 60:
                    b["d60"] += it["remain"]
                elif age <= 90:
                    b["d90"] += it["remain"]
                else:
                    b["over90"] += it["remain"]
            if info["over"]:
                b["notdue"] -= info["over"]
            b["total"] = b["notdue"] + b["d30"] + b["d60"] + b["d90"] + b["over90"]
        else:
            b = {"name": prow["name"], "pid": prow["id"], "cur": 0, "d30": 0, "d60": 0, "over90": 0, "total": 0}
            for it in info["items"]:
                if it["remain"] <= 0:
                    continue
                age = (as_of - date.fromisoformat(it["date"])).days
                if age <= 30:
                    b["cur"] += it["remain"]
                elif age <= 60:
                    b["d30"] += it["remain"]
                elif age <= 90:
                    b["d60"] += it["remain"]
                else:
                    b["over90"] += it["remain"]
            if info["over"]:
                b["cur"] -= info["over"]
            b["total"] = b["cur"] + b["d30"] + b["d60"] + b["over90"]
        buckets.append(b)

    partners = sorted([v for v in buckets if v["total"] != 0], key=lambda x: -abs(x["total"]))
    bucket_keys = ("notdue", "d30", "d60", "d90", "over90") if basis == "due" else ("cur", "d30", "d60", "over90")
    totals = {k: 0 for k in bucket_keys}
    totals["total"] = 0
    for p in partners:
        for k in bucket_keys:
            totals[k] += p[k]
        totals["total"] += p["total"]

    return templates.TemplateResponse(request=request, name="erp/aging_report.html", context={
        "request": request, "page_title": "채권/채무 연령분석", "user_name": u["name"],
        "view": view, "basis": basis, "partners": partners, "totals": totals,
    })


@router.get("/budget/edit", response_class=HTMLResponse)
async def budget_edit(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    year = int(request.query_params.get("year", date.today().year))
    dept = request.query_params.get("dept", "")
    book = u["book"]
    accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 AND category='비용' ORDER BY code",
        (book,),
    ).fetchall()
    depts = [r["name"] for r in conn.execute(
        "SELECT name FROM depts WHERE company_id=? AND is_active=1 ORDER BY sort, name", (book,)
    ).fetchall()]
    if dept and dept not in depts:
        dept = ""

    saved = {}
    for row in conn.execute(
        "SELECT account_code, month, amount FROM budgets WHERE company_id=? AND year=? AND dept=?",
        (book, year, dept),
    ).fetchall():
        saved[(row["account_code"], row["month"])] = row["amount"]

    grid = []
    for a in accounts:
        months = [saved.get((a["code"], m), 0) for m in range(1, 13)]
        grid.append({"code": a["code"], "name": a["name"], "months": months, "total": sum(months)})

    return templates.TemplateResponse(request=request, name="erp/budget_edit.html", context={
        "request": request, "page_title": "예산 편성", "user_name": u["name"],
        "year": year, "dept": dept, "depts": depts, "grid": grid,
    })


@router.post("/budget/save")
async def budget_save(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    if u["role"] not in ("client", "platform_staff"):
        return JSONResponse({"error": "예산 편성 권한이 없습니다."}, status_code=403)

    book = u["book"]
    form = await request.form()
    year = int(form.get("year", date.today().year))
    dept = form.get("dept", "")
    if dept and not conn.execute("SELECT 1 FROM depts WHERE company_id=? AND name=?", (book, dept)).fetchone():
        dept = ""

    for key, val in form.items():
        if not key.startswith("b_"):
            continue
        parts = key.split("_")
        if len(parts) != 3:
            continue
        code, month = parts[1], int(parts[2])
        amount = int(str(val or 0).replace(",", ""))
        conn.execute(
            "INSERT INTO budgets (company_id, dept, account_code, year, month, amount, created_by) "
            "VALUES (?,?,?,?,?,?,?) "
            "ON CONFLICT(company_id, dept, account_code, year, month) DO UPDATE SET amount=?, created_by=?",
            (book, dept, code, year, month, amount, u["id"], amount, u["id"]),
        )
    conn.commit()
    return RedirectResponse(f"/budget/edit?year={year}&dept={quote(dept)}", status_code=303)


@router.get("/budget/vs-actual", response_class=HTMLResponse)
async def budget_vs_actual(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    year = int(request.query_params.get("year", date.today().year))
    month = int(request.query_params.get("month", date.today().month))
    dept = request.query_params.get("dept", "")
    book = u["book"]
    d_from = f"{year}-{month:02d}-01"
    d_to = f"{year+1}-01-01" if month == 12 else f"{year}-{month+1:02d}-01"

    depts = [r["name"] for r in conn.execute(
        "SELECT name FROM depts WHERE company_id=? AND is_active=1 ORDER BY sort, name", (book,)
    ).fetchall()]
    if dept and dept not in depts:
        dept = ""

    bgt = {}
    for row in conn.execute(
        "SELECT account_code, amount FROM budgets WHERE company_id=? AND year=? AND month=? AND dept=?",
        (book, year, month, dept),
    ).fetchall():
        bgt[row["account_code"]] = row["amount"]

    # 부서 미지정(전사) 조회는 전체 실적과 비교, 부서 선택 시 해당 부서 전표만 집계
    act_sql = """SELECT sl.account_code, SUM(sl.debit) - SUM(sl.credit) AS v
           FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
           WHERE d.company_id=? AND d.status='done' AND d.slip_date >= ? AND d.slip_date < ?"""
    act_params = [book, d_from, d_to]
    if dept:
        act_sql += " AND COALESCE(sl.dept, '')=?"
        act_params.append(dept)
    act_sql += " GROUP BY sl.account_code"

    act_map = {}
    for row in conn.execute(act_sql, act_params).fetchall():
        act_map[row["account_code"]] = row["v"]

    accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 AND category='비용' ORDER BY code", (book,)
    ).fetchall()
    rows = []
    for a in accounts:
        b = bgt.get(a["code"], 0)
        act = act_map.get(a["code"], 0)
        if b == 0 and act == 0:
            continue
        pct = round(act / b * 100) if b > 0 else 0
        rows.append({"code": a["code"], "name": a["name"], "budget": b, "actual": act, "var": b - act, "pct": pct})

    tot_b = sum(r["budget"] for r in rows)
    tot_a = sum(r["actual"] for r in rows)

    return templates.TemplateResponse(request=request, name="erp/budget_vs_actual.html", context={
        "request": request, "page_title": "예산 대비 실적", "user_name": u["name"],
        "year": year, "month": month, "dept": dept, "depts": depts,
        "rows": rows, "tot_b": tot_b, "tot_a": tot_a,
    })
