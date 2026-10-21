import calendar
import csv
import io
import re
import secrets
import time
from datetime import date
from urllib.parse import quote

import msoffcrypto
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, Depends, HTTPException, Form, Request
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response, StreamingResponse

from core.tz import now_kst
from core.auth import verify_pw
from core.deps import get_db, require_staff, templates
from core.books import (
    INTERNAL, book_emp, book_label, company_info, emp_book, hr_book_ok, require_hr_book, require_hr_chief,
)
from core.constants import BANK_CODES, QUOTA, SUBSIDY_RATES, DISABILITY_TYPE_CODES
from core.export import (
    esc_cell, mask_rno, mask_acct, excel_response, styled_workbook, styled_sheet,
    HEADER_FONT, HEADER_FILL, CELL_FONT, THIN_BORDER,
)
from core.crypto import dec_rno, enc_rno
from core.payroll import (
    below_min_wage, build_payslip, build_payslip_detail, calc_employer_insurance, calc_ordinary_hourly,
    calc_premium_pay, get_rates, is_five_plus,
)
from core.pay_items import get_items, PREMIUM_CODES
from core.attendance import is_month_closed, get_month_work_stats, get_worker_month_leaves
from core.severance import calc_ordinary_daily
from core.leave import get_leave_balance
from core.appointments import log_changes
from core.notification import create_bulk_notifications, create_notification
from core import audit, books
from core.rate_limit import get_client_ip
from routers.erp_accounting import _create_payroll_slip

router = APIRouter()


@router.post("/hr-book")
async def set_hr_book(
    request: Request,
    book: int = Form(...),
    next: str = Form("/payroll"),
    u: dict = Depends(require_staff),
    conn=Depends(get_db),
):
    if not next.startswith("/") or next.startswith("//"):
        next = "/payroll"
    request.session["hr_book"] = book if hr_book_ok(conn, u, book) else 0
    return RedirectResponse(url=next, status_code=303)


def _payroll_month(request):
    now = now_kst()
    try:
        year = int(request.query_params.get("year", now.year))
        month = int(request.query_params.get("month", now.month))
    except ValueError:
        year, month = now.year, now.month
    return year, month


def _prev_month(year, month):
    if month > 1:
        return year, month - 1
    return year - 1, 12


def _pay_date(conn, book, year, month):
    co = conn.execute("SELECT pay_day FROM client_companies WHERE id=?", (book,)).fetchone()
    pay_day = co["pay_day"] if co and co["pay_day"] else 25
    return f"{year}-{month:02d}-{pay_day:02d}"


# 급여 이체

BANK_ALIASES = {
    "KB": "국민", "IBK": "기업", "NH": "농협", "KEB": "하나", "SC제일": "SC",
    "iM": "대구", "새마을금고": "새마을", "케이": "케이",
}


def bank_code(name):
    name = (name or "").replace(" ", "")
    for word in ("은행", "뱅크"):
        name = name.replace(word, "")
    for alias, real in BANK_ALIASES.items():
        name = name.replace(alias, real)
    matched = [k for k in BANK_CODES if k in name]
    if not matched:
        return ""
    return BANK_CODES[max(matched, key=len)]


def _transfer_rows(conn, book, year, month):
    rows = conn.execute(
        """SELECT e.id AS emp_id, e.name, e.bank_name, e.bank_account, p.id AS payslip_id, p.net_pay
           FROM employees e JOIN payslips p ON e.id = p.employee_id
           WHERE e.company_id IS ? AND p.pay_year=? AND p.pay_month=? AND p.status='confirmed'
           ORDER BY e.employee_no, e.name""",
        (emp_book(book), year, month),
    ).fetchall()
    rows = [dict(r) for r in rows]
    for r in rows:
        r["missing"] = not r["bank_name"] or not r["bank_account"]
        r["code"] = bank_code(r["bank_name"])
        r["acct_masked"] = mask_acct(r["bank_account"])
    return rows


@router.get("/payroll/transfer", response_class=HTMLResponse)
async def payroll_transfer(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    year, month = _payroll_month(request)
    book = user["book"]

    rows = _transfer_rows(conn, book, year, month)
    missing_cnt = sum(1 for r in rows if r["missing"])
    code_issue_cnt = sum(1 for r in rows if not r["missing"] and not r["code"])

    last_at = conn.execute(
        """SELECT MAX(p.transfer_file_at) AS at FROM employees e JOIN payslips p ON e.id = p.employee_id
           WHERE e.company_id IS ? AND p.pay_year=? AND p.pay_month=? AND p.status='confirmed'""",
        (emp_book(book), year, month),
    ).fetchone()["at"]

    accounts = conn.execute(
        "SELECT id, bank, account_no, name FROM bank_accounts WHERE company_id=? AND is_active=1 ORDER BY bank",
        (book,),
    ).fetchall()

    pm = month - 1 if month > 1 else 12
    py_ = year if month > 1 else year - 1
    nm = month + 1 if month < 12 else 1
    ny = year if month < 12 else year + 1

    return templates.TemplateResponse(
        request=request, name="hr/payroll_transfer.html", context={
            "request": request,
            "page_title": "급여 이체",
            "user_name": user["name"],
            "user_role": user["role"],
            "rows": rows,
            "year": year,
            "month": month,
            "book_name": book_label(conn, book),
            "total_count": len(rows),
            "missing_cnt": missing_cnt,
            "code_issue_cnt": code_issue_cnt,
            "last_at": last_at,
            "accounts": accounts,
            "pm": pm, "py": py_, "nm": nm, "ny": ny,
        }
    )


@router.get("/payroll/transfer/csv")
async def payroll_transfer_csv(
    request: Request,
    account_id: int = 0,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    reason: str = Depends(audit.require_reason),
):
    year, month = _payroll_month(request)
    book = user["book"]

    if account_id and not conn.execute(
        "SELECT 1 FROM bank_accounts WHERE id=? AND company_id=?", (account_id, book)
    ).fetchone():
        raise HTTPException(status_code=404)

    rows = [r for r in _transfer_rows(conn, book, year, month) if not r["missing"]]

    if book == INTERNAL:
        deposit_name = company_info(conn, book)["name"]
    else:
        deposit_name = conn.execute("SELECT name FROM client_companies WHERE id=?", (book,)).fetchone()["name"]
    withdraw_name = f"{year}년 {month}월 급여"

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["입금은행코드", "입금계좌번호", "입금액", "예금주", "입금통장표시", "출금통장표시"])
    for r in rows:
        w.writerow([esc_cell(v) for v in [
            r["code"], re.sub(r"\D", "", r["bank_account"]), r["net_pay"], r["name"], deposit_name, withdraw_name,
        ]])

    now = now_kst().strftime("%Y-%m-%d %H:%M:%S")
    for r in rows:
        conn.execute("UPDATE payslips SET transfer_file_at=? WHERE id=?", (now, r["payslip_id"]))
    total = sum(r["net_pay"] for r in rows)
    audit.log(conn, user["id"], "payroll_transfer", "payslip", None,
              f"{year}년 {month}월 {len(rows)}명 합계 {total:,}", company_id=book)
    audit.log_pii_export(conn, user, request, reason, "급여 이체파일", len(rows), company_id=book)

    buf.seek(0)
    return Response(
        content=buf.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f"attachment; filename=payroll_transfer_{year}{month:02d}.csv"},
    )


# 고용장려금 산정

SEVERE_GRADE_KEYWORDS = ("중증", "심한", "1급", "2급", "3급")
RATE_ORDER = ("경증남", "경증여", "중증남", "중증여")


def _is_severe_grade(grade):
    g = grade or ""
    return any(k in g for k in SEVERE_GRADE_KEYWORDS)


def _rno_gender(resident_no):
    digits = re.sub(r"\D", "", resident_no or "")
    if len(digits) < 7:
        return None
    return "남" if digits[6] in "13579" else "여"


def _base_days(conn, emp_id, year, month):
    # 임금지급 기초일수(실근로일+유급휴가). 근태 미기록 근로자는 판정 보류(None)
    ym = f"{year}-{month:02d}"
    has_record = conn.execute(
        "SELECT 1 FROM attendance WHERE employee_id=? AND work_date LIKE ? LIMIT 1",
        (emp_id, f"{ym}%"),
    ).fetchone()
    if not has_record:
        return None
    stats = get_month_work_stats(conn, emp_id, ym)
    leaves = get_worker_month_leaves(conn, emp_id, ym)
    return stats["work_days"] + len(leaves)


def _is_regular_worker(hours, severe, base_days):
    # 기초일수 16일 미만은 적용제외. 소정시간 60시간 미만도 적용제외하되 중증장애인은 예외
    if base_days is None or base_days < 16:
        return False
    if hours < 60 and not severe:
        return False
    return True


def calc_subsidy(conn, book, year, month):
    sector = "민간"
    if book != INTERNAL:
        row = conn.execute("SELECT sector FROM client_companies WHERE id=?", (book,)).fetchone()
        sector = (row["sector"] if row and row["sector"] else "민간")

    first_day = f"{year}-{month:02d}-01"
    _, last_d = calendar.monthrange(year, month)
    last_day = f"{year}-{month:02d}-{last_d:02d}"

    emps = [dict(e) for e in conn.execute(
        """SELECT id, name, hire_date, resign_date, monthly_work_hours,
                  disability_type, disability_grade, resident_no, base_salary
           FROM employees WHERE company_id IS ?""",
        (emp_book(book),),
    ).fetchall()]

    present = [e for e in emps
               if (not e["hire_date"] or e["hire_date"] <= last_day)
               and (not e["resign_date"] or e["resign_date"] >= first_day)]

    no_hours_cnt = sum(1 for e in present if not e["monthly_work_hours"])

    for e in present:
        e["severe"] = _is_severe_grade(e["disability_grade"]) if e["disability_type"] else False
        e["hours"] = e["monthly_work_hours"] or 0
        e["base_days"] = _base_days(conn, e["id"], year, month)

    all_cnt = len(present)
    pending = [e for e in present if e["base_days"] is None]
    checked = [e for e in present if e["base_days"] is not None]
    regular = [e for e in checked if _is_regular_worker(e["hours"], e["severe"], e["base_days"])]
    regular_cnt = len(regular)
    pending_cnt = len(pending)
    excluded_cnt = len(checked) - regular_cnt

    # 부동소수 오차 방지. 천분율 정수로 변환 후 버림(의무고용인원)/올림(장려금기준인원)
    permille = round(QUOTA[sector] * 10)
    quota_cnt = regular_cnt * permille // 1000
    reward_base_cnt = -(-(regular_cnt * permille) // 1000)

    rate_year = year if year in SUBSIDY_RATES else max(SUBSIDY_RATES)
    rates = SUBSIDY_RATES[rate_year]

    rows = []
    for e in regular:
        if not e["disability_type"]:
            continue
        severe = e["severe"]
        gender = _rno_gender(dec_rno(e["resident_no"]))
        ps = conn.execute(
            "SELECT gross_pay FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=? ORDER BY id DESC LIMIT 1",
            (e["id"], year, month),
        ).fetchone()
        wage = ps["gross_pay"] if ps else (e["base_salary"] or 0)
        if gender is None:
            rows.append({"name": e["name"], "severe": severe, "gender": None,
                         "monthly_hours": e["hours"], "wage": wage,
                         "unit_price": 0, "category": "제외(주민번호 없음)", "capped_amount": 0})
            continue
        rows.append({"name": e["name"], "severe": severe, "gender": gender,
                     "monthly_hours": e["hours"], "wage": wage,
                     "unit_price": rates[("중증" if severe else "경증") + gender],
                     "category": None, "capped_amount": 0})

    # 단가 낮은 순으로 의무고용 먼저 충당해야 고단가 대상이 장려금으로 남음
    eligible = [r for r in rows if r["gender"] is not None]
    eligible.sort(key=lambda r: (RATE_ORDER.index(("중증" if r["severe"] else "경증") + r["gender"]), r["wage"]))

    remain = quota_cnt
    for r in eligible:
        if remain > 0:
            # 잔여가 중증 1인분보다 적어도 분할하지 않고 통째로 충당 처리
            r["category"] = "의무고용 충당"
            remain = max(0, remain - (2 if r["severe"] else 1))
        else:
            r["category"] = "장려금 대상"

    for r in rows:
        if r["category"] == "장려금 대상":
            r["capped_amount"] = min(r["unit_price"], r["wage"] * 60 // 100)

    target_rows = [r for r in rows if r["category"] == "장려금 대상"]
    excess_cnt = sum(2 if r["severe"] else 1 for r in target_rows)
    subsidy_total = sum(r["capped_amount"] for r in target_rows)

    # 별지 제15호서식 ⑱~⑲ 장애인 상시근로자 구분(2배수 가중 전 headcount)
    all_disabled_cnt = sum(1 for e in regular if e["disability_type"])
    severe_ge60_cnt = sum(1 for e in regular if e["disability_type"] and e["severe"] and e["hours"] >= 60)
    severe_lt60_cnt = sum(1 for e in regular if e["disability_type"] and e["severe"] and e["hours"] < 60)
    mild_cnt_form = all_disabled_cnt - severe_ge60_cnt - severe_lt60_cnt

    # ㉑ 고용장려금 제외인원(타 장려금 중복수급)은 시스템에서 추적하지 않아 0으로 고정
    exempt_cnt = 0
    pay_cnt = max(0, all_disabled_cnt - exempt_cnt - reward_base_cnt)
    shortfall = all_disabled_cnt + severe_ge60_cnt - quota_cnt

    return {
        "sector": sector, "rate_year": rate_year,
        "all_cnt": all_cnt, "excluded_cnt": excluded_cnt,
        "regular_cnt": regular_cnt, "quota_cnt": quota_cnt, "reward_base_cnt": reward_base_cnt,
        "excess_cnt": excess_cnt,
        "sev_m": sum(1 for r in rows if r["severe"] and r["gender"] == "남"),
        "sev_f": sum(1 for r in rows if r["severe"] and r["gender"] == "여"),
        "mild_m": sum(1 for r in rows if not r["severe"] and r["gender"] == "남"),
        "mild_f": sum(1 for r in rows if not r["severe"] and r["gender"] == "여"),
        "all_disabled_cnt": all_disabled_cnt, "mild_cnt_form": mild_cnt_form,
        "severe_ge60_cnt": severe_ge60_cnt, "severe_lt60_cnt": severe_lt60_cnt,
        "exempt_cnt": exempt_cnt, "pay_cnt": pay_cnt, "shortfall": shortfall,
        "subsidy_total": subsidy_total, "no_hours_cnt": no_hours_cnt,
        "over100": regular_cnt >= 100, "rows": rows,
        "pending_cnt": pending_cnt, "pending_workers": [e["name"] for e in pending],
    }


@router.get("/hr/subsidy", response_class=HTMLResponse)
async def subsidy_page(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    year, month = _payroll_month(request)
    book = user["book"]
    result = calc_subsidy(conn, book, year, month)

    pm = month - 1 if month > 1 else 12
    py_ = year if month > 1 else year - 1
    nm = month + 1 if month < 12 else 1
    ny = year if month < 12 else year + 1

    return templates.TemplateResponse(
        request=request, name="hr/subsidy.html", context={
            "request": request,
            "page_title": "고용장려금 산정",
            "user_name": user["name"],
            "user_role": user["role"],
            "book_name": book_label(conn, book),
            "year": year, "month": month,
            "quota_pct": QUOTA[result["sector"]],
            "result": result,
            "pm": pm, "py": py_, "nm": nm, "ny": ny,
        }
    )


@router.get("/hr/subsidy/excel")
async def subsidy_excel(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    reason: str = Depends(audit.require_reason),
):
    year, month = _payroll_month(request)
    book = user["book"]
    result = calc_subsidy(conn, book, year, month)

    headers = ["성명", "중증여부", "성별", "월소정시간", "월임금", "단가", "60%상한적용액", "구분"]
    rows = [[
        r["name"], "중증" if r["severe"] else "경증", r["gender"] or "-",
        r["monthly_hours"], r["wage"], r["unit_price"], r["capped_amount"], r["category"],
    ] for r in result["rows"]]
    wb = styled_workbook("고용장려금", headers, rows, [10, 8, 6, 10, 12, 12, 14, 16])
    filename = f"고용장려금산정_{year}{month:02d}.xlsx"
    audit.log_pii_export(conn, user, request, reason, "고용장려금 엑셀", len(rows), company_id=book or None)
    return excel_response(wb, filename)


# 장애인 고용계획 및 실시상황 보고서 (장애인고용촉진법 시행규칙 별지 제15호서식)

def calc_disability_report(conn, book, year):
    months = []
    sector = "민간"
    pending_workers = set()
    for m in range(1, 13):
        r = calc_subsidy(conn, book, year, m)
        sector = r["sector"]
        disabled_cnt = r["all_disabled_cnt"] + r["severe_ge60_cnt"]
        rate = round(disabled_cnt / r["regular_cnt"] * 100, 2) if r["regular_cnt"] else 0.0
        pending_workers.update(r["pending_workers"])
        months.append({
            "month": m,
            "all_cnt": r["all_cnt"], "excluded_cnt": r["excluded_cnt"],
            "regular_cnt": r["regular_cnt"], "quota_cnt": r["quota_cnt"],
            "reward_base_cnt": r["reward_base_cnt"],
            "all_disabled_cnt": r["all_disabled_cnt"], "mild_cnt_form": r["mild_cnt_form"],
            "severe_ge60_cnt": r["severe_ge60_cnt"], "severe_lt60_cnt": r["severe_lt60_cnt"],
            "shortfall": r["shortfall"],
            "shortfall_unemployed": r["shortfall"] if r["all_disabled_cnt"] == 0 else None,
            "shortfall_employed": r["shortfall"] if r["all_disabled_cnt"] > 0 else None,
            "exempt_cnt": r["exempt_cnt"], "pay_cnt": r["pay_cnt"],
            "disabled_cnt": disabled_cnt, "gap": disabled_cnt - r["quota_cnt"], "rate": rate,
            "pending_cnt": r["pending_cnt"], "pending_workers": r["pending_workers"],
        })
    avg_regular = round(sum(x["regular_cnt"] for x in months) / 12, 1)
    avg_disabled = round(sum(x["disabled_cnt"] for x in months) / 12, 1)
    return {
        "months": months, "sector": sector, "quota_pct": QUOTA.get(sector, QUOTA["민간"]),
        "avg_regular": avg_regular, "avg_disabled": avg_disabled,
        "avg_rate": round(avg_disabled / avg_regular * 100, 2) if avg_regular else 0.0,
        "pending_workers": sorted(pending_workers),
    }


# 첨부서류 - 장애인 근로자 명부 (별지 제15호서식 첨부 양식)

def calc_disability_roster(conn, book, year=None):
    info = company_info(conn, book)
    sql = """SELECT name, resident_no, disability_class, disability_type, veteran_grade,
                     disability_grade, disability_date, hire_date, resign_date,
                     job_title, job_code, base_salary, monthly_work_hours
              FROM employees
              WHERE company_id IS ? AND disability_type IS NOT NULL AND disability_type != ''"""
    params = [emp_book(book)]
    if year:
        sql += " AND (hire_date IS NULL OR hire_date <= ?) AND (resign_date IS NULL OR resign_date = '' OR resign_date >= ?)"
        params += [f"{year}-12-31", f"{year}-01-01"]
    sql += " ORDER BY employee_no, name"
    emps = conn.execute(sql, params).fetchall()

    rows = []
    for e in emps:
        severe = _is_severe_grade(e["disability_grade"])
        if severe:
            double_label = "60시간 이상" if (e["monthly_work_hours"] or 0) >= 60 else "60시간 미만"
        else:
            double_label = "-"
        rows.append({
            "biz_name": info["name"], "biz_no": info["biz_no"],
            "name": e["name"], "resident_no": dec_rno(e["resident_no"]) or "",
            "disability_class": e["disability_class"] or "",
            "disability_type": DISABILITY_TYPE_CODES.get(e["disability_type"], ""),
            "veteran_grade": e["veteran_grade"] or "",
            "severe_label": "중증장애인" if severe else "경증장애인",
            "double_label": double_label,
            "disability_date": e["disability_date"] or "",
            "hire_date": e["hire_date"] or "",
            "resign_date": e["resign_date"] or "",
            "job_code": e["job_code"] or "",
            "wage": e["base_salary"] or 0,
            "subsidy_period": "",  # 시스템에 데이터 없음. 담당자 직접 기재
        })
    return rows


@router.get("/hr/disability-report", response_class=HTMLResponse)
async def disability_report_page(
    request: Request,
    year: int = 0,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    year = year or now_kst().year
    book = user["book"]
    result = calc_disability_report(conn, book, year)

    return templates.TemplateResponse(
        request=request, name="hr/disability_report.html", context={
            "request": request,
            "page_title": "장애인 고용현황 신고 자료",
            "user_name": user["name"],
            "user_role": user["role"],
            "book_name": book_label(conn, book),
            "year": year,
            "result": result,
        }
    )


@router.get("/hr/disability-report/excel")
async def disability_report_excel(
    request: Request,
    year: int = 0,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    reason: str = Depends(audit.require_reason),
):
    year = year or now_kst().year
    book = user["book"]
    result = calc_disability_report(conn, book, year)

    headers = [
        "월", "전체근로자수", "적용제외근로자수", "상시근로자수", "의무고용인원", "고용장려금기준인원",
        "장애인근로자수(전체)", "장애인근로자수(경증)", "중증장애인(60시간이상)", "중증장애인(60시간미만)",
        "미달고용인원(미고용월)", "미달고용인원(고용월)", "고용장려금제외인원(담당자 기재)", "고용장려금지급인원",
    ]
    rows = [[
        f"{m['month']}월", m["all_cnt"], m["excluded_cnt"], m["regular_cnt"], m["quota_cnt"], m["reward_base_cnt"],
        m["all_disabled_cnt"], m["mild_cnt_form"], m["severe_ge60_cnt"], m["severe_lt60_cnt"],
        m["shortfall_unemployed"] if m["shortfall_unemployed"] is not None else "",
        m["shortfall_employed"] if m["shortfall_employed"] is not None else "",
        "", m["pay_cnt"],
    ] for m in result["months"]]
    wb = styled_workbook("장애인고용현황", headers, rows, [8, 11, 12, 11, 11, 13, 13, 13, 15, 15, 14, 12, 13, 13])

    roster = calc_disability_roster(conn, book)
    roster_headers = [
        "연번", "사업장명", "사업자등록번호", "성명", "주민등록번호", "장애인정구분", "장애유형", "상이등급",
        "중증(경증)여부", "중증2배수인정여부", "장애인정일", "입사일", "퇴사일", "근무직무", "임금",
        "장려금 지급받은 기간(담당자 기재)",
    ]
    roster_rows = [[
        i, r["biz_name"], r["biz_no"], r["name"], mask_rno(r["resident_no"]), r["disability_class"], r["disability_type"],
        r["veteran_grade"], r["severe_label"], r["double_label"], r["disability_date"], r["hire_date"],
        r["resign_date"], r["job_code"], r["wage"], r["subsidy_period"],
    ] for i, r in enumerate(roster, 1)]
    styled_sheet(wb, "장애인근로자명부", roster_headers, roster_rows,
                 [6, 16, 14, 10, 14, 12, 10, 10, 12, 14, 12, 12, 12, 12, 12, 18])

    filename = f"장애인고용현황_{year}.xlsx"
    audit.log_pii_export(conn, user, request, reason, "장애인고용 신고 엑셀", len(roster_rows), company_id=book or None)
    return excel_response(wb, filename)


# 공단 제출용 원본은 서버에 파일로 남기지 않음. 다중 워커 대응 위해 세션에 토큰과 암호화한 파일 비밀번호만 두고 다운로드 시 재생성
SUBMISSION_TTL = 300


def _roster_xlsx(roster):
    headers = [
        "연번", "사업장명", "사업자등록번호", "성명", "주민등록번호", "장애인정구분", "장애유형", "상이등급",
        "중증(경증)여부", "중증2배수인정여부", "장애인정일", "입사일", "퇴사일", "근무직무", "임금",
        "장려금 지급받은 기간(담당자 기재)",
    ]
    rows = [[
        i, r["biz_name"], r["biz_no"], r["name"], r["resident_no"], r["disability_class"], r["disability_type"],
        r["veteran_grade"], r["severe_label"], r["double_label"], r["disability_date"], r["hire_date"],
        r["resign_date"], r["job_code"], r["wage"], r["subsidy_period"],
    ] for i, r in enumerate(roster, 1)]
    return styled_workbook("장애인근로자명부(원본)", headers, rows,
                            [6, 16, 14, 10, 14, 12, 10, 10, 12, 14, 12, 12, 12, 12, 12, 18])


@router.get("/hr/disability-report/submission", response_class=HTMLResponse)
async def disability_submission_form(
    request: Request,
    year: int = 0,
    user: dict = Depends(require_hr_chief),
    conn=Depends(get_db),
):
    year = year or now_kst().year
    return templates.TemplateResponse(
        request=request, name="hr/disability_submission.html", context={
            "request": request,
            "page_title": "공단 제출용 장애인 근로자명부",
            "user_name": user["name"],
            "year": year,
        }
    )


@router.post("/hr/disability-report/submission", response_class=HTMLResponse)
async def disability_submission_issue(
    request: Request,
    year: int = Form(...),
    reason: str = Form(...),
    password: str = Form(...),
    user: dict = Depends(require_hr_chief),
    conn=Depends(get_db),
):
    if not reason.strip():
        return templates.TemplateResponse(
            request=request, name="hr/disability_submission.html", context={
                "request": request, "page_title": "공단 제출용 장애인 근로자명부",
                "user_name": user["name"], "year": year, "error": "사유를 입력하세요.",
            }
        )
    row = conn.execute("SELECT password FROM users WHERE id=?", (user["id"],)).fetchone()
    if not row or not verify_pw(password, row["password"]):
        return templates.TemplateResponse(
            request=request, name="hr/disability_submission.html", context={
                "request": request, "page_title": "공단 제출용 장애인 근로자명부",
                "user_name": user["name"], "year": year, "error": "비밀번호가 일치하지 않습니다.",
            }
        )

    book = user["book"]
    roster = calc_disability_roster(conn, book, year=year)
    file_password = secrets.token_urlsafe(16)
    token = secrets.token_hex(16)
    request.session["pii_submission"] = {
        "token": token, "pw": enc_rno(file_password), "year": year, "exp": time.time() + SUBMISSION_TTL,
    }

    audit.log_pii_export(conn, user, request, reason, "장애인 근로자명부 공단제출용(원본)",
                          len(roster), company_id=book, action="download")

    chiefs = conn.execute(
        """SELECT u.id FROM client_users cu JOIN users u ON u.id = cu.user_id
           WHERE cu.company_id=? AND cu.client_role='all'""",
        (book,),
    ).fetchall()
    chief_ids = [r["id"] for r in chiefs]
    if chief_ids:
        create_bulk_notifications(
            conn, chief_ids, "장애인 근로자명부 원본 다운로드",
            f"{user['name']}님이 {year}년 장애인 근로자명부 원본({len(roster)}명분)을 다운로드했습니다. 사유: {reason}",
            "/hr/disability-report",
        )

    return templates.TemplateResponse(
        request=request, name="hr/disability_submission.html", context={
            "request": request, "page_title": "공단 제출용 장애인 근로자명부",
            "user_name": user["name"], "year": year,
            "issued": True, "file_password": file_password, "token": token, "count": len(roster),
        }
    )


@router.get("/hr/disability-report/submission/download")
async def disability_submission_download(
    request: Request,
    user: dict = Depends(require_hr_chief),
    conn=Depends(get_db),
):
    token = request.query_params.get("token", "")
    sub = request.session.pop("pii_submission", None)
    if not token or not sub or token != sub["token"]:
        raise HTTPException(status_code=403, detail="다운로드 권한이 없습니다.")
    if sub["exp"] < time.time():
        raise HTTPException(status_code=400, detail="만료된 다운로드입니다. 다시 발급하세요.")

    buf = io.BytesIO()
    _roster_xlsx(calc_disability_roster(conn, user["book"], year=sub["year"])).save(buf)
    buf.seek(0)
    enc_buf = io.BytesIO()
    msoffcrypto.OfficeFile(buf).encrypt(dec_rno(sub["pw"]), enc_buf)
    enc_buf.seek(0)
    filename = f"장애인근로자명부_공단제출용_{sub['year']}.xlsx"
    return StreamingResponse(
        enc_buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# 급여 관리 (근로자별 명세 입력/확정)

@router.get("/payroll", response_class=HTMLResponse)
async def payroll_list(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    year, month = _payroll_month(request)
    book = user["book"]

    workers = conn.execute(
        "SELECT id, name, employee_no, dept, position, monthly_work_hours FROM employees "
        "WHERE company_id IS ? AND (status='active' OR resign_date >= ?)",
        (emp_book(book), f"{year:04d}-{month:02d}-01"),
    ).fetchall()
    workers = [dict(w) for w in workers]

    for w in workers:
        ps = conn.execute(
            "SELECT id, status, gross_pay, total_deduction, net_pay, paid_at FROM payslips "
            "WHERE employee_id=? AND pay_year=? AND pay_month=?",
            (w["id"], year, month),
        ).fetchone()
        if ps:
            w["payslip_status"] = ps["status"]
            w["gross_pay"] = ps["gross_pay"]
            w["net_pay"] = ps["net_pay"]
            w["total_deduction"] = ps["total_deduction"]
            w["paid_at"] = ps["paid_at"]
            ordinary_total = conn.execute(
                """SELECT COALESCE(SUM(pi.amount),0) FROM payslip_items pi
                   JOIN pay_items it ON it.id=pi.item_id
                   WHERE pi.payslip_id=? AND it.ordinary=1""",
                (ps["id"],),
            ).fetchone()[0]
            w["below_min_wage"] = below_min_wage(ordinary_total, w["monthly_work_hours"], year)
        else:
            w["payslip_status"] = None

    has_draft = any(w["payslip_status"] == "draft" for w in workers)
    all_confirmed = all(w["payslip_status"] == "confirmed" for w in workers if w["payslip_status"])
    has_unpaid = any(w["payslip_status"] == "confirmed" and not w.get("paid_at") for w in workers)
    min_wage_cnt = sum(1 for w in workers if w.get("below_min_wage"))
    confirm_msg = f"{year}년 {month:02d}월 급여를 확정합니다. 확정 후에는 수정할 수 없습니다."
    if min_wage_cnt:
        confirm_msg += f" 최저임금 미달 {min_wage_cnt}명"

    emp_ids = [w["id"] for w in workers]
    pay_batches = []
    if emp_ids:
        placeholders = ",".join("?" * len(emp_ids))
        pay_batches = conn.execute(
            f"""SELECT pay_doc_id AS doc_id, COUNT(*) AS cnt, SUM(net_pay) AS total, MIN(paid_at) AS paid_at
                FROM payslips WHERE employee_id IN ({placeholders}) AND pay_year=? AND pay_month=?
                AND pay_doc_id IS NOT NULL GROUP BY pay_doc_id""",
            (*emp_ids, year, month),
        ).fetchall()
        pay_batches = [dict(b) for b in pay_batches]

    # 원 내부(book 0)는 근태 마감 대상이 아닌 wone 자체 서비스. 항상 마감된 것으로 취급
    att_closed = True if book == INTERNAL else is_month_closed(conn, year, month, book)

    accounts = conn.execute(
        "SELECT id, bank, account_no, name FROM bank_accounts WHERE company_id=? AND is_active=1 ORDER BY bank",
        (book,),
    ).fetchall()

    pm = month - 1 if month > 1 else 12
    py_ = year if month > 1 else year - 1
    nm = month + 1 if month < 12 else 1
    ny = year if month < 12 else year + 1

    return templates.TemplateResponse(
        request=request, name="client/payroll.html", context={
            "request": request,
            "page_title": "급여 관리",
            "user_name": user["name"],
            "user_role": user["role"],
            "workers": workers,
            "year": year,
            "month": month,
            "book_name": book_label(conn, book),
            "has_draft": has_draft,
            "all_confirmed": all_confirmed,
            "has_unpaid": has_unpaid,
            "min_wage_cnt": min_wage_cnt,
            "confirm_msg": confirm_msg,
            "pay_batches": pay_batches,
            "att_closed": att_closed,
            "accounts": accounts,
            "default_pay_date": _pay_date(conn, book, year, month),
            "pm": pm, "py": py_, "nm": nm, "ny": ny,
            "error": request.query_params.get("error", ""),
        }
    )


# 급여대장

def _load_ledger(conn, year, month, book, filter_disabled=False):
    first_day = f"{year}-{month:02d}-01"
    _, last_d = calendar.monthrange(year, month)
    last_day = f"{year}-{month:02d}-{last_d:02d}"

    emp_sql = """
        SELECT e.id, e.name, e.employee_no, e.resident_no, e.hire_date, e.job_title,
               e.monthly_work_hours, e.disability_type, e.disability_grade,
               p.id AS payslip_id,
               p.gross_pay, p.national_pension, p.health_insurance,
               p.employment_insurance AS ei_deduction, p.income_tax, p.resident_tax,
               p.total_deduction, p.net_pay, p.status AS pay_status,
               p.employer_pension, p.employer_health, p.employer_employment, p.employer_accident
        FROM employees e
        LEFT JOIN payslips p ON e.id = p.employee_id AND p.pay_year=? AND p.pay_month=?
        WHERE (e.status = 'active' OR e.resign_date >= ?) AND e.company_id IS ?
    """
    params = [year, month, first_day, emp_book(book)]

    if filter_disabled:
        emp_sql += " AND e.disability_type IS NOT NULL AND e.disability_type != ''"

    emp_sql += " ORDER BY e.employee_no, e.name"
    rows = conn.execute(emp_sql, params).fetchall()
    rows = [dict(r) for r in rows]
    for r in rows:
        r["resident_no"] = dec_rno(r["resident_no"])

    if not rows:
        return rows, {}

    payslip_ids = [r["payslip_id"] for r in rows if r["payslip_id"]]
    item_amounts = {}
    if payslip_ids:
        ph_p = ",".join("?" * len(payslip_ids))
        item_rows = conn.execute(
            f"""SELECT pi.payslip_id, it.code, pi.amount
                FROM payslip_items pi JOIN pay_items it ON it.id = pi.item_id
                WHERE pi.payslip_id IN ({ph_p})""",
            payslip_ids,
        ).fetchall()
        for r in item_rows:
            item_amounts.setdefault(r["payslip_id"], {})[r["code"]] = r["amount"]

    for r in rows:
        amounts = item_amounts.get(r["payslip_id"], {})
        r["base_salary"] = amounts.get("BASE", 0)
        r["overtime_pay"] = sum(amounts.get(c, 0) for c in ("OT", "NIGHT", "HOLIDAY"))
        r["disability_allowance"] = amounts.get("DISABLE", 0)
        r["meal_allowance"] = amounts.get("MEAL", 0)

    emp_ids = [r["id"] for r in rows]
    ph = ",".join("?" * len(emp_ids))
    att_rows = conn.execute(
        f"""SELECT employee_id,
                   COUNT(*) AS work_days,
                   SUM(work_minutes) AS total_work_mins,
                   SUM(overtime_minutes) AS total_ot_mins
            FROM attendance
            WHERE work_date BETWEEN ? AND ?
              AND status IN ('normal','late','early_leave','late_early')
              AND employee_id IN ({ph})
            GROUP BY employee_id""",
        [first_day, last_day, *emp_ids],
    ).fetchall()

    att_map = {r["employee_id"]: dict(r) for r in att_rows}

    year_month = f"{year}-{month:02d}"
    for eid in emp_ids:
        stats = get_month_work_stats(conn, eid, year_month)
        att_map.setdefault(eid, {})
        att_map[eid]["night_mins"] = stats["night_mins"]
        att_map[eid]["holiday_mins"] = stats["holiday_mins"]

    return rows, att_map


@router.get("/payroll/ledger", response_class=HTMLResponse)
async def payroll_ledger(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    filter_val = request.query_params.get("filter", "all")
    filter_disabled = filter_val == "disabled"
    book = user["book"]

    rows, att_map = _load_ledger(conn, year, month, book, filter_disabled)

    for r in rows:
        att = att_map.get(r["id"], {})
        r["work_days"] = att.get("work_days") or 0
        r["total_work_mins"] = att.get("total_work_mins") or 0
        r["total_ot_mins"] = att.get("total_ot_mins") or 0
        r["work_hours"] = r["total_work_mins"] // 60
        r["ot_hours"] = r["total_ot_mins"] // 60
        r["night_hours"] = att.get("night_mins", 0) // 60
        r["holiday_hours"] = att.get("holiday_mins", 0) // 60
        r["resident_no_masked"] = mask_rno(r.get("resident_no"))

    total_count = len(rows)
    sum_base = sum(r["base_salary"] or 0 for r in rows)
    sum_ot = sum(r["overtime_pay"] or 0 for r in rows)
    sum_dis = sum(r["disability_allowance"] or 0 for r in rows)
    sum_meal = sum(r["meal_allowance"] or 0 for r in rows)
    sum_gross = sum(r["gross_pay"] or 0 for r in rows)
    sum_pension = sum(r["national_pension"] or 0 for r in rows)
    sum_health = sum(r["health_insurance"] or 0 for r in rows)
    sum_ei = sum(r["ei_deduction"] or 0 for r in rows)
    sum_itax = sum(r["income_tax"] or 0 for r in rows)
    sum_rtax = sum(r["resident_tax"] or 0 for r in rows)
    sum_deduction = sum(r["total_deduction"] or 0 for r in rows)
    sum_net = sum(r["net_pay"] or 0 for r in rows)
    sum_er_pension = sum(r["employer_pension"] or 0 for r in rows)
    sum_er_health = sum(r["employer_health"] or 0 for r in rows)
    sum_er_employment = sum(r["employer_employment"] or 0 for r in rows)
    sum_er_accident = sum(r["employer_accident"] or 0 for r in rows)
    sum_er_total = sum_er_pension + sum_er_health + sum_er_employment + sum_er_accident

    return templates.TemplateResponse(
        request=request,
        name="hr/payroll_ledger.html",
        context={
            "request": request,
            "page_title": "급여대장",
            "user_name": user["name"],
            "user_role": user["role"],
            "rows": rows,
            "year": year,
            "month": month,
            "book_name": book_label(conn, book),
            "filter": filter_val,
            "total_count": total_count,
            "sum_base": sum_base,
            "sum_ot": sum_ot,
            "sum_dis": sum_dis,
            "sum_meal": sum_meal,
            "sum_gross": sum_gross,
            "sum_pension": sum_pension,
            "sum_health": sum_health,
            "sum_ei": sum_ei,
            "sum_itax": sum_itax,
            "sum_rtax": sum_rtax,
            "sum_deduction": sum_deduction,
            "sum_net": sum_net,
            "sum_er_pension": sum_er_pension,
            "sum_er_health": sum_er_health,
            "sum_er_employment": sum_er_employment,
            "sum_er_accident": sum_er_accident,
            "sum_er_total": sum_er_total,
        },
    )


@router.get("/payroll/ledger/export")
async def payroll_ledger_export(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    reason: str = Depends(audit.require_reason),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    filter_val = request.query_params.get("filter", "all")
    filter_disabled = filter_val == "disabled"
    book = user["book"]

    rows, att_map = _load_ledger(conn, year, month, book, filter_disabled)

    for r in rows:
        att = att_map.get(r["id"], {})
        r["work_days"] = att.get("work_days") or 0
        r["total_work_mins"] = att.get("total_work_mins") or 0
        r["total_ot_mins"] = att.get("total_ot_mins") or 0
        r["work_hours"] = r["total_work_mins"] // 60
        r["ot_hours"] = r["total_ot_mins"] // 60
        r["night_hours"] = att.get("night_mins", 0) // 60
        r["holiday_hours"] = att.get("holiday_mins", 0) // 60

    book_name = book_label(conn, book)

    wb = Workbook()
    ws = wb.active
    ws.title = "급여대장"

    title_text = f"급여대장 - {year}년 {month}월 ({book_name})"

    headers = [
        "순번", "근로자명", "주민등록번호", "고용연월일", "종사업무",
        "소정근로시간", "근로일수", "근로시간수", "연장근로시간", "야간근로시간", "휴일근로시간",
        "기본급", "초과근무수당", "장애수당", "식대", "지급합계",
        "국민연금", "건강보험", "고용보험", "소득세", "지방소득세",
        "공제합계", "실지급액",
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="맑은 고딕", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    money_cols = set(range(12, len(headers) + 1))

    for ri, r in enumerate(rows, 3):
        vals = [
            ri - 2,
            esc_cell(r["name"]),
            esc_cell(mask_rno(r.get("resident_no"))),
            esc_cell(r.get("hire_date") or ""),
            esc_cell(r.get("job_title") or ""),
            r.get("monthly_work_hours") or 0,
            r["work_days"],
            r["work_hours"],
            r["ot_hours"],
            r["night_hours"],
            r["holiday_hours"],
            r.get("base_salary") or 0,
            r.get("overtime_pay") or 0,
            r.get("disability_allowance") or 0,
            r.get("meal_allowance") or 0,
            r.get("gross_pay") or 0,
            r.get("national_pension") or 0,
            r.get("health_insurance") or 0,
            r.get("ei_deduction") or 0,
            r.get("income_tax") or 0,
            r.get("resident_tax") or 0,
            r.get("total_deduction") or 0,
            r.get("net_pay") or 0,
        ]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            if ci in money_cols and isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="center")

    # 합계행
    sum_row = len(rows) + 3
    sum_vals = [
        "", "합계", "", "", "", "", "", "", "", "", "",
        sum(r.get("base_salary") or 0 for r in rows),
        sum(r.get("overtime_pay") or 0 for r in rows),
        sum(r.get("disability_allowance") or 0 for r in rows),
        sum(r.get("meal_allowance") or 0 for r in rows),
        sum(r.get("gross_pay") or 0 for r in rows),
        sum(r.get("national_pension") or 0 for r in rows),
        sum(r.get("health_insurance") or 0 for r in rows),
        sum(r.get("ei_deduction") or 0 for r in rows),
        sum(r.get("income_tax") or 0 for r in rows),
        sum(r.get("resident_tax") or 0 for r in rows),
        sum(r.get("total_deduction") or 0 for r in rows),
        sum(r.get("net_pay") or 0 for r in rows),
    ]
    for ci, val in enumerate(sum_vals, 1):
        cell = ws.cell(row=sum_row, column=ci, value=val)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        if ci in money_cols and isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="center")

    col_widths = [6, 10, 16, 12, 12, 10, 8, 8, 8, 8, 8, 12, 12, 10, 10, 12, 10, 10, 10, 10, 10, 12, 12]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    filename = f"급여대장_{year}{month:02d}.xlsx"
    audit.log_pii_export(conn, user, request, reason, "급여대장 엑셀", len(rows), company_id=book or None)
    return excel_response(wb, filename)

def _item_rows_for_edit(conn, employee_id, book, year, month, worker, payslip):
    items = get_items(conn, book)
    five_plus = is_five_plus(conn, book)

    saved = {}
    if payslip:
        for r in conn.execute(
            "SELECT item_id, amount, overridden FROM payslip_items WHERE payslip_id=?", (payslip["id"],)
        ):
            saved[r["item_id"]] = {"amount": r["amount"], "overridden": bool(r["overridden"])}

    py, pm = _prev_month(year, month)
    prev_ps = conn.execute(
        "SELECT id FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (employee_id, py, pm),
    ).fetchone()
    prev_amounts = {}
    if prev_ps:
        for r in conn.execute(
            "SELECT item_id, amount FROM payslip_items WHERE payslip_id=?", (prev_ps["id"],)
        ):
            prev_amounts[r["item_id"]] = r["amount"]

    monthly_hours = worker.get("monthly_work_hours") or 0
    ordinary_total = 0
    for it in items:
        if not it["ordinary"]:
            continue
        if it["id"] in saved:
            ordinary_total += saved[it["id"]]["amount"]
        elif it["code"] == "BASE":
            ordinary_total += worker.get("base_salary") or 0
        else:
            ordinary_total += prev_amounts.get(it["id"], 0)
    hourly = calc_ordinary_hourly(ordinary_total, monthly_hours)

    stats = get_month_work_stats(conn, employee_id, f"{year}-{month:02d}")
    premium = calc_premium_pay(hourly, stats["ot_mins"], stats["night_mins"], stats["holiday_day_mins"], five_plus)

    # 퇴사월 한정. 미사용 연차수당을 ANNUAL 추천값으로 제시
    leave_days = None
    resign_date = worker.get("resign_date")
    if resign_date:
        rd = date.fromisoformat(resign_date)
        if rd.year == year and rd.month == month:
            leave_days = max(0, get_leave_balance(conn, employee_id, rd.year)["remaining"])
            leave_daily = calc_ordinary_daily(conn, employee_id, resign_date)

    rows = []
    for it in items:
        suggested = premium.get(it["code"]) if it["code"] in PREMIUM_CODES else None
        row_leave_days = None
        if it["code"] == "ANNUAL" and leave_days is not None:
            suggested = round(leave_days * leave_daily)
            row_leave_days = leave_days
        if it["id"] in saved:
            amount = saved[it["id"]]["amount"]
            overridden = saved[it["id"]]["overridden"]
        elif it["code"] == "BASE":
            amount = worker.get("base_salary") or 0
            overridden = False
        elif suggested is not None:
            amount = suggested
            overridden = False
        else:
            amount = prev_amounts.get(it["id"], 0)
            overridden = False
        rows.append({
            "item_id": it["id"], "code": it["code"], "name": it["name"],
            "amount": amount, "suggested": suggested, "overridden": overridden,
            "tax_free_limit": it["tax_free_limit"], "leave_days": row_leave_days,
        })
    return rows, five_plus, stats, ordinary_total


@router.get("/payroll/{emp_id}", response_class=HTMLResponse)
async def payroll_edit(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    year, month = _payroll_month(request)
    book = user["book"]
    worker = book_emp(conn, book, emp_id)

    payslip = conn.execute(
        "SELECT * FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (emp_id, year, month),
    ).fetchone()
    payslip = dict(payslip) if payslip else None

    is_locked = payslip and payslip.get("status") == "confirmed"

    item_rows, five_plus, stats, ordinary_total = _item_rows_for_edit(conn, emp_id, book, year, month, worker, payslip)
    below_min = below_min_wage(ordinary_total, worker.get("monthly_work_hours") or 0, year)

    att_closed = True if book == INTERNAL else is_month_closed(conn, year, month, book)

    pm = month - 1 if month > 1 else 12
    py_ = year if month > 1 else year - 1
    nm = month + 1 if month < 12 else 1
    ny = year if month < 12 else year + 1

    return templates.TemplateResponse(
        request=request, name="client/payroll_edit.html", context={
            "request": request,
            "page_title": f"{worker['name']} 급여 입력",
            "user_name": user["name"],
            "worker": worker,
            "payslip": payslip,
            "item_rows": item_rows,
            "five_plus": five_plus,
            "below_min_wage": below_min,
            "year": year,
            "month": month,
            "is_locked": is_locked,
            "att_closed": att_closed,
            "pm": pm, "py": py_, "nm": nm, "ny": ny,
        }
    )


@router.post("/payroll/{emp_id}/reopen")
async def payroll_reopen(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    year: int = Form(...),
    month: int = Form(...),
    reason: str = Form(...),
    password: str = Form(...),
):
    book = user["book"]
    worker = book_emp(conn, book, emp_id)

    row = conn.execute("SELECT password FROM users WHERE id=?", (user["id"],)).fetchone()
    if not row or not verify_pw(password, row["password"]):
        return JSONResponse({"error": "비밀번호가 일치하지 않습니다."}, status_code=400)
    if not reason.strip():
        return JSONResponse({"error": "취소 사유를 입력해 주세요."}, status_code=400)

    payslip = conn.execute(
        "SELECT * FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (emp_id, year, month),
    ).fetchone()
    if not payslip or payslip["status"] != "confirmed":
        return JSONResponse({"error": "확정된 급여가 아닙니다."}, status_code=400)
    if payslip["pay_doc_id"]:
        return JSONResponse({"error": "이미 지급 처리된 급여는 확정을 취소할 수 없습니다."}, status_code=400)
    if payslip["transfer_file_at"]:
        return JSONResponse({"error": "이체파일이 생성된 급여는 확정을 취소할 수 없습니다."}, status_code=400)

    doc = conn.execute(
        """SELECT * FROM erp_docs WHERE payslip_id=?
           AND NOT EXISTS (SELECT 1 FROM erp_docs r WHERE r.reversal_of=erp_docs.id)""",
        (payslip["id"],),
    ).fetchone()
    if doc:
        if books.period_closed(conn, book, doc["slip_date"]):
            return JSONResponse({"error": "마감된 회계기간의 급여는 확정을 취소할 수 없습니다."}, status_code=400)
        books.reverse(conn, doc["id"], user["id"], doc["slip_date"])

    conn.execute("UPDATE payslips SET status='draft', paid_at=NULL WHERE id=?", (payslip["id"],))
    audit.log(conn, user["id"], "payroll_reopen", "payslip", payslip["id"],
              f"{worker['name']} {year}년 {month}월 급여확정 취소: {reason.strip()}", company_id=book)
    conn.commit()

    return RedirectResponse(url=f"/payroll/{emp_id}?year={year}&month={month}", status_code=303)


@router.post("/payroll/confirm")
async def payroll_confirm(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    form = await request.form()
    year = int(form.get("year", now_kst().year))
    month = int(form.get("month", now_kst().month))
    book = user["book"]

    # 원 내부(book 0)는 근태 마감 대상이 아닌 wone 자체 서비스. 마감 체크 생략
    if book != INTERNAL and not is_month_closed(conn, year, month, book):
        return RedirectResponse(
            url=f"/payroll?year={year}&month={month}&error=att_not_closed",
            status_code=303,
        )

    emp_ids = [r["id"] for r in conn.execute(
        "SELECT id FROM employees WHERE company_id IS ? AND (status='active' OR resign_date >= ?)",
        (emp_book(book), f"{year:04d}-{month:02d}-01"),
    ).fetchall()]
    if emp_ids:
        placeholders = ",".join("?" * len(emp_ids))
        to_confirm = conn.execute(
            f"""SELECT p.id, e.name AS employee_name, e.user_id FROM payslips p JOIN employees e ON p.employee_id=e.id
                WHERE p.employee_id IN ({placeholders}) AND p.pay_year=? AND p.pay_month=? AND p.status='draft'""",
            (*emp_ids, year, month),
        ).fetchall()
        conn.execute(
            f"UPDATE payslips SET status='confirmed' WHERE employee_id IN ({placeholders}) AND pay_year=? AND pay_month=? AND status='draft'",
            (*emp_ids, year, month),
        )
        for ps_row in to_confirm:
            audit.log(conn, user["id"], "payroll_confirm", "payslip", ps_row["id"],
                      f"{ps_row['employee_name']} {year}년 {month}월 급여확정", company_id=book)
        uid = user["id"]
        confirmed = conn.execute(
            f"""SELECT p.*, e.name as employee_name
                FROM payslips p JOIN employees e ON p.employee_id=e.id
                WHERE p.employee_id IN ({placeholders}) AND p.pay_year=? AND p.pay_month=? AND p.status='confirmed'""",
            (*emp_ids, year, month),
        ).fetchall()
        for ps in confirmed:
            ps = dict(ps)
            existing = conn.execute(
                """SELECT id FROM erp_docs WHERE payslip_id=?
                   AND NOT EXISTS (SELECT 1 FROM erp_docs r WHERE r.reversal_of=erp_docs.id)""",
                (ps["id"],),
            ).fetchone()
            if not existing:
                try:
                    _create_payroll_slip(conn, book, ps, ps["employee_name"], uid)
                except ValueError as e:
                    conn.rollback()
                    return RedirectResponse(url=f"/payroll?year={year}&month={month}&error={quote(str(e))}", status_code=303)

        # 알림 함수 내부 commit 존재. 전표 생성 성공 후 호출
        notify_ids = [r["user_id"] for r in to_confirm if r["user_id"]]
        if notify_ids:
            create_bulk_notifications(conn, notify_ids, f"{year}년 {month}월 급여명세서 확정",
                                      f"{year}년 {month}월 급여명세서를 확인할 수 있습니다", "/payslip")
        conn.commit()

    return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)


@router.post("/payroll/pay")
async def payroll_pay(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    year: int = Form(...),
    month: int = Form(...),
    pay_date: str = Form(""),
    account_id: int = Form(0),
):
    book = user["book"]
    uid = user["id"]

    withdraw_code = "103"
    if account_id:
        acct = conn.execute(
            "SELECT account_code FROM bank_accounts WHERE id=? AND company_id=?", (account_id, book)
        ).fetchone()
        if not acct:
            raise HTTPException(status_code=404)
        withdraw_code = acct["account_code"]

    targets = conn.execute(
        """SELECT p.*, e.name AS employee_name FROM payslips p JOIN employees e ON p.employee_id=e.id
           WHERE e.company_id IS ? AND p.pay_year=? AND p.pay_month=? AND p.status='confirmed' AND p.paid_at IS NULL""",
        (emp_book(book), year, month),
    ).fetchall()

    if not targets:
        return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)

    slip_date = pay_date or _pay_date(conn, book, year, month)
    try:
        date.fromisoformat(slip_date)
    except ValueError:
        return RedirectResponse(url=f"/payroll?year={year}&month={month}&error={quote('지급일을 확인해 주세요.')}",
                                status_code=303)
    lines = []
    total = 0
    for ps in targets:
        lines.append({"code": "254", "debit": ps["net_pay"], "credit": 0,
                      "summary": f"{ps['employee_name']} {month}월 급여 지급"})
        total += ps["net_pay"]
    lines.append({"code": withdraw_code, "debit": 0, "credit": total,
                  "summary": f"{year}년 {month}월 급여 지급"})

    title = f"급여 지급 ({year}.{month:02d})"
    content = f"{year}년 {month}월 급여 지급 {len(targets)}명"
    try:
        doc_id = books.post(conn, book, uid, slip_date, lines, title, "payroll_pay", None, content=content)
    except ValueError as e:
        conn.rollback()
        return RedirectResponse(url=f"/payroll?year={year}&month={month}&error={quote(str(e))}", status_code=303)

    now = now_kst().strftime("%Y-%m-%d %H:%M:%S")
    for ps in targets:
        conn.execute("UPDATE payslips SET paid_at=?, pay_doc_id=? WHERE id=?", (now, doc_id, ps["id"]))
    audit.log(conn, uid, "payroll_pay", "payslip", None,
              f"{year}년 {month}월 {len(targets)}명 급여지급 합계 {total:,}", company_id=book)
    conn.commit()

    return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)


@router.post("/payroll/pay/{doc_id}/cancel")
async def payroll_pay_cancel(
    request: Request,
    doc_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    year: int = Form(...),
    month: int = Form(...),
    reason: str = Form(...),
    password: str = Form(...),
):
    book = user["book"]
    err_url = f"/payroll?year={year}&month={month}&error="

    row = conn.execute("SELECT password FROM users WHERE id=?", (user["id"],)).fetchone()
    if not row or not verify_pw(password, row["password"]):
        return RedirectResponse(url=err_url + "wrong_password", status_code=303)
    if not reason.strip():
        return RedirectResponse(url=err_url + "reason_required", status_code=303)

    doc = conn.execute(
        """SELECT * FROM erp_docs WHERE id=? AND company_id=? AND source_type='payroll_pay'
           AND NOT EXISTS (SELECT 1 FROM erp_docs r WHERE r.reversal_of=erp_docs.id)""",
        (doc_id, book),
    ).fetchone()
    if not doc:
        raise HTTPException(status_code=404)
    if books.period_closed(conn, book, doc["slip_date"]):
        return RedirectResponse(url=err_url + "period_closed", status_code=303)

    cnt = conn.execute("SELECT COUNT(*) FROM payslips WHERE pay_doc_id=?", (doc_id,)).fetchone()[0]
    books.reverse(conn, doc_id, user["id"], doc["slip_date"])
    conn.execute("UPDATE payslips SET paid_at=NULL, pay_doc_id=NULL WHERE pay_doc_id=?", (doc_id,))
    audit.log(conn, user["id"], "payroll_pay_cancel", "erp_docs", doc_id,
              f"{year}년 {month}월 급여지급 취소 ({cnt}명): {reason.strip()}", company_id=book)
    conn.commit()

    return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)


def _save_payslip_items(conn, payslip_id, items, amounts, premium):
    conn.execute("DELETE FROM payslip_items WHERE payslip_id=?", (payslip_id,))
    for it in items:
        amount = amounts.get(it["id"], 0)
        if amount == 0 and it["code"] != "BASE":
            continue
        overridden = it["code"] in PREMIUM_CODES and amount != premium.get(it["code"])
        conn.execute(
            "INSERT INTO payslip_items (payslip_id, item_id, amount, overridden) VALUES (?,?,?,?)",
            (payslip_id, it["id"], amount, 1 if overridden else 0),
        )


@router.post("/payroll/bulk-generate")
async def payroll_bulk_generate(
    request: Request,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    year: int = Form(...),
    month: int = Form(...),
):
    book = user["book"]

    workers = conn.execute(
        """SELECT id, base_salary, monthly_work_hours, tax_dependent_count,
                  tax_child_count, tax_rate
           FROM employees
           WHERE company_id IS ? AND (status='active' OR resign_date >= ?)""",
        (emp_book(book), f"{year:04d}-{month:02d}-01"),
    ).fetchall()

    pay_date = _pay_date(conn, book, year, month)
    rates = get_rates(conn, year, month)
    five_plus = is_five_plus(conn, book)
    items = get_items(conn, book)
    cnt = 0

    for w in workers:
        existing = conn.execute(
            "SELECT id FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
            (w["id"], year, month),
        ).fetchone()
        if existing:
            continue

        base = w["base_salary"] or 0
        monthly_hours = w["monthly_work_hours"] or 0
        hourly = calc_ordinary_hourly(base, monthly_hours)
        stats = get_month_work_stats(conn, w["id"], f"{year}-{month:02d}")
        premium = calc_premium_pay(hourly, stats["ot_mins"], stats["night_mins"], stats["holiday_day_mins"], five_plus)

        amounts = {}
        for it in items:
            if it["code"] == "BASE":
                amounts[it["id"]] = base
            elif it["code"] in premium:
                amounts[it["id"]] = premium[it["code"]]
            else:
                amounts[it["id"]] = 0

        build_items = [
            {"amount": amounts[it["id"]], "taxable": it["taxable"], "tax_free_limit": it["tax_free_limit"]}
            for it in items
        ]
        dependents = w["tax_dependent_count"] or 1
        children = w["tax_child_count"] or 0
        tax_rate = w["tax_rate"] or 100
        ps = build_payslip(build_items, rates, dependents, children, tax_rate)
        er = calc_employer_insurance(ps)
        premium_rule = "5인이상" if five_plus else "5인미만"

        cur = conn.execute("""
            INSERT INTO payslips (employee_id, pay_year, pay_month, gross_pay, national_pension,
                health_insurance, long_care, employment_insurance, income_tax, resident_tax,
                total_deduction, net_pay, pay_date, status, created_at, premium_rule,
                employer_pension, employer_health, employer_employment, employer_accident)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,'draft',datetime('now','localtime'),?,?,?,?,?)
        """, (w["id"], year, month, ps["gross_pay"], ps["national_pension"], ps["health_insurance"],
              ps["long_care"], ps["employment_insurance"], ps["income_tax"], ps["resident_tax"],
              ps["total_deduction"], ps["net_pay"], pay_date, premium_rule,
              er["employer_pension"], er["employer_health"], er["employer_employment"], er["employer_accident"]))
        _save_payslip_items(conn, cur.lastrowid, items, amounts, premium)
        cnt += 1

    conn.commit()

    return RedirectResponse(
        url=f"/payroll?year={year}&month={month}&bulk_done={cnt}",
        status_code=303,
    )


@router.post("/payroll/{emp_id}")
async def payroll_save(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    year: int = Form(...),
    month: int = Form(...),
):
    book = user["book"]
    worker = book_emp(conn, book, emp_id)

    existing = conn.execute(
        "SELECT * FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (emp_id, year, month),
    ).fetchone()
    if existing and existing["status"] == "confirmed":
        return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)

    items = get_items(conn, book)
    form = await request.form()
    amounts = {}
    for it in items:
        raw = form.get(f"amount_{it['id']}", "0")
        try:
            amt = int(raw)
        except ValueError:
            amt = 0
        if amt < 0:
            return JSONResponse({"error": "금액은 0 이상이어야 합니다."}, status_code=400)
        amounts[it["id"]] = amt

    five_plus = is_five_plus(conn, book)
    ordinary_total = sum(amounts[it["id"]] for it in items if it["ordinary"])
    hourly = calc_ordinary_hourly(ordinary_total, worker["monthly_work_hours"] or 0)
    stats = get_month_work_stats(conn, emp_id, f"{year}-{month:02d}")
    premium = calc_premium_pay(hourly, stats["ot_mins"], stats["night_mins"], stats["holiday_day_mins"], five_plus)

    build_items = [
        {"amount": amounts[it["id"]], "taxable": it["taxable"], "tax_free_limit": it["tax_free_limit"]}
        for it in items
    ]
    dependents = worker["tax_dependent_count"] or 1
    children = worker["tax_child_count"] or 0
    tax_rate = worker["tax_rate"] or 100
    rates = get_rates(conn, year, month)
    ps = build_payslip(build_items, rates, dependents, children, tax_rate)
    er = calc_employer_insurance(ps)
    pay_date = _pay_date(conn, book, year, month)
    premium_rule = "5인이상" if five_plus else "5인미만"

    new_vals = {"gross_pay": ps["gross_pay"], "national_pension": ps["national_pension"],
                "health_insurance": ps["health_insurance"], "long_care": ps["long_care"],
                "employment_insurance": ps["employment_insurance"], "income_tax": ps["income_tax"],
                "resident_tax": ps["resident_tax"], "total_deduction": ps["total_deduction"],
                "net_pay": ps["net_pay"], "pay_date": pay_date, "premium_rule": premium_rule,
                "employer_pension": er["employer_pension"], "employer_health": er["employer_health"],
                "employer_employment": er["employer_employment"], "employer_accident": er["employer_accident"]}

    if existing:
        payslip_id = existing["id"]
        conn.execute("""
            UPDATE payslips SET gross_pay=?, national_pension=?, health_insurance=?,
                long_care=?, employment_insurance=?, income_tax=?, resident_tax=?,
                total_deduction=?, net_pay=?, pay_date=?, status='draft', premium_rule=?,
                employer_pension=?, employer_health=?, employer_employment=?, employer_accident=?
            WHERE id=?
        """, (ps["gross_pay"], ps["national_pension"], ps["health_insurance"], ps["long_care"],
              ps["employment_insurance"], ps["income_tax"], ps["resident_tax"], ps["total_deduction"],
              ps["net_pay"], pay_date, premium_rule,
              er["employer_pension"], er["employer_health"], er["employer_employment"], er["employer_accident"],
              payslip_id))
        changed = audit.diff(dict(existing), new_vals)
        audit.log(conn, user["id"], "payroll_edit", "payslip", payslip_id,
                  f"{worker['name']} {year}년 {month}월 급여수정", detail=changed, company_id=book)
    else:
        cur = conn.execute("""
            INSERT INTO payslips (employee_id, pay_year, pay_month, gross_pay, national_pension,
                health_insurance, long_care, employment_insurance, income_tax, resident_tax,
                total_deduction, net_pay, pay_date, status, created_at, premium_rule,
                employer_pension, employer_health, employer_employment, employer_accident)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,'draft',datetime('now','localtime'),?,?,?,?,?)
        """, (emp_id, year, month, ps["gross_pay"], ps["national_pension"], ps["health_insurance"],
              ps["long_care"], ps["employment_insurance"], ps["income_tax"], ps["resident_tax"],
              ps["total_deduction"], ps["net_pay"], pay_date, premium_rule,
              er["employer_pension"], er["employer_health"], er["employer_employment"], er["employer_accident"]))
        payslip_id = cur.lastrowid
        audit.log(conn, user["id"], "payroll_edit", "payslip", payslip_id,
                  f"{worker['name']} {year}년 {month}월 급여입력", company_id=book)

    _save_payslip_items(conn, payslip_id, items, amounts, premium)
    conn.commit()

    return RedirectResponse(url=f"/payroll?year={year}&month={month}", status_code=303)


@router.get("/payroll/{emp_id}/print", response_class=HTMLResponse)
async def payslip_print(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    now = now_kst()
    book = user["book"]
    worker = book_emp(conn, book, emp_id)
    try:
        year = int(request.query_params.get("year", now.year))
        month = int(request.query_params.get("month", now.month))
    except ValueError:
        return RedirectResponse(url="/payroll/ledger", status_code=303)

    row = conn.execute(
        "SELECT * FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (emp_id, year, month),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/payroll/ledger", status_code=303)

    payslip = dict(row)
    work_stats, calc_notes, items = build_payslip_detail(
        conn, emp_id, worker.get("monthly_work_hours") or 0, payslip
    )

    return templates.TemplateResponse(
        request=request, name="hr/payslip_print.html", context={
            "payslip": payslip,
            "employee": worker,
            "company_name": book_label(conn, book),
            "work_stats": work_stats,
            "calc_notes": calc_notes,
            "items": items,
            "viewer_name": user["name"],
            "viewed_at": now.strftime("%Y-%m-%d %H:%M"),
        }
    )


# 급여 마스터 정보 (사원 화면에서 분리된 지급 기준값)

@router.get("/payroll/{emp_id}/info", response_class=HTMLResponse)
async def payroll_info_form(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
):
    worker = book_emp(conn, user["book"], emp_id)
    return templates.TemplateResponse(
        request=request, name="client/payroll_info.html", context={
            "request": request,
            "page_title": f"{worker['name']} - 급여 마스터 정보",
            "user_name": user["name"],
            "worker": worker,
            "acct_masked": mask_acct(worker["bank_account"]) if worker["bank_account"] else "",
            "error": "",
        }
    )


@router.post("/payroll/{emp_id}/info", response_class=HTMLResponse)
async def payroll_info_save(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_hr_book),
    conn=Depends(get_db),
    base_salary: int = Form(0),
    monthly_work_hours: int = Form(0),
    employment_insurance: int = Form(1),
    tax_dependent_count: int = Form(1),
    tax_child_count: int = Form(0),
    tax_rate: int = Form(100),
    bank_name: str = Form(""),
    bank_account: str = Form(""),
):
    worker = book_emp(conn, user["book"], emp_id)
    if tax_rate not in (80, 100, 120):
        return templates.TemplateResponse(
            request=request, name="client/payroll_info.html", context={
                "request": request,
                "page_title": f"{worker['name']} - 급여 마스터 정보",
                "user_name": user["name"],
                "worker": worker,
                "acct_masked": mask_acct(worker["bank_account"]) if worker["bank_account"] else "",
                "error": "원천징수 비율은 80/100/120 중 하나여야 합니다",
            }
        )
    new_acct = bank_account.strip() or worker["bank_account"]
    conn.execute(
        """UPDATE employees SET base_salary=?, monthly_work_hours=?, employment_insurance=?,
               tax_dependent_count=?, tax_child_count=?, tax_rate=?, bank_name=?, bank_account=?,
               updated_at=datetime('now','localtime')
           WHERE id=?""",
        (base_salary, monthly_work_hours, employment_insurance, max(1, tax_dependent_count),
         max(0, tax_child_count), tax_rate, bank_name, new_acct, emp_id),
    )
    # 급여 입금 계좌 변경은 횡령 경로이므로 이력 기록 후 본인에게 통지
    if (bank_name, new_acct) != (worker["bank_name"] or "", worker["bank_account"] or ""):
        audit.log(conn, user["id"], "bank_account_change", "employee", emp_id, worker["name"],
                  detail={"before": f"{worker['bank_name'] or ''} {mask_acct(worker['bank_account'])}",
                          "after": f"{bank_name} {mask_acct(new_acct)}"},
                  company_id=worker["company_id"], ip=get_client_ip(request))
        if worker["user_id"]:
            create_notification(conn, worker["user_id"], "급여 계좌 변경",
                                f"{bank_name} {mask_acct(new_acct)} 담당자 변경. 본인 요청이 아니면 회사에 문의", "/mypage/profile")
    if base_salary != (worker["base_salary"] or 0):
        log_changes(conn, worker["company_id"], emp_id, {"base_salary": (worker["base_salary"], base_salary)}, user["id"])
    conn.commit()
    return RedirectResponse(url=f"/payroll/{emp_id}/info", status_code=303)



