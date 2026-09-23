import calendar

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from core.tz import now_kst
from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse
from core.deps import get_db, require_login, templates
from core.export import esc_cell, excel_response, HEADER_FONT, HEADER_FILL, CELL_FONT, THIN_BORDER

router = APIRouter()


def _load_ledger(conn, year, month, company_id, filter_disabled=False):
    first_day = f"{year}-{month:02d}-01"
    _, last_d = calendar.monthrange(year, month)
    last_day = f"{year}-{month:02d}-{last_d:02d}"

    emp_sql = """
        SELECT e.id, e.name, e.employee_no, e.resident_no, e.hire_date, e.job_title,
               e.monthly_work_hours, e.disability_type, e.disability_grade,
               p.base_salary, p.overtime_pay, p.disability_allowance, p.meal_allowance,
               p.gross_pay, p.national_pension, p.health_insurance,
               p.employment_insurance AS ei_deduction, p.income_tax, p.resident_tax,
               p.total_deduction, p.net_pay, p.status AS pay_status,
               p.employer_pension, p.employer_health, p.employer_employment, p.employer_accident
        FROM employees e
        LEFT JOIN payslips p ON e.id = p.employee_id AND p.pay_year=? AND p.pay_month=?
        WHERE e.status = 'active'
    """
    params = [year, month]

    if company_id:
        emp_sql += " AND e.company_id=?"
        params.append(company_id)
    if filter_disabled:
        emp_sql += " AND e.disability_type IS NOT NULL AND e.disability_type != ''"

    emp_sql += " ORDER BY e.employee_no, e.name"
    rows = conn.execute(emp_sql, params).fetchall()
    rows = [dict(r) for r in rows]

    if not rows:
        return rows, {}

    emp_ids = [r["id"] for r in rows]
    ph = ",".join("?" * len(emp_ids))
    att_rows = conn.execute(
        f"""SELECT employee_id,
                   COUNT(*) AS work_days,
                   SUM(work_minutes) AS total_work_mins,
                   SUM(overtime_minutes) AS total_ot_mins
            FROM attendance
            WHERE work_date BETWEEN ? AND ?
              AND status IN ('normal','late','early_leave','late_early')
              AND employee_id IN ({ph})
            GROUP BY employee_id""",
        [first_day, last_day, *emp_ids],
    ).fetchall()

    att_map = {r["employee_id"]: dict(r) for r in att_rows}
    return rows, att_map


def _mask_rno(resident_no):
    if not resident_no:
        return "-"
    rno = resident_no.replace("-", "")
    if len(rno) >= 7:
        return rno[:6] + "-" + "*" * 7
    return resident_no


@router.get("/payroll/payslip/{emp_id}/print", response_class=HTMLResponse)
async def hr_payslip_print(
    request: Request,
    emp_id: int,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    now = now_kst()
    try:
        year = int(request.query_params.get("year", now.year))
        month = int(request.query_params.get("month", now.month))
    except ValueError:
        return RedirectResponse(url="/hr/payroll/ledger", status_code=303)

    row = conn.execute(
        "SELECT * FROM payslips WHERE employee_id=? AND pay_year=? AND pay_month=?",
        (emp_id, year, month),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/hr/payroll/ledger", status_code=303)

    emp_row = conn.execute(
        "SELECT name, employee_no, dept, position, company_id FROM employees WHERE id=?",
        (emp_id,),
    ).fetchone()
    company_name = ""
    if emp_row:
        comp = conn.execute(
            "SELECT name FROM client_companies WHERE id=?", (emp_row["company_id"],)
        ).fetchone()
        if comp:
            company_name = comp["name"]

    return templates.TemplateResponse(
        request=request, name="hr/payslip_print.html", context={
            "payslip": dict(row),
            "employee": dict(emp_row) if emp_row else {},
            "company_name": company_name,
        }
    )


@router.get("/payroll/ledger", response_class=HTMLResponse)
async def payroll_ledger(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None
    filter_val = request.query_params.get("filter", "all")
    filter_disabled = filter_val == "disabled"

    rows, att_map = _load_ledger(conn, year, month, company_id, filter_disabled)

    for r in rows:
        att = att_map.get(r["id"], {})
        r["work_days"] = att.get("work_days") or 0
        r["total_work_mins"] = att.get("total_work_mins") or 0
        r["total_ot_mins"] = att.get("total_ot_mins") or 0
        r["work_hours"] = r["total_work_mins"] // 60
        r["ot_hours"] = r["total_ot_mins"] // 60
        r["resident_no_masked"] = _mask_rno(r.get("resident_no"))

    total_count = len(rows)
    sum_base = sum(r["base_salary"] or 0 for r in rows)
    sum_ot = sum(r["overtime_pay"] or 0 for r in rows)
    sum_dis = sum(r["disability_allowance"] or 0 for r in rows)
    sum_meal = sum(r["meal_allowance"] or 0 for r in rows)
    sum_gross = sum(r["gross_pay"] or 0 for r in rows)
    sum_pension = sum(r["national_pension"] or 0 for r in rows)
    sum_health = sum(r["health_insurance"] or 0 for r in rows)
    sum_ei = sum(r["ei_deduction"] or 0 for r in rows)
    sum_itax = sum(r["income_tax"] or 0 for r in rows)
    sum_rtax = sum(r["resident_tax"] or 0 for r in rows)
    sum_deduction = sum(r["total_deduction"] or 0 for r in rows)
    sum_net = sum(r["net_pay"] or 0 for r in rows)
    sum_er_pension = sum(r["employer_pension"] or 0 for r in rows)
    sum_er_health = sum(r["employer_health"] or 0 for r in rows)
    sum_er_employment = sum(r["employer_employment"] or 0 for r in rows)
    sum_er_accident = sum(r["employer_accident"] or 0 for r in rows)
    sum_er_total = sum_er_pension + sum_er_health + sum_er_employment + sum_er_accident

    companies = conn.execute(
        "SELECT id, name FROM client_companies WHERE status='active' ORDER BY name"
    ).fetchall()

    company_name = ""
    if company_id:
        cn = conn.execute("SELECT name FROM client_companies WHERE id=?", (company_id,)).fetchone()
        if cn:
            company_name = cn["name"]

    return templates.TemplateResponse(
        request=request,
        name="hr/payroll_ledger.html",
        context={
            "request": request,
            "page_title": "급여대장",
            "user_name": user["name"],
            "user_role": user["role"],
            "rows": rows,
            "year": year,
            "month": month,
            "company_id": company_id or "",
            "company_name": company_name,
            "companies": companies,
            "filter": filter_val,
            "total_count": total_count,
            "sum_base": sum_base,
            "sum_ot": sum_ot,
            "sum_dis": sum_dis,
            "sum_meal": sum_meal,
            "sum_gross": sum_gross,
            "sum_pension": sum_pension,
            "sum_health": sum_health,
            "sum_ei": sum_ei,
            "sum_itax": sum_itax,
            "sum_rtax": sum_rtax,
            "sum_deduction": sum_deduction,
            "sum_net": sum_net,
            "sum_er_pension": sum_er_pension,
            "sum_er_health": sum_er_health,
            "sum_er_employment": sum_er_employment,
            "sum_er_accident": sum_er_accident,
            "sum_er_total": sum_er_total,
        },
    )


@router.get("/payroll/ledger/export")
async def payroll_ledger_export(
    request: Request,
    user: dict = Depends(require_login),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    company_id = request.query_params.get("company_id", "")
    company_id = int(company_id) if company_id else None
    filter_val = request.query_params.get("filter", "all")
    filter_disabled = filter_val == "disabled"

    rows, att_map = _load_ledger(conn, year, month, company_id, filter_disabled)

    for r in rows:
        att = att_map.get(r["id"], {})
        r["work_days"] = att.get("work_days") or 0
        r["total_work_mins"] = att.get("total_work_mins") or 0
        r["total_ot_mins"] = att.get("total_ot_mins") or 0
        r["work_hours"] = r["total_work_mins"] // 60
        r["ot_hours"] = r["total_ot_mins"] // 60

    company_name = ""
    if company_id:
        cn = conn.execute("SELECT name FROM client_companies WHERE id=?", (company_id,)).fetchone()
        if cn:
            company_name = cn["name"]

    wb = Workbook()
    ws = wb.active
    ws.title = "급여대장"

    title_text = f"급여대장 - {year}년 {month}월"
    if company_name:
        title_text += f" ({company_name})"

    headers = [
        "순번", "근로자명", "주민등록번호", "고용연월일", "종사업무",
        "소정근로시간", "근로일수", "근로시간수", "연장근로시간",
        "기본급", "초과근무수당", "장애수당", "식대", "지급합계",
        "국민연금", "건강보험", "고용보험", "소득세", "지방소득세",
        "공제합계", "실지급액",
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="맑은 고딕", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    money_cols = set(range(10, len(headers) + 1))

    for ri, r in enumerate(rows, 3):
        vals = [
            ri - 2,
            esc_cell(r["name"]),
            esc_cell(_mask_rno(r.get("resident_no"))),
            esc_cell(r.get("hire_date") or ""),
            esc_cell(r.get("job_title") or ""),
            r.get("monthly_work_hours") or 0,
            r["work_days"],
            r["work_hours"],
            r["ot_hours"],
            r.get("base_salary") or 0,
            r.get("overtime_pay") or 0,
            r.get("disability_allowance") or 0,
            r.get("meal_allowance") or 0,
            r.get("gross_pay") or 0,
            r.get("national_pension") or 0,
            r.get("health_insurance") or 0,
            r.get("ei_deduction") or 0,
            r.get("income_tax") or 0,
            r.get("resident_tax") or 0,
            r.get("total_deduction") or 0,
            r.get("net_pay") or 0,
        ]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            if ci in money_cols and isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="center")

    # 합계행
    sum_row = len(rows) + 3
    sum_vals = [
        "", "합계", "", "", "", "", "", "", "",
        sum(r.get("base_salary") or 0 for r in rows),
        sum(r.get("overtime_pay") or 0 for r in rows),
        sum(r.get("disability_allowance") or 0 for r in rows),
        sum(r.get("meal_allowance") or 0 for r in rows),
        sum(r.get("gross_pay") or 0 for r in rows),
        sum(r.get("national_pension") or 0 for r in rows),
        sum(r.get("health_insurance") or 0 for r in rows),
        sum(r.get("ei_deduction") or 0 for r in rows),
        sum(r.get("income_tax") or 0 for r in rows),
        sum(r.get("resident_tax") or 0 for r in rows),
        sum(r.get("total_deduction") or 0 for r in rows),
        sum(r.get("net_pay") or 0 for r in rows),
    ]
    for ci, val in enumerate(sum_vals, 1):
        cell = ws.cell(row=sum_row, column=ci, value=val)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        if ci in money_cols and isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="center")

    col_widths = [6, 10, 16, 12, 12, 10, 8, 8, 8, 12, 12, 10, 10, 12, 10, 10, 10, 10, 10, 12, 12]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    filename = f"급여대장_{year}{month:02d}.xlsx"
    return excel_response(wb, filename)
