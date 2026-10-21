import calendar
import json
from datetime import date

from core.tz import now_kst, today_kst

from fastapi import APIRouter, Request, Depends, Form, File, UploadFile, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse

from core.deps import get_db, require_client_hr_with_company, require_client_with_company, templates
from core.constants import APPEAL_TYPE_LABELS, COMPANY_INFO, DISABILITY_TYPES, JOB_CODES
from core.work_schedule import get_schedule
from core.leave import get_leave_balance
from core.severance import calc_severance, calc_sev_tax, calc_ordinary_daily
from core.appointments import log_changes
from core.attendance import (
    get_company_workers,
    get_attendance_records,
    get_attendance_summary,
    get_worker_detail,
    get_worker_month_records,
    get_worker_month_leaves,
    merge_records_leaves,
    build_calendar_data,
    worker_month_summary,
    prev_next_month,
    available_months,
    STATUS_LABELS,
)
from core.export import excel_response, pdf_response, csv_response, mask_rno
from core.crypto import enc_rno, dec_rno, valid_rno
from core import audit, employee_files
from core.auth import hash_pw
from routers.hr_employee import _save_cert_file
from routers.platform_company import gen_temp_pw

router = APIRouter(prefix="/client")

DISABILITY_TYPE_LABELS = {
    "physical": "지체", "visual": "시각", "hearing": "청각",
    "intellectual": "지적", "mental": "정신", "brain": "뇌병변",
    "autism": "자폐성", "kidney": "신장", "heart": "심장",
    "respiratory": "호흡기", "liver": "간", "facial": "안면",
    "intestinal": "장루요루", "epilepsy": "뇌전증",
}

DISABILITY_GRADE_LABELS = {
    "severe": "중증", "mild": "경증",
}


@router.get("/workers", response_class=HTMLResponse)
async def worker_list(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    status_filter = request.query_params.get("status", "")
    search_query = request.query_params.get("q", "")

    company_id = user["company_id"]
    sql = """SELECT id, name, employee_no, dept, position, hire_date, status,
                    disability_type, disability_grade
             FROM employees WHERE company_id=?"""
    params = [company_id]
    if status_filter:
        sql += " AND status=?"
        params.append(status_filter)
    if search_query:
        sql += " AND (name LIKE ? OR employee_no LIKE ?)"
        params += [f"%{search_query}%", f"%{search_query}%"]
    sql += " ORDER BY name"
    rows = conn.execute(sql, params).fetchall()
    workers = [dict(r) for r in rows]
    for w in workers:
        w["disability_type_label"] = DISABILITY_TYPE_LABELS.get(w["disability_type"], w["disability_type"] or "")
        w["disability_grade_label"] = DISABILITY_GRADE_LABELS.get(w["disability_grade"], w["disability_grade"] or "")
    total_count = len(workers)
    active_count = sum(1 for w in workers if w["status"] == "active")
    inactive_count = total_count - active_count

    month_start = today_kst().replace(day=1).isoformat()
    new_hire_count = sum(
        1 for w in workers
        if w.get("hire_date") and w["hire_date"] >= month_start
    )

    return templates.TemplateResponse(
        request=request, name="client/workers.html", context={
            "request": request,
            "page_title": "근로자 관리",
            "user_name": user["name"],
            "workers": workers,
            "total_count": total_count,
            "active_count": active_count,
            "inactive_count": inactive_count,
            "new_hire_count": new_hire_count,
            "status_filter": status_filter,
            "search_query": search_query,
            "disability_type_labels": DISABILITY_TYPE_LABELS,
        }
    )


@router.get("/workers/new", response_class=HTMLResponse)
async def worker_new_form(
    request: Request,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    return templates.TemplateResponse(
        request=request, name="hr/employee_form.html", context={
            "request": request,
            "page_title": "근로자 신규 등록",
            "user_name": user["name"],
            "employee": None,
            "companies": [],
            "disability_types": DISABILITY_TYPES,
            "job_codes": JOB_CODES,
            "error": "",
            "client_mode": True,
            "form_action": "/client/workers",
            "cancel_url": "/client/workers",
        },
    )


@router.post("/workers", response_class=HTMLResponse)
async def worker_create(
    request: Request,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
    name: str = Form(...),
    employee_no: str = Form(""),
    dept: str = Form("경영지원팀"),
    position: str = Form("사원"),
    hire_date: str = Form(""),
    status: str = Form("active"),
    gender: str = Form(""),
    birth_date: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    address: str = Form(""),
    disability_type: str = Form(""),
    disability_grade: str = Form(""),
    emergency_contact: str = Form(""),
    emergency_phone: str = Form(""),
    notes: str = Form(""),
    resident_no: str = Form(""),
    job_title: str = Form(""),
    job_code: str = Form(""),
    disability_cert_file: UploadFile = File(None),
    create_account: str = Form(""),
    login_username: str = Form(""),
    renew_date: str = Form(""),
    contract_end: str = Form(""),
    career: str = Form(""),
    education: str = Form(""),
    military: str = Form(""),
    license: str = Form(""),
    training: str = Form(""),
    disability_class: str = Form(""),
    veteran_grade: str = Form(""),
    disability_date: str = Form(""),
):
    company_id = user["company_id"]  # 폼 값 불신. 세션 소속사 기준
    cert_path = await _save_cert_file(disability_cert_file)

    error = ""
    if resident_no.strip() and not valid_rno(resident_no):
        error = "주민등록번호 형식이 올바르지 않습니다."
    if not error and create_account and login_username:
        if conn.execute("SELECT 1 FROM users WHERE username=?", (login_username,)).fetchone():
            error = "이미 사용 중인 로그인 아이디입니다."

    if not error:
        try:
            cur = conn.execute(
                """INSERT INTO employees
                   (name, employee_no, dept, position, hire_date, status,
                    gender, birth_date, phone, email, address,
                    disability_type, disability_grade,
                    emergency_contact, emergency_phone,
                    notes, company_id,
                    resident_no, job_title, job_code, disability_cert_file,
                    renew_date, contract_end, career, education, military, license, training,
                    disability_class, veteran_grade, disability_date)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (name, employee_no or None, dept, position, hire_date, status,
                 gender, birth_date, phone, email, address,
                 disability_type, disability_grade,
                 emergency_contact, emergency_phone,
                 notes, company_id,
                 enc_rno(resident_no), job_title, job_code or None, cert_path,
                 renew_date, contract_end, career, education, military, license, training,
                 disability_class, veteran_grade, disability_date),
            )
            new_id = cur.lastrowid
            temp_pw_show = None
            if create_account and login_username:
                temp_pw = gen_temp_pw()
                uid = conn.execute(
                    "INSERT INTO users (username, password, name, phone, role, must_change_password) "
                    "VALUES (?,?,?,?,'worker',1)",
                    (login_username.strip(), hash_pw(temp_pw), name, phone),
                ).lastrowid
                conn.execute("UPDATE employees SET user_id=? WHERE id=?", (uid, new_id))
                temp_pw_show = {"username": login_username.strip(), "temp_pw": temp_pw}
            audit.log(conn, user["id"], "worker_create", "employee", new_id, name, company_id=company_id)
            conn.commit()
            if temp_pw_show:
                request.session["temp_pw_show"] = temp_pw_show
        except Exception:
            error = "사번이 중복되었거나 필수 항목이 누락되었습니다."

    if error:
        return JSONResponse({"error": error}, status_code=400)
    return RedirectResponse(url=f"/client/workers/{new_id}", status_code=303)


@router.get("/workers/{employee_id}", response_class=HTMLResponse)
async def worker_detail(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        """SELECT e.*, c.name as company_name
           FROM employees e
           LEFT JOIN client_companies c ON e.company_id = c.id
           WHERE e.id=?""",
        (employee_id,),
    ).fetchone()
    if not row or row["company_id"] != company_id:
        return RedirectResponse(url="/client/workers", status_code=303)
    worker = dict(row)
    worker["disability_type_label"] = DISABILITY_TYPE_LABELS.get(worker["disability_type"], worker["disability_type"] or "")
    worker["disability_grade_label"] = DISABILITY_GRADE_LABELS.get(worker["disability_grade"], worker["disability_grade"] or "")

    recent_attendance = conn.execute(
        """SELECT work_date, clock_in, clock_out, status, notes
           FROM attendance WHERE employee_id=?
           ORDER BY work_date DESC LIMIT 5""",
        (employee_id,),
    ).fetchall()
    recent_attendance = [dict(r) for r in recent_attendance]

    active_accom_count = conn.execute(
        """SELECT COUNT(*) FROM accommodation_requests
           WHERE employee_id=? AND status='pending'""",
        (employee_id,),
    ).fetchone()[0]

    bal = get_leave_balance(conn, employee_id, now_kst().year)
    leave_total = bal["total"]
    leave_used = bal["used"]
    leave_remaining = max(0, leave_total - leave_used)
    leave_pct = int(leave_used / leave_total * 100) if leave_total else 0

    appointments = conn.execute(
        """SELECT a.*, u.name AS created_by_name FROM hr_appointments a
           LEFT JOIN users u ON u.id = a.created_by
           WHERE a.employee_id=? ORDER BY a.id DESC""",
        (employee_id,),
    ).fetchall()

    return templates.TemplateResponse(
        request=request, name="client/worker_detail.html", context={
            "request": request,
            "page_title": f"{worker['name']} 프로필",
            "user_name": user["name"],
            "worker": worker,
            "recent_attendance": recent_attendance,
            "active_accom_count": active_accom_count,
            "leave_total": leave_total,
            "leave_used": leave_used,
            "leave_remaining": leave_remaining,
            "leave_pct": leave_pct,
            "files": employee_files.list_files(conn, employee_id),
            "file_kinds": employee_files.KINDS,
            "appointments": appointments,
            "temp_pw_show": request.session.pop("temp_pw_show", None),
        }
    )


@router.get("/workers/{employee_id}/edit", response_class=HTMLResponse)
async def worker_edit_form(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    employee = conn.execute("SELECT * FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not employee or employee["company_id"] != user["company_id"]:
        raise HTTPException(status_code=404)
    if employee["resident_no"]:
        audit.log(conn, user["id"], "resident_no_view", "employee", employee_id, employee["name"],
                  company_id=employee["company_id"])
        conn.commit()
    employee = dict(employee)
    employee["resident_no_masked"] = mask_rno(dec_rno(employee["resident_no"]))
    employee["resident_no"] = ""
    return templates.TemplateResponse(
        request=request, name="hr/employee_form.html", context={
            "request": request,
            "page_title": f"{employee['name']} - 정보 수정",
            "user_name": user["name"],
            "employee": employee,
            "companies": [],
            "disability_types": DISABILITY_TYPES,
            "job_codes": JOB_CODES,
            "error": "",
            "client_mode": True,
            "form_action": f"/client/workers/{employee_id}/edit",
            "cancel_url": f"/client/workers/{employee_id}",
        },
    )


@router.post("/workers/{employee_id}/edit", response_class=HTMLResponse)
async def worker_update(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
    name: str = Form(...),
    employee_no: str = Form(""),
    dept: str = Form("경영지원팀"),
    position: str = Form("사원"),
    hire_date: str = Form(""),
    status: str = Form("active"),
    gender: str = Form(""),
    birth_date: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    address: str = Form(""),
    disability_type: str = Form(""),
    disability_grade: str = Form(""),
    emergency_contact: str = Form(""),
    emergency_phone: str = Form(""),
    notes: str = Form(""),
    resident_no: str = Form(""),
    job_title: str = Form(""),
    job_code: str = Form(""),
    disability_cert_file: UploadFile = File(None),
    renew_date: str = Form(""),
    contract_end: str = Form(""),
    career: str = Form(""),
    education: str = Form(""),
    military: str = Form(""),
    license: str = Form(""),
    training: str = Form(""),
    disability_class: str = Form(""),
    veteran_grade: str = Form(""),
    disability_date: str = Form(""),
):
    company_id = user["company_id"]
    old_row = conn.execute("SELECT * FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not old_row or old_row["company_id"] != company_id:
        raise HTTPException(status_code=404)

    if resident_no.strip() and not valid_rno(resident_no):
        return JSONResponse({"error": "주민등록번호 형식이 올바르지 않습니다."}, status_code=400)
    cert_path = await _save_cert_file(disability_cert_file)
    rno_val = enc_rno(resident_no) if resident_no.strip() else old_row["resident_no"]
    try:
        sql = """UPDATE employees SET
               name=?, employee_no=?, dept=?, position=?, hire_date=?, status=?,
               gender=?, birth_date=?, phone=?, email=?, address=?,
               disability_type=?, disability_grade=?,
               emergency_contact=?, emergency_phone=?,
               notes=?,
               resident_no=?, job_title=?, job_code=?,
               renew_date=?, contract_end=?, career=?, education=?, military=?, license=?, training=?,
               disability_class=?, veteran_grade=?, disability_date=?,
               updated_at=datetime('now','localtime')"""
        params = [
            name, employee_no or None, dept, position, hire_date, status,
            gender, birth_date, phone, email, address,
            disability_type, disability_grade,
            emergency_contact, emergency_phone,
            notes,
            rno_val, job_title, job_code or None,
            renew_date, contract_end, career, education, military, license, training,
            disability_class, veteran_grade, disability_date,
        ]
        new_vals = {"name": name, "employee_no": employee_no or None, "dept": dept, "position": position,
                    "hire_date": hire_date, "status": status, "gender": gender, "birth_date": birth_date,
                    "phone": phone, "email": email, "address": address, "disability_type": disability_type,
                    "disability_grade": disability_grade, "emergency_contact": emergency_contact,
                    "emergency_phone": emergency_phone, "notes": notes, "job_title": job_title,
                    "job_code": job_code or None,
                    "renew_date": renew_date, "contract_end": contract_end, "career": career,
                    "education": education, "military": military, "license": license, "training": training,
                    "disability_class": disability_class, "veteran_grade": veteran_grade,
                    "disability_date": disability_date}
        if cert_path:
            sql += ", disability_cert_file=?"
            params.append(cert_path)
            new_vals["disability_cert_file"] = cert_path
        sql += " WHERE id=? AND company_id=?"
        params += [employee_id, company_id]
        conn.execute(sql, params)
        changed = audit.diff(dict(old_row), new_vals)
        audit.log(conn, user["id"], "worker_update", "employee", employee_id, name,
                  detail=changed, company_id=company_id)
        log_changes(conn, company_id, employee_id, changed, user["id"])
        conn.commit()
    except Exception:
        return JSONResponse({"error": "사번이 중복되었거나 필수 항목이 누락되었습니다."}, status_code=400)
    return RedirectResponse(url=f"/client/workers/{employee_id}", status_code=303)


@router.get("/workers/{employee_id}/resign", response_class=HTMLResponse)
async def worker_resign_form(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    employee = conn.execute("SELECT * FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not employee or employee["company_id"] != user["company_id"]:
        raise HTTPException(status_code=404)
    return templates.TemplateResponse(
        request=request, name="hr/resign_form.html", context={
            "request": request,
            "page_title": f"{employee['name']} - 퇴사 처리",
            "user_name": user["name"],
            "employee": employee,
            "form_action": f"/client/workers/{employee_id}/resign",
            "cancel_url": f"/client/workers/{employee_id}",
        },
    )


@router.post("/workers/{employee_id}/resign", response_class=HTMLResponse)
async def worker_resign_submit(
    employee_id: int,
    resign_date: str = Form(...),
    resign_reason: str = Form(""),
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    emp = conn.execute("SELECT * FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not emp or emp["company_id"] != user["company_id"]:
        raise HTTPException(status_code=404)

    conn.execute(
        "UPDATE employees SET resign_date=?, resign_reason=?, status='retired', "
        "updated_at=datetime('now','localtime') WHERE id=?",
        (resign_date, resign_reason, employee_id),
    )
    audit.log(conn, user["id"], "worker_resign", "employee", employee_id, emp["name"], company_id=emp["company_id"])
    conn.commit()

    sev = calc_severance(conn, employee_id, resign_date=resign_date)
    if sev:
        tax = calc_sev_tax(sev["amount"], emp["hire_date"], resign_date)
        bal = get_leave_balance(conn, employee_id, date.fromisoformat(resign_date).year)
        leave_days = max(0, bal["remaining"])
        leave_pay = round(leave_days * calc_ordinary_daily(conn, employee_id, resign_date))

        conn.execute(
            """INSERT INTO severance_records
               (employee_id, calc_date, avg_salary, tenure_days, amount,
                income_tax, local_tax, net_amount, leave_days, leave_pay, status, created_by)
               VALUES (?,?,?,?,?,?,?,?,?,?,'pending',?)""",
            (employee_id, resign_date, sev["avg_salary"], sev["tenure_days"], sev["amount"],
             tax["income_tax"], tax["local_tax"], tax["net"], leave_days, leave_pay, user["id"]),
        )
        conn.commit()

    return RedirectResponse(url=f"/severance/{employee_id}", status_code=303)


@router.post("/workers/{employee_id}/files", response_class=HTMLResponse)
async def worker_file_upload(
    employee_id: int,
    kind: str = Form(...),
    file: UploadFile = File(...),
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    employee = conn.execute("SELECT id, company_id FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not employee or employee["company_id"] != user["company_id"]:
        raise HTTPException(status_code=404)
    saved = await employee_files.save_file(conn, employee_id, employee["company_id"], kind, file, user["id"])
    if isinstance(saved, str):
        audit.log(conn, user["id"], "employee_file_upload", "employee", employee_id, kind,
                  company_id=employee["company_id"])
        conn.commit()
        return RedirectResponse(url=f"/client/workers/{employee_id}", status_code=303)
    return RedirectResponse(url=f"/client/workers/{employee_id}?file_err=1", status_code=303)


@router.post("/workers/{employee_id}/files/{file_id}/delete", response_class=HTMLResponse)
async def worker_file_delete(
    employee_id: int,
    file_id: int,
    user: dict = Depends(require_client_hr_with_company),
    conn=Depends(get_db),
):
    employee = conn.execute("SELECT id, company_id FROM employees WHERE id=?", (employee_id,)).fetchone()
    if not employee or employee["company_id"] != user["company_id"]:
        raise HTTPException(status_code=404)
    if employee_files.delete_file(conn, employee_id, file_id):
        audit.log(conn, user["id"], "employee_file_delete", "employee", employee_id, str(file_id),
                  company_id=employee["company_id"])
        conn.commit()
    return RedirectResponse(url=f"/client/workers/{employee_id}", status_code=303)


@router.get("/attendance", response_class=HTMLResponse)
async def attendance(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    status_filter = request.query_params.get("status", "")

    company_id = user["company_id"]
    workers = get_company_workers(conn, company_id) if company_id else []
    emp_ids = [w["id"] for w in workers]
    records = get_attendance_records(conn, emp_ids, month)
    summary = get_attendance_summary(records, len(workers))
    if status_filter:
        statuses = status_filter.split(",")
        records = [r for r in records if r["status"] in statuses]

    return templates.TemplateResponse(
        request=request, name="client/attendance.html", context={
            "request": request,
            "page_title": "근태 관리",
            "user_name": user["name"],
            "records": records,
            "summary": summary,
            "workers": workers,
            "current_month": month,
            "status_filter": status_filter,
            "months": available_months(),
        }
    )


def _build_client_sheet(conn, year, month, company_id):
    _, days_in_month = calendar.monthrange(year, month)
    ym = f"{year}-{month:02d}"

    employees = conn.execute(
        "SELECT id, name, employee_no, monthly_work_hours FROM employees WHERE status='active' AND company_id=? ORDER BY employee_no, name",
        (company_id,),
    ).fetchall()
    emp_ids = [e["id"] for e in employees]

    if not emp_ids:
        return [], days_in_month

    ph = ",".join("?" * len(emp_ids))
    att_rows = conn.execute(
        f"""SELECT employee_id, work_date, status, clock_in, clock_out,
                   work_minutes, overtime_minutes, break_minutes, notes
            FROM attendance
            WHERE employee_id IN ({ph}) AND work_date LIKE ?
            ORDER BY work_date""",
        [*emp_ids, f"{ym}%"],
    ).fetchall()

    att_map = {}
    for r in att_rows:
        att_map.setdefault(r["employee_id"], {})[r["work_date"]] = dict(r)

    rows = []
    for emp in employees:
        eid = emp["id"]
        days = att_map.get(eid, {})
        cells = []
        work_days = 0
        late_cnt = 0
        early_cnt = 0
        absent_cnt = 0
        holiday_cnt = 0
        total_mins = 0
        ot_mins = 0
        for d in range(1, days_in_month + 1):
            dt = f"{ym}-{d:02d}"
            rec = days.get(dt)
            if rec:
                cells.append(rec)
                if rec["status"] in ("normal", "late", "early_leave", "late_early"):
                    work_days += 1
                    total_mins += rec.get("work_minutes") or 0
                    ot_mins += rec.get("overtime_minutes") or 0
                if rec["status"] in ("late", "late_early"):
                    late_cnt += 1
                if rec["status"] in ("early_leave", "late_early"):
                    early_cnt += 1
                if rec["status"] == "absent":
                    absent_cnt += 1
                elif rec["status"] == "holiday":
                    holiday_cnt += 1
            else:
                cells.append(None)
        rows.append({
            "employee": dict(emp),
            "cells": cells,
            "work_days": work_days,
            "late": late_cnt,
            "early_leave": early_cnt,
            "absent": absent_cnt,
            "holiday_days": holiday_cnt,
            "monthly_work_hours": emp["monthly_work_hours"] or 0,
            "total_mins": total_mins,
            "total_hours": total_mins // 60,
            "total_mins_r": total_mins % 60,
            "ot_hours": ot_mins // 60,
            "ot_mins_r": ot_mins % 60,
        })
    return rows, days_in_month


@router.get("/attendance/sheet", response_class=HTMLResponse)
async def attendance_sheet(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    page = int(request.query_params.get("page", 1))
    company_id = user["company_id"]

    rows, days_in_month = _build_client_sheet(conn, year, month, company_id)

    ci_row = conn.execute(
        "SELECT name, business_no, representative, address, contact_name FROM client_companies WHERE id=?",
        (company_id,),
    ).fetchone()
    company_name = ci_row["name"] if ci_row else ""
    company_info = dict(ci_row) if ci_row else None

    first_weekday = calendar.monthrange(year, month)[0]
    weekdays = [(first_weekday + d) % 7 for d in range(days_in_month)]

    # 배너 집계
    weekday_count = sum(1 for wd in weekdays if wd < 5)
    total_emp = len(rows)
    sum_work = sum(r["work_days"] for r in rows)
    sum_late = sum(r["late"] for r in rows)
    sum_absent = sum(r["absent"] for r in rows)
    expected = weekday_count * total_emp
    banner = {
        "total_emp": total_emp,
        "work": sum_work,
        "late": sum_late,
        "absent": sum_absent,
        "rate": round(sum_work / expected * 100, 1) if expected else 0,
    }

    # 페이지네이션
    per_page = 50
    total_pages = max(1, (total_emp + per_page - 1) // per_page)
    if page < 1:
        page = 1
    if page > total_pages:
        page = total_pages
    rows_page = rows[(page - 1) * per_page: page * per_page]

    return templates.TemplateResponse(
        request=request,
        name="client/attendance_sheet.html",
        context={
            "request": request,
            "page_title": "출근부",
            "user_name": user["name"],
            "rows": rows_page,
            "all_rows": rows,
            "year": year,
            "month": month,
            "days_in_month": days_in_month,
            "weekdays": weekdays,
            "company_name": company_name,
            "company_info": company_info,
            "banner": banner,
            "page": page,
            "total_pages": total_pages,
            "status_labels": STATUS_LABELS,
        },
    )


@router.get("/attendance/sheet/export")
async def attendance_sheet_export(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    now = now_kst()
    year = int(request.query_params.get("year", now.year))
    month = int(request.query_params.get("month", now.month))
    company_id = user["company_id"]
    fmt = request.query_params.get("format", "xlsx")

    rows, days_in_month = _build_client_sheet(conn, year, month, company_id)

    company_row = conn.execute("SELECT name FROM client_companies WHERE id=?", (company_id,)).fetchone()
    company_name = company_row["name"] if company_row else ""

    ym_str = f"{year}년 {month}월"
    filename_base = f"출근부_{year}{month:02d}_{company_name}"

    if fmt == "pdf":
        html = templates.get_template("hr/attendance_sheet_pdf.html").render(
            rows=rows, year=year, month=month, days_in_month=days_in_month,
            company_name=company_name, ym_str=ym_str,
            company_info=COMPANY_INFO, status_labels=STATUS_LABELS,
        )
        return pdf_response(html, f"{filename_base}.pdf")

    from openpyxl.styles import Font, Alignment
    from core.export import esc_cell, HEADER_FONT, HEADER_FILL, CELL_FONT, THIN_BORDER
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "출근부"

    sum_cols = ["출근", "지각", "조퇴", "결근", "휴일", "연장", "총근무"]
    title_text = f"출근부 - {ym_str} ({company_name})"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_in_month + 2 + len(sum_cols))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="맑은 고딕", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 3
    headers = ["사번", "이름"] + [str(d) for d in range(1, days_in_month + 1)] + sum_cols
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    status_short = {"normal": "O", "late": "△", "early_leave": "▽", "late_early": "△▽", "absent": "X", "holiday": "-"}
    for ri, r in enumerate(rows, header_row + 1):
        emp = r["employee"]
        ws.cell(row=ri, column=1, value=esc_cell(emp.get("employee_no", ""))).font = CELL_FONT
        ws.cell(row=ri, column=1).border = THIN_BORDER
        ws.cell(row=ri, column=2, value=esc_cell(emp["name"])).font = CELL_FONT
        ws.cell(row=ri, column=2).border = THIN_BORDER
        for di, cell_data in enumerate(r["cells"], 3):
            cell = ws.cell(row=ri, column=di)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.font = CELL_FONT
            if cell_data:
                cell.value = status_short.get(cell_data["status"], "")
            else:
                cell.value = ""
        base_col = days_in_month + 3
        vals = [
            r["work_days"], r["late"], r["early_leave"], r["absent"],
            r["holiday_days"],
            f"{r['ot_hours']}:{r['ot_mins_r']:02d}",
            f"{r['total_hours']}:{r['total_mins_r']:02d}",
        ]
        for offset, val in enumerate(vals):
            cell = ws.cell(row=ri, column=base_col + offset, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 10
    for ci in range(3, days_in_month + 3):
        col_letter = get_column_letter(ci)
        ws.column_dimensions[col_letter].width = 4
    for ci in range(days_in_month + 3, days_in_month + 3 + len(sum_cols)):
        col_letter = get_column_letter(ci)
        ws.column_dimensions[col_letter].width = 9

    return excel_response(wb, f"{filename_base}.xlsx")


APPEAL_SORTS = {"date": "a.work_date", "worker": "e.name", "appeal": "a.appeal"}
APPEAL_STATUS_LABEL = {"pending": "대기", "approved": "인정", "rejected": "불인정"}


@router.get("/attendance/appeals", response_class=HTMLResponse)
async def appeal_list(
    request: Request,
    sort: str = "",
    dir: str = "",
    export: str = "",
    status: str = "",
    date_from: str = "",
    date_to: str = "",
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    order = APPEAL_SORTS.get(sort)
    order_sql = f"ORDER BY {order} {'DESC' if dir == 'desc' else 'ASC'}" if order else \
        "ORDER BY CASE a.appeal WHEN 'pending' THEN 0 ELSE 1 END, a.work_date DESC"
    sql = ("SELECT a.*, e.name AS worker_name "
           "FROM attendance a JOIN employees e ON a.employee_id=e.id "
           "WHERE e.company_id=? AND a.appeal IS NOT NULL")
    params = [company_id]
    if status:
        sql += " AND a.appeal = ?"
        params.append(status)
    if date_from:
        sql += " AND a.work_date >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND a.work_date <= ?"
        params.append(date_to)
    rows = conn.execute(sql + " " + order_sql, params).fetchall()
    appeals = []
    for r in rows:
        d = dict(r)
        d["appeal_type_label"] = APPEAL_TYPE_LABELS.get(d.get("appeal_type", ""), d.get("appeal_type", ""))
        appeals.append(d)

    if export == "csv":
        data = [[a["work_date"], a["worker_name"], a["appeal_type_label"], a["appeal_reason"] or "-",
                 APPEAL_STATUS_LABEL.get(a["appeal"], a["appeal"])] for a in appeals]
        return csv_response(["날짜", "근로자", "사유구분", "소명내용", "상태"], data, "appeals.csv")

    pending_cnt = sum(1 for a in appeals if a["appeal"] == "pending")
    return templates.TemplateResponse(
        request=request, name="client/appeal_list.html", context={
            "request": request,
            "page_title": "소명 관리",
            "user_name": user["name"],
            "appeals": appeals,
            "pending_cnt": pending_cnt,
            "sort": sort,
            "dir": dir,
            "status_filter": status,
            "date_from": date_from,
            "date_to": date_to,
        }
    )


@router.get("/attendance/appeals/{att_id}", response_class=HTMLResponse)
async def appeal_detail(
    request: Request,
    att_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT a.*, e.name AS worker_name, u.name AS reviewer_name "
        "FROM attendance a JOIN employees e ON a.employee_id=e.id "
        "LEFT JOIN users u ON a.appeal_reviewed_by=u.id "
        "WHERE a.id=? AND e.company_id=? AND a.appeal IS NOT NULL",
        (att_id, company_id),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/client/attendance/appeals", status_code=303)
    a = dict(row)
    a["appeal_type_label"] = APPEAL_TYPE_LABELS.get(a.get("appeal_type", ""), a.get("appeal_type", ""))
    return templates.TemplateResponse(
        request=request, name="client/appeal_detail.html", context={
            "request": request,
            "page_title": "소명 상세",
            "user_name": user["name"],
            "a": a,
        }
    )


@router.post("/attendance/{att_id}/appeal-review")
async def appeal_review(
    request: Request,
    att_id: int,
    action: str = Form(...),
    reject_reason: str = Form(""),
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT a.id, a.appeal_type FROM attendance a JOIN employees e ON a.employee_id=e.id "
        "WHERE a.id=? AND e.company_id=? AND a.appeal='pending'",
        (att_id, company_id),
    ).fetchone()
    if not row:
        return RedirectResponse(url="/client/attendance/appeals", status_code=303)

    result = "approved" if action == "approve" else "rejected"
    if result == "approved":
        note = f"소명 승인 ({row['appeal_type'] or ''})"
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=?, "
            "status='normal', notes=CASE WHEN notes IS NULL OR notes='' THEN ? "
            "ELSE notes || ' / ' || ? END WHERE id=?",
            (result, user["id"], now_kst().isoformat(), note, note, att_id),
        )
    else:
        conn.execute(
            "UPDATE attendance SET appeal=?, appeal_reviewed_by=?, appeal_at=?, "
            "appeal_reject_reason=? WHERE id=?",
            (result, user["id"], now_kst().isoformat(), reject_reason.strip(), att_id),
        )
    conn.commit()
    return RedirectResponse(url="/client/attendance/appeals", status_code=303)


@router.get("/attendance/{employee_id}", response_class=HTMLResponse)
async def attendance_detail(
    request: Request,
    employee_id: int,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    month = request.query_params.get("month", now_kst().strftime("%Y-%m"))
    company_id = user["company_id"]
    worker = get_worker_detail(conn, employee_id)
    if not worker or worker.get("company_id") != company_id:
        return RedirectResponse(url="/client/attendance", status_code=303)
    records = get_worker_month_records(conn, employee_id, month)
    leaves = get_worker_month_leaves(conn, employee_id, month)
    summary = worker_month_summary(records, leaves)
    sched = get_schedule(conn, company_id)
    cal = build_calendar_data(records, month, leaves, sched)
    merged = merge_records_leaves(records, leaves)
    prev_m, next_m = prev_next_month(month)

    return templates.TemplateResponse(
        request=request,
        name="client/attendance_detail.html",
        context={
            "request": request,
            "page_title": f"{worker['name']} 근태 상세",
            "user_name": user["name"],
            "worker": worker,
            "records": merged,
            "summary": summary,
            "cal": cal,
            "current_month": month,
            "prev_month": prev_m,
            "next_month": next_m,
            "back_url": "/client/attendance",
        }
    )


@router.get("/settings/appeal", response_class=HTMLResponse)
async def appeal_settings_form(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    row = conn.execute(
        "SELECT appeal_guide, appeal_types, appeal_require_file FROM client_companies WHERE id=?",
        (company_id,),
    ).fetchone()
    guide = row["appeal_guide"] or "" if row else ""
    types_json = row["appeal_types"] if row else '["transit","medical","accident","disability","family","etc"]'
    enabled_types = json.loads(types_json) if types_json else []
    require_file = (row["appeal_require_file"] or 0) if row else 0

    return templates.TemplateResponse(
        request=request, name="client/appeal_settings.html", context={
            "request": request,
            "page_title": "소명 정책 설정",
            "user_name": user["name"],
            "guide": guide,
            "enabled_types": enabled_types,
            "all_types": APPEAL_TYPE_LABELS,
            "require_file": require_file,
        }
    )


@router.post("/settings/appeal")
async def appeal_settings_save(
    request: Request,
    user: dict = Depends(require_client_with_company()),
    conn=Depends(get_db),
):
    company_id = user["company_id"]
    form = await request.form()
    guide = form.get("appeal_guide", "").strip()
    selected = form.getlist("appeal_types")
    require_file = 1 if form.get("appeal_require_file") else 0

    conn.execute(
        "UPDATE client_companies SET appeal_guide=?, appeal_types=?, appeal_require_file=? WHERE id=?",
        (guide, json.dumps(selected), require_file, company_id),
    )
    conn.commit()
    return RedirectResponse(url="/client/settings/appeal", status_code=303)
