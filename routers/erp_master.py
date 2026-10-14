import io
import json
import pathlib
import sqlite3
from datetime import date, timedelta
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from starlette.responses import JSONResponse
from core.deps import get_db, templates
from core.task_link import stamp_task
from core.erp_log import log_activity
from core.tz import today_kst
from core import audit, stock
from core.books import require_book, book_users
from core import credit
from core import partner as ptn
from routers._helpers import save_upload
from routers.erp_purchase import PO_STATUS
from routers.erp_sales import QT_STATUS, SO_STATUS
from core.crm import ACT_KINDS, STAGES, STAGE_CLASS

router = APIRouter()

PER_PAGE = 20

ITEM_TYPES = {
    "material": "원재료",
    "product": "완제품",
    "semi": "반제품",
    "consumable": "소모품",
}


@router.get("/items", response_class=HTMLResponse)
async def item_list(request: Request, q: str = "", item_type: str = "", page: int = 1,
                    u: dict = Depends(require_book), conn=Depends(get_db)):
    where = "WHERE company_id=? AND is_active=1"
    params = [u["book"]]
    if q:
        where += " AND (item_code LIKE ? OR item_name LIKE ?)"
        params += [f"%{q}%", f"%{q}%"]
    if item_type:
        where += " AND item_type=?"
        params.append(item_type)

    total = conn.execute(f"SELECT COUNT(*) FROM items {where}", params).fetchone()[0]
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
    page = max(1, min(page, total_pages))

    sql = f"""SELECT *, (SELECT name FROM partners WHERE id=items.partner_id) AS partner_name FROM items
              {where} ORDER BY item_code LIMIT {PER_PAGE} OFFSET {(page-1)*PER_PAGE}"""
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    for r in rows:
        r["item_type_label"] = ITEM_TYPES.get(r["item_type"], r["item_type"])
    return templates.TemplateResponse(request=request, name="erp/item_list.html", context={
        "request": request, "page_title": "품목 관리",
        "user_name": u["name"],
        "items": rows, "q": q, "item_type": item_type,
        "item_types": ITEM_TYPES,
        "page": page, "total_pages": total_pages, "total_count": total,
        "base_url": f"/items?q={q}&item_type={item_type}",
    })


def _suppliers(conn, book):
    return [dict(r) for r in conn.execute(
        """SELECT id, code, name FROM partners WHERE company_id=? AND is_active=1
           AND partner_type IN ('supplier','both') ORDER BY code""", (book,)
    ).fetchall()]


@router.get("/items/new", response_class=HTMLResponse)
async def item_form_new(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    return templates.TemplateResponse(request=request, name="erp/item_form.html", context={
        "request": request, "page_title": "품목 등록",
        "user_name": u["name"],
        "item": None, "item_types": ITEM_TYPES, "suppliers": _suppliers(conn, u["book"]),
    })


@router.get("/items/{item_id}/edit", response_class=HTMLResponse)
async def item_form_edit(request: Request, item_id: int,
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    row = conn.execute("SELECT * FROM items WHERE id=? AND company_id=?", (item_id, u["book"])).fetchone()
    if not row:
        return RedirectResponse("/items", status_code=303)
    return templates.TemplateResponse(request=request, name="erp/item_form.html", context={
        "request": request, "page_title": "품목 수정",
        "user_name": u["name"],
        "item": dict(row), "item_types": ITEM_TYPES, "suppliers": _suppliers(conn, u["book"]),
    })


@router.post("/items/save", response_class=HTMLResponse)
async def item_save(request: Request,
                    item_id: int = Form(0),
                    item_code: str = Form(...),
                    item_name: str = Form(...),
                    item_type: str = Form("material"),
                    unit: str = Form("EA"),
                    unit_price: int = Form(0),
                    safety_stock: int = Form(0),
                    category: str = Form(""),
                    purch_group: str = Form(""),
                    partner_id: int = Form(0),
                    lead_days: int = Form(0),
                    moq: int = Form(0),
                    order_unit: int = Form(1),
                    reorder_pt: int = Form(0),
                    u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    if partner_id and partner_id not in {p["id"] for p in _suppliers(conn, book)}:
        return JSONResponse({"error": "주거래처를 찾을 수 없습니다."}, status_code=400)
    if min(lead_days, moq, reorder_pt) < 0 or order_unit < 1:
        return JSONResponse({"error": "조달기간, 최소발주량, 재발주점은 0 이상, 발주단위는 1 이상입니다."}, status_code=400)
    purch = (purch_group.strip(), partner_id or None, lead_days, moq, order_unit, reorder_pt)
    is_new = not bool(item_id)
    if item_id:
        conn.execute(
            """UPDATE items SET item_code=?, item_name=?, item_type=?, unit=?, unit_price=?, safety_stock=?, category=?,
               purch_group=?, partner_id=?, lead_days=?, moq=?, order_unit=?, reorder_pt=?
               WHERE id=? AND company_id=?""",
            (item_code, item_name, item_type, unit, unit_price, safety_stock, category, *purch, item_id, book),
        )
    else:
        task_id = stamp_task(conn, request, u["id"], "master")
        cur = conn.execute(
            """INSERT INTO items (company_id, item_code, item_name, item_type, unit, unit_price, safety_stock, category,
                   purch_group, partner_id, lead_days, moq, order_unit, reorder_pt, created_by, task_id)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (book, item_code, item_name, item_type, unit, unit_price, safety_stock, category, *purch, u["id"], task_id),
        )
        item_id = cur.lastrowid
    if is_new:
        log_activity(conn, u["id"], "create", "master", item_id, item_code, f"품목 등록 {item_name} ({item_code})")
    else:
        log_activity(conn, u["id"], "update", "master", item_id, item_code, f"품목 수정 {item_name} ({item_code})")
    conn.commit()
    return RedirectResponse("/items", status_code=303)



ITEM_IMPORT_COLS = [
    ("품목코드", "item_code"), ("품목명", "item_name"), ("유형", "item_type"),
    ("단위", "unit"), ("단가", "unit_price"), ("안전재고", "safety_stock"), ("분류", "category"),
]


@router.get("/items/import", response_class=HTMLResponse)
async def item_import_form(request: Request, u: dict = Depends(require_book)):
    return templates.TemplateResponse(request=request, name="erp/item_import.html", context={
        "request": request, "page_title": "품목 엑셀 등록", "user_name": u["name"],
        "cols": [c[0] for c in ITEM_IMPORT_COLS], "item_types": ITEM_TYPES, "errors": [], "result": None,
    })


@router.get("/items/import/template")
async def item_import_template(u: dict = Depends(require_book), conn=Depends(get_db)):
    rows = conn.execute("SELECT * FROM items WHERE company_id=? ORDER BY item_code", (u["book"],)).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "품목"
    ws.append([c[0] for c in ITEM_IMPORT_COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        line = []
        for _, key in ITEM_IMPORT_COLS:
            v = ITEM_TYPES.get(r["item_type"], r["item_type"]) if key == "item_type" else r[key]
            line.append(v)
        ws.append(line)
    # 품목코드 앞자리 0 보존용 텍스트 서식
    for c in ws["A"]:
        c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('품목등록양식.xlsx')}"},
    )


@router.post("/items/import", response_class=HTMLResponse)
async def item_import(request: Request, file: UploadFile = File(...),
                      u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    ctx = {"request": request, "page_title": "품목 엑셀 등록", "user_name": u["name"],
           "cols": [c[0] for c in ITEM_IMPORT_COLS], "item_types": ITEM_TYPES, "errors": [], "result": None}
    try:
        ws = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True).active
        sheet = list(ws.iter_rows(values_only=True))
    except Exception:
        ctx["errors"] = ["엑셀(xlsx) 파일을 읽을 수 없습니다."]
        return templates.TemplateResponse(request=request, name="erp/item_import.html", context=ctx)

    head = [str(h).strip() if h is not None else "" for h in (sheet[0] if sheet else [])]
    if "품목명" not in head:
        ctx["errors"] = ["필수 열이 없습니다: 품목명"]
        return templates.TemplateResponse(request=request, name="erp/item_import.html", context=ctx)

    types = {**{v: k for k, v in ITEM_TYPES.items()}, **{k: k for k in ITEM_TYPES}}
    existing = {r["item_code"]: r["id"] for r in conn.execute(
        "SELECT id, item_code FROM items WHERE company_id=?", (book,))}
    last = conn.execute(
        "SELECT item_code FROM items WHERE company_id=? AND item_code GLOB '[0-9][0-9][0-9][0-9][0-9]' "
        "ORDER BY item_code DESC LIMIT 1", (book,),
    ).fetchone()
    seq = int(last["item_code"]) if last else 0

    errors, rows, seen = [], [], set()
    for no, raw in enumerate(sheet[1:], 2):
        cells = {}
        for h, key in ITEM_IMPORT_COLS:
            v = raw[head.index(h)] if h in head and head.index(h) < len(raw) else None
            # 숫자 셀로 저장된 코드의 소수점 제거
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            cells[key] = "" if v is None else str(v).strip()
        if not any(cells.values()):
            continue
        err = []
        if not cells["item_name"]:
            err.append("품목명 누락")
        item_type = types.get(cells["item_type"] or "material")
        if not item_type:
            err.append(f"유형 오류({cells['item_type']})")
        nums = {}
        for key, label in (("unit_price", "단가"), ("safety_stock", "안전재고")):
            v = cells[key].replace(",", "") or "0"
            if not v.isdigit():
                err.append(f"{label} 숫자 오류({cells[key]})")
            nums[key] = int(v) if v.isdigit() else 0
        if cells["item_code"] in seen:
            err.append(f"파일 내 코드 중복({cells['item_code']})")
        if cells["item_code"]:
            seen.add(cells["item_code"])
        if err:
            errors.append(f"{no}행: {', '.join(err)}")
            continue
        rows.append((cells, item_type, nums))

    if errors:
        ctx["errors"] = errors
        return templates.TemplateResponse(request=request, name="erp/item_import.html", context=ctx)

    added = updated = 0
    for cells, item_type, nums in rows:
        unit = cells["unit"] or "EA"
        iid = existing.get(cells["item_code"])
        if iid:
            conn.execute(
                """UPDATE items SET item_name=?, item_type=?, unit=?, unit_price=?, safety_stock=?, category=?
                   WHERE id=?""",
                (cells["item_name"], item_type, unit, nums["unit_price"], nums["safety_stock"], cells["category"], iid),
            )
            updated += 1
        else:
            code = cells["item_code"]
            if not code:
                seq += 1
                while f"{seq:05d}" in existing:
                    seq += 1
                code = f"{seq:05d}"
            existing[code] = conn.execute(
                """INSERT INTO items (company_id, item_code, item_name, item_type, unit, unit_price,
                   safety_stock, category, created_by) VALUES (?,?,?,?,?,?,?,?,?)""",
                (book, code, cells["item_name"], item_type, unit, nums["unit_price"], nums["safety_stock"],
                 cells["category"], u["id"]),
            ).lastrowid
            added += 1
    log_activity(conn, u["id"], "create", "master", 0, "", f"품목 엑셀 등록 신규 {added}건, 수정 {updated}건")
    conn.commit()
    ctx["result"] = {"added": added, "updated": updated}
    return templates.TemplateResponse(request=request, name="erp/item_import.html", context=ctx)


PARTNER_TYPES = {
    "supplier": "공급처",
    "customer": "고객",
    "both": "공급/고객",
}

FILE_KINDS = {"bizreg": "사업자등록증", "bankbook": "통장사본", "contract": "계약서", "etc": "기타"}
CONTACT_FIELDS = ("dept", "position", "phone", "mobile", "email", "fax", "birthday", "memo")
REF_TABLES_CONTACT = ("crm_activities", "quotations", "sales_orders", "purchase_orders", "shipments", "tax_invoices")
MERGE_PARTNER_TABLES = ("quotations", "sales_orders", "purchase_orders", "tax_invoices", "crm_opps",
                        "crm_activities", "crm_leads")

PARTNER_SORTS = {
    "code": "p.code", "name": "p.name", "type": "p.partner_type", "group": "g.name", "grade": "gr.name",
    "rep": "ru.name", "trade_status": "p.trade_status", "pay_terms": "p.pay_terms",
}


IMPORT_COLS = [
    ("거래처코드", "code"), ("거래처명", "name"), ("유형", "partner_type"), ("사업자번호", "biz_no"),
    ("대표자", "representative"), ("업태", "biz_type"), ("종목", "biz_item"), ("과세유형", "tax_type"),
    ("주소", "address"), ("전화", "phone"), ("팩스", "fax"), ("이메일", "email"),
    ("담당자명", "contact_name"), ("담당자연락처", "contact_phone"), ("담당자이메일", "contact_email"),
    ("은행", "bank_name"), ("계좌번호", "bank_account"), ("예금주", "bank_holder"),
    ("거래처그룹", "group"), ("여신한도", "credit_limit"), ("여신기간", "credit_days"),
    ("판매단가그룹", "sale_pg"), ("구매단가그룹", "buy_pg"), ("비고", "notes"),
]

# 사업자번호 10자리 검증번호 가중치. 엑셀 등록은 reg_kind 구분이 없어 기존 방식(10/13자리 통과) 유지
BIZ_W = (1, 3, 7, 1, 3, 7, 1, 3, 5)


def norm_biz_no(v):
    d = "".join(c for c in v if c.isdigit())
    if not d:
        return ""
    if len(d) == 13:
        return v.strip()
    if len(d) != 10:
        return None
    n = sum(int(d[i]) * BIZ_W[i] for i in range(9)) + int(d[8]) * 5 // 10
    if (10 - n % 10) % 10 != int(d[9]):
        return None
    return f"{d[:3]}-{d[3:5]}-{d[5:]}"


def dup_partners(conn, book, biz_no, pid=0):
    d = "".join(c for c in biz_no if c.isdigit())
    if not d:
        return []
    return conn.execute(
        """SELECT id, code, name FROM partners
           WHERE company_id=? AND id!=? AND REPLACE(biz_no, '-', '')=?""", (book, pid, d)
    ).fetchall()


def _masters(conn, book):
    return {
        "groups": conn.execute("SELECT id, name FROM partner_groups WHERE company_id=? ORDER BY name",
                               (book,)).fetchall(),
        "pgs": conn.execute("SELECT id, name, rate FROM price_groups WHERE company_id=? ORDER BY name",
                            (book,)).fetchall(),
        "fields": conn.execute("SELECT id, label, kind FROM partner_fields WHERE company_id=? ORDER BY sort, id",
                               (book,)).fetchall(),
        "grades": conn.execute("SELECT id, name, min_sales FROM partner_grades WHERE company_id=? ORDER BY min_sales DESC",
                               (book,)).fetchall(),
        "reps": book_users(conn, book),
    }


# 최근 1년 매출 세금계산서 공급가액 기준(수정, 취소분 차감 반영). 기준금액 높은 등급부터 적용
def regrade(conn, book, pid=None):
    since = (today_kst() - timedelta(days=365)).isoformat()
    grades = conn.execute("SELECT id, min_sales FROM partner_grades WHERE company_id=? ORDER BY min_sales DESC",
                          (book,)).fetchall()
    sql = """SELECT p.id, COALESCE((SELECT SUM(t.supply_amt) FROM tax_invoices t
                  WHERE t.partner_id=p.id AND t.company_id=p.company_id AND t.kind='매출'
                    AND t.issue_date>=?), 0) AS sales
             FROM partners p WHERE p.company_id=? AND p.grade_fixed=0"""
    args = [since, book]
    if pid:
        sql += " AND p.id=?"
        args.append(pid)
    rows = conn.execute(sql, args).fetchall()
    for r in rows:
        gid = next((g["id"] for g in grades if r["sales"] >= g["min_sales"]), None)
        conn.execute("UPDATE partners SET grade_id=? WHERE id=?", (gid, r["id"]))
    return len(rows)


def _own(conn, table, rid, book):
    if rid and conn.execute(f"SELECT 1 FROM {table} WHERE id=? AND company_id=?", (rid, book)).fetchone():
        return rid
    return None


def _contact_referenced(conn, cid):
    return any(conn.execute(f"SELECT 1 FROM {t} WHERE contact_id=?", (cid,)).fetchone() for t in REF_TABLES_CONTACT)


def _bank_summary(conn, pid):
    rows = conn.execute(
        "SELECT bank, account_no, holder FROM partner_banks WHERE partner_id=? ORDER BY is_default DESC, sort, id",
        (pid,)).fetchall()
    return "; ".join(f"{r['bank']} {r['account_no']}" + (f" ({r['holder']})" if r["holder"] else "") for r in rows)


# 변경이력 기록용 표시문자열 변환. old/new 스냅샷을 같은 기준으로 만들어 넘겨야 함
def _partner_display(conn, book, row, m):
    reps = {r["id"]: r["name"] for r in m["reps"]}
    groups = {g["id"]: g["name"] for g in m["groups"]}
    return {
        "code": row.get("code") or "", "name": row.get("name") or "",
        "kind": ptn.KINDS.get(row.get("kind"), row.get("kind") or ""),
        "reg_kind": ptn.REG_KINDS.get(row.get("reg_kind"), row.get("reg_kind") or ""),
        "biz_no": row.get("biz_no") or "", "tax_type": row.get("tax_type") or "",
        "rep_id": reps.get(row.get("rep_id"), ""), "sub_rep_id": reps.get(row.get("sub_rep_id"), ""),
        "credit_limit": f"{row.get('credit_limit') or 0:,}", "credit_days": str(row.get("credit_days") or 0),
        "pay_terms": ptn.PAY_TERMS.get(row.get("pay_terms"), ""),
        "trade_status": ptn.TRADE_STATUS.get(row.get("trade_status"), ""),
        "biz_status": ptn.BIZ_STATUS.get(row.get("biz_status"), ""),
        "group_id": groups.get(row.get("group_id"), ""),
        "is_active": "사용" if row.get("is_active") else "사용중지",
        "address": row.get("address") or "",
        "bank": _bank_summary(conn, row["id"]) if row.get("id") else "",
    }


def _partner_filter(book, q, ptype, group, rep, grade, trade_status, include_inactive):
    where = "WHERE p.company_id=?"
    params = [book]
    if not include_inactive:
        where += " AND p.is_active=1 AND p.merged_into IS NULL"
    if q:
        dq = ptn.digits(q)
        if len(dq) >= 3:
            where += """ AND (p.code LIKE ? OR p.name LIKE ? OR p.biz_no LIKE ?
                         OR REPLACE(p.phone,'-','') LIKE ?
                         OR EXISTS (SELECT 1 FROM partner_contacts c WHERE c.partner_id=p.id
                                    AND (REPLACE(c.phone,'-','') LIKE ? OR REPLACE(c.mobile,'-','') LIKE ?)))"""
            params += [f"%{q}%", f"%{q}%", f"%{q}%", f"%{dq}%", f"%{dq}%", f"%{dq}%"]
        else:
            where += " AND (p.code LIKE ? OR p.name LIKE ? OR p.biz_no LIKE ?)"
            params += [f"%{q}%", f"%{q}%", f"%{q}%"]
    if ptype:
        where += " AND p.partner_type=?"
        params.append(ptype)
    if group:
        where += " AND p.group_id=?"
        params.append(group)
    if rep:
        where += " AND p.rep_id=?"
        params.append(rep)
    if grade:
        where += " AND p.grade_id=?"
        params.append(grade)
    if trade_status:
        where += " AND p.trade_status=?"
        params.append(trade_status)
    return where, params


@router.get("/partners", response_class=HTMLResponse)
async def partner_list(request: Request, q: str = "", ptype: str = "", group: int = 0, page: int = 1,
                       include_inactive: int = 0, rep: int = 0, grade: int = 0, trade_status: str = "",
                       alert: str = "", sort: str = "code", dir: str = "asc",
                       u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    where, params = _partner_filter(book, q, ptype, group, rep, grade, trade_status, include_inactive)
    order = PARTNER_SORTS.get(sort, "p.code")
    updown = "DESC" if dir == "desc" else "ASC"

    base_sql = f"""SELECT p.*, g.name AS group_name, gr.name AS grade_name, ru.name AS rep_name,
                     (SELECT name FROM partner_contacts WHERE partner_id=p.id AND is_active=1
                      ORDER BY is_primary DESC, sort, id LIMIT 1) AS contact
              FROM partners p
              LEFT JOIN partner_groups g ON g.id=p.group_id
              LEFT JOIN partner_grades gr ON gr.id=p.grade_id
              LEFT JOIN users ru ON ru.id=p.rep_id
              {where} ORDER BY {order} {updown}, p.code"""

    alert_map = ptn.alerts(conn, book, today_kst().isoformat())

    if alert:
        rows = [dict(r) for r in conn.execute(base_sql, params).fetchall()]
        rows = [r for r in rows if (alert_map.get(r["id"]) if alert == "1" else alert in alert_map.get(r["id"], []))]
        total = len(rows)
        total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
        page = max(1, min(page, total_pages))
        rows = rows[(page - 1) * PER_PAGE:page * PER_PAGE]
    else:
        total = conn.execute(f"SELECT COUNT(*) FROM partners p {where}", params).fetchone()[0]
        total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
        page = max(1, min(page, total_pages))
        rows = [dict(r) for r in conn.execute(
            f"{base_sql} LIMIT {PER_PAGE} OFFSET {(page - 1) * PER_PAGE}", params).fetchall()]

    can_edit = ptn.can_edit(u)
    for r in rows:
        r["type_label"] = PARTNER_TYPES.get(r.get("partner_type", ""), "공급처")
        r["trade_label"] = ptn.TRADE_STATUS.get(r["trade_status"], "")
        r["trade_class"] = ptn.TRADE_CLASS.get(r["trade_status"], "")
        r["pay_label"] = ptn.PAY_TERMS.get(r["pay_terms"], "") + (
            f" {r['pay_days']}일" if r["pay_terms"] in ("days", "eom") else "")
        r["alert_codes"] = alert_map.get(r["id"], [])
        r["biz_no_display"] = r["biz_no"] if (can_edit or r["reg_kind"] == "biz") else ptn.mask_reg(r["biz_no"])

    m = _masters(conn, book)
    return templates.TemplateResponse(request=request, name="erp/partner_list.html", context={
        "request": request, "page_title": "거래처 관리", "user_name": u["name"], "user_id": u["id"],
        "can_edit": can_edit,
        "partners": rows, "q": q, "ptype": ptype, "group": group, "include_inactive": include_inactive,
        "rep": rep, "grade": grade, "trade_status": trade_status, "alert": alert, "sort": sort, "dir": dir,
        "partner_types": PARTNER_TYPES, "trade_statuses": ptn.TRADE_STATUS, "alert_labels": ptn.ALERTS,
        "groups": m["groups"], "reps": m["reps"], "grades": m["grades"],
        "page": page, "total_pages": total_pages, "total_count": total,
        "base_url": f"/partners?q={quote(q)}&ptype={ptype}&group={group}&rep={rep}&grade={grade}"
                    f"&trade_status={trade_status}&alert={alert}&sort={sort}&dir={dir}"
                    f"&include_inactive={include_inactive}",
    })


@router.get("/partners/excel")
async def partner_excel(q: str = "", ptype: str = "", group: int = 0, include_inactive: int = 0,
                        rep: int = 0, grade: int = 0, trade_status: str = "",
                        u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    can_edit = ptn.can_edit(u)
    where, params = _partner_filter(book, q, ptype, group, rep, grade, trade_status, include_inactive)
    rows = conn.execute(
        f"""SELECT p.code, p.name, p.partner_type, p.reg_kind, g.name AS group_name, gr.name AS grade_name,
                   ru.name AS rep_name, p.biz_no, p.trade_status, p.pay_terms, p.pay_days, p.phone
            FROM partners p
            LEFT JOIN partner_groups g ON g.id=p.group_id
            LEFT JOIN partner_grades gr ON gr.id=p.grade_id
            LEFT JOIN users ru ON ru.id=p.rep_id
            {where} ORDER BY p.code""", params,
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "거래처"
    ws.append(["거래처코드", "거래처명", "유형", "그룹", "등급", "담당사원", "사업자번호", "거래상태", "결제조건", "전화"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        pay = ptn.PAY_TERMS.get(r["pay_terms"], "") + (f" {r['pay_days']}일" if r["pay_terms"] in ("days", "eom") else "")
        biz_no = r["biz_no"] if (can_edit or r["reg_kind"] == "biz") else ptn.mask_reg(r["biz_no"])
        ws.append([r["code"], r["name"], PARTNER_TYPES.get(r["partner_type"], r["partner_type"]), r["group_name"],
                   r["grade_name"], r["rep_name"], biz_no, ptn.TRADE_STATUS.get(r["trade_status"], ""),
                   pay, r["phone"]])
    for col in ("A", "G"):
        for c in ws[col]:
            c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('거래처목록.xlsx')}"},
    )


@router.post("/partners/bulk")
async def partner_bulk(ids: list[int] = Form(...), action: str = Form(...), value: str = Form(""),
                       u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    if action not in ("rep", "group", "trade_status"):
        return JSONResponse({"error": "잘못된 처리입니다."}, status_code=400)
    valid_ids = [r["id"] for r in conn.execute(
        f"SELECT id FROM partners WHERE company_id=? AND id IN ({','.join('?' * len(ids))})", (book, *ids))]
    if not valid_ids:
        return JSONResponse({"error": "대상 거래처가 없습니다."}, status_code=400)

    m = _masters(conn, book)
    if action == "rep":
        new_val = int(value) if value else None
        if new_val and new_val not in {r["id"] for r in m["reps"]}:
            return JSONResponse({"error": "담당사원을 확인해 주세요."}, status_code=400)
        field, flabel = "rep_id", "담당사원"
        label = next((r["name"] for r in m["reps"] if r["id"] == new_val), "미지정")
    elif action == "group":
        new_val = _own(conn, "partner_groups", int(value), book) if value else None
        field, flabel = "group_id", "거래처그룹"
        label = next((g["name"] for g in m["groups"] if g["id"] == new_val), "미지정")
    else:
        if value not in ptn.TRADE_STATUS:
            return JSONResponse({"error": "거래상태를 확인해 주세요."}, status_code=400)
        field, flabel = "trade_status", "거래상태"
        new_val, label = value, ptn.TRADE_STATUS[value]

    changed = 0
    for pid in valid_ids:
        old = conn.execute(f"SELECT {field} FROM partners WHERE id=?", (pid,)).fetchone()[0]
        if old == new_val:
            continue
        conn.execute(f"UPDATE partners SET {field}=? WHERE id=?", (new_val, pid))
        if field == "rep_id":
            old_label = next((r["name"] for r in m["reps"] if r["id"] == old), "미지정")
        elif field == "group_id":
            old_label = next((g["name"] for g in m["groups"] if g["id"] == old), "미지정")
        else:
            old_label = ptn.TRADE_STATUS.get(old, "")
        conn.execute(
            "INSERT INTO partner_history (company_id, partner_id, field, old, new, user_id) VALUES (?,?,?,?,?,?)",
            (book, pid, field, old_label, label, u["id"]),
        )
        changed += 1
    log_activity(conn, u["id"], "update", "master", 0, "", f"거래처 일괄변경 {flabel} -> {label} ({changed}건)")
    conn.commit()
    return RedirectResponse("/partners?saved=1", status_code=303)


@router.get("/partners/contacts", response_class=HTMLResponse)
async def partner_contact_list(request: Request, q: str = "", role: str = "", group: int = 0,
                               include_resigned: int = 0, page: int = 1,
                               u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    where = "WHERE p.company_id=?"
    params = [book]
    if not include_resigned:
        where += " AND c.is_active=1"
    if q:
        dq = ptn.digits(q)
        if len(dq) >= 3:
            where += """ AND (c.name LIKE ? OR c.email LIKE ? OR p.name LIKE ?
                         OR REPLACE(c.phone,'-','') LIKE ? OR REPLACE(c.mobile,'-','') LIKE ?)"""
            params += [f"%{q}%", f"%{q}%", f"%{q}%", f"%{dq}%", f"%{dq}%"]
        else:
            where += " AND (c.name LIKE ? OR c.email LIKE ? OR p.name LIKE ?)"
            params += [f"%{q}%", f"%{q}%", f"%{q}%"]
    if role:
        where += " AND (',' || c.roles || ',') LIKE ?"
        params.append(f"%,{role},%")
    if group:
        where += " AND p.group_id=?"
        params.append(group)

    total = conn.execute(
        f"SELECT COUNT(*) FROM partner_contacts c JOIN partners p ON p.id=c.partner_id {where}", params
    ).fetchone()[0]
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
    page = max(1, min(page, total_pages))

    sql = f"""SELECT c.*, p.id AS partner_id, p.name AS partner_name, p.code AS partner_code,
                     (SELECT MAX(act_date) FROM crm_activities WHERE contact_id=c.id) AS last_contact
              FROM partner_contacts c JOIN partners p ON p.id=c.partner_id
              {where} ORDER BY p.name, c.sort, c.id LIMIT {PER_PAGE} OFFSET {(page - 1) * PER_PAGE}"""
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    for r in rows:
        r["role_labels"] = [ptn.ROLES.get(x, x) for x in (r["roles"] or "").split(",") if x]
    m = _masters(conn, book)
    return templates.TemplateResponse(request=request, name="erp/partner_contacts.html", context={
        "request": request, "page_title": "담당자 목록", "user_name": u["name"],
        "contacts": rows, "q": q, "role": role, "group": group, "include_resigned": include_resigned,
        "roles": ptn.ROLES, "groups": m["groups"],
        "page": page, "total_pages": total_pages, "total_count": total,
        "base_url": f"/partners/contacts?q={quote(q)}&role={role}&group={group}&include_resigned={include_resigned}",
    })


@router.get("/partners/contacts/excel")
async def partner_contact_excel(q: str = "", role: str = "", group: int = 0, include_resigned: int = 0,
                                u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    where = "WHERE p.company_id=?"
    params = [book]
    if not include_resigned:
        where += " AND c.is_active=1"
    if q:
        dq = ptn.digits(q)
        if len(dq) >= 3:
            where += """ AND (c.name LIKE ? OR c.email LIKE ? OR p.name LIKE ?
                         OR REPLACE(c.phone,'-','') LIKE ? OR REPLACE(c.mobile,'-','') LIKE ?)"""
            params += [f"%{q}%", f"%{q}%", f"%{q}%", f"%{dq}%", f"%{dq}%"]
        else:
            where += " AND (c.name LIKE ? OR c.email LIKE ? OR p.name LIKE ?)"
            params += [f"%{q}%", f"%{q}%", f"%{q}%"]
    if role:
        where += " AND (',' || c.roles || ',') LIKE ?"
        params.append(f"%,{role},%")
    if group:
        where += " AND p.group_id=?"
        params.append(group)

    rows = conn.execute(
        f"""SELECT p.code, p.name AS partner_name, c.name, c.dept, c.position, c.phone, c.mobile, c.email,
                   c.roles, c.is_active
            FROM partner_contacts c JOIN partners p ON p.id=c.partner_id {where} ORDER BY p.name, c.sort, c.id""",
        params,
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "담당자"
    ws.append(["거래처코드", "거래처명", "이름", "부서", "직위", "전화", "휴대폰", "이메일", "역할", "재직"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        roles = ",".join(ptn.ROLES.get(x, x) for x in (r["roles"] or "").split(",") if x)
        ws.append([r["code"], r["partner_name"], r["name"], r["dept"], r["position"], r["phone"], r["mobile"],
                   r["email"], roles, "재직" if r["is_active"] else "퇴사"])
    for c in ws["A"]:
        c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('담당자목록.xlsx')}"},
    )


@router.get("/partners/new", response_class=HTMLResponse)
async def partner_form_new(request: Request, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    return templates.TemplateResponse(request=request, name="erp/partner_form.html", context={
        "request": request, "page_title": "거래처 등록", "user_name": u["name"],
        "partner": None, "partner_types": PARTNER_TYPES, "extra": {}, "contacts": [], "addrs": [], "banks": [],
        "files": [], "file_kinds": FILE_KINDS,
        "kinds": ptn.KINDS, "reg_kinds": ptn.REG_KINDS, "biz_statuses": ptn.BIZ_STATUS,
        "trade_statuses": ptn.TRADE_STATUS, "pay_terms_opts": ptn.PAY_TERMS, "roles": ptn.ROLES,
        **_masters(conn, u["book"]),
    })


@router.get("/partners/settings", response_class=HTMLResponse)
async def partner_settings(request: Request, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    co = conn.execute("SELECT credit_limit_mode, credit_days_mode FROM client_companies WHERE id=?",
                      (u["book"],)).fetchone()
    return templates.TemplateResponse(request=request, name="erp/partner_settings.html", context={
        "request": request, "page_title": "거래처 설정", "user_name": u["name"],
        "co": co, "modes": credit.MODES, "alert_cfg": ptn.alert_cfg(conn, u["book"]), "alert_labels": ptn.ALERTS,
        **_masters(conn, u["book"]),
    })


@router.post("/partners/settings/credit")
async def partner_credit_mode(credit_limit_mode: str = Form("off"), credit_days_mode: str = Form("off"),
                              u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    if credit_limit_mode in credit.MODES and credit_days_mode in credit.MODES:
        conn.execute("UPDATE client_companies SET credit_limit_mode=?, credit_days_mode=? WHERE id=?",
                     (credit_limit_mode, credit_days_mode, u["book"]))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/settings/alerts")
async def partner_alert_cfg_save(over_limit: int = Form(0), overdue: int = Form(30), no_sale: int = Form(90),
                                 no_act: int = Form(60), closed_bal: int = Form(0),
                                 u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    vals = {"over_limit": 1 if over_limit else 0, "overdue": overdue, "no_sale": no_sale,
            "no_act": no_act, "closed_bal": 1 if closed_bal else 0}
    if any(not (0 <= vals[k] <= 3650) for k in ("overdue", "no_sale", "no_act")):
        return RedirectResponse("/partners/settings?error=alert", status_code=303)
    conn.execute("UPDATE client_companies SET partner_alerts=? WHERE id=?",
                 (json.dumps(vals, ensure_ascii=False), u["book"]))
    conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/groups/add")
async def partner_group_add(name: str = Form(...), u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    name = name.strip()
    if name:
        conn.execute("INSERT OR IGNORE INTO partner_groups (company_id, name) VALUES (?,?)", (u["book"], name))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/groups/{gid}/delete")
async def partner_group_delete(gid: int, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    if _own(conn, "partner_groups", gid, u["book"]):
        conn.execute("UPDATE partners SET group_id=NULL WHERE group_id=?", (gid,))
        conn.execute("DELETE FROM partner_groups WHERE id=?", (gid,))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/price-groups/save")
async def price_group_save(pg_id: int = Form(0), name: str = Form(...), rate: float = Form(0),
                           u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    name = name.strip()
    if not name or not -100 < rate <= 1000:
        return RedirectResponse("/partners/settings?error=pg", status_code=303)
    try:
        if pg_id:
            if _own(conn, "price_groups", pg_id, u["book"]):
                conn.execute("UPDATE price_groups SET name=?, rate=? WHERE id=?", (name, rate, pg_id))
        else:
            conn.execute("INSERT INTO price_groups (company_id, name, rate) VALUES (?,?,?)", (u["book"], name, rate))
    except sqlite3.IntegrityError:
        return RedirectResponse("/partners/settings?error=dup", status_code=303)
    conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/price-groups/{pg_id}/delete")
async def price_group_delete(pg_id: int, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    if _own(conn, "price_groups", pg_id, u["book"]):
        conn.execute("UPDATE partners SET sale_pg_id=NULL WHERE sale_pg_id=?", (pg_id,))
        conn.execute("UPDATE partners SET buy_pg_id=NULL WHERE buy_pg_id=?", (pg_id,))
        conn.execute("DELETE FROM price_groups WHERE id=?", (pg_id,))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/fields/add")
async def partner_field_add(label: str = Form(...), kind: str = Form("text"),
                            u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    label = label.strip()
    if label and kind in ("text", "num"):
        sort = conn.execute("SELECT COALESCE(MAX(sort),0)+1 FROM partner_fields WHERE company_id=?",
                            (u["book"],)).fetchone()[0]
        conn.execute("INSERT INTO partner_fields (company_id, label, kind, sort) VALUES (?,?,?,?)",
                     (u["book"], label, kind, sort))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/fields/{fid}/delete")
async def partner_field_delete(fid: int, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    # 입력값은 extra에 보존. 같은 항목 재등록 시 id 상이로 미노출
    if _own(conn, "partner_fields", fid, u["book"]):
        conn.execute("DELETE FROM partner_fields WHERE id=?", (fid,))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/grades/save")
async def partner_grade_save(name: str = Form(...), min_sales: str = Form("0"),
                             u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    name, amt = name.strip(), min_sales.replace(",", "")
    if not name or not amt.isdigit():
        return RedirectResponse("/partners/settings?error=grade", status_code=303)
    try:
        conn.execute("INSERT INTO partner_grades (company_id, name, min_sales) VALUES (?,?,?)",
                     (u["book"], name, int(amt)))
    except sqlite3.IntegrityError:
        return RedirectResponse("/partners/settings?error=grade_dup", status_code=303)
    conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/grades/{gid}/delete")
async def partner_grade_delete(gid: int, u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    if _own(conn, "partner_grades", gid, u["book"]):
        conn.execute("UPDATE partners SET grade_id=NULL, grade_fixed=0 WHERE grade_id=?", (gid,))
        conn.execute("DELETE FROM partner_grades WHERE id=?", (gid,))
        conn.commit()
    return RedirectResponse("/partners/settings?saved=1", status_code=303)


@router.post("/partners/grades/regrade")
async def partner_regrade(u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    cnt = regrade(conn, u["book"])
    log_activity(conn, u["id"], "update", "master", 0, "", f"고객등급 재산정 {cnt}건")
    conn.commit()
    return RedirectResponse(f"/partners/settings?regraded={cnt}", status_code=303)


@router.get("/api/partners/biz-no")
async def partner_biz_no_check(biz_no: str = "", pid: int = 0,
                               u: dict = Depends(require_book), conn=Depends(get_db)):
    v = norm_biz_no(biz_no)
    if v is None:
        return {"ok": False, "dups": []}
    return {"ok": True, "biz_no": v, "dups": [dict(r) for r in dup_partners(conn, u["book"], v, pid)]}


@router.get("/partners/import", response_class=HTMLResponse)
async def partner_import_form(request: Request, u: dict = Depends(ptn.require_editor)):
    return templates.TemplateResponse(request=request, name="erp/partner_import.html", context={
        "request": request, "page_title": "거래처 엑셀 등록",
        "user_name": u["name"],
        "cols": [c[0] for c in IMPORT_COLS], "errors": [], "result": None,
    })


@router.get("/partners/import/template")
async def partner_import_template(u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    rows = conn.execute(
        """SELECT p.*, g.name AS "group", sp.name AS sale_pg, bp.name AS buy_pg,
                  c.name AS c_name, c.phone AS c_phone, c.email AS c_email,
                  pb.bank AS bank_name_pb, pb.account_no AS bank_account_pb, pb.holder AS bank_holder_pb
           FROM partners p
           LEFT JOIN partner_groups g ON g.id=p.group_id
           LEFT JOIN price_groups sp ON sp.id=p.sale_pg_id
           LEFT JOIN price_groups bp ON bp.id=p.buy_pg_id
           LEFT JOIN partner_contacts c ON c.id=(SELECT id FROM partner_contacts WHERE partner_id=p.id
                                                AND is_active=1 ORDER BY is_primary DESC, sort, id LIMIT 1)
           LEFT JOIN partner_banks pb ON pb.id=(SELECT id FROM partner_banks WHERE partner_id=p.id
                                                ORDER BY is_default DESC, sort, id LIMIT 1)
           WHERE p.company_id=? ORDER BY p.code""", (u["book"],)
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "거래처"
    ws.append([c[0] for c in IMPORT_COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        line = []
        for _, key in IMPORT_COLS:
            v = r[f"{key}_pb"] if key in ("bank_name", "bank_account", "bank_holder") else r[key.replace("contact_", "c_")]
            if key == "partner_type":
                v = PARTNER_TYPES.get(v, v)
            line.append(v)
        ws.append(line)
    # 코드, 사업자번호 앞자리 0 보존용 텍스트 서식
    for col in ("A", "D"):
        for c in ws[col]:
            c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('거래처등록양식.xlsx')}"},
    )


@router.post("/partners/import", response_class=HTMLResponse)
async def partner_import(request: Request, file: UploadFile = File(...),
                         u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    ctx = {"request": request, "page_title": "거래처 엑셀 등록", "user_name": u["name"],
           "cols": [c[0] for c in IMPORT_COLS], "errors": [], "result": None}
    try:
        ws = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True).active
        sheet = list(ws.iter_rows(values_only=True))
    except Exception:
        ctx["errors"] = ["엑셀(xlsx) 파일을 읽을 수 없습니다."]
        return templates.TemplateResponse(request=request, name="erp/partner_import.html", context=ctx)

    head = [str(h).strip() if h is not None else "" for h in (sheet[0] if sheet else [])]
    if "거래처명" not in head:
        ctx["errors"] = ["필수 열이 없습니다: 거래처명"]
        return templates.TemplateResponse(request=request, name="erp/partner_import.html", context=ctx)

    m = _masters(conn, book)
    groups = {g["name"]: g["id"] for g in m["groups"]}
    pgs = {g["name"]: g["id"] for g in m["pgs"]}
    types = {**{v: k for k, v in PARTNER_TYPES.items()}, **{k: k for k in PARTNER_TYPES}}
    existing = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM partners WHERE company_id=?", (book,))}
    last = conn.execute(
        "SELECT code FROM partners WHERE company_id=? AND code GLOB '[0-9][0-9][0-9][0-9][0-9]' ORDER BY code DESC LIMIT 1",
        (book,),
    ).fetchone()
    seq = int(last["code"]) if last else 0

    errors, rows, seen = [], [], set()
    for no, raw in enumerate(sheet[1:], 2):
        cells = {}
        for h, key in IMPORT_COLS:
            v = raw[head.index(h)] if h in head and head.index(h) < len(raw) else None
            # 숫자 셀로 저장된 코드, 사업자번호의 소수점 제거
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            cells[key] = "" if v is None else str(v).strip()
        if not any(cells.values()):
            continue
        err = []
        if not cells["name"]:
            err.append("거래처명 누락")
        ptype = types.get(cells["partner_type"] or "supplier")
        if not ptype:
            err.append(f"유형 오류({cells['partner_type']})")
        tax = cells["tax_type"] or "과세"
        if tax not in ("과세", "영세", "면세"):
            err.append(f"과세유형 오류({tax})")
        biz = norm_biz_no(cells["biz_no"])
        if biz is None:
            err.append(f"사업자번호 오류({cells['biz_no']})")
        cells["biz_no"] = biz or ""
        ids = {}
        for key, src, label in (("group", groups, "거래처그룹"), ("sale_pg", pgs, "판매단가그룹"),
                                ("buy_pg", pgs, "구매단가그룹")):
            ids[key] = src.get(cells[key]) if cells[key] else None
            if cells[key] and not ids[key]:
                err.append(f"없는 {label}({cells[key]})")
        nums = {}
        for key, label in (("credit_limit", "여신한도"), ("credit_days", "여신기간")):
            v = cells[key].replace(",", "") or "0"
            if not v.isdigit():
                err.append(f"{label} 숫자 오류({cells[key]})")
            nums[key] = int(v) if v.isdigit() else 0
        if cells["code"] in seen:
            err.append(f"파일 내 코드 중복({cells['code']})")
        if cells["code"]:
            seen.add(cells["code"])
        if err:
            errors.append(f"{no}행: {', '.join(err)}")
            continue
        rows.append((cells, ptype, tax, ids, nums))

    if errors:
        ctx["errors"] = errors
        return templates.TemplateResponse(request=request, name="erp/partner_import.html", context=ctx)

    added = updated = 0
    for cells, ptype, tax, ids, nums in rows:
        vals = (cells["name"], ptype, cells["biz_no"], cells["representative"], cells["biz_type"],
                cells["biz_item"], tax, cells["address"], cells["phone"], cells["email"], cells["fax"],
                cells["notes"], ids["group"], ids["sale_pg"], ids["buy_pg"], nums["credit_limit"], nums["credit_days"])
        pid = existing.get(cells["code"])
        if pid:
            conn.execute(
                """UPDATE partners SET name=?, partner_type=?, biz_no=?, representative=?, biz_type=?, biz_item=?,
                   tax_type=?, address=?, phone=?, email=?, fax=?,
                   notes=?, group_id=?, sale_pg_id=?, buy_pg_id=?, credit_limit=?, credit_days=? WHERE id=?""",
                (*vals, pid),
            )
            updated += 1
        else:
            code = cells["code"]
            if not code:
                seq += 1
                while f"{seq:05d}" in existing:
                    seq += 1
                code = f"{seq:05d}"
            pid = existing[code] = conn.execute(
                """INSERT INTO partners (name, partner_type, biz_no, representative, biz_type, biz_item,
                   tax_type, address, phone, email, fax,
                   notes, group_id, sale_pg_id, buy_pg_id, credit_limit, credit_days, company_id, code, created_by)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (*vals, book, code, u["id"]),
            ).lastrowid
            added += 1
        # 엑셀 담당자 열은 대표 담당자에 반영. 계좌는 기본계좌에 반영
        if cells["contact_name"]:
            primary = conn.execute("SELECT id FROM partner_contacts WHERE partner_id=? AND is_primary=1",
                                   (pid,)).fetchone()
            c = (cells["contact_name"], cells["contact_phone"], cells["contact_email"])
            if primary:
                conn.execute("UPDATE partner_contacts SET name=?, phone=?, email=? WHERE id=?", (*c, primary["id"]))
            else:
                conn.execute(
                    """INSERT INTO partner_contacts (company_id, partner_id, name, phone, email, is_primary,
                       is_active, sort) VALUES (?,?,?,?,?,1,1,-1)""", (book, pid, *c))
        if cells["bank_name"] or cells["bank_account"]:
            default_bank = conn.execute("SELECT id FROM partner_banks WHERE partner_id=? AND is_default=1",
                                        (pid,)).fetchone()
            b = (cells["bank_name"], cells["bank_account"], cells["bank_holder"])
            if default_bank:
                conn.execute("UPDATE partner_banks SET bank=?, account_no=?, holder=? WHERE id=?",
                             (*b, default_bank["id"]))
            else:
                conn.execute(
                    """INSERT INTO partner_banks (company_id, partner_id, bank, account_no, holder, is_default, sort)
                       VALUES (?,?,?,?,?,1,0)""", (book, pid, *b))
    log_activity(conn, u["id"], "create", "master", 0, "", f"거래처 엑셀 등록 신규 {added}건, 수정 {updated}건")
    conn.commit()
    ctx["result"] = {"added": added, "updated": updated}
    return templates.TemplateResponse(request=request, name="erp/partner_import.html", context=ctx)


@router.get("/api/partners/{partner_id}/credit")
async def partner_credit(partner_id: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    p = conn.execute("SELECT id, name, credit_limit, credit_days FROM partners WHERE id=? AND company_id=?",
                     (partner_id, u["book"])).fetchone()
    if not p:
        raise HTTPException(status_code=404)
    return {"balance": credit.ar_balance(conn, u["book"], p), "limit": p["credit_limit"], "days": p["credit_days"]}


@router.get("/api/partners/{pid}/contacts")
async def api_partner_contacts(pid: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    if not conn.execute("SELECT 1 FROM partners WHERE id=? AND company_id=?", (pid, u["book"])).fetchone():
        raise HTTPException(status_code=404)
    rows = conn.execute(
        """SELECT id, name, dept, position, phone, mobile, email, roles, is_primary FROM partner_contacts
           WHERE partner_id=? AND is_active=1 ORDER BY is_primary DESC, sort, id""", (pid,)).fetchall()
    return [{"id": r["id"], "name": r["name"], "dept": r["dept"], "position": r["position"],
             "phone": r["phone"], "mobile": r["mobile"], "email": r["email"],
             "roles": [x for x in (r["roles"] or "").split(",") if x], "is_primary": bool(r["is_primary"])}
            for r in rows]


@router.get("/api/partners/{pid}/addrs")
async def api_partner_addrs(pid: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    if not conn.execute("SELECT 1 FROM partners WHERE id=? AND company_id=?", (pid, u["book"])).fetchone():
        raise HTTPException(status_code=404)
    rows = conn.execute(
        """SELECT id, name, postcode, address, receiver, phone, is_default FROM partner_addrs
           WHERE partner_id=? ORDER BY is_default DESC, sort, id""", (pid,)).fetchall()
    return [{"id": r["id"], "name": r["name"], "postcode": r["postcode"], "address": r["address"],
             "receiver": r["receiver"], "phone": r["phone"], "is_default": bool(r["is_default"])} for r in rows]


PRICE_KINDS = {"sale": "출고단가", "buy": "입고단가"}


@router.get("/api/partners/{partner_id}/prices")
async def partner_price_map(partner_id: int, kind: str = "sale", on: str = "",
                            u: dict = Depends(require_book), conn=Depends(get_db)):
    on = on or today_kst().isoformat()
    rows = conn.execute(
        """SELECT item_id, price FROM partner_prices WHERE partner_id=? AND company_id=? AND kind=?
           AND (valid_from='' OR valid_from<=?) AND (valid_to='' OR valid_to>=?)""",
        (partner_id, u["book"], kind, on, on)).fetchall()
    return {str(r["item_id"]): r["price"] for r in rows}


@router.get("/partners/{partner_id}/prices", response_class=HTMLResponse)
async def partner_prices(request: Request, partner_id: int, kind: str = "sale",
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    p = conn.execute(
        """SELECT p.id, p.code, p.name, sp.name AS sale_pg, sp.rate AS sale_rate, bp.name AS buy_pg, bp.rate AS buy_rate
           FROM partners p
           LEFT JOIN price_groups sp ON sp.id=p.sale_pg_id
           LEFT JOIN price_groups bp ON bp.id=p.buy_pg_id
           WHERE p.id=? AND p.company_id=?""", (partner_id, book)).fetchone()
    if not p:
        raise HTTPException(status_code=404)
    if kind not in PRICE_KINDS:
        kind = "sale"
    rate = (p["sale_rate"] if kind == "sale" else p["buy_rate"]) or 0
    rows = [dict(r) for r in conn.execute(
        """SELECT pp.id, pp.price, pp.updated_at, pp.valid_from, pp.valid_to, pp.lead_days,
                  i.id AS item_id, i.item_code, i.item_name, i.unit, i.unit_price
           FROM partner_prices pp JOIN items i ON i.id=pp.item_id
           WHERE pp.partner_id=? AND pp.company_id=? AND pp.kind=? ORDER BY i.item_code""",
        (partner_id, book, kind)).fetchall()]
    today = today_kst().isoformat()
    for r in rows:
        r["group_price"] = round(r["unit_price"] * (1 + rate / 100))
        r["expired"] = bool(r["valid_to"]) and r["valid_to"] < today
    return templates.TemplateResponse(request=request, name="erp/partner_prices.html", context={
        "request": request, "page_title": "거래처별 단가",
        "user_name": u["name"], "can_edit": ptn.can_edit(u),
        "partner": p, "kind": kind, "kinds": PRICE_KINDS, "rows": rows, "rate": rate,
    })


@router.post("/partners/{partner_id}/prices")
async def partner_price_save(partner_id: int, kind: str = Form(...), item_id: int = Form(0), price: str = Form(""),
                             valid_from: str = Form(""), valid_to: str = Form(""), lead_days: int = Form(0),
                             u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    amt = price.replace(",", "")
    if kind not in PRICE_KINDS or not amt.isdigit():
        return JSONResponse({"error": "단가를 0 이상 숫자로 입력해 주세요."}, status_code=400)
    try:
        for d in (valid_from, valid_to):
            if d:
                date.fromisoformat(d)
    except ValueError:
        return JSONResponse({"error": "적용기간 날짜를 확인해 주세요."}, status_code=400)
    if valid_from and valid_to and valid_from > valid_to:
        return JSONResponse({"error": "적용 시작일이 종료일보다 늦습니다."}, status_code=400)
    if lead_days < 0:
        return JSONResponse({"error": "조달기간은 0 이상입니다."}, status_code=400)
    p = conn.execute("SELECT name FROM partners WHERE id=? AND company_id=?", (partner_id, book)).fetchone()
    item = conn.execute("SELECT item_code FROM items WHERE id=? AND company_id=?", (item_id, book)).fetchone()
    if not p or not item:
        return JSONResponse({"error": "품목을 선택해 주세요."}, status_code=400)
    conn.execute(
        """INSERT INTO partner_prices (company_id, partner_id, item_id, kind, price, valid_from, valid_to, lead_days)
           VALUES (?,?,?,?,?,?,?,?)
           ON CONFLICT(partner_id, item_id, kind) DO UPDATE SET price=excluded.price, valid_from=excluded.valid_from,
           valid_to=excluded.valid_to, lead_days=excluded.lead_days, updated_at=datetime('now','localtime')""",
        (book, partner_id, item_id, kind, int(amt), valid_from, valid_to, lead_days),
    )
    log_activity(conn, u["id"], "update", "master", partner_id, item["item_code"],
                 f"거래처별 {PRICE_KINDS[kind]} {p['name']} {item['item_code']} {int(amt):,}원")
    conn.commit()
    return RedirectResponse(f"/partners/{partner_id}/prices?kind={kind}", status_code=303)


@router.post("/partners/{partner_id}/prices/{pp_id}/delete")
async def partner_price_delete(partner_id: int, pp_id: int,
                               u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    row = conn.execute("SELECT kind FROM partner_prices WHERE id=? AND partner_id=? AND company_id=?",
                       (pp_id, partner_id, u["book"])).fetchone()
    if not row:
        raise HTTPException(status_code=404)
    conn.execute("DELETE FROM partner_prices WHERE id=?", (pp_id,))
    conn.commit()
    return RedirectResponse(f"/partners/{partner_id}/prices?kind={row['kind']}", status_code=303)


@router.get("/partners/{partner_id}", response_class=HTMLResponse)
async def partner_detail(request: Request, partner_id: int,
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    row = conn.execute(
        """SELECT p.*, g.name AS group_name, sp.name AS sale_pg, sp.rate AS sale_rate,
                  bp.name AS buy_pg, bp.rate AS buy_rate, gr.name AS grade_name, ru.name AS rep_name,
                  su.name AS sub_rep_name
           FROM partners p
           LEFT JOIN partner_groups g ON g.id=p.group_id
           LEFT JOIN price_groups sp ON sp.id=p.sale_pg_id
           LEFT JOIN price_groups bp ON bp.id=p.buy_pg_id
           LEFT JOIN partner_grades gr ON gr.id=p.grade_id
           LEFT JOIN users ru ON ru.id=p.rep_id
           LEFT JOIN users su ON su.id=p.sub_rep_id
           WHERE p.id=? AND p.company_id=?""", (partner_id, book)).fetchone()
    if not row:
        raise HTTPException(status_code=404)
    partner = dict(row)
    partner["type_label"] = PARTNER_TYPES.get(partner.get("partner_type", ""), "공급처")
    partner["trade_label"] = ptn.TRADE_STATUS.get(partner["trade_status"], "")
    partner["trade_class"] = ptn.TRADE_CLASS.get(partner["trade_status"], "")
    partner["biz_label"] = ptn.BIZ_STATUS.get(partner["biz_status"], "")
    partner["pay_label"] = ptn.PAY_TERMS.get(partner["pay_terms"], "") + (
        f" {partner['pay_days']}일" if partner["pay_terms"] in ("days", "eom") else "")
    can_edit = ptn.can_edit(u)
    partner["biz_no_display"] = partner["biz_no"] if (can_edit or partner["reg_kind"] == "biz") \
        else ptn.mask_reg(partner["biz_no"])

    match = "(sl.partner_id=? OR (sl.partner_id IS NULL AND sl.partner=?))"
    args = [book, partner_id, partner["name"]]

    ar_bal = credit.ar_balance(conn, book, partner)
    ap_bal = conn.execute(
        f"""SELECT COALESCE(SUM(sl.credit),0) - COALESCE(SUM(sl.debit),0)
            FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
            WHERE d.company_id=? AND d.status='done' AND {match} AND sl.account_code='251'""",
        args,
    ).fetchone()[0]
    credit_pct = round(ar_bal / partner["credit_limit"] * 100) if partner["credit_limit"] else None

    today_iso = today_kst().isoformat()
    open_ar = credit.open_items(conn, book, "110", partner_id).get(partner_id, {"items": [], "over": 0})
    overdue_amt = next30_amt = 0
    limit30 = (date.fromisoformat(today_iso) + timedelta(days=30)).isoformat()
    for it in open_ar["items"]:
        due = ptn.due_date(it["date"], partner["pay_terms"], partner["pay_days"])
        if due < today_iso:
            overdue_amt += it["remain"]
        elif due <= limit30:
            next30_amt += it["remain"]

    last_sale = conn.execute(
        "SELECT MAX(issue_date) FROM tax_invoices WHERE company_id=? AND partner_id=? AND kind='매출'",
        (book, partner_id)).fetchone()[0]
    last_act_date = conn.execute(
        "SELECT MAX(act_date) FROM crm_activities WHERE company_id=? AND partner_id=?",
        (book, partner_id)).fetchone()[0]
    open_opp = conn.execute(
        """SELECT COUNT(*), COALESCE(SUM(amount),0) FROM crm_opps
           WHERE company_id=? AND partner_id=? AND stage NOT IN ('won','lost')""",
        (book, partner_id)).fetchone()

    my_alerts = ptn.alerts(conn, book, today_iso).get(partner_id, [])

    slips = conn.execute(
        f"""SELECT d.id as doc_id, d.slip_date, d.doc_number, sl.summary, sl.debit, sl.credit
            FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
            WHERE d.company_id=? AND d.status='done' AND {match}
            ORDER BY d.slip_date DESC, d.id DESC LIMIT 10""",
        args,
    ).fetchall()

    invoices = conn.execute(
        """SELECT id, issue_date, kind, evidence, total FROM tax_invoices
           WHERE company_id=? AND partner_id=? ORDER BY issue_date DESC, id DESC LIMIT 10""",
        (book, partner_id),
    ).fetchall()

    sales = conn.execute(
        """SELECT 'qt' AS kind, id, quote_number AS no, quote_date AS dt, total_amount + tax_amount AS total, status
           FROM quotations WHERE company_id=? AND partner_id=?
           UNION ALL
           SELECT 'so', id, so_number, order_date, total_amount + tax_amount, status
           FROM sales_orders WHERE company_id=? AND partner_id=?
           ORDER BY dt DESC, id DESC LIMIT 10""",
        (book, partner_id, book, partner_id),
    ).fetchall()
    buys = conn.execute(
        """SELECT id, po_number, po_date, total_amount + tax_amount AS total, status
           FROM purchase_orders WHERE company_id=? AND partner_id=? ORDER BY po_date DESC, id DESC LIMIT 10""",
        (book, partner_id),
    ).fetchall()

    extra = json.loads(partner["extra"] or "{}")
    fields = [(f["label"], extra.get(str(f["id"]), ""), f["kind"]) for f in _masters(conn, book)["fields"]]

    contacts = [dict(r) for r in conn.execute(
        "SELECT * FROM partner_contacts WHERE partner_id=? ORDER BY is_primary DESC, sort, id", (partner_id,))]
    for c in contacts:
        c["role_list"] = [ptn.ROLES.get(x, x) for x in (c["roles"] or "").split(",") if x]
    addrs = conn.execute("SELECT * FROM partner_addrs WHERE partner_id=? ORDER BY is_default DESC, sort, id",
                        (partner_id,)).fetchall()
    banks = conn.execute("SELECT * FROM partner_banks WHERE partner_id=? ORDER BY is_default DESC, sort, id",
                        (partner_id,)).fetchall()
    files = conn.execute(
        """SELECT f.*, us.name AS created_by_name FROM partner_files f LEFT JOIN users us ON us.id=f.created_by
           WHERE f.partner_id=? ORDER BY f.id DESC""", (partner_id,)).fetchall()

    history = [dict(r) for r in conn.execute(
        """SELECT h.*, us.name AS user_name FROM partner_history h LEFT JOIN users us ON us.id=h.user_id
           WHERE h.partner_id=? ORDER BY h.id DESC LIMIT 100""", (partner_id,)).fetchall()]
    for h in history:
        h["field_label"] = ptn.FIELDS.get(h["field"], h["field"])
        if h["field"] == "biz_no" and not can_edit:
            h["old"], h["new"] = ptn.mask_reg(h["old"]), ptn.mask_reg(h["new"])
    rep_history = [h for h in history if h["field"] in ("rep_id", "sub_rep_id")]

    acts = conn.execute(
        """SELECT a.*, ru.name AS rep_name, o.title AS opp_title FROM crm_activities a
           LEFT JOIN users ru ON ru.id=a.rep_id
           LEFT JOIN crm_opps o ON o.id=a.opp_id
           WHERE a.company_id=? AND a.partner_id=? ORDER BY a.act_date DESC, a.id DESC LIMIT 20""",
        (book, partner_id)).fetchall()
    opps = conn.execute(
        """SELECT o.*, ru.name AS rep_name FROM crm_opps o LEFT JOIN users ru ON ru.id=o.rep_id
           WHERE o.company_id=? AND o.partner_id=? ORDER BY o.id DESC""", (book, partner_id)).fetchall()
    price_cnt = conn.execute("SELECT COUNT(*) FROM partner_prices WHERE partner_id=? AND company_id=?",
                             (partner_id, book)).fetchone()[0]

    return templates.TemplateResponse(request=request, name="erp/partner_detail.html", context={
        "request": request, "page_title": "거래처 정보", "user_name": u["name"], "can_edit": can_edit,
        "partner": partner, "ar_bal": ar_bal, "ap_bal": ap_bal, "credit_pct": credit_pct,
        "credit_left": partner["credit_limit"] - ar_bal if partner["credit_limit"] else None,
        "overdue_amt": overdue_amt, "next30_amt": next30_amt,
        "last_sale": last_sale, "last_act_date": last_act_date,
        "open_opp_cnt": open_opp[0], "open_opp_amt": open_opp[1], "my_alerts": my_alerts,
        "alert_labels": ptn.ALERTS,
        "slips": slips, "invoices": invoices,
        "sales": sales, "buys": buys, "fields": fields,
        "status_labels": {"qt": QT_STATUS, "so": SO_STATUS, "po": PO_STATUS},
        "contacts": contacts, "addrs": addrs, "banks": banks, "files": files,
        "history": history, "rep_history": rep_history,
        "acts": acts, "opps": opps, "price_cnt": price_cnt,
        "act_kinds": ACT_KINDS, "stages": STAGES, "stage_class": STAGE_CLASS, "roles": ptn.ROLES,
    })


@router.get("/partners/{partner_id}/edit", response_class=HTMLResponse)
async def partner_form_edit(request: Request, partner_id: int,
                            u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    row = conn.execute("SELECT * FROM partners WHERE id=? AND company_id=?", (partner_id, book)).fetchone()
    if not row:
        return RedirectResponse("/partners", status_code=303)
    contacts = [dict(r) for r in conn.execute(
        "SELECT * FROM partner_contacts WHERE partner_id=? ORDER BY is_primary DESC, sort, id", (partner_id,))]
    for c in contacts:
        c["roles"] = [x for x in (c["roles"] or "").split(",") if x]
    addrs = [dict(r) for r in conn.execute(
        "SELECT * FROM partner_addrs WHERE partner_id=? ORDER BY is_default DESC, sort, id", (partner_id,))]
    banks = [dict(r) for r in conn.execute(
        "SELECT * FROM partner_banks WHERE partner_id=? ORDER BY is_default DESC, sort, id", (partner_id,))]
    files = [dict(r) for r in conn.execute(
        "SELECT * FROM partner_files WHERE partner_id=? ORDER BY id DESC", (partner_id,))]
    return templates.TemplateResponse(request=request, name="erp/partner_form.html", context={
        "request": request, "page_title": "거래처 수정", "user_name": u["name"],
        "partner": dict(row), "partner_types": PARTNER_TYPES, "extra": json.loads(row["extra"] or "{}"),
        "contacts": contacts, "addrs": addrs, "banks": banks, "files": files, "file_kinds": FILE_KINDS,
        "kinds": ptn.KINDS, "reg_kinds": ptn.REG_KINDS, "biz_statuses": ptn.BIZ_STATUS,
        "trade_statuses": ptn.TRADE_STATUS, "pay_terms_opts": ptn.PAY_TERMS, "roles": ptn.ROLES,
        **_masters(conn, book),
    })


@router.get("/partners/{partner_id}/merge", response_class=HTMLResponse)
async def partner_merge_form(request: Request, partner_id: int, src: int = 0,
                             u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    a = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=? AND merged_into IS NULL",
                     (partner_id, book)).fetchone()
    if not a:
        raise HTTPException(status_code=404)
    others = conn.execute(
        """SELECT id, code, name FROM partners WHERE company_id=? AND id!=? AND is_active=1
           AND merged_into IS NULL ORDER BY code""", (book, partner_id)).fetchall()
    b, counts = None, None
    if src:
        b = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=? AND merged_into IS NULL",
                         (src, book)).fetchone()
        if b:
            counts = {
                "견적": conn.execute("SELECT COUNT(*) FROM quotations WHERE partner_id=? AND company_id=?",
                                   (src, book)).fetchone()[0],
                "수주": conn.execute("SELECT COUNT(*) FROM sales_orders WHERE partner_id=? AND company_id=?",
                                   (src, book)).fetchone()[0],
                "발주": conn.execute("SELECT COUNT(*) FROM purchase_orders WHERE partner_id=? AND company_id=?",
                                   (src, book)).fetchone()[0],
                "세금계산서": conn.execute("SELECT COUNT(*) FROM tax_invoices WHERE partner_id=? AND company_id=?",
                                       (src, book)).fetchone()[0],
                "영업활동": conn.execute("SELECT COUNT(*) FROM crm_activities WHERE partner_id=? AND company_id=?",
                                      (src, book)).fetchone()[0],
                "영업기회": conn.execute("SELECT COUNT(*) FROM crm_opps WHERE partner_id=? AND company_id=?",
                                      (src, book)).fetchone()[0],
            }
    return templates.TemplateResponse(request=request, name="erp/partner_merge.html", context={
        "request": request, "page_title": "거래처 통합", "user_name": u["name"],
        "partner": dict(a), "others": [dict(r) for r in others], "src": src, "b": dict(b) if b else None,
        "counts": counts,
    })


@router.post("/partners/{partner_id}/merge")
async def partner_merge(partner_id: int, src_id: int = Form(...), confirm: int = Form(0),
                        u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    if not confirm:
        return JSONResponse({"error": "통합 확인 체크가 필요합니다."}, status_code=400)
    if partner_id == src_id:
        return JSONResponse({"error": "같은 거래처는 통합할 수 없습니다."}, status_code=400)
    a = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=? AND merged_into IS NULL",
                     (partner_id, book)).fetchone()
    b = conn.execute("SELECT id, code, name FROM partners WHERE id=? AND company_id=? AND merged_into IS NULL",
                     (src_id, book)).fetchone()
    if not a or not b:
        return JSONResponse({"error": "거래처를 확인해 주세요."}, status_code=400)

    for t in MERGE_PARTNER_TABLES:
        conn.execute(f"UPDATE {t} SET partner_id=? WHERE partner_id=? AND company_id=?", (partner_id, src_id, book))

    a_primary = conn.execute("SELECT 1 FROM partner_contacts WHERE partner_id=? AND is_primary=1",
                             (partner_id,)).fetchone()
    conn.execute("UPDATE partner_contacts SET partner_id=?, is_primary=0 WHERE partner_id=? AND company_id=?",
                 (partner_id, src_id, book))
    if not a_primary:
        first = conn.execute("SELECT id FROM partner_contacts WHERE partner_id=? AND is_active=1 ORDER BY sort, id LIMIT 1",
                             (partner_id,)).fetchone()
        if first:
            conn.execute("UPDATE partner_contacts SET is_primary=1 WHERE id=?", (first["id"],))

    a_def_addr = conn.execute("SELECT 1 FROM partner_addrs WHERE partner_id=? AND is_default=1",
                              (partner_id,)).fetchone()
    conn.execute("UPDATE partner_addrs SET partner_id=?, is_default=0 WHERE partner_id=? AND company_id=?",
                 (partner_id, src_id, book))
    if not a_def_addr:
        first = conn.execute("SELECT id FROM partner_addrs WHERE partner_id=? ORDER BY sort, id LIMIT 1",
                             (partner_id,)).fetchone()
        if first:
            conn.execute("UPDATE partner_addrs SET is_default=1 WHERE id=?", (first["id"],))

    a_def_bank = conn.execute("SELECT 1 FROM partner_banks WHERE partner_id=? AND is_default=1",
                              (partner_id,)).fetchone()
    conn.execute("UPDATE partner_banks SET partner_id=?, is_default=0 WHERE partner_id=? AND company_id=?",
                 (partner_id, src_id, book))
    if not a_def_bank:
        first = conn.execute("SELECT id FROM partner_banks WHERE partner_id=? ORDER BY sort, id LIMIT 1",
                             (partner_id,)).fetchone()
        if first:
            conn.execute("UPDATE partner_banks SET is_default=1 WHERE id=?", (first["id"],))

    conn.execute("UPDATE partner_files SET partner_id=? WHERE partner_id=? AND company_id=?",
                 (partner_id, src_id, book))

    for pp in conn.execute("SELECT id, item_id, kind FROM partner_prices WHERE partner_id=? AND company_id=?",
                           (src_id, book)).fetchall():
        if conn.execute("SELECT 1 FROM partner_prices WHERE partner_id=? AND item_id=? AND kind=?",
                        (partner_id, pp["item_id"], pp["kind"])).fetchone():
            conn.execute("DELETE FROM partner_prices WHERE id=?", (pp["id"],))
        else:
            conn.execute("UPDATE partner_prices SET partner_id=? WHERE id=?", (partner_id, pp["id"]))

    # 신규 라인(partner_id=B) + 이름만 남은 과거 라인(B명) 모두 A로 이전. slip_lines는 company_id가 없어 문서로 제한
    conn.execute(
        """UPDATE slip_lines SET partner_id=?, partner=?
           WHERE doc_id IN (SELECT id FROM erp_docs WHERE company_id=?)
             AND (partner_id=? OR (partner_id IS NULL AND partner=?))""",
        (partner_id, a["name"], book, src_id, b["name"]),
    )
    # 출고/입고는 수주/발주를 따라가므로 거래처 컬럼이 없어 별도 이전 불필요

    conn.execute("UPDATE partners SET is_active=0, trade_status='stop', merged_into=? WHERE id=?",
                 (partner_id, src_id))

    conn.execute(
        "INSERT INTO partner_history (company_id, partner_id, field, old, new, user_id) VALUES (?,?,?,?,?,?)",
        (book, partner_id, "merge", "", f"{b['name']}({b['code']}) 통합", u["id"]))
    conn.execute(
        "INSERT INTO partner_history (company_id, partner_id, field, old, new, user_id) VALUES (?,?,?,?,?,?)",
        (book, src_id, "merge", "", f"-> {a['name']}({a['code']})", u["id"]))

    log_activity(conn, u["id"], "update", "master", partner_id, a["code"],
                 f"거래처 통합 {b['name']}({b['code']}) -> {a['name']}({a['code']})")
    audit.log(conn, u["id"], "partner_merge", "partner", partner_id,
              f"{b['name']}({b['code']}) -> {a['name']}({a['code']})", company_id=book)
    regrade(conn, book, partner_id)
    conn.commit()
    return RedirectResponse(f"/partners/{partner_id}?saved=1", status_code=303)


@router.post("/partners/save", response_class=HTMLResponse)
async def partner_save(request: Request,
                       partner_id: int = Form(0),
                       code: str = Form(""),
                       name: str = Form(...),
                       partner_type: str = Form("supplier"),
                       kind: str = Form("general"),
                       reg_kind: str = Form("biz"),
                       biz_no: str = Form(""),
                       representative: str = Form(""),
                       biz_type: str = Form(""),
                       biz_item: str = Form(""),
                       tax_type: str = Form("과세"),
                       biz_status: str = Form("normal"),
                       biz_checked_at: str = Form(""),
                       trade_status: str = Form("active"),
                       postcode: str = Form(""),
                       address: str = Form(""),
                       phone: str = Form(""),
                       email: str = Form(""),
                       fax: str = Form(""),
                       notes: str = Form(""),
                       is_active: int = Form(1),
                       group_id: int = Form(0),
                       sale_pg_id: int = Form(0),
                       buy_pg_id: int = Form(0),
                       credit_limit: str = Form("0"),
                       credit_days: int = Form(0),
                       pay_terms: str = Form(""),
                       pay_days: int = Form(0),
                       rep_id: int = Form(0),
                       sub_rep_id: int = Form(0),
                       grade_id: int = Form(0),
                       u: dict = Depends(ptn.require_editor), conn=Depends(get_db)):
    book = u["book"]
    name = name.strip()
    if not name:
        return JSONResponse({"error": "거래처명을 입력해 주세요."}, status_code=400)
    if partner_type not in PARTNER_TYPES:
        return JSONResponse({"error": "거래처유형을 확인해 주세요."}, status_code=400)
    if kind not in ptn.KINDS:
        kind = "general"
    if reg_kind not in ptn.REG_KINDS:
        reg_kind = "biz"
    if tax_type not in ("과세", "영세", "면세"):
        tax_type = "과세"
    if biz_status not in ptn.BIZ_STATUS:
        biz_status = "normal"
    if trade_status not in ptn.TRADE_STATUS:
        trade_status = "active"
    if pay_terms not in ptn.PAY_TERMS:
        pay_terms = ""
    limit = credit_limit.replace(",", "")
    if not limit.isdigit() or credit_days < 0 or pay_days < 0:
        return JSONResponse({"error": "여신한도, 여신기간, 결제일수는 0 이상 숫자로 입력해 주세요."}, status_code=400)
    biz_no, reg_err = ptn.validate_reg_no(reg_kind, biz_no)
    if reg_err:
        return JSONResponse({"error": reg_err}, status_code=400)

    old_row = None
    if partner_id:
        old_row = conn.execute("SELECT * FROM partners WHERE id=? AND company_id=?", (partner_id, book)).fetchone()
        if not old_row:
            raise HTTPException(status_code=404)

    m = _masters(conn, book)
    old_disp = _partner_display(conn, book, dict(old_row), m) if old_row else None

    rep_ids = {r["id"] for r in m["reps"]}
    rep_id = rep_id if rep_id in rep_ids else None
    sub_rep_id = sub_rep_id if sub_rep_id in rep_ids else None
    group_id = _own(conn, "partner_groups", group_id, book)
    sale_pg_id = _own(conn, "price_groups", sale_pg_id, book)
    buy_pg_id = _own(conn, "price_groups", buy_pg_id, book)
    grade_id = _own(conn, "partner_grades", grade_id, book)

    form = await request.form()
    extra = {}
    for f in m["fields"]:
        v = form.get(f"f_{f['id']}", "").strip()
        if v:
            extra[str(f["id"])] = v.replace(",", "") if f["kind"] == "num" else v

    phone = ptn.fmt_phone(phone)
    fax = ptn.fmt_phone(fax)

    is_new = not bool(partner_id)
    code = code.strip()
    if not is_new and not code:
        return JSONResponse({"error": "거래처코드를 입력해 주세요."}, status_code=400)
    if is_new and not code:
        last = conn.execute(
            "SELECT code FROM partners WHERE company_id=? AND code GLOB '[0-9][0-9][0-9][0-9][0-9]' ORDER BY code DESC LIMIT 1",
            (book,),
        ).fetchone()
        code = f"{(int(last['code']) + 1) if last else 1:05d}"

    more_keys = ("kind", "reg_kind", "biz_no", "representative", "biz_type", "biz_item", "tax_type",
                 "biz_status", "biz_checked_at", "trade_status", "postcode", "address", "phone", "email",
                 "fax", "notes", "is_active", "group_id", "sale_pg_id", "buy_pg_id", "credit_limit",
                 "credit_days", "pay_terms", "pay_days", "extra", "rep_id", "sub_rep_id")
    more = (kind, reg_kind, biz_no, representative, biz_type, biz_item, tax_type, biz_status, biz_checked_at,
            trade_status, postcode, address, phone, email, fax, notes, is_active,
            group_id, sale_pg_id, buy_pg_id, int(limit), credit_days, pay_terms, pay_days,
            json.dumps(extra, ensure_ascii=False), rep_id, sub_rep_id)
    new_vals = {"code": code, "name": name, "partner_type": partner_type, **dict(zip(more_keys, more))}

    try:
        if partner_id:
            conn.execute(
                """UPDATE partners SET code=?, name=?, partner_type=?, kind=?, reg_kind=?, biz_no=?,
                   representative=?, biz_type=?, biz_item=?, tax_type=?, biz_status=?, biz_checked_at=?,
                   trade_status=?, postcode=?, address=?, phone=?, email=?, fax=?, notes=?, is_active=?,
                   group_id=?, sale_pg_id=?, buy_pg_id=?, credit_limit=?, credit_days=?, pay_terms=?, pay_days=?,
                   extra=?, rep_id=?, sub_rep_id=?
                   WHERE id=? AND company_id=?""",
                (code, name, partner_type, *more, partner_id, book),
            )
            pid = partner_id
        else:
            task_id = stamp_task(conn, request, u["id"], "master")
            cur = conn.execute(
                """INSERT INTO partners (company_id, code, name, partner_type, kind, reg_kind, biz_no,
                   representative, biz_type, biz_item, tax_type, biz_status, biz_checked_at, trade_status,
                   postcode, address, phone, email, fax, notes, is_active, group_id, sale_pg_id, buy_pg_id,
                   credit_limit, credit_days, pay_terms, pay_days, extra, rep_id, sub_rep_id, created_by, task_id)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (book, code, name, partner_type, *more, u["id"], task_id),
            )
            pid = cur.lastrowid
    except sqlite3.IntegrityError:
        conn.rollback()
        return JSONResponse({"error": f"거래처코드 {code}가 이미 사용 중입니다."}, status_code=400)

    # 담당자. c_id가 해당 거래처 소속인 경우만 수정 대상
    existing_contacts = {r["id"] for r in conn.execute(
        "SELECT id FROM partner_contacts WHERE partner_id=? AND company_id=?", (pid, book))}
    c_rows = []
    for n in form.get("c_rows", "").split(","):
        n = n.strip()
        if not n:
            continue
        cname = form.get(f"c_name_{n}", "").strip()
        if not cname:
            continue
        cid = int(form.get(f"c_id_{n}", 0) or 0)
        if cid not in existing_contacts:
            cid = 0
        roles = ",".join(code_ for code_ in ptn.ROLES if form.get(f"c_role_{n}_{code_}"))
        vals_c = {k: (ptn.fmt_phone(form.get(f"c_{k}_{n}", "")) if k in ("phone", "mobile", "fax")
                      else form.get(f"c_{k}_{n}", "").strip()) for k in CONTACT_FIELDS}
        c_rows.append({"n": n, "cid": cid, "name": cname, "roles": roles,
                       "active": 1 if form.get(f"c_active_{n}") else 0, **vals_c})

    kept_ids = {r["cid"] for r in c_rows if r["cid"]}
    # 참조 중인 담당자는 비활성화만. 참조 없으면 삭제
    for rid in existing_contacts - kept_ids:
        if _contact_referenced(conn, rid):
            conn.execute("UPDATE partner_contacts SET is_active=0 WHERE id=?", (rid,))
        else:
            conn.execute("DELETE FROM partner_contacts WHERE id=?", (rid,))

    for i, r in enumerate(c_rows):
        if r["cid"]:
            conn.execute(
                """UPDATE partner_contacts SET name=?, dept=?, position=?, phone=?, mobile=?, email=?, fax=?,
                   birthday=?, memo=?, roles=?, is_active=?, sort=? WHERE id=?""",
                (r["name"], r["dept"], r["position"], r["phone"], r["mobile"], r["email"], r["fax"],
                 r["birthday"], r["memo"], r["roles"], r["active"], i, r["cid"]),
            )
        else:
            cur = conn.execute(
                """INSERT INTO partner_contacts (company_id, partner_id, name, dept, position, phone, mobile,
                   email, fax, birthday, memo, roles, is_active, sort) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (book, pid, r["name"], r["dept"], r["position"], r["phone"], r["mobile"], r["email"], r["fax"],
                 r["birthday"], r["memo"], r["roles"], r["active"], i),
            )
            r["cid"] = cur.lastrowid

    primary_n = form.get("c_primary", "")
    primary_id = next((r["cid"] for r in c_rows if r["n"] == primary_n), None)
    if not primary_id:
        active_rows = [r for r in c_rows if r["active"]]
        primary_id = active_rows[0]["cid"] if active_rows else None
    conn.execute("UPDATE partner_contacts SET is_primary=0 WHERE partner_id=?", (pid,))
    if primary_id:
        conn.execute("UPDATE partner_contacts SET is_primary=1 WHERE id=?", (primary_id,))

    # 배송지. 문서에는 ship_to 텍스트로 스냅샷되므로 삭제 후 재등록 처리
    conn.execute("DELETE FROM partner_addrs WHERE partner_id=?", (pid,))
    a_rows = []
    for n in form.get("a_rows", "").split(","):
        n = n.strip()
        if not n:
            continue
        addr_v = form.get(f"a_address_{n}", "").strip()
        if not addr_v:
            continue
        a_rows.append({
            "n": n, "name": form.get(f"a_name_{n}", "").strip(), "postcode": form.get(f"a_postcode_{n}", "").strip(),
            "address": addr_v, "receiver": form.get(f"a_receiver_{n}", "").strip(),
            "phone": ptn.fmt_phone(form.get(f"a_phone_{n}", "")),
        })
    a_default_n = form.get("a_default", "")
    a_default_idx = next((i for i, r in enumerate(a_rows) if r["n"] == a_default_n), 0 if a_rows else None)
    for i, r in enumerate(a_rows):
        conn.execute(
            """INSERT INTO partner_addrs (company_id, partner_id, name, postcode, address, receiver, phone,
               is_default, sort) VALUES (?,?,?,?,?,?,?,?,?)""",
            (book, pid, r["name"], r["postcode"], r["address"], r["receiver"], r["phone"],
             1 if i == a_default_idx else 0, i),
        )

    # 계좌. 배송지와 동일하게 삭제 후 재등록
    conn.execute("DELETE FROM partner_banks WHERE partner_id=?", (pid,))
    b_rows = []
    for n in form.get("b_rows", "").split(","):
        n = n.strip()
        if not n:
            continue
        bank_v = form.get(f"b_bank_{n}", "").strip()
        acct_v = form.get(f"b_account_{n}", "").strip()
        if not bank_v or not acct_v:
            continue
        b_rows.append({"n": n, "bank": bank_v, "account_no": acct_v, "holder": form.get(f"b_holder_{n}", "").strip()})
    b_default_n = form.get("b_default", "")
    b_default_idx = next((i for i, r in enumerate(b_rows) if r["n"] == b_default_n), 0 if b_rows else None)
    for i, r in enumerate(b_rows):
        conn.execute(
            """INSERT INTO partner_banks (company_id, partner_id, bank, account_no, holder, is_default, sort)
               VALUES (?,?,?,?,?,?,?)""",
            (book, pid, r["bank"], r["account_no"], r["holder"], 1 if i == b_default_idx else 0, i),
        )

    # 첨부파일. 삭제 선택분은 소유 확인 후 디스크에서도 제거
    del_files = [int(x) for x in form.getlist("del_file") if x]
    if del_files:
        placeholders = ",".join("?" * len(del_files))
        for rf in conn.execute(
            f"SELECT id, file FROM partner_files WHERE partner_id=? AND company_id=? AND id IN ({placeholders})",
            (pid, book, *del_files)).fetchall():
            safe = rf["file"].split("|")[-1]
            (pathlib.Path("var/uploads/partners") / safe).unlink(missing_ok=True)
            conn.execute("DELETE FROM partner_files WHERE id=?", (rf["id"],))

    for fkind in FILE_KINDS:
        up = form.get(f"file_{fkind}")
        if up is not None and getattr(up, "filename", ""):
            saved = await save_upload(up, "var/uploads/partners")
            if isinstance(saved, JSONResponse):
                conn.rollback()
                return saved
            if saved:
                conn.execute(
                    "INSERT INTO partner_files (company_id, partner_id, kind, file, created_by) VALUES (?,?,?,?,?)",
                    (book, pid, fkind, saved, u["id"]),
                )

    # 등급 직접 지정 시 자동 재산정 대상 제외
    if grade_id:
        conn.execute("UPDATE partners SET grade_id=?, grade_fixed=1 WHERE id=?", (grade_id, pid))
    else:
        conn.execute("UPDATE partners SET grade_fixed=0 WHERE id=?", (pid,))
        regrade(conn, book, pid)

    if is_new:
        log_activity(conn, u["id"], "create", "master", pid, biz_no, f"거래처 등록 {name}")
        audit.log(conn, u["id"], "partner_create", "partner", pid, name, company_id=book)
    else:
        new_row = conn.execute("SELECT * FROM partners WHERE id=?", (pid,)).fetchone()
        new_disp = _partner_display(conn, book, dict(new_row), m)
        ptn.log_changes(conn, book, pid, u["id"], old_disp, new_disp)
        log_activity(conn, u["id"], "update", "master", pid, biz_no, f"거래처 수정 {name}")
        changed = audit.diff(dict(old_row), new_vals)
        audit.log(conn, u["id"], "partner_update", "partner", pid, name, detail=changed, company_id=book)
    conn.commit()
    return RedirectResponse(f"/partners/{pid}?saved=1", status_code=303)


@router.get("/bank-accounts", response_class=HTMLResponse)
async def bank_account_list(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM bank_accounts WHERE company_id=? AND is_active=1 ORDER BY bank, account_no", (u["book"],)
    ).fetchall()]
    return templates.TemplateResponse(request=request, name="erp/bank_account_list.html", context={
        "request": request, "page_title": "계좌 관리",
        "user_name": u["name"],
        "accounts": rows,
    })


@router.get("/bank-accounts/new", response_class=HTMLResponse)
async def bank_account_form_new(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    gl_accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()
    return templates.TemplateResponse(request=request, name="erp/bank_account_form.html", context={
        "request": request, "page_title": "계좌 등록",
        "user_name": u["name"],
        "account": None, "gl_accounts": gl_accounts,
    })


@router.get("/bank-accounts/{account_id}/edit", response_class=HTMLResponse)
async def bank_account_form_edit(request: Request, account_id: int, u: dict = Depends(require_book), conn=Depends(get_db)):
    row = conn.execute("SELECT * FROM bank_accounts WHERE id=? AND company_id=?", (account_id, u["book"])).fetchone()
    if not row:
        return RedirectResponse("/bank-accounts", status_code=303)
    gl_accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()
    return templates.TemplateResponse(request=request, name="erp/bank_account_form.html", context={
        "request": request, "page_title": "계좌 수정",
        "user_name": u["name"],
        "account": dict(row), "gl_accounts": gl_accounts,
    })


@router.post("/bank-accounts/save", response_class=HTMLResponse)
async def bank_account_save(request: Request,
                            account_id: int = Form(0),
                            bank: str = Form(...),
                            account_no: str = Form(...),
                            name: str = Form(...),
                            account_code: str = Form(...),
                            is_active: int = Form(1),
                            u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    if not conn.execute("SELECT 1 FROM accounts WHERE company_id=? AND code=?", (book, account_code)).fetchone():
        return JSONResponse({"error": "연결 계정을 확인해 주세요."}, status_code=400)
    is_new = not bool(account_id)
    if account_id:
        conn.execute(
            "UPDATE bank_accounts SET bank=?, account_no=?, name=?, account_code=?, is_active=? WHERE id=? AND company_id=?",
            (bank, account_no, name, account_code, is_active, account_id, book),
        )
        aid = account_id
    else:
        cur = conn.execute(
            "INSERT INTO bank_accounts (company_id, bank, account_no, name, account_code, created_by) VALUES (?,?,?,?,?,?)",
            (book, bank, account_no, name, account_code, u["id"]),
        )
        aid = cur.lastrowid
    if is_new:
        log_activity(conn, u["id"], "create", "master", aid, account_no, f"계좌 등록 {bank} {account_no}")
    else:
        log_activity(conn, u["id"], "update", "master", aid, account_no, f"계좌 수정 {bank} {account_no}")
    conn.commit()
    return RedirectResponse("/bank-accounts", status_code=303)


@router.get("/warehouses", response_class=HTMLResponse)
async def warehouse_list(request: Request,
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM warehouses WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()]
    for r in rows:
        r["item_count"] = conn.execute(
            "SELECT COUNT(DISTINCT item_id) FROM stock WHERE warehouse_id=? AND quantity>0",
            (r["id"],),
        ).fetchone()[0]
        r["total_qty"] = conn.execute(
            "SELECT COALESCE(SUM(quantity),0) FROM stock WHERE warehouse_id=?",
            (r["id"],),
        ).fetchone()[0]
    return templates.TemplateResponse(request=request, name="erp/warehouse_list.html", context={
        "request": request, "page_title": "창고 관리",
        "user_name": u["name"],
        "warehouses": rows,
    })


@router.get("/warehouses/new", response_class=HTMLResponse)
async def warehouse_form_new(request: Request, u: dict = Depends(require_book)):
    return templates.TemplateResponse(request=request, name="erp/warehouse_form.html", context={
        "request": request, "page_title": "창고 등록",
        "user_name": u["name"],
        "warehouse": None,
    })


@router.get("/warehouses/{wh_id}/edit", response_class=HTMLResponse)
async def warehouse_form_edit(request: Request, wh_id: int,
                              u: dict = Depends(require_book), conn=Depends(get_db)):
    row = conn.execute("SELECT * FROM warehouses WHERE id=? AND company_id=?", (wh_id, u["book"])).fetchone()
    if not row:
        return RedirectResponse("/warehouses", status_code=303)
    return templates.TemplateResponse(request=request, name="erp/warehouse_form.html", context={
        "request": request, "page_title": "창고 수정",
        "user_name": u["name"],
        "warehouse": dict(row),
    })


@router.post("/warehouses/save", response_class=HTMLResponse)
async def warehouse_save(request: Request,
                         wh_id: int = Form(0),
                         code: str = Form(...),
                         name: str = Form(...),
                         location: str = Form(""),
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    is_new = not bool(wh_id)
    if wh_id:
        conn.execute(
            "UPDATE warehouses SET code=?, name=?, location=? WHERE id=? AND company_id=?",
            (code, name, location, wh_id, book),
        )
    else:
        cur = conn.execute(
            "INSERT INTO warehouses (company_id, code, name, location, created_by) VALUES (?,?,?,?,?)",
            (book, code, name, location, u["id"]),
        )
        wh_id = cur.lastrowid
    if is_new:
        log_activity(conn, u["id"], "create", "master", wh_id, code, f"창고 등록 {name} ({code})")
    else:
        log_activity(conn, u["id"], "update", "master", wh_id, code, f"창고 수정 {name} ({code})")
    conn.commit()
    return RedirectResponse("/warehouses", status_code=303)


@router.get("/memo-codes", response_class=HTMLResponse)
async def memo_code_list(request: Request, q: str = "",
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    where = "WHERE company_id=? AND is_active=1"
    params = [u["book"]]
    if q:
        where += " AND (code LIKE ? OR text LIKE ?)"
        params += [f"%{q}%", f"%{q}%"]
    rows = [dict(r) for r in conn.execute(
        f"SELECT * FROM memo_codes {where} ORDER BY code", params
    ).fetchall()]
    return templates.TemplateResponse(request=request, name="erp/memo_code_list.html", context={
        "request": request, "page_title": "적요 관리",
        "user_name": u["name"],
        "memos": rows, "q": q,
    })


@router.get("/memo-codes/new", response_class=HTMLResponse)
async def memo_code_form_new(request: Request, u: dict = Depends(require_book), conn=Depends(get_db)):
    gl_accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()
    return templates.TemplateResponse(request=request, name="erp/memo_code_form.html", context={
        "request": request, "page_title": "적요 등록",
        "user_name": u["name"],
        "memo": None, "gl_accounts": gl_accounts,
    })


@router.get("/memo-codes/{memo_id}/edit", response_class=HTMLResponse)
async def memo_code_form_edit(request: Request, memo_id: int,
                              u: dict = Depends(require_book), conn=Depends(get_db)):
    row = conn.execute("SELECT * FROM memo_codes WHERE id=? AND company_id=?", (memo_id, u["book"])).fetchone()
    if not row:
        return RedirectResponse("/memo-codes", status_code=303)
    gl_accounts = conn.execute(
        "SELECT code, name FROM accounts WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()
    return templates.TemplateResponse(request=request, name="erp/memo_code_form.html", context={
        "request": request, "page_title": "적요 수정",
        "user_name": u["name"],
        "memo": dict(row), "gl_accounts": gl_accounts,
    })


@router.post("/memo-codes/save", response_class=HTMLResponse)
async def memo_code_save(request: Request,
                         memo_id: int = Form(0),
                         code: str = Form(...),
                         text: str = Form(...),
                         account_code: str = Form(""),
                         is_active: int = Form(1),
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    account_code = account_code or None
    is_new = not bool(memo_id)
    if memo_id:
        conn.execute(
            "UPDATE memo_codes SET code=?, text=?, account_code=?, is_active=? WHERE id=? AND company_id=?",
            (code, text, account_code, is_active, memo_id, book),
        )
        mid = memo_id
    else:
        cur = conn.execute(
            "INSERT INTO memo_codes (company_id, code, text, account_code, created_by) VALUES (?,?,?,?,?)",
            (book, code, text, account_code, u["id"]),
        )
        mid = cur.lastrowid
    if is_new:
        log_activity(conn, u["id"], "create", "master", mid, code, f"적요 등록 {code} ({text})")
    else:
        log_activity(conn, u["id"], "update", "master", mid, code, f"적요 수정 {code} ({text})")
    conn.commit()
    return RedirectResponse("/memo-codes?saved=1", status_code=303)


@router.get("/api/memo_codes")
async def api_memo_codes(q: str = "", account: str = "",
                         u: dict = Depends(require_book), conn=Depends(get_db)):
    where = "WHERE company_id=? AND is_active=1"
    params = [u["book"]]
    if q:
        where += " AND (code LIKE ? OR text LIKE ?)"
        params += [f"%{q}%", f"%{q}%"]
    if account:
        where += " AND account_code=?"
        params.append(account)
    rows = conn.execute(
        f"SELECT id, code, text AS name, account_code FROM memo_codes {where} ORDER BY code", params
    ).fetchall()
    return [dict(r) for r in rows]


@router.get("/stock", response_class=HTMLResponse)
async def stock_list(request: Request, q: str = "", wh: int = 0, low: int = 0, page: int = 1,
                     u: dict = Depends(require_book), conn=Depends(get_db)):
    base_where = "WHERE i.company_id=? AND i.is_active=1 AND w.is_active=1"
    params = [u["book"]]
    if q:
        base_where += " AND (i.item_code LIKE ? OR i.item_name LIKE ?)"
        params += [f"%{q}%", f"%{q}%"]
    if wh:
        base_where += " AND s.warehouse_id=?"
        params.append(wh)
    if low:
        base_where += " AND s.quantity <= i.safety_stock"

    total = conn.execute(
        f"SELECT COUNT(*) FROM stock s JOIN items i ON s.item_id=i.id JOIN warehouses w ON s.warehouse_id=w.id {base_where}",
        params,
    ).fetchone()[0]
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
    page = max(1, min(page, total_pages))

    sql = f"""
        SELECT s.id, s.quantity, i.item_code, i.item_name, i.item_type, i.unit,
               i.safety_stock, w.code as wh_code, w.name as wh_name, i.id as item_id
        FROM stock s
        JOIN items i ON s.item_id=i.id
        JOIN warehouses w ON s.warehouse_id=w.id
        {base_where}
        ORDER BY i.item_code, w.code
        LIMIT {PER_PAGE} OFFSET {(page-1)*PER_PAGE}
    """
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    for r in rows:
        r["item_type_label"] = ITEM_TYPES.get(r["item_type"], r["item_type"])
        r["is_low"] = r["quantity"] <= r["safety_stock"]
    warehouses = [dict(r) for r in conn.execute(
        "SELECT id, code, name FROM warehouses WHERE company_id=? AND is_active=1 ORDER BY code", (u["book"],)
    ).fetchall()]
    return templates.TemplateResponse(request=request, name="erp/stock_list.html", context={
        "request": request, "page_title": "재고 현황",
        "user_name": u["name"],
        "stocks": rows, "q": q, "wh": wh, "low": low,
        "warehouses": warehouses, "item_types": ITEM_TYPES,
        "page": page, "total_pages": total_pages, "total_count": total,
        "base_url": f"/stock?q={q}&wh={wh}&low={low}",
    })


@router.get("/stock/valuation", response_class=HTMLResponse)
async def stock_valuation(request: Request, date_to: str = "",
                          u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    date_to = date_to or today_kst().isoformat()
    rows = stock.valuation(conn, book, date_to)
    for r in rows:
        r["avg"] = round(r["amt"] / r["qty"]) if r["qty"] else None
    total = sum(r["amt"] for r in rows)
    gl = conn.execute(
        """SELECT COALESCE(SUM(sl.debit - sl.credit), 0) FROM slip_lines sl JOIN erp_docs d ON sl.doc_id=d.id
           WHERE d.company_id=? AND d.status='done' AND d.slip_date<=? AND sl.account_code='120'""",
        (book, date_to),
    ).fetchone()[0]
    return templates.TemplateResponse(request=request, name="erp/stock_valuation.html", context={
        "request": request, "page_title": "재고평가",
        "user_name": u["name"],
        "rows": rows, "total": total, "gl": gl, "date_to": date_to,
    })


REF_TYPE_LABELS = {
    "open": "기초재고",
    "gr": "입고",
    "ship": "출고",
    "wo_material": "생산출고",
    "wo_product": "생산입고",
    "adj": "재고조정",
    "transfer": "재고이동",
}

TRANS_TYPE_CLASS = {
    "in": "bg-success",
    "out": "bg-danger",
    "adjust": "bg-warning text-dark",
}


@router.get("/stock/ledger", response_class=HTMLResponse)
async def stock_ledger(request: Request, item_id: int = 0, wh: int = 0,
                       date_from: str = "", date_to: str = "", page: int = 1,
                       u: dict = Depends(require_book), conn=Depends(get_db)):
    book = u["book"]
    warehouses = [dict(r) for r in conn.execute(
        "SELECT id, code, name FROM warehouses WHERE company_id=? AND is_active=1 ORDER BY code", (book,)
    ).fetchall()]

    if not date_from:
        date_from = f"{today_kst().year}-01-01"
    if not date_to:
        date_to = today_kst().isoformat()

    base_url = f"/stock/ledger?item_id={item_id}&wh={wh}&date_from={date_from}&date_to={date_to}"

    if not item_id:
        where = "WHERE i.company_id=? AND st.trans_date BETWEEN ? AND ?"
        params = [book, date_from, date_to]
        if wh:
            where += " AND st.warehouse_id=?"
            params.append(wh)

        sql = f"""
            SELECT i.id as item_id, i.item_code, i.item_name, i.unit,
                   SUM(CASE WHEN st.trans_type='in' THEN st.quantity ELSE 0 END) as total_in,
                   SUM(CASE WHEN st.trans_type='out' THEN st.quantity ELSE 0 END) as total_out
            FROM stock_transactions st
            JOIN items i ON st.item_id=i.id
            {where}
            GROUP BY i.id, i.item_code, i.item_name, i.unit
            ORDER BY i.item_code
        """
        all_rows = [dict(r) for r in conn.execute(sql, params).fetchall()]

        for r in all_rows:
            bp = [date_from]
            bw = ""
            if wh:
                bw = " AND warehouse_id=?"
                bp.append(wh)
            before = conn.execute(
                f"""SELECT COALESCE(SUM(CASE WHEN trans_type='in' THEN quantity ELSE 0 END)
                    - SUM(CASE WHEN trans_type='out' THEN quantity ELSE 0 END), 0)
                    FROM stock_transactions WHERE item_id=? AND trans_date < ?{bw}""",
                [r["item_id"]] + bp,
            ).fetchone()[0]
            r["opening"] = before
            r["closing"] = before + r["total_in"] - r["total_out"]

        total = len(all_rows)
        total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
        page = max(1, min(page, total_pages))
        rows = all_rows[(page - 1) * PER_PAGE:page * PER_PAGE]

        return templates.TemplateResponse(request=request, name="erp/stock_ledger.html", context={
            "request": request, "page_title": "수불부",
            "user_name": u["name"],
            "mode": "summary", "rows": rows,
            "warehouses": warehouses, "wh": wh,
            "date_from": date_from, "date_to": date_to,
            "page": page, "total_pages": total_pages, "total_count": total, "base_url": base_url,
        })

    # 품목 상세
    item = conn.execute("SELECT * FROM items WHERE id=? AND company_id=?", (item_id, book)).fetchone()
    if not item:
        return RedirectResponse("/stock/ledger", status_code=303)
    item = dict(item)

    where = "WHERE st.item_id=? AND st.trans_date BETWEEN ? AND ?"
    params = [item_id, date_from, date_to]
    if wh:
        where += " AND st.warehouse_id=?"
        params.append(wh)

    total = conn.execute(f"SELECT COUNT(*) FROM stock_transactions st {where}", params).fetchone()[0]
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)
    page = max(1, min(page, total_pages))

    sql = f"""
        SELECT st.*, w.name as wh_name
        FROM stock_transactions st
        JOIN warehouses w ON st.warehouse_id=w.id
        {where}
        ORDER BY st.trans_date, st.id
        LIMIT {PER_PAGE} OFFSET {(page - 1) * PER_PAGE}
    """
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    for r in rows:
        r["type_label"] = REF_TYPE_LABELS.get(r["ref_type"], r["trans_type"])
        r["type_class"] = TRANS_TYPE_CLASS.get(r["trans_type"], "")

    bp = [item_id, date_from]
    bw = ""
    if wh:
        bw = " AND warehouse_id=?"
        bp.append(wh)
    opening = conn.execute(
        f"""SELECT COALESCE(SUM(CASE WHEN trans_type='in' THEN quantity ELSE 0 END)
            - SUM(CASE WHEN trans_type='out' THEN quantity ELSE 0 END), 0)
            FROM stock_transactions WHERE item_id=? AND trans_date < ?{bw}""",
        bp,
    ).fetchone()[0]

    sp = [item_id, date_from, date_to]
    if wh:
        sp.append(wh)
    summary = conn.execute(
        f"""SELECT COALESCE(SUM(CASE WHEN trans_type='in' THEN quantity ELSE 0 END), 0) as total_in,
                   COALESCE(SUM(CASE WHEN trans_type='out' THEN quantity ELSE 0 END), 0) as total_out
            FROM stock_transactions WHERE item_id=? AND trans_date BETWEEN ? AND ?{bw}""",
        sp,
    ).fetchone()

    return templates.TemplateResponse(request=request, name="erp/stock_ledger.html", context={
        "request": request, "page_title": f"수불부 - {item['item_name']}",
        "user_name": u["name"],
        "mode": "detail", "item": item, "rows": rows,
        "opening": opening, "total_in": summary["total_in"], "total_out": summary["total_out"],
        "closing": opening + summary["total_in"] - summary["total_out"],
        "warehouses": warehouses, "wh": wh,
        "date_from": date_from, "date_to": date_to,
        "page": page, "total_pages": total_pages, "total_count": total, "base_url": base_url,
    })
